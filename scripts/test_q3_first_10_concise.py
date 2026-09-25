import asyncio
import json
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q3_QUESTION_TEXT = "การใช้ลิ้งค์ลิสต์เป็นสแตกและคิวจะแตกต่างจากการใช้อาร์เรย์อย่างไร และมีข้อดีข้อเสียอย่างไร (1 คะแนน)"

Q3_ANSWER_KEY = (
    "เฉลยและแนวคำตอบมาตรฐาน:\n"
    "1. ความแตกต่าง (Difference):\n"
    "   - Array: มีขนาดคงที่ (Fixed-size / Static) ต้องจองพื้นที่ล่วงหน้า การทำ Stack/Queue ต้องคุม index (top, front, rear)\n"
    "   - Linked List: มีขนาดปรับเปลี่ยนได้ตามจริง (Dynamic / ยืดหยุ่น) ไม่ต้องกำหนดขนาดล่วงหน้า จัดการผ่าน Node และ Pointer (head, tail)\n"
    "2. ข้อดี-ข้อเสีย (Pros & Cons):\n"
    "   - Linked List: ข้อดีคือไม่จำกัดขนาด ไม่เกิด Stack Overflow ได้ง่าย เพิ่ม/ลดโหนดได้สะดวก; "
    "ข้อเสียคือเข้าถึงข้อมูลช้ากว่า (Sequential access) และเปลืองหน่วยความจำเก็บพอยน์เตอร์\n"
    "   - Array: ข้อดีคือเข้าถึงข้อมูลได้รวดเร็วทันทีด้วย Index (O(1)); "
    "ข้อเสียคือขนาดคงที่เสี่ยงเกิด Overflow หากจองน้อยไป หรือเปลืองพื้นที่หากจองมากเกินไป"
)

# เกณฑ์กระชับ เป็นกลาง ไร้รหัสนิสิต และไม่อ้างอิงประโยคเฉพาะบุคคล
Q3_RUBRIC_NAME = "การเปรียบเทียบ Linked List vs Array สำหรับ Stack และ Queue พร้อมข้อดีข้อเสีย"

Q3_RUBRIC_DESC = (
    "เกณฑ์การประเมิน 5 ระดับ (คะแนนเต็ม 1.00):\n"
    "• 1.00 คะแนน (สมบูรณ์): ตอบครบ 2 ส่วนหลักอย่างชัดเจน ได้แก่ (1) อธิบายความแตกต่างของโครงสร้าง/ขนาด (Dynamic vs Static/Fixed หรือไม่ต้องประกาศขนาดล่วงหน้า) หรือเชื่อมโยงการทำงานของ Stack/Queue และ (2) ระบุทั้ง 'ข้อดี' และ 'ข้อเสีย' อย่างครบถ้วน (อนุโลมสำนวนภาษาตามระดับนักศึกษา ไม่ต้องเคร่งครัดศัพท์ pointer/index หากเนื้อหาสื่อข้อดีข้อเสียได้สมเหตุสมผล)\n"
    "• 0.75 คะแนน (ดีมาก): ตอบความแตกต่างได้ถูกต้อง แต่ระบุข้อดีหรือข้อเสียเพียงด้านเดียว (เช่น มีข้อดีแต่ไม่มีข้อเสีย) หรืออธิบายข้อดีข้อเสียกว้างๆ ในเชิงการเขียนโปรแกรม\n"
    "• 0.50 คะแนน (มาตรฐาน): ตอบถูกเพียงบางส่วน ได้แก่ (1) เปรียบเทียบโครงสร้างทั่วไปแต่ไม่ได้แจกแจงแยกข้อดีข้อเสีย, (2) เขียนรวมเป็นความเรียงเดียว, (3) ตอบเฉพาะข้อดีหรือข้อเสียเพียงด้านเดียวสั้นๆ, หรือ (4) อธิบายเฉพาะพฤติกรรม LIFO/FIFO ของ Stack/Queue โดยข้อดีข้อเสียไม่ชัดเจน\n"
    "• 0.25 คะแนน (เล็กน้อย): คำตอบสั้นมาก คลุมเครือ ไม่ระบุชื่อโครงสร้าง แต่มีคีย์เวิร์ดที่ถูกต้องตามหลักการเพียงจุดเดียว (เช่น ยืดหยุ่น หรือ static)\n"
    "• 0.00 คะแนน (ไม่ได้คะแนน): ตอบไม่ตรงประเด็น (เช่น ช่วยประหยัดพลังงาน, รถไฟ), ตอบสั้นเพียงประโยคเดียวโดยไม่มีข้อดีข้อเสีย, ระบุข้อดีข้อเสียที่ผิดหลักการอย่างมีนัยสำคัญ, หรือไม่ตอบ\n\n"
    "(ต้องระบุ teacher_feedback แจกแจงจุดให้/หักคะแนนสำหรับผู้สอน และ student_feedback เป็นคำแนะนำพัฒนาการเรียนรู้สำหรับนักเรียน)"
)

Q3_RUBRICS = [
    {
        "name": Q3_RUBRIC_NAME,
        "score": 1.0,
        "description": Q3_RUBRIC_DESC,
    }
]

async def main():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws_data = wb["ชุดข้อสอบ_dataset"]
    
    # ดึง 10 คนแรกของข้อ 3: แถว 74 ถึง 83 (DS-069 ถึง DS-078)
    first_10 = []
    for r in range(74, 84):
        sid = ws_data.cell(r, 1).value
        q_no = ws_data.cell(r, 2).value
        assert q_no == 3
        ans = str(ws_data.cell(r, 6).value or "")
        h = float(ws_data.cell(r, 7).value or 0.0)
        old_ai = float(ws_data.cell(r, 8).value or 0.0)
        first_10.append({
            "row": r,
            "sample_id": sid,
            "student_ans": ans,
            "human_score": h,
            "old_ai_score": old_ai
        })
        
    print(f"=== กำลังส่งทดสอบ 10 ข้อแรก (DS-069 ถึง DS-078) ด้วยเกณฑ์กระชับฉบับมาตรฐาน ===")
    print(f"จำนวนที่ส่งตรวจ: {len(first_10)} ข้อ\n")
    
    sem = asyncio.Semaphore(5)
    
    async def grade_one(item):
        async with sem:
            res = await score_with_openai(
                question_text=Q3_QUESTION_TEXT,
                answer_text=item["student_ans"],
                max_score=1.0,
                answer_key=Q3_ANSWER_KEY,
                rubrics=Q3_RUBRICS,
                image_bytes_list=None,
                strict_rubric_enforcement=True,
            )
            new_ai = float(res.get("score", 0.0))
            diff = round(new_ai - item["human_score"], 2)
            match = (new_ai == item["human_score"])
            
            return {
                "row": item["row"],
                "sample_id": item["sample_id"],
                "student_ans": item["student_ans"],
                "human_score": item["human_score"],
                "old_ai_score": item["old_ai_score"],
                "new_ai_score": new_ai,
                "confidence": res.get("confidence", "high"),
                "diff": diff,
                "is_exact": match,
                "teacher_feedback": res.get("teacher_feedback", ""),
                "student_feedback": res.get("student_feedback", "")
            }

    tasks = [grade_one(item) for item in first_10]
    results = await asyncio.gather(*tasks)
    results = sorted(results, key=lambda x: x["row"])
    
    print("-" * 80)
    print(f"{'Sample ID':<10} | {'คะแนนอาจารย์':<12} | {'คะแนน AI':<10} | {'ผลต่าง':<8} | {'สถานะ':<10}")
    print("-" * 80)
    for r in results:
        status = "🟢 ตรงเป๊ะ" if r["is_exact"] else f"🟠 ต่าง ({r['diff']:+0.2f})"
        print(f"{r['sample_id']:<10} | {r['human_score']:<12.2f} | {r['new_ai_score']:<10.2f} | {r['diff']:<+8.2f} | {status}")
        
    exact_count = sum(1 for x in results if x["is_exact"])
    w025_count = sum(1 for x in results if abs(x["diff"]) <= 0.2501)
    mae = sum(abs(x["diff"]) for x in results) / len(results)
    
    print("-" * 80)
    print(f"สรุป 10 ข้อแรก:")
    print(f"• ตรงกันเป๊ะ (Exact Match): {exact_count} / 10 ({exact_count * 10}%)")
    print(f"• ต่างไม่เกิน ±0.25 pt:       {w025_count} / 10 ({w025_count * 10}%)")
    print(f"• MAE เฉลี่ย:                {mae:.4f} คะแนน")
    print("=" * 80)
    
    out_file = ROOT / "artifacts" / "q3_first10_concise_test_results.json"
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"\nบันทึกรายละเอียดคำตอบและ Feedback ลงใน {out_file.name}")

if __name__ == "__main__":
    asyncio.run(main())
