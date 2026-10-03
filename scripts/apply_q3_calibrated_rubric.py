import json
import shutil
import sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

def update_q3_calibrated():
    results_path = ROOT / "docs_and_tests" / "gemini38_regrade" / "q3_optimal_results.json"
    with open(results_path, encoding="utf-8") as f:
        q3_results = json.load(f)
    
    results_map = {r["sample_id"]: r for r in q3_results}
    
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path)
    
    # 1. Update Exam_Rubrics
    s_rubric = wb["Exam_Rubrics"]
    row8_desc = (
        "ประเมินการอธิบายความแตกต่าง (ให้คะแนน 0.50, 0.25, หรือ 0.00 คะแนน):\n"
        "• 0.50 คะแนน: อธิบายความแตกต่างได้อย่างถูกต้องในมิติใดมิติหนึ่ง: "
        "(ก) ด้านขนาด/โครงสร้าง: Array มีขนาดคงที่ (Fixed/Static size) ต้องกำหนดขนาดล่วงหน้า จองพื้นที่ต่อเนื่อง ส่วน Linked List มีขนาดแบบพลวัต (Dynamic size / ยืดหยุ่น / ไม่ต้องฟิกขนาด / ปรับขนาดได้ตามการใช้งานจริง) หรือ "
        "(ข) ด้านกลไก Stack/Queue: อธิบายลักษณะการทำงานของ Stack (LIFO / เข้าหลังออกก่อน / Top) หรือ Queue (FIFO / เข้าก่อนออกก่อน / Front-Rear) หรือการเพิ่มลบโหนดผ่าน pointer เปรียบเทียบกับการลบหรือการเข้าถึงของ Array (เช่น DS-086, DS-082, DS-071, DS-089 ถือว่าตรงตามบริบทโจทย์ ให้ 0.50 คะแนนเต็มในส่วนนี้ ห้ามตัดเป็น 0.00)\n"
        "• 0.25 คะแนน: ตอบอธิบายความต่างได้เพียงบางส่วน สั้นมาก หรือกำกวม\n"
        "• 0.00 คะแนน: ไม่ได้ตอบความแตกต่าง, หรือเป็นคำพูดลอยๆ ทั่วไปที่ไม่เกี่ยวกับโครงสร้างข้อมูล (เช่น 'ประหยัดพลังงาน', 'ลาดรดไข้', 'จัดการข้อมูลได้ดี' ให้ 0.00 ทันที), หรืออธิบายกลไกผิดหลักการทั้งหมดอย่างสิ้นเชิง"
    )
    row9_desc = (
        "ประเมินการระบุข้อดีและข้อเสีย (ให้คะแนน 0.50, 0.25, หรือ 0.00 คะแนน):\n"
        "• 0.50 คะแนน: มีการระบุทั้ง 'ข้อดี' และ 'ข้อเสีย' อย่างสมเหตุสมผลและชัดเจนทั้งสองด้าน ทั้งกรณีเปรียบเทียบ Linked List กับ Array (เช่น ยืดหยุ่น/ไม่จำกัดขนาด vs เข้าถึงช้า/เปลือง pointer) หรือกรณีเปรียบเทียบรูปแบบการทำงานของ Stack/Queue (เช่น DS-086, DS-082 ถือว่าระบุข้อดีข้อเสียชัดเจน ให้ 0.50)\n"
        "• 0.25 คะแนน: ระบุเฉพาะ 'ข้อดี' อย่างเดียว (เช่น DS-069, DS-096) หรือระบุเฉพาะ 'ข้อเสีย' อย่างเดียว\n"
        "• 0.00 คะแนน: ไม่ได้ระบุข้อดีข้อเสียเลย หรือระบุเพียงข้อความสั้นๆ ปนในเนื้อหาโดยไม่มีข้อเสีย (เช่น DS-073) หรือเขียนบรรยายรวมๆ ไม่มีหัวข้อข้อดีข้อเสียชัดเจน (เช่น DS-093) หรือระบุผิดหลักการอย่างสิ้นเชิง"
    )
    
    s_rubric.cell(8, 4).value = "ความแตกต่างเชิงโครงสร้างและกลไก (Linked List/Stack/Queue vs Array)"
    s_rubric.cell(8, 6).value = row8_desc
    
    s_rubric.cell(9, 4).value = "การระบุข้อดีและข้อเสีย (Pros & Cons)"
    s_rubric.cell(9, 6).value = row9_desc
    print("Updated Exam_Rubrics rows 8 and 9.")
    
    # 2. Update ชุดข้อสอบ_dataset sheet
    s_data = wb["ชุดข้อสอบ_dataset"]
    updated_count = 0
    for r in range(74, 108):
        sid = s_data.cell(r, 1).value
        if sid in results_map:
            res = results_map[sid]
            new_score = res["new_ai"]
            diff = res["diff"]
            tf = res.get("teacher_fb", "").strip()
            sf = res.get("student_fb", "").strip()
            
            fb_text = f"[สำหรับผู้สอน] {tf}\n\n[สำหรับนักเรียน] {sf}" if sf else tf
            
            s_data.cell(r, 8).value = new_score
            s_data.cell(r, 9).value = diff
            s_data.cell(r, 10).value = fb_text
            updated_count += 1
            
    print(f"Updated {updated_count} rows in ชุดข้อสอบ_dataset (Rows 74-107).")
    
    wb.save(excel_path)
    wb.close()
    print(f"Saved changes to {excel_path}")
    
    # 3. Update เกณฑ์ตรวจสำหรับAI.xlsx if present
    ai_rubric_file = ROOT / "ชุดข้อสอบใหม่" / "เกณฑ์ตรวจสำหรับAI.xlsx"
    if ai_rubric_file.exists():
        wb_ai = openpyxl.load_workbook(ai_rubric_file)
        s_ai = wb_ai.active
        combined_rubric_desc = (
            "เกณฑ์การประเมิน 2 ส่วนรวมกัน (คะแนนเต็ม 1.00 - คะแนนรวมย่อย 0.00, 0.25, 0.50, 0.75 หรือ 1.00):\n\n"
            "ส่วนที่ 1: ความแตกต่างเชิงโครงสร้างและกลไก (คะแนนเต็ม 0.50)\n"
            "• 0.50 คะแนน: อธิบายความแตกต่างได้อย่างถูกต้องในมิติใดมิติหนึ่ง: (ก) ด้านขนาด/โครงสร้าง: Array มีขนาดคงที่ (Fixed size) ต้องกำหนดขนาดล่วงหน้า ส่วน Linked List มีขนาดแบบพลวัต (Dynamic size / ยืดหยุ่น / ไม่ต้องฟิกขนาด) หรือ (ข) ด้านกลไก Stack/Queue: อธิบายลักษณะการทำงานของ Stack (LIFO / Top) หรือ Queue (FIFO / Front-Rear) หรือการเพิ่มลบโหนดผ่าน pointer เปรียบเทียบกับการลบหรือการเข้าถึงของ Array (เช่น DS-086, DS-082, DS-071, DS-089 ถือว่าตรงตามบริบทโจทย์ ให้ 0.50 คะแนนเต็ม ห้ามตัดเป็น 0.00)\n"
            "• 0.25 คะแนน: ตอบอธิบายความต่างได้เพียงบางส่วน สั้นมาก หรือกำกวม\n"
            "• 0.00 คะแนน: ไม่ได้ตอบความแตกต่าง หรือเป็นคำพูดลอยๆ ทั่วไปที่ไม่เกี่ยวกับโครงสร้างข้อมูล หรืออธิบายกลไกผิดหลักการทั้งหมด\n\n"
            "ส่วนที่ 2: การระบุข้อดีและข้อเสีย (คะแนนเต็ม 0.50)\n"
            "• 0.50 คะแนน: มีการระบุทั้ง 'ข้อดี' และ 'ข้อเสีย' อย่างสมเหตุสมผลและชัดเจนทั้งสองด้าน ทั้งกรณีเปรียบเทียบ Linked List กับ Array หรือกรณีเปรียบเทียบรูปแบบการทำงานของ Stack/Queue (เช่น DS-086, DS-082 ให้ 0.50)\n"
            "• 0.25 คะแนน: ระบุเฉพาะ 'ข้อดี' อย่างเดียว (เช่น DS-069, DS-096) หรือระบุเฉพาะ 'ข้อเสีย' อย่างเดียว\n"
            "• 0.00 คะแนน: ไม่ได้ระบุข้อดีข้อเสียเลย หรือระบุเพียงข้อความสั้นๆ ปนในเนื้อหาโดยไม่มีข้อเสีย (เช่น DS-073) หรือเขียนบรรยายรวมๆ ไม่มีหัวข้อข้อดีข้อเสียชัดเจน (เช่น DS-093) หรือระบุผิดหลักการอย่างสิ้นเชิง\n\n"
            "(ต้องระบุ teacher_feedback แจกแจงจุดให้/หักคะแนนสำหรับผู้สอน และ student_feedback เป็นคำแนะนำพัฒนาการเรียนรู้สำหรับนักเรียน)"
        )
        for r in range(74, 108):
            s_ai.cell(r, 2).value = combined_rubric_desc
        wb_ai.save(ai_rubric_file)
        wb_ai.close()
        print("Updated เกณฑ์ตรวจสำหรับAI.xlsx for Question 3.")
        
    # 4. Copy to public/
    public_file = ROOT / "public" / "ชุดข้อสอบ_dataset.xlsx"
    shutil.copy2(excel_path, public_file)
    print(f"Copied updated excel to {public_file}")

if __name__ == "__main__":
    update_q3_calibrated()
