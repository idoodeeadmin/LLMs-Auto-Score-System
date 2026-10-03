import shutil
import sys
from pathlib import Path
import openpyxl
import json

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding='utf-8')

EXCEL_PATH = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
BACKUP_PATH = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset_backup_before_optimal34.xlsx"
JSON_PATH = ROOT / "scratch" / "optimal_34_graded.json"

def main():
    print(f"1. Creating backup: {BACKUP_PATH.name}...")
    shutil.copyfile(EXCEL_PATH, BACKUP_PATH)

    print("2. Loading optimal 34 graded data...")
    with open(JSON_PATH, "r", encoding="utf-8") as f:
        data = json.load(f)

    assert len(data) == 34, f"Expected 34 items, got {len(data)}"

    print(f"3. Loading workbook: {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb["ชุดข้อสอบ_dataset"]

    print("4. Updating rows 6 to 39 for Question 1...")
    for i, item in enumerate(data):
        row = i + 6
        expected_id = f"DS-{i+1:03d}"
        
        ws.cell(row, 1, expected_id)
        ws.cell(row, 2, 1)
        ws.cell(row, 3, "text")
        ws.cell(row, 4, "อธิบายความต่างของ Row-major vs Column-major")
        ws.cell(row, 5, "text")
        ws.cell(row, 6, item["student_answer"])
        ws.cell(row, 7, float(item["human_score"]))
        ws.cell(row, 8, float(item["ai_score"]))
        ws.cell(row, 9, item.get("confidence", "high"))
        
        fb = item.get("feedback", "")
        if not fb:
            tf = item.get("teacher_feedback", "")
            sf = item.get("student_feedback", "")
            fb = f"[สำหรับผู้สอน]\n{tf}\n\n[สำหรับนักเรียน]\n{sf}"
        ws.cell(row, 10, fb.strip())

    print(f"5. Saving workbook...")
    wb.save(EXCEL_PATH)
    print("Update complete!")

    # Verify check
    print("\n6. Verifying updated Excel (Q1 rows 6-39):")
    wb_check = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws_check = wb_check["ชุดข้อสอบ_dataset"]
    
    exact = 0
    total = 0
    for r in range(6, 40):
        sid = ws_check.cell(r, 1).value
        h = float(ws_check.cell(r, 7).value or 0)
        ai = float(ws_check.cell(r, 8).value or 0)
        ans = (ws_check.cell(r, 6).value or "").replace("\n", " ")[:35]
        match = "MATCH" if h == ai else f"DIFF({ai-h:+.1f})"
        print(f"Row {r:02d} | {sid} | H:{h:.1f} AI:{ai:.1f} | {match} | {ans}")
        total += 1
        if h == ai:
            exact += 1

    print(f"\nQ1 Exact Match: {exact}/{total} ({exact/total*100:.1f}%)")
    
    print("\n7. Check Row 40 (Question 2 untouched):")
    print(f"Row 40 | {ws_check.cell(40, 1).value} | Q{ws_check.cell(40, 2).value} | H:{ws_check.cell(40, 7).value} | AI:{ws_check.cell(40, 8).value}")

if __name__ == "__main__":
    main()
