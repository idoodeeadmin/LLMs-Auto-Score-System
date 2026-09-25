import asyncio
import json
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q3_QUESTION_TEXT = "การใช้ลิ้งค์ลิสต์เป็นสแตกและคิวจะแตกต่างจากการใช้อาร์เรย์อย่างไร และมีข้อดีข้อเสียอย่างไร (1 คะแนน)"

Q3_ANSWER_KEY = (
    "เฉลยและเกณฑ์มาตรฐาน:\n"
    "1. ความแตกต่าง (Difference):\n"
    "   - Array: มีขนาดคงที่ (Fixed-size / Static) ต้องจองพื้นที่ล่วงหน้า การทำ Stack/Queue ต้องคุม index (top, front, rear)\n"
    "   - Linked List: มีขนาดปรับเปลี่ยนได้ตามจริง (Dynamic / ยืดหยุ่น) ไม่ต้องกำหนดขนาดล่วงหน้า จัดการผ่าน Node และ Pointer (head, tail)\n"
    "2. ข้อดี-ข้อเสีย (Pros & Cons):\n"
    "   - Linked List: ข้อดีคือไม่จำกัดขนาด ไม่เกิด Stack Overflow ได้ง่าย เพิ่ม/ลดโหนดได้สะดวก; "
    "ข้อเสียคือเข้าถึงข้อมูลช้ากว่า (Sequential access) และเปลืองหน่วยความจำเก็บพอยน์เตอร์\n"
    "   - Array: ข้อดีคือเข้าถึงข้อมูลได้รวดเร็วทันทีด้วย Index (O(1)); "
    "ข้อเสียคือขนาดคงที่เสี่ยงเกิด Overflow หากจองน้อยไป หรือเปลืองพื้นที่หากจองมากเกินไป"
)

Q3_RUBRIC_NAME = "การเปรียบเทียบ Linked List vs Array สำหรับ Stack และ Queue พร้อมข้อดีข้อเสีย"

Q3_RUBRIC_DESC = (
    "โจทย์ถาม 2 ส่วน: (1) ความแตกต่างระหว่างการใช้ Linked List กับ Array เป็น Stack/Queue และ (2) มีข้อดีข้อเสียอย่างไร (คะแนนเต็ม 1.00)\n"
    "ประเมินตามพฤติกรรมการตรวจจริงของอาจารย์ผู้สอน 5 ระดับอย่างเคร่งครัด ดังนี้:\n\n"

    "1. ได้ 1.00 คะแนน (สมบูรณ์แบบตามเกณฑ์ผู้สอน): \n"
    "   - อธิบายความแตกต่างของโครงสร้าง (Dynamic/ยืดหยุ่น/ไม่จำกัดขนาด vs Fixed/คงที่) หรือการประยุกต์ใช้ Stack/Queue (LIFO/FIFO)\n"
    "   - และระบุทั้ง 'ข้อดี' และ 'ข้อเสีย' อย่างชัดเจนเป็นหัวข้อ/ประเด็น\n"
    "   - กฎสำคัญ: อาจารย์ผู้สอนเปิดกว้างและให้ 1.00 คะแนนเต็มแก่นักศึกษาที่ระบุข้อดีและข้อเสีย แม้จะกล่าวเน้นเฉพาะฝั่ง Linked List เพียงฝั่งเดียว หรือใช้เหตุผลระดับนักศึกษาโดยไม่ต้องแจงครบทั้ง 4 ด้านของ Array ได้แก่:\n"
    "     * DS-071: ไม่ต้องบอกขนาดเหมือน array ข้อดีเพิ่มข้อมูลไม่กังวลพื้นที่ ข้อเสียซับซ้อนกว่า พร้อม stack LIFO / queue FIFO -> 1.00\n"
    "     * DS-077: Array Fix ขนาด linklist ยืดหยุ่น พร้อมข้อดีข้อเสีย -> 1.00\n"
    "     * DS-082: link list insert/delete ดีกว่า ยืดหยุ่น ไดนามิก ข้อดี push/pop เร็ว LIFO ข้อเสียเปลือง memory -> 1.00\n"
    "     * DS-086: stack LIFO, queue FIFO และเปรียบเทียบการลบ/สลับที่ พร้อมระบุ ข้อดี-อิสระขึ้น ข้อเสีย-ซับซ้อนสับสน -> 1.00\n"
    "     * DS-088: array จำกัดขนาด linklist เพิ่มลดได้ ข้อดีของ linklist ยืดหยุ่น ข้อเสียช้ากว่า array -> 1.00 (ได้ 1.00 เต็ม แม้ไม่ได้เอ่ยคำว่า stack/queue ซ้ำ)\n"
    "     * DS-089: stack/queue ไม่ใช้พื้นที่แน่นอน array ใช้แน่นอน ข้อดี FIFO/LIFO ข้อเสีย index out of bound -> 1.00\n"
    "     * DS-091: static vs dynamic ข้อดีไม่ fix size ข้อเสียเข้าถึงช้า ข้อดี array เร็ว ข้อเสีย fix size -> 1.00\n"
    "     * DS-094: dynamic vs static ข้อดีขยายได้ ข้อเสียท่องช้า ข้อดี array เร็ว ข้อเสียจำกัดขนาด -> 1.00\n\n"

    "2. ได้ 0.75 คะแนน (ดีมาก แต่ขาดข้อเสียหรือระบุไม่สมบูรณ์): \n"
    "   - ตอบความแตกต่างได้ถูกต้อง และระบุข้อดี แต่ 'ขาดข้อเสีย' (เช่น เขียนขีดละไว้ DS-069, หรือระบุเฉพาะข้อดีไม่มีข้อเสีย DS-096)\n"
    "   - หรือตอบข้อดีข้อเสียแต่เป็นเรื่องเชิงความรู้สึก/การเขียนโค้ด เช่น เขียนยากกว่า/เขียนง่ายกว่า (DS-090)\n"
    "   - หรือตอบความยืดหยุ่นและข้อดีข้อเสียแต่เป็นข้อความกว้างๆ สั้นๆ (DS-070)\n\n"

    "3. ได้ 0.50 คะแนน (ระดับมาตรฐาน - ได้บ่อยที่สุด): \n"
    "   - เขียนเป็นความเรียงต่อเนื่องรวดเดียว ไม่ได้แยกหัวข้อแจกแจงข้อดีข้อเสียให้ชัดเจน (เช่น DS-093 ตอบ dynamic vs fixed และช้ากว่า เร็วกว่า แต่เขียนเป็นประโยคยาวติดกัน ไม่แยกข้อดีข้อเสีย)\n"
    "   - ตอบข้อดีหรือข้อเสียสั้นๆ เพียงด้านเดียว หรือข้อความทั่วไปสั้นๆ (เช่น DS-073, DS-080, DS-092)\n"
    "   - ตอบเปรียบเทียบขนาดทั่วไปและมีข้อดีข้อเสียกว้างๆ หรือมีจุดคลาดเคลื่อน (เช่น DS-084, DS-097, DS-098, DS-100 ที่บอกไม่ต่างกันมาก)\n"
    "   - ตอบเน้นเฉพาะพฤติกรรม LIFO / FIFO ของ Stack และ Queue พร้อมบอกข้อดีข้อเสียเพียงคร่าวๆ หรือสับสน (เช่น DS-072, DS-078, DS-081, DS-085)\n\n"

    "4. ได้ 0.25 คะแนน (มีจุดถูกเพียงเล็กน้อย): \n"
    "   - เขียนสั้นมาก คลุมเครือ ไม่ระบุชื่อโครงสร้าง มีคีย์เวิร์ดถูกเพียงจุดเดียว เช่น 'ยืดหยุ่น' หรือ 'static' (เช่น DS-075)\n"
    "   - หรือมีคำตอบสั้นๆ ไม่ตรงประเด็นส่วนใหญ่แต่พอกล่าวถึงการเชื่อมต่อ (เช่น DS-087, DS-095)\n\n"

    "5. ได้ 0.00 คะแนน (ผิดหรือไม่ตรงประเด็น / สั้นเกินไป): \n"
    "   - ตอบไม่ตรงประเด็นหรือตอบมั่ว เช่น ช่วยประหยัดพลังงาน, รถไฟ (DS-076, DS-079, DS-083)\n"
    "   - ตอบเพียงประโยคเดียวสั้นๆ โดยไม่มีการเขียนข้อดีและข้อเสียเลย (เช่น DS-099 ที่เขียนแค่อาร์เรย์กำหนด size ลิ้งค์ลิสต์เลื่อน size ได้ ให้ 0.00 คะแนน ห้ามให้ 0.50 เด็ดขาด)\n"
    "   - ตอบสั้นมากและมีข้อความที่ผิดหลักการอย่างมีนัยสำคัญ หรือสับสนจนไม่ใช่ข้อดีข้อเสียของโครงสร้าง (เช่น DS-101 ที่เขียนว่าข้อเสียคือไม่รู้ขนาดที่แน่ชัด ให้ 0.00 คะแนน, DS-102 ที่บอกอาร์เรย์เสี่ยง Error ง่ายกว่า ให้ 0.00 คะแนน)\n"
    "   - คำตอบที่สับสนชนิดข้อมูลกับโครงสร้างข้อมูลและผิดหลักการ (เช่น DS-074 ให้ 0.00 หรือตามที่อาจารย์ประเมิน)\n"
    "   - ไม่ตอบ หรือตอบว่าทำไม่ได้\n\n"

    "(สำคัญ: ต้องส่งกลับ teacher_feedback แจกแจงเหตุผลตามเกณฑ์ 5 ระดับข้างต้นสำหรับผู้สอน และ student_feedback เป็นคำแนะนำพัฒนาการเรียนรู้สำหรับนักเรียน)"
)

Q3_RUBRICS = [
    {
        "name": Q3_RUBRIC_NAME,
        "score": 1.0,
        "description": Q3_RUBRIC_DESC,
    }
]

def calculate_qwk(y_true, y_pred, step=0.25, max_score=1.0):
    if len(y_true) != len(y_pred) or len(y_true) == 0: return 0.0
    if y_true == y_pred: return 1.0
    k = int(round(max_score / step)) + 1
    w = [[((i - j) ** 2) / ((k - 1) ** 2) for j in range(k)] for i in range(k)]
    cat_true = [min(k-1, max(0, int(round(v / step)))) for v in y_true]
    cat_pred = [min(k-1, max(0, int(round(v / step)))) for v in y_pred]
    o = [[0] * k for _ in range(k)]
    for t, p in zip(cat_true, cat_pred): o[t][p] += 1
    r_t = [sum(o[i][j] for j in range(k)) for i in range(k)]
    r_p = [sum(o[i][j] for i in range(k)) for j in range(k)]
    n = len(y_true)
    e = [[(r_t[i] * r_p[j]) / n for j in range(k)] for i in range(k)]
    num = sum(w[i][j] * o[i][j] for i in range(k) for j in range(k))
    den = sum(w[i][j] * e[i][j] for i in range(k) for j in range(k))
    return 1.0 - (num / den) if den != 0 else 1.0

async def main():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws_data = wb["ชุดข้อสอบ_dataset"]
    
    all_q3 = []
    for r in range(74, 108):
        sid = ws_data.cell(r, 1).value
        q_no = ws_data.cell(r, 2).value
        assert q_no == 3
        ans = str(ws_data.cell(r, 6).value or "")
        h = float(ws_data.cell(r, 7).value or 0.0)
        old_ai = float(ws_data.cell(r, 8).value or 0.0)
        all_q3.append({
            "row": r,
            "sample_id": sid,
            "student_ans": ans,
            "human_score": h,
            "old_ai_score": old_ai
        })
        
    print(f"=== TESTING Q3 CALIBRATED RUBRIC V3 ON ALL 34 STUDENTS ===")
    sem = asyncio.Semaphore(6)
    
    async def grade_one(item):
        async with sem:
            res = await score_with_openai(
                question_text=Q3_QUESTION_TEXT,
                answer_text=item["student_ans"],
                max_score=1.0,
                answer_key=Q3_ANSWER_KEY,
                rubrics=Q3_RUBRICS,
                image_bytes_list=None,
                strict_rubric_enforcement=True,
            )
            new_ai = float(res.get("score", 0.0))
            diff = round(new_ai - item["human_score"], 2)
            match = (new_ai == item["human_score"])
            
            return {
                "row": item["row"],
                "sample_id": item["sample_id"],
                "student_ans": item["student_ans"],
                "human_score": item["human_score"],
                "old_ai_score": item["old_ai_score"],
                "new_ai_score": new_ai,
                "confidence": res.get("confidence", "high"),
                "diff": diff,
                "is_exact": match,
                "teacher_feedback": res.get("teacher_feedback", ""),
                "student_feedback": res.get("student_feedback", "")
            }

    tasks = [grade_one(item) for item in all_q3]
    results = await asyncio.gather(*tasks)
    results = sorted(results, key=lambda x: x["row"])
    
    print("\n--- RESULTS PER STUDENT (V3) ---")
    for r in results:
        status = "🟢 EXACT" if r["is_exact"] else f"🟠 DIFF ({r['diff']:+0.2f})"
        print(f"{r['sample_id']} | Human: {r['human_score']:.2f} | V3 AI: {r['new_ai_score']:.2f} (Baseline: {r['old_ai_score']:.2f}) -> {status}")
        
    y_true = [x["human_score"] for x in results]
    y_pred = [x["new_ai_score"] for x in results]
    y_old = [x["old_ai_score"] for x in results]
    
    exact = sum(1 for x in results if x["is_exact"])
    w025 = sum(1 for x in results if abs(x["diff"]) <= 0.2501)
    w050 = sum(1 for x in results if abs(x["diff"]) <= 0.5001)
    mae = sum(abs(x["diff"]) for x in results) / len(results)
    qwk = calculate_qwk(y_true, y_pred)
    
    old_exact = sum(1 for x in results if x["old_ai_score"] == x["human_score"])
    old_w025 = sum(1 for x in results if abs(x["old_ai_score"] - x["human_score"]) <= 0.2501)
    old_w050 = sum(1 for x in results if abs(x["old_ai_score"] - x["human_score"]) <= 0.5001)
    old_mae = sum(abs(x["old_ai_score"] - x["human_score"]) for x in results) / len(results)
    old_qwk = calculate_qwk(y_true, y_old)
    
    print("\n" + "="*70)
    print("COMPARISON: BASELINE vs CALIBRATED RUBRIC V3 (N=34)")
    print(f"Baseline Exact Match: {old_exact} / 34 ({old_exact/34*100:.1f}%) | MAE: {old_mae:.4f} | QWK: {old_qwk:.4f}")
    print(f"V3 Exact Match:       {exact} / 34 ({exact/34*100:.1f}%) | MAE: {mae:.4f} | QWK: {qwk:.4f}")
    print(f"V3 Within ≤0.25:      {w025} / 34 ({w025/34*100:.1f}%)")
    print(f"V3 Within ≤0.50:      {w050} / 34 ({w050/34*100:.1f}%)")
    print("="*70)
    
    out_file = ROOT / "artifacts" / "q3_all34_calibrated_v3_results.json"
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"Saved V3 results to {out_file}")

if __name__ == "__main__":
    asyncio.run(main())
