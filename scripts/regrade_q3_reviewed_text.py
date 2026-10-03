"""Regrade the 34 user-reviewed Q3 text answers with the production grader."""
import asyncio
import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")
from scripts.regrade_corrected_q2_q3_answers import OPENAI_MODEL, PROMPT_VERSION, score_with_openai

BOOK = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx"
OUT = ROOT / "docs_and_tests/q3_transcription_review/q3_reviewed_text_regrade.json"
ANSWER_KEY = (
    "Array ต้องกำหนดขนาดล่วงหน้าและใช้ตำแหน่ง index ส่วน Linked List ปรับขนาดได้และใช้โหนดเชื่อมกัน "
    "ข้อดีของ Linked List คือเพิ่มหรือลบข้อมูลได้โดยไม่ต้องจองขนาดคงที่ "
    "ข้อเสียคือใช้หน่วยความจำสำหรับตัวชี้และเข้าถึงตำแหน่งโดยตรงไม่ได้"
)
ALLOWED = [0.0, 0.25, 0.5, 0.75, 1.0]


async def main():
    wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    ws = wb["ชุดข้อสอบ_dataset"]
    rs = wb["Exam_Rubrics"]
    rubric = [
        {"name": rs.cell(r, 4).value, "score": 0.5,
         "description": rs.cell(r, 6).value, "allowed_scores": [0.0, 0.25, 0.5]}
        for r in (8, 9)
    ]
    items = []
    for row_no in range(74, 108):
        values = [ws.cell(row_no, c).value for c in range(1, 9)]
        assert values[0] == f"DS-{row_no - 5:03d}" and values[1] == 3
        assert isinstance(values[5], str) and values[5].strip()
        items.append({"sample_id": values[0], "question": values[3], "answer": values[5],
                      "human_score": float(values[6]), "previous_ai_score": float(values[7])})
    wb.close()
    assert len(items) == 34
    report = {"started_at": datetime.now(timezone.utc).isoformat(),
              "model": OPENAI_MODEL, "prompt_version": PROMPT_VERSION,
              "grading_function": "server.services.openai_grading.score_with_openai",
              "source_sha256": hashlib.sha256(BOOK.read_bytes()).hexdigest(),
              "input_modality": "text_only", "answer_key": ANSWER_KEY,
              "rubric": rubric, "allowed_scores": ALLOWED, "results": []}
    sem = asyncio.Semaphore(3)

    async def grade(item):
        async with sem:
            result = await score_with_openai(
                question_text=item["question"], answer_text=item["answer"], max_score=1.0,
                answer_key=ANSWER_KEY, rubrics=rubric, allowed_scores=ALLOWED)
            success = (not result.get("metrics", {}).get("manual_review_required", True)
                       and float(result.get("score", -1)) in ALLOWED
                       and bool(result.get("teacher_feedback", "").strip())
                       and bool(result.get("student_feedback", "").strip()))
            report["results"].append({**item, "success": success, "result": result})
            OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
            print(item["sample_id"], "score", result.get("score"), "previous", item["previous_ai_score"],
                  "success", success, flush=True)

    print("model", OPENAI_MODEL, "count", len(items), "output", OUT, flush=True)
    await asyncio.gather(*(grade(item) for item in items))
    report["results"].sort(key=lambda x: x["sample_id"])
    valid = [x for x in report["results"] if x["success"]]
    report["summary"] = {
        "completed": len(valid), "total": len(items),
        "new_human_exact": sum(x["result"]["score"] == x["human_score"] for x in valid),
        "old_human_exact": sum(x["previous_ai_score"] == x["human_score"] for x in valid),
        "changed_from_previous": sum(x["result"]["score"] != x["previous_ai_score"] for x in valid),
        "new_mae": sum(abs(x["result"]["score"] - x["human_score"]) for x in valid) / len(valid) if valid else None,
    }
    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report["summary"]), flush=True)


if __name__ == "__main__":
    asyncio.run(main())
