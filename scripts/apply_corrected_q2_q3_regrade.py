"""Apply five verified AI regrades without rewriting workbook images."""

import hashlib
import json
import os
import shutil
import zipfile
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx"
BACKUP = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset_before_q2_q3_regrade.xlsx"
REPORT = ROOT / "docs_and_tests/q2_q3_review/regrade_corrected_results.json"
data = json.loads(REPORT.read_text(encoding="utf-8"))
assert hashlib.sha256(BOOK.read_bytes()).hexdigest() == data["source_sha256"]
items = data["items"]
assert len(items) == 5 and {item["sample_id"] for item in items} == {
    "DS-051", "DS-055", "DS-057", "DS-089", "DS-101"}
assert all(item["success"] for item in items)

wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
sheet = wb["ชุดข้อสอบ_dataset"]
updates = {}
for item in items:
    row = item["row"]
    result = item["result"]
    assert sheet.cell(row, 1).value == item["sample_id"]
    assert sheet.cell(row, 6).value == item["answer"]
    assert sheet.cell(row, 8).value == item["prior_ai_score"]
    assert not result["metrics"]["manual_review_required"]
    assert result["teacher_feedback"].strip() and result["student_feedback"].strip()
    updates[row] = (float(result["score"]), result["confidence"], result["feedback"])
wb.close()
if not BACKUP.exists():
    shutil.copy2(BOOK, BACKUP)

ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ET.register_namespace("", ns)
with zipfile.ZipFile(BOOK) as source:
    root = ET.fromstring(source.read("xl/worksheets/sheet1.xml"))
    cells = {cell.attrib["r"]: cell for cell in root.iter(f"{{{ns}}}c")}
    for row, (score, confidence, feedback) in updates.items():
        for column, value in (("H", score), ("I", confidence), ("J", feedback)):
            cell = cells[f"{column}{row}"]
            for child in list(cell):
                if child.tag in {f"{{{ns}}}v", f"{{{ns}}}is", f"{{{ns}}}f"}:
                    cell.remove(child)
            if column == "H":
                cell.attrib.pop("t", None)
                ET.SubElement(cell, f"{{{ns}}}v").text = str(value)
            else:
                cell.attrib["t"] = "inlineStr"
                inline = ET.SubElement(cell, f"{{{ns}}}is")
                text = ET.SubElement(inline, f"{{{ns}}}t")
                text.attrib["{http://www.w3.org/XML/1998/namespace}space"] = "preserve"
                text.text = value
    xml = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with NamedTemporaryFile(dir=BOOK.parent, suffix=".xlsx", delete=False) as temp:
        temp_name = temp.name
    with zipfile.ZipFile(temp_name, "w") as target:
        for entry in source.infolist():
            target.writestr(entry, xml if entry.filename == "xl/worksheets/sheet1.xml" else source.read(entry.filename))

try:
    with zipfile.ZipFile(temp_name) as archive, zipfile.ZipFile(BACKUP) as original:
        for name in original.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert archive.read(name) == original.read(name), name
    check = openpyxl.load_workbook(temp_name, read_only=True, data_only=True)
    try:
        sheet = check["ชุดข้อสอบ_dataset"]
        assert all(tuple(sheet.cell(row, col).value for col in (8, 9, 10)) == values
                   for row, values in updates.items())
    finally:
        check.close()
    os.replace(temp_name, BOOK)
finally:
    if os.path.exists(temp_name):
        os.unlink(temp_name)
print("Applied five corrected-answer regrades")
