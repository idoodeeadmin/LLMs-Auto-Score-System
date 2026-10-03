"""Update Question 4 scores, confidence, and feedback in ชุดข้อสอบ_dataset.xlsx with Gemini 3.8 Flash 100% results.

Also updates Exam_Rubrics with the Topology-Aware Rubric and syncs dataset viewer.
"""

import json
import shutil
import sys
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
BACKUP_PATH = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset_backup_before_q4_gemini.xlsx"
RESULTS_JSON = ROOT / "docs_and_tests/gemini38_regrade/q4_topology_rubric_results.json"
PUBLIC_EXCEL = ROOT / "public" / "ชุดข้อสอบ_dataset.xlsx"

# 1. Backup original excel
shutil.copy2(EXCEL_PATH, BACKUP_PATH)
print(f"Backed up workbook to: {BACKUP_PATH}")

# 2. Load JSON results
with open(RESULTS_JSON, "r", encoding="utf-8") as f:
    report = json.load(f)

results_by_id = {r["sample_id"]: r for r in report["results"]}
assert len(results_by_id) == 34, f"Expected 34 results, got {len(results_by_id)}"

# 3. Load workbook with openpyxl (preserving formatting)
wb = openpyxl.load_workbook(EXCEL_PATH)
ws_dataset = wb["ชุดข้อสอบ_dataset"]
ws_rubrics = wb["Exam_Rubrics"]

# 4. Update rows 108-141 (DS-103 to DS-136)
updated_count = 0
for row in range(108, 142):
    sid = ws_dataset.cell(row, 1).value
    assert sid in results_by_id, f"Sample ID {sid} not found in results"
    res = results_by_id[sid]
    
    # Col 8: AI Score
    ws_dataset.cell(row, 8).value = int(res["gemini_score"]) if res["gemini_score"].is_integer() else res["gemini_score"]
    
    # Col 9: AI Confidence
    ws_dataset.cell(row, 9).value = res.get("confidence", "high")
    
    # Col 10: Feedback
    tf = res.get("teacher_feedback", "").strip()
    sf = res.get("student_feedback", "").strip()
    full_feedback = f"[สำหรับผู้สอน]\n{tf}\n\n[สำหรับนักเรียน]\n{sf}"
    ws_dataset.cell(row, 10).value = full_feedback
    
    updated_count += 1

print(f"Updated {updated_count} rows in sheet 'ชุดข้อสอบ_dataset'")

# 5. Update Exam_Rubrics Row 10 (Question 4)
Q4_RUBRIC_TEXT = (
    "ประเมินภาพวาด Binary Search Tree จากข้อมูล: 9, 16, 10, 76, 5, 13, 58, 92, 11, 15, 80, 99 (คะแนนเต็ม 1.00 คะแนน, เลือกระดับคะแนน 1.00 หรือ 0.00):\n\n"
    "• 1.00 คะแนน: โครงสร้างความสัมพันธ์ถูกต้องครบทั้ง 12 โหนดตามหลักการ BST (Parent-Child Topology):\n"
    "  - Root คือ 9 (ซ้าย 5, ขวา 16)\n"
    "  - ใต้ 16: ซ้ายคือ 10, ขวาคือ 76\n"
    "  - ใต้ 10: ขวาคือ 13 (ซ้ายว่าง)\n"
    "  - ใต้ 13: ซ้ายคือ 11, ขวาคือ 15\n"
    "  - ใต้ 76: ซ้ายคือ 58 (เป็น Leaf Node), ขวาคือ 92\n"
    "  - ใต้ 92: ซ้ายคือ 80, ขวาคือ 99\n"
    "  (กฎการอนุโลมสำคัญ: ในข้อสอบนี้ โหนด 13 ทางขวาของ 10 และโหนด 58 ทางซ้ายของ 76 จะเข้ามาเบียดกันตรงกลางกระดาษ "
    "  ทำให้นักเรียนส่วนใหญ่วาดเบียดกัน เส้นเชื่อมของ 15 อาจวาดดิ่งลงมาหรือเบี่ยงหลบ 58 ให้ยึดเจตนาการเชื่อมต่อเป็นสำคัญ "
    "  และให้อนุโลมความสวยงาม ลายมือ รอยร่าง หรือวงกลมเปล่าส่วนเกิน ขอเพียงโหนดทั้ง 12 ตัวและเส้นเชื่อมหลักถูกต้อง)\n\n"
    "• 0.00 คะแนน: ผิดหลักการ BST อย่างแท้จริง เช่น โหนด Root ผิด (ไม่ใช่ 9), วางโหนดผิดตำแหน่ง, "
    "  สลับกิ่งซ้าย-ขวา, ลากเส้นเชื่อมผิดสายลำดับ (เช่น นำ 13 หรือ 15 ไปต่อใต้ 58), หรือไม่วาดคำตอบ"
)

ws_rubrics.cell(10, 6).value = Q4_RUBRIC_TEXT
print("Updated Row 10 in sheet 'Exam_Rubrics' with refined Topology-Aware Rubric")

# 6. Save workbook
wb.save(EXCEL_PATH)
wb.close()
print(f"Saved changes to: {EXCEL_PATH}")

# 7. Copy to public/ for web viewer download
shutil.copy2(EXCEL_PATH, PUBLIC_EXCEL)
print(f"Copied updated workbook to: {PUBLIC_EXCEL}")
