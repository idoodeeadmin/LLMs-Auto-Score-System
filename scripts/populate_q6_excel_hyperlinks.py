import openpyxl
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"

wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb["ชุดข้อสอบ_dataset"]

count = 0
for r in range(176, 210):
    std_idx = r - 175
    # col 4: question_content
    if not ws.cell(r, 4).value:
        ws.cell(r, 4, "จงแปลง tree ต่อไปนี้ให้เป็น Binary Tree (1 คะแนน)\n[รูปภาพโจทย์: โจทphoto3/LINE_ALBUM_โจทphoto6_260918_1.jpg]")
        count += 1
    # col 6: student_answer
    if not ws.cell(r, 6).value or ws.cell(r, 6).value == "None":
        img_name = f"LINE_ALBUM_Photo2.2_260918_{std_idx}.jpg"
        ws.cell(r, 6, f'=HYPERLINK("photo_clean_ชุดที่3/{img_name}", "เปิดภาพคำตอบที่ลบคะแนนแล้ว")')
        count += 1

wb.save(EXCEL_PATH)
print(f"Successfully populated {count} missing cells in Q6 rows (176-209) in {EXCEL_PATH}!")
