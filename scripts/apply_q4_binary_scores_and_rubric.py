import json
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[1]
RESULTS_FILE = ROOT / 'artifacts' / 'q4_binary_eval_results.json'
EXCEL_FILE = ROOT / 'ชุดข้อสอบใหม่' / 'ชุดข้อสอบ_dataset.xlsx'
RUBRIC_AI_FILE = ROOT / 'ชุดข้อสอบใหม่' / 'เกณฑ์ตรวจสำหรับAI.xlsx'

def apply_q4():
    print(f"Loading results from {RESULTS_FILE.name}...")
    with open(RESULTS_FILE, encoding='utf-8') as f:
        data = json.load(f)

    results_map = {r['sample_id']: r for r in data['results']}
    print(f"Loaded {len(results_map)} student results for Question 4.")

    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb['ชุดข้อสอบ_dataset']

    thin_border = Border(
        left=Side(style='thin', color='00BFBFBF'),
        right=Side(style='thin', color='00BFBFBF'),
        top=Side(style='thin', color='00BFBFBF'),
        bottom=Side(style='thin', color='00BFBFBF')
    )
    font_aptos = Font(name='Aptos', size=10)
    align_center = Alignment(horizontal='center', vertical='top')
    align_feedback = Alignment(horizontal='left', vertical='top', wrap_text=True)

    # 1. Update ชุดข้อสอบ_dataset rows 108 to 141
    updated_count = 0
    for row in range(108, 142):
        sid = str(ws.cell(row=row, column=1).value or '').strip()
        qno = ws.cell(row=row, column=2).value
        if qno != 4 or not sid or sid not in results_map:
            continue

        res = results_map[sid]
        score = float(res['ai'])
        conf = str(res.get('confidence', 'high'))
        fb_teacher = res.get('teacher_feedback', '').strip()
        fb_student = res.get('student_feedback', '').strip()
        full_fb = f"[สำหรับผู้สอน]\n{fb_teacher}\n\n[สำหรับนักเรียน]\n{fb_student}"

        c8 = ws.cell(row=row, column=8, value=score)
        c8.font = font_aptos
        c8.alignment = align_center
        c8.border = thin_border

        c9 = ws.cell(row=row, column=9, value=conf)
        c9.font = font_aptos
        c9.alignment = align_center
        c9.border = thin_border

        c10 = ws.cell(row=row, column=10, value=full_fb)
        c10.font = font_aptos
        c10.alignment = align_feedback
        c10.border = thin_border

        updated_count += 1

    print(f"Updated {updated_count} student scores and dual feedbacks in 'ชุดข้อสอบ_dataset'")

    # 2. Update Exam_Rubrics
    if 'Exam_Rubrics' in wb.sheetnames:
        ws_rubric = wb['Exam_Rubrics']
        r10_desc = (
            "ประเมินภาพวาด Binary Search Tree จากข้อมูล: 9, 16, 10, 76, 5, 13, 58, 92, 11, 15, 80, 99 (ให้เลือกได้เฉพาะ 1.0 หรือ 0.0 คะแนนเท่านั้น):\n\n"
            "• 1.0 คะแนน: โครงสร้าง BST ถูกต้องครบทั้ง 12 โหนดตามหลักการ BST (โหนดซ้าย < โหนดแม่ < โหนดขวา):\n"
            "  - Root คือ 9 (ซ้าย 5, ขวา 16)\n"
            "  - ใต้ 16 มี 10 ทางซ้าย, 76 ทางขวา\n"
            "  - ใต้ 10 มี 13 ทางขวา\n"
            "  - ใต้ 13 มี 11 ทางซ้าย, 15 ทางขวา\n"
            "  - ใต้ 76 มี 58 ทางซ้าย, 92 ทางขวา\n"
            "  - ใต้ 92 มี 80 ทางซ้าย, 99 ทางขวา\n"
            "  (ให้อนุโลมความสวยงาม ลายมือ ความยาว/มุมเอียงของกิ่ง และรอยร่างดินสอส่วนเกิน ขอเพียงโครงสร้างและค่าของโหนดสื่อสารได้ถูกต้อง)\n\n"
            "• 0.0 คะแนน: ผิดหลักการ BST เช่น วางตำแหน่งโหนดผิด, สลับกิ่งซ้าย-ขวา, โหนดแตกกิ่งเกิน 2 กิ่ง, ตัวเลขตกหล่นไม่ครบ 12 โหนด, หรือไม่วาดคำตอบ"
        )
        ws_rubric.cell(row=10, column=6, value=r10_desc)
        print("Updated Question 4 rubric in 'Exam_Rubrics' sheet (Row 10).")

    wb.save(EXCEL_FILE)
    print(f"Saved {EXCEL_FILE.name} successfully.")

    # 3. Update เกณฑ์ตรวจสำหรับAI.xlsx
    if RUBRIC_AI_FILE.exists():
        wb_ai = openpyxl.load_workbook(RUBRIC_AI_FILE)
        ws_ai = wb_ai.worksheets[0]
        q4_ai_criteria = (
            "ข้อ 4 (1 คะแนน - ให้เฉพาะ 1 หรือ 0 เท่านั้น):\n"
            "• 1 คะแนน: สร้างโครงสร้าง Binary Search Tree (BST) จากข้อมูล 9, 16, 10, 76, 5, 13, 58, 92, 11, 15, 80, 99 ได้ถูกต้องครบทั้ง 12 โหนดตามหลัก BST "
            "(Root 9, ซ้าย 5, ขวา 16 -> ใต้ 16 มี 10 ซ้าย, 76 ขวา -> ใต้ 10 มี 13 ขวา -> ใต้ 13 มี 11 ซ้าย, 15 ขวา -> ใต้ 76 มี 58 ซ้าย, 92 ขวา -> ใต้ 92 มี 80 ซ้าย, 99 ขวา) "
            "อนุโลมลายมือ รอยร่าง หรือความยาวกิ่ง\n"
            "• 0 คะแนน: วางโหนดผิดตำแหน่ง ผิดหลัก BST ไม่ครบ 12 โหนด หรือไม่วาดคำตอบ"
        )
        ai_updated = 0
        for r in range(2, ws_ai.max_row + 1):
            q_val = ws_ai.cell(row=r, column=1).value
            if q_val and 'Binary search tree' in str(q_val) and 'Array 1 มิติ' not in str(q_val):
                ws_ai.cell(row=r, column=2, value=q4_ai_criteria)
                ai_updated += 1
        wb_ai.save(RUBRIC_AI_FILE)
        print(f"Updated {ai_updated} rows in '{RUBRIC_AI_FILE.name}'.")

    print("\nAll Q4 updates completed successfully!")

if __name__ == '__main__':
    apply_q4()
