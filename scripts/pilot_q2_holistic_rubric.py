"""Blind pilot of the proposed holistic four-level Q2 rubric."""

import asyncio
import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")
from scripts.regrade_corrected_q2_q3_answers import OPENAI_MODEL, PROMPT_VERSION, score_with_openai

BOOK = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx"
OUT = ROOT / "docs_and_tests/q2_q3_review/q2_holistic_pilot.json"
RUBRIC = [{
    "name": "การเปรียบเทียบ O(n log n) และ O(n^2) พร้อมตัวอย่าง",
    "score": 2.0,
    "allowed_scores": [0.0, 1.0, 1.5, 2.0],
    "description": (
        "ให้คะแนนองค์รวมโดยพิจารณาเหตุผลทางเทคนิค ไม่ใช้เพียงการนับคำสำคัญ:\n"
        "2.0: อธิบายถูกว่าเมื่อข้อมูลใหญ่ O(n log n) ใช้เวลาหรือจำนวนขั้นตอนเพิ่มช้ากว่า O(n^2) "
        "และยกตัวอย่างอัลกอริทึมหรือกลไกที่ถูกต้องอย่างน้อยหนึ่งอย่าง\n"
        "1.5: สื่อเหตุผลเปรียบเทียบได้ถูกทิศทาง แต่ยังไม่ชัดหรือครบ พร้อมตัวอย่างที่ถูกต้อง "
        "หรืออธิบายกลไกเปรียบเทียบได้ชัดแม้ไม่ระบุชื่ออัลกอริทึม\n"
        "1.0: บอกได้เพียงว่า O(n log n) เร็วกว่า/เหมาะกับข้อมูลใหญ่กว่า โดยเหตุผลยังไม่ชัด "
        "หรือยกตัวอย่างได้ถูกต้องแต่ไม่อธิบายเหตุผล\n"
        "0.0: ไม่ตอบ ไม่เกี่ยวข้อง หรือคำอธิบายหลักผิดจนไม่สื่อความแตกต่างของความซับซ้อนทั้งสอง\n"
        "ตัวอย่างไม่จำเป็นต้องมีทั้งสองฝั่ง: Merge Sort, Heap Sort, Quick Sort กรณีเฉลี่ย "
        "หรือ nested loop ที่วนตาม n สองชั้นเป็นตัวอย่างที่ใช้ได้ตามบริบท; "
        "การเรียกซ้ำหรือ factorial เพียงอย่างเดียวไม่ใช่ตัวอย่าง O(n log n) และไม่ทำให้ได้คะแนนเพิ่ม"
    ),
}]
KEY = (
    "O(n log n) เติบโตช้ากว่า O(n^2) เมื่อขนาดข้อมูลเพิ่ม จึงมักใช้เวลาน้อยกว่าสำหรับข้อมูลใหญ่ "
    "เช่น Merge Sort มี O(n log n), Bubble Sort มี O(n^2), หรือ nested loop ที่วนตาม n สองชั้นแสดง O(n^2). "
    "การใช้ recursion โดยตัวมันเองไม่ได้กำหนดความซับซ้อนเป็น O(n log n)."
)


async def main():
    wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    sheet = wb["ชุดข้อสอบ_dataset"]
    items = []
    for row in range(40, 74):
        assert sheet.cell(row, 2).value == 2
        items.append({"sample_id": sheet.cell(row, 1).value, "row": row,
                      "question": sheet.cell(row, 4).value, "answer": sheet.cell(row, 6).value,
                      "human_score": float(sheet.cell(row, 7).value),
                      "current_ai_score": float(sheet.cell(row, 8).value)})
    wb.close()
    report = {"started_at": datetime.now(timezone.utc).isoformat(), "model": OPENAI_MODEL,
              "prompt_version": PROMPT_VERSION, "source_sha256": hashlib.sha256(BOOK.read_bytes()).hexdigest(),
              "rubric": RUBRIC, "results": []}
    semaphore = asyncio.Semaphore(4)

    async def grade(item):
        async with semaphore:
            result = await score_with_openai(question_text=item["question"], answer_text=item["answer"],
                                             max_score=2.0, answer_key=KEY, rubrics=RUBRIC,
                                             allowed_scores=[0.0, 1.0, 1.5, 2.0],
                                             strict_rubric_enforcement=True)
            success = (not result.get("metrics", {}).get("manual_review_required", True)
                       and float(result.get("score", -1)) in (0.0, 1.0, 1.5, 2.0)
                       and bool(result.get("teacher_feedback", "").strip()))
            return {**item, "success": success, "result": result}

    tasks = [asyncio.create_task(grade(item)) for item in items]
    for task in asyncio.as_completed(tasks):
        item = await task
        report["results"].append(item)
        OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(item["sample_id"], item["result"].get("score"), "success", item["success"], flush=True)
    print("report", OUT)


if __name__ == "__main__":
    asyncio.run(main())
