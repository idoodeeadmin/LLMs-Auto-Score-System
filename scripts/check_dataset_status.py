import openpyxl
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws_data = wb["ชุดข้อสอบ_dataset"]

status = {}
for r in range(6, ws_data.max_row + 1):
    sid = ws_data.cell(r, 1).value
    if not sid:
        continue
    q_no = ws_data.cell(r, 2).value
    h_score = ws_data.cell(r, 7).value
    ai_score = ws_data.cell(r, 8).value
    ai_fb = ws_data.cell(r, 10).value
    
    if q_no not in status:
        status[q_no] = {"total": 0, "scored": 0, "exact": 0, "diffs": [], "has_feedback": 0}
    status[q_no]["total"] += 1
    if ai_score is not None and str(ai_score).strip() != "":
        status[q_no]["scored"] += 1
        h = float(h_score) if h_score is not None else 0.0
        a = float(ai_score)
        if round(h, 2) == round(a, 2):
            status[q_no]["exact"] += 1
        status[q_no]["diffs"].append(abs(h - a))
    if ai_fb is not None and str(ai_fb).strip() != "":
        status[q_no]["has_feedback"] += 1

print("--- EXCEL DATASET STATUS ---")
for q, s in sorted(status.items()):
    pct = (s["scored"] / s["total"]) * 100 if s["total"] else 0
    exact_pct = (s["exact"] / s["scored"]) * 100 if s["scored"] else 0
    mae = sum(s["diffs"]) / len(s["diffs"]) if s["diffs"] else 0
    print(f"Q{q}: Scored {s['scored']}/{s['total']} ({pct:.1f}%) | Exact: {s['exact']}/{s['scored']} ({exact_pct:.1f}%) | MAE: {mae:.4f} | FB: {s['has_feedback']}/{s['total']}")
