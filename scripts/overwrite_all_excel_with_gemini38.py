"""Overwrite all 204 rows in ชุดข้อสอบ_dataset.xlsx with Gemini 3.8 Flash results.

- Q1, Q2, Q3, Q5, Q6: from docs_and_tests/gemini38_regrade/all_204_results.json
- Q4: from docs_and_tests/gemini38_regrade/q4_topology_rubric_results.json (100% 34/34 match)
"""

import json
import shutil
import sys
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
BACKUP_PATH = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset_backup_before_full_gemini.xlsx"
ALL_204_JSON = ROOT / "docs_and_tests/gemini38_regrade/all_204_results.json"
Q4_TOPOLOGY_JSON = ROOT / "docs_and_tests/gemini38_regrade/q4_topology_rubric_results.json"
PUBLIC_EXCEL = ROOT / "public" / "ชุดข้อสอบ_dataset.xlsx"

# 1. Backup workbook
shutil.copy2(EXCEL_PATH, BACKUP_PATH)
print(f"Backed up workbook to: {BACKUP_PATH}")

# 2. Load result data
d204 = json.load(open(ALL_204_JSON, encoding="utf-8"))
dq4 = json.load(open(Q4_TOPOLOGY_JSON, encoding="utf-8"))

q4_by_id = {r["sample_id"]: r for r in dq4["results"]}
all_by_id = {r["sample_id"]: r for r in d204["results"]}

# 3. Load workbook
wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb["ชุดข้อสอบ_dataset"]

updated_count = 0
matches_by_q = {q: 0 for q in range(1, 7)}
total_by_q = {q: 0 for q in range(1, 7)}

for row in range(6, 210):
    sid = ws.cell(row, 1).value
    q_no = int(ws.cell(row, 2).value)
    h_score = float(ws.cell(row, 7).value)
    
    total_by_q[q_no] += 1
    
    if q_no == 4:
        item = q4_by_id[sid]
        score = item["gemini_score"]
        conf = item.get("confidence", "high").upper()
        tf = item.get("teacher_feedback", "").strip()
        sf = item.get("student_feedback", "").strip()
    else:
        item = all_by_id[sid]["result"]
        score = item["score"]
        conf = item.get("confidence", "high").upper()
        tf = item.get("teacher_feedback", "").strip()
        sf = item.get("student_feedback", "").strip()
        
    if score == h_score:
        matches_by_q[q_no] += 1
        
    # Write to Excel
    # Col 8: AI Score (numeric format)
    ws.cell(row, 8).value = int(score) if float(score).is_integer() else round(float(score), 2)
    # Col 9: AI Confidence
    ws.cell(row, 9).value = conf
    # Col 10: Feedback
    ws.cell(row, 10).value = f"[สำหรับผู้สอน]\n{tf}\n\n[สำหรับนักเรียน]\n{sf}"
    
    updated_count += 1

print(f"\nUpdated {updated_count} rows across all 6 questions in sheet 'ชุดข้อสอบ_dataset':")
total_exact = sum(matches_by_q.values())
for q in range(1, 7):
    print(f"  ข้อ {q}: {matches_by_q[q]} / {total_by_q[q]} ({matches_by_q[q]/total_by_q[q]*100:.1f}%)")
print(f"Total Exact Matches: {total_exact} / 204 ({total_exact/204*100:.2f}%)")

# Save workbook
wb.save(EXCEL_PATH)
wb.close()
print(f"\nSaved updated workbook to: {EXCEL_PATH}")

# Copy to public
shutil.copy2(EXCEL_PATH, PUBLIC_EXCEL)
print(f"Copied updated workbook to: {PUBLIC_EXCEL}")
