import openpyxl
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb["ชุดข้อสอบ_dataset"]

samples = []
for r in range(6, ws.max_row + 1):
    if ws.cell(r, 2).value == 3:
        sid = ws.cell(r, 1).value
        ans = str(ws.cell(r, 6).value or "")
        h = float(ws.cell(r, 7).value or 0.0)
        ai = float(ws.cell(r, 8).value or 0.0)
        samples.append({"sid": sid, "ans": ans, "h": h, "ai": ai})

by_h = {}
for s in samples:
    by_h.setdefault(s["h"], []).append(s)

for score in sorted(by_h.keys()):
    items = by_h[score]
    print(f"\n{'='*70}")
    print(f"TEACHER SCORE = {score:.2f} ({len(items)} students)")
    print(f"{'='*70}")
    for item in items:
        diff = item["ai"] - item["h"]
        diff_str = "EXACT" if diff == 0 else f"{diff:+0.2f}"
        print(f"[{item['sid']}] AI: {item['ai']:.2f} | Diff: {diff_str}")
        print(f"Ans: {item['ans']}")
        print("-" * 50)
