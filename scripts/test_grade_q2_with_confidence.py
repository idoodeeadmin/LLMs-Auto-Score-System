import os
import sys
import json
import asyncio
from pathlib import Path
from datetime import datetime, timezone

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.append(str(ROOT / '.venv/Lib/site-packages'))

from dotenv import load_dotenv
load_dotenv()
sys.stdout.reconfigure(encoding='utf-8')

import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

from scripts.run_dataset_benchmark import (
    load_dataset,
    calculate_qwk,
    calculate_mae,
    calculate_confusion_matrix,
    get_agreement_interpretation,
    EXAM_QUESTIONS,
)
from server.services.openai_grading import score_with_openai, OPENAI_MODEL, PROMPT_VERSION

EXCEL_FILE = ROOT / 'ชุดข้อสอบใหม่' / 'ชุดข้อสอบ_dataset.xlsx'

async def main():
    items = [x for x in load_dataset() if x['question_no'] == 2]
    assert len(items) == 34 and all(x['answer_type'] == 'text' for x in items)

    q = EXAM_QUESTIONS[2]
    rubrics = q['rubrics']
    
    timestamp = datetime.now().strftime('%Y%m%d-%H%M%S')
    folder = ROOT / 'artifacts' / f'q2-confidence-eval-{timestamp}'
    folder.mkdir(parents=True, exist_ok=True)

    print(f"=== Grading Question 2 ({len(items)} students) ===")
    print(f"Model: {OPENAI_MODEL} | Prompt Version: {PROMPT_VERSION}")
    print(f"Topic: {q['topic']}")
    print(f"Rubrics: {[r['name'] + ' (' + str(r['score']) + ')' for r in rubrics]}")
    print("-" * 65)

    results = []
    sem = asyncio.Semaphore(4)

    async def grade_one(item):
        async with sem:
            res = await score_with_openai(
                question_text=item['question_content'],
                answer_text=item['student_answer'],
                max_score=q['max_score'],
                answer_key=q['answer_key'],
                rubrics=rubrics,
            )
            h_score = item['human_score']
            ai_score = res.get('score', 0.0)
            conf = res.get('confidence', 'unknown')
            review = res.get('metrics', {}).get('manual_review_required', False)
            diff = round(ai_score - h_score, 2)
            
            entry = {
                'sample_id': item['sample_id'],
                'row': item['row'],
                'answer': item['student_answer'],
                'human_score': h_score,
                'ai_score': ai_score,
                'diff': diff,
                'exact': ai_score == h_score,
                'confidence': conf,
                'manual_review_required': review,
                'feedback': res.get('feedback', ''),
            }
            results.append(entry)
            status = "MATCH" if ai_score == h_score else f"DIFF ({diff:+0.1f})"
            print(f"[{entry['sample_id']}] Human={h_score:0.1f} | AI={ai_score:0.1f} | Conf={conf:6s} | {status} | Review={review}")
            return entry

    # Run all students
    tasks = [grade_one(item) for item in items]
    await asyncio.gather(*tasks)

    # Sort by row
    results.sort(key=lambda x: x['row'])

    truth = [r['human_score'] for r in results]
    predicted = [r['ai_score'] for r in results]
    exact_count = sum(r['exact'] for r in results)
    tol_count = sum(abs(r['diff']) <= 0.5 for r in results)
    mae = calculate_mae(truth, predicted)
    qwk = calculate_qwk(truth, predicted, step=0.5, max_score=2.0)

    conf_dist = {'high': 0, 'medium': 0, 'low': 0, 'other': 0}
    for r in results:
        conf_dist[r['confidence']] = conf_dist.get(r['confidence'], 0) + 1

    summary = {
        'total': len(results),
        'exact_count': exact_count,
        'exact_pct': round((exact_count / len(results)) * 100, 2),
        'tolerance_0_5_count': tol_count,
        'tolerance_0_5_pct': round((tol_count / len(results)) * 100, 2),
        'mae': round(mae, 4),
        'qwk': round(qwk, 4),
        'qwk_interpretation': get_agreement_interpretation(qwk),
        'confidence_distribution': conf_dist,
    }

    full_output = {
        'timestamp': datetime.now(timezone.utc).isoformat(),
        'question_no': 2,
        'topic': q['topic'],
        'rubrics': rubrics,
        'summary': summary,
        'results': results,
    }

    output_path = folder / 'q2_results.json'
    output_path.write_text(json.dumps(full_output, ensure_ascii=False, indent=2), encoding='utf-8')

    print("\n" + "=" * 65)
    print("=== SUMMARY RESULTS QUESTION 2 ===")
    print(f"Total Students:        {summary['total']}")
    print(f"Exact Match:           {summary['exact_count']}/{summary['total']} ({summary['exact_pct']}%)")
    print(f"Tolerance <= 0.5:      {summary['tolerance_0_5_count']}/{summary['total']} ({summary['tolerance_0_5_pct']}%)")
    print(f"MAE (Mean Abs Error):  {summary['mae']}")
    print(f"QWK (Quadratic Kappa): {summary['qwk']} ({summary['qwk_interpretation']['level_th']})")
    print(f"Confidence Breakdown:  High={conf_dist.get('high', 0)}, Medium={conf_dist.get('medium', 0)}, Low={conf_dist.get('low', 0)}")
    print(f"JSON Report saved to: {output_path}")

    # Record directly into Excel
    print(f"\nRecording Question 2 results into Excel: {EXCEL_FILE.name}...")
    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb['ชุดข้อสอบ_dataset']

    thin_border = Border(
        left=Side(style='thin', color='00BFBFBF'),
        right=Side(style='thin', color='00BFBFBF'),
        top=Side(style='thin', color='00BFBFBF'),
        bottom=Side(style='thin', color='00BFBFBF')
    )
    font_aptos = Font(name='Aptos', size=10)
    align_center = Alignment(horizontal='center', vertical='top')
    align_feedback = Alignment(horizontal='left', vertical='top', wrap_text=True)

    results_by_sid = {r['sample_id']: r for r in results}
    updated_count = 0

    for row in range(6, ws.max_row + 1):
        sid = ws.cell(row=row, column=1).value
        qno = ws.cell(row=row, column=2).value
        if qno != 2 or not sid:
            continue
        
        sid_clean = str(sid).strip()
        if sid_clean in results_by_sid:
            res = results_by_sid[sid_clean]
            
            c8 = ws.cell(row=row, column=8, value=float(res['ai_score']))
            c8.font = font_aptos
            c8.alignment = align_center
            c8.border = thin_border
            
            c9 = ws.cell(row=row, column=9, value=str(res['confidence']))
            c9.font = font_aptos
            c9.alignment = align_center
            c9.border = thin_border
            
            c10 = ws.cell(row=row, column=10, value=str(res['feedback']))
            c10.font = font_aptos
            c10.alignment = align_feedback
            c10.border = thin_border
            
            updated_count += 1

    wb.save(EXCEL_FILE)
    print(f"Successfully recorded {updated_count} rows for Question 2 in {EXCEL_FILE}")

if __name__ == '__main__':
    asyncio.run(main())
