"""Apply the verified DS-051 score under the clarified Q2 rubric."""

import hashlib
import json
import os
import shutil
import zipfile
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx"
BACKUP = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset_before_ds051_clarified_score.xlsx"
REPORT = ROOT / "docs_and_tests/q2_q3_review/regrade_ds051_clarified_rubric.json"
data = json.loads(REPORT.read_text(encoding="utf-8"))
assert hashlib.sha256(BOOK.read_bytes()).hexdigest() == data["source_sha256"]
assert len(data["items"]) == 1
item = data["items"][0]
result = item["result"]
assert item["sample_id"] == "DS-051" and item["row"] == 56 and item["success"]
assert float(result["score"]) == 1.5 and result["confidence"] in ("high", "medium")

wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
sheet = wb["ชุดข้อสอบ_dataset"]
assert sheet.cell(56, 6).value == item["answer"]
assert sheet.cell(56, 8).value == 0
assert "คำอธิบายเพิ่มเติมในการแยกคะแนนบางส่วน" in wb["Exam_Rubrics"].cell(7, 6).value
wb.close()
if not BACKUP.exists():
    shutil.copy2(BOOK, BACKUP)

ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ET.register_namespace("", ns)
with zipfile.ZipFile(BOOK) as source:
    root = ET.fromstring(source.read("xl/worksheets/sheet1.xml"))
    cells = {cell.attrib["r"]: cell for cell in root.iter(f"{{{ns}}}c")}
    for column, value in (("H", 1.5), ("I", result["confidence"]), ("J", result["feedback"])):
        cell = cells[f"{column}56"]
        for child in list(cell):
            if child.tag in {f"{{{ns}}}v", f"{{{ns}}}is", f"{{{ns}}}f"}:
                cell.remove(child)
        if column == "H":
            cell.attrib.pop("t", None)
            ET.SubElement(cell, f"{{{ns}}}v").text = str(value)
        else:
            cell.attrib["t"] = "inlineStr"
            inline = ET.SubElement(cell, f"{{{ns}}}is")
            text = ET.SubElement(inline, f"{{{ns}}}t")
            text.attrib["{http://www.w3.org/XML/1998/namespace}space"] = "preserve"
            text.text = value
    xml = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with NamedTemporaryFile(dir=BOOK.parent, suffix=".xlsx", delete=False) as temp:
        temp_name = temp.name
    with zipfile.ZipFile(temp_name, "w") as target:
        for entry in source.infolist():
            target.writestr(entry, xml if entry.filename == "xl/worksheets/sheet1.xml" else source.read(entry.filename))

try:
    with zipfile.ZipFile(temp_name) as current, zipfile.ZipFile(BACKUP) as original:
        for name in original.namelist():
            if name != "xl/worksheets/sheet1.xml":
                assert current.read(name) == original.read(name), name
    check = openpyxl.load_workbook(temp_name, read_only=True, data_only=True)
    try:
        assert check["ชุดข้อสอบ_dataset"].cell(56, 8).value == 1.5
    finally:
        check.close()
    os.replace(temp_name, BOOK)
finally:
    if os.path.exists(temp_name):
        os.unlink(temp_name)
print("Updated DS-051 score and feedback")
