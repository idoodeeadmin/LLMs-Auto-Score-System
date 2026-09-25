import asyncio
import json
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q4_RUBRIC_NAME = "ความถูกต้องของโครงสร้างต้นไม้ค้นหาทวิภาค (Binary Search Tree)"
Q4_RUBRIC_DESC = (
    "ประเมินตามโครงสร้างและความถูกต้องของ Binary Search Tree (BST) จากข้อมูล 9, 16, 10, 76, 5, 13, 58, 92, 11, 15, 80, 99 (รวม 12 โหนด) ดังนี้:\n\n"
    "- ได้ 1.00 คะแนน: โครงสร้าง BST ถูกต้องสมบูรณ์ครบทั้ง 12 โหนด โดย Root คือ 9 (ซ้าย 5, ขวา 16) -> "
    "ใต้ 16 (ซ้าย 10, ขวา 76) -> ใต้ 10 (ขวา 13, ซ้ายว่าง) -> ใต้ 13 (ซ้าย 11, ขวา 15) -> "
    "ใต้ 76 (ซ้าย 58, ขวา 92) -> ใต้ 92 (ซ้าย 80, ขวา 99) "
    "ทุกโหนดรักษาคุณสมบัติ โหนดซ้าย < โหนดแม่ < โหนดขวา อย่างถูกต้องครบถ้วน "
    "(อนุโลมความสวยงาม ความยาวของกิ่ง หรือลายมือ หากความสัมพันธ์และทิศทางซ้าย/ขวาสื่อสารได้ถูกต้อง)\n"
    "- ได้ 0.25 คะแนน: โครงสร้างหลักถูกต้องเกิน 70% (เช่น แกน Root 9, 5, 16 และกิ่งขวา 76, 58, 92, 80, 99 ถูกต้องครบถ้วน) "
    "แต่มีข้อผิดพลาดเฉพาะที่ตำแหน่งโหนดในกิ่งย่อยเพียง 1 จุด (เช่น สับสนการวาง 11 ใต้ 10) หรือมีเส้นกิ่งว่างเปล่าส่วนเกิน\n"
    "- ได้ 0.00 คะแนน: ผิดพลาดอย่างมีนัยสำคัญ เช่น โหนดตกหล่นไม่ครบ 12 โหนด (เช่น ลืมใส่โหนด 10), "
    "โหนดใดโหนดหนึ่งแตกกิ่งเกิน 2 กิ่ง (ไม่ใช่ Binary Tree), วางโหนดผิดตำแหน่งหลักตั้งแต่ระดับบน (เช่น วาง 10 ไว้ทางซ้ายของ 9), "
    "วาง 58 ใต้ 13, หรือวาดผิดหลักการเป็นเส้นตรงซิกแซก\n\n"
    "(ต้องระบุ teacher_feedback แจกแจงเหตุผลและจุดถูก/ผิดตามเกณฑ์สำหรับผู้สอนอย่างละเอียด และระบุ student_feedback เป็นคำแนะนำพัฒนาการเรียนรู้สำหรับนักเรียน)"
)

Q4_RUBRICS = [
    {
        "name": Q4_RUBRIC_NAME,
        "score": 1.0,
        "description": Q4_RUBRIC_DESC,
    }
]

Q4_QUESTION_TEXT = (
    "จากข้อมูลต่อไปนี้จงนำไปสร้างเป็น Binary search tree (1 คะแนน)\n"
    "9 16 10 76 5 13 58 92 11 15 80 99"
)

Q4_ANSWER_KEY = (
    "เฉลยโครงสร้าง Binary Search Tree (BST) ที่ถูกต้องสมบูรณ์:\n"
    "- Root คือ 9\n"
    "  - กิ่งซ้าย: 5\n"
    "  - กิ่งขวา: 16\n"
    "    - ใต้ 16 กิ่งซ้าย: 10\n"
    "      - ใต้ 10 กิ่งขวา: 13 (กิ่งซ้ายว่าง)\n"
    "        - ใต้ 13 กิ่งซ้าย: 11\n"
    "        - ใต้ 13 กิ่งขวา: 15\n"
    "    - ใต้ 16 กิ่งขวา: 76\n"
    "      - ใต้ 76 กิ่งซ้าย: 58\n"
    "      - ใต้ 76 กิ่งขวา: 92\n"
    "        - ใต้ 92 กิ่งซ้าย: 80\n"
    "        - ใต้ 92 กิ่งขวา: 99\n"
    "รวมทั้งหมด 12 โหนด โดยทุกโหนดลูกทางซ้ายต้องมีค่าน้อยกว่าโหนดแม่ และลูกทางขวาต้องมีค่ามากกว่าโหนดแม่"
)

KEY_IMG_PATH = ROOT / "public" / "answer-keys" / "q4-bst-answer-key.png"
key_img_bytes = KEY_IMG_PATH.read_bytes() if KEY_IMG_PATH.exists() else None

def calculate_qwk(y_true, y_pred, step=0.25, max_score=1.0):
    if len(y_true) != len(y_pred) or len(y_true) == 0:
        return 0.0
    if y_true == y_pred:
        return 1.0
    k = int(round(max_score / step)) + 1
    w = [[((i - j) ** 2) / ((k - 1) ** 2) for j in range(k)] for i in range(k)]
    cat_true = [int(round(v / step)) for v in y_true]
    cat_pred = [int(round(v / step)) for v in y_pred]
    o = [[0] * k for _ in range(k)]
    for t, p in zip(cat_true, cat_pred):
        o[t][p] += 1
    r_t = [sum(o[i][j] for j in range(k)) for i in range(k)]
    r_p = [sum(o[i][j] for i in range(k)) for j in range(k)]
    n = len(y_true)
    e = [[(r_t[i] * r_p[j]) / n for j in range(k)] for i in range(k)]
    num = sum(w[i][j] * o[i][j] for i in range(k) for j in range(k))
    den = sum(w[i][j] * e[i][j] for i in range(k) for j in range(k))
    if den == 0:
        return 1.0
    return 1.0 - (num / den)

async def main():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws_dataset = wb["ชุดข้อสอบ_dataset"]
    
    rows = list(range(108, 118)) # First 10 rows for Q4 (DS-103 to DS-112)
    assert len(rows) == 10
    
    to_evaluate = []
    for r in rows:
        sid = ws_dataset.cell(r, 1).value
        human_score = float(ws_dataset.cell(r, 7).value or 0.0)
        std_idx = r - 107
        img_name = f"LINE_ALBUM_Photo1_260917_{std_idx}.jpg"
        img_path = ROOT / "ชุดข้อสอบใหม่" / "photo_clean_ชุดที่1" / img_name
        to_evaluate.append({
            "row": r,
            "sample_id": sid,
            "student_idx": std_idx,
            "human_score": human_score,
            "img_name": img_name,
            "img_path": img_path
        })
        
    print(f"============================================================")
    print(f"Testing Question 4 (First 10 Students: DS-103 to DS-112)")
    print(f"Target Ground Truth Scores: {[x['human_score'] for x in to_evaluate]}")
    print(f"============================================================")
    
    sem = asyncio.Semaphore(4)
    
    async def grade_one(item):
        async with sem:
            img_bytes = item["img_path"].read_bytes()
            print(f"Grading {item['sample_id']} (idx={item['student_idx']:2d}) ...", flush=True)
            res = await score_with_openai(
                question_text=Q4_QUESTION_TEXT,
                answer_text="",
                max_score=1.0,
                answer_key=Q4_ANSWER_KEY,
                rubrics=Q4_RUBRICS,
                image_bytes_list=[img_bytes],
                image_mime_list=["image/jpeg"],
                answer_key_image_bytes_list=[key_img_bytes] if key_img_bytes else None,
                answer_key_image_mime_list=["image/png"] if key_img_bytes else None,
                strict_rubric_enforcement=True,
            )
            ai_score = float(res.get("score", 0.0))
            conf = res.get("confidence", "unknown")
            teacher_fb = res.get("teacher_feedback", "")
            student_fb = res.get("student_feedback", "")
            combined_fb = res.get("feedback", "")
            diff = round(ai_score - item["human_score"], 2)
            match = (ai_score == item["human_score"])
            
            result_item = {
                "row": item["row"],
                "sample_id": item["sample_id"],
                "student_idx": item["student_idx"],
                "img_name": item["img_name"],
                "human_score": item["human_score"],
                "ai_score": ai_score,
                "diff": diff,
                "match": match,
                "confidence": conf,
                "teacher_feedback": teacher_fb,
                "student_feedback": student_fb,
                "feedback": combined_fb,
            }
            status_str = "MATCH 🟢" if match else f"DIFF 🟠 ({diff:+.2f})"
            print(f"  -> {item['sample_id']} (idx={item['student_idx']:2d}): Human={item['human_score']:.2f} | AI={ai_score:.2f} | {status_str}")
            return result_item

    tasks = [grade_one(item) for item in to_evaluate]
    results = await asyncio.gather(*tasks)
    results.sort(key=lambda x: x["row"])
    
    # Save results to artifacts
    out_json = ROOT / "artifacts" / "q4_first10_test_results.json"
    out_json.parent.mkdir(parents=True, exist_ok=True)
    with open(out_json, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"\nSaved test results to {out_json}")
    
    # Statistics
    y_true = [x["human_score"] for x in results]
    y_pred = [x["ai_score"] for x in results]
    exacts = sum(1 for x in results if x["match"])
    mae = sum(abs(t - p) for t, p in zip(y_true, y_pred)) / len(y_true)
    qwk = calculate_qwk(y_true, y_pred, step=0.25, max_score=1.0)
    
    print("\n" + "="*80)
    print("EVALUATION SUMMARY REPORT FOR QUESTION 4 (FIRST 10 SAMPLES)")
    print("="*80)
    print(f"Total Evaluated: {len(results)}")
    print(f"Exact Matches:   {exacts} / {len(results)} ({exacts / len(results) * 100:.1f}%)")
    print(f"MAE:             {mae:.4f}")
    print(f"QWK (Kappa):     {qwk:.4f}")
    print("="*80)
    
    print(f"\n{'Sample ID':<10} | {'Idx':<4} | {'Human':<6} | {'AI':<6} | {'Diff':<7} | {'Status':<10} | {'Confidence':<10}")
    print("-" * 65)
    for x in results:
        status_text = "EXACT" if x["match"] else f"{x['diff']:+.2f}"
        print(f"{x['sample_id']:<10} | {x['student_idx']:<4d} | {x['human_score']:<6.2f} | {x['ai_score']:<6.2f} | {x['diff']:<+7.2f} | {status_text:<10} | {x['confidence']:<10}")
        
    print("\nSample Feedbacks:")
    for x in results[:3]:
        print(f"\n[{x['sample_id']} - Score: Human={x['human_score']}, AI={x['ai_score']}]")
        print(f"Teacher FB: {x['teacher_feedback'][:160]}...")
        print(f"Student FB: {x['student_feedback'][:160]}...")

if __name__ == "__main__":
    asyncio.run(main())
