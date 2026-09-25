import openpyxl
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[1]

# 1. Update Exam_Rubrics in ชุดข้อสอบ_dataset.xlsx
DATASET_PATH = ROOT / 'ชุดข้อสอบใหม่' / 'ชุดข้อสอบ_dataset.xlsx'
wb = openpyxl.load_workbook(DATASET_PATH)
ws = wb['Exam_Rubrics']

q2_rubric_5level = (
    "ประเมินตามระดับคะแนน 5 ระดับ (2.0, 1.5, 1.0, 0.5, 0.0) อย่างเคร่งครัด:\n"
    "• 2.0 คะแนน: อธิบายได้ว่าทำไม O(n log n) ดีกว่าเมื่อข้อมูลใหญ่ เช่น เร็วกว่า, จำนวนรอบ/การทำงานโตช้ากว่า, n² โตเร็วมาก และ มีตัวอย่าง Algorithm หรืออธิบายกลไกชัดเจน\n"
    "• 1.5 คะแนน: เข้าใจแก่นว่า O(n log n) มีประสิทธิภาพกว่า แต่คำอธิบายยังไม่ครบ/ไม่ยกตัวอย่าง หรือมีตัวอย่างแต่เหตุผลยังไม่ลงลึก\n"
    "• 1.0 คะแนน: รู้เพียงว่า O(n log n) “เร็วกว่า/ดีกว่า/ซ้ำซ้อนน้อยกว่า” แต่ไม่ได้อธิบายกลไกอย่างชัดเจน หรือคำอธิบายคลุมเครือ และไม่ได้ยกตัวอย่าง (หรือมีเพียงตัวอย่างคำนวณ)\n"
    "• 0.5 คะแนน: ระบุได้เพียงตัวอย่าง Algorithm ที่เกี่ยวข้องอย่างถูกต้อง หรือกล่าวถึงความซับซ้อน/ประสิทธิภาพเพียงฝั่งเดียวอย่างสั้นๆ โดยยังไม่สื่อชัดว่า O(n log n) เหมาะกับข้อมูลใหญ่กว่า O(n²) อย่างไร\n"
    "• 0.0 คะแนน: ไม่สามารถอธิบายความสัมพันธ์ของ O(n log n) กับ O(n²) ได้อย่างมีสาระ ตอบผิดหลักการสำคัญ หรือไม่ตอบ\n"
    "(ระดับคะแนนที่ให้ได้คือ 2.0, 1.5, 1.0, 0.5 หรือ 0.0 คะแนนเท่านั้น ห้ามให้เศษทศนิยมอื่น)"
)

# Row 7 is Q2
ws.cell(7, 6, value=q2_rubric_5level)
wb.save(DATASET_PATH)
print(f"1. Successfully updated Exam_Rubrics Row 7 in {DATASET_PATH.name}")

# 2. Update write_rubrics_to_excel.py
SCRIPT_RUBRIC = ROOT / 'scripts' / 'write_rubrics_to_excel.py'
if SCRIPT_RUBRIC.exists():
    text = SCRIPT_RUBRIC.read_text(encoding='utf-8')
    # Update Q2_RUBRIC_TEXT if needed
    text = text.replace(
        '(ระดับคะแนนที่ให้ได้คือ 2.0, 1.5, 1.0 หรือ 0.0 คะแนนเท่านั้น)',
        '(ระดับคะแนนที่ให้ได้คือ 2.0, 1.5, 1.0, 0.5 หรือ 0.0 คะแนนเท่านั้น)'
    )
    text = text.replace(
        '(ระดับคะแนนที่ให้ได้คือ 2.0, 1.5, 1.0 หรือ 0.0 คะแนนเท่านั้น ห้ามให้เศษทศนิยมอื่น)',
        '(ระดับคะแนนที่ให้ได้คือ 2.0, 1.5, 1.0, 0.5 หรือ 0.0 คะแนนเท่านั้น ห้ามให้เศษทศนิยมอื่น)'
    )
    SCRIPT_RUBRIC.write_text(text, encoding='utf-8')
    print(f"2. Updated {SCRIPT_RUBRIC.name}")

# 3. Update เกณฑ์ตรวจสำหรับAI.xlsx
AI_RUBRIC_PATH = ROOT / 'ชุดข้อสอบใหม่' / 'เกณฑ์ตรวจสำหรับAI.xlsx'
if AI_RUBRIC_PATH.exists():
    wb_ai = openpyxl.load_workbook(AI_RUBRIC_PATH)
    ws_ai = wb_ai['ข้อมูลสำหรับ API']
    updated_ai = 0
    for r in range(2, ws_ai.max_row + 1):
        q_val = str(ws_ai.cell(r, 1).value or '')
        if 'O(n log n)' in q_val:
            ws_ai.cell(r, 2, value=q2_rubric_5level)
            updated_ai += 1
    wb_ai.save(AI_RUBRIC_PATH)
    print(f"3. Updated {updated_ai} rows in {AI_RUBRIC_PATH.name}")

print("\nAll files synchronized with 5-Level Q2 Rubric (supporting 0.5)!")
