import asyncio
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env")
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q2_TWO_RUBRICS = [
    {
        "name": "เหตุผลเชิงเปรียบเทียบความซับซ้อน (Time Complexity Reasoning)",
        "score": 1.0,
        "description": (
            "ประเมินเหตุผลว่าทำไม O(n log n) เหมาะกับข้อมูลขนาดใหญ่กว่า O(n^2) (คะแนนเต็ม 1.00 คะแนน):\n"
            "• ได้ 1.00 คะแนน: อธิบายถึงเหตุผลที่สื่อถึงความเข้าใจว่า O(n log n) ทำงานน้อยกว่า/เร็วกว่า/ตัดแบ่งครึ่งข้อมูล/ลูปน้อยกว่า O(n^2) เมื่อข้อมูลมีขนาดใหญ่ (ให้อนุโลมภาษาพูด ไม่หักคะแนนหากนิสิตเขียนเรื่อง memory/recursive/CPU)\n"
            "• ได้ 0.50 คะแนน: ตอบสั้นๆ แบบกว้างๆ เพียงว่ามีประสิทธิภาพดีกว่า หรือเร็วกว่า โดยไม่อธิบายการวนลูปหรือการเติบโต\n"
            "• ได้ 0.00 คะแนน: ไม่ได้ตอบ หรือตอบผิดหลักการทั้งหมด (เช่น บอกว่าเพราะเสถียรกว่า หรือเขียนเรื่องลูกเต๋าโยนเหรียญ)"
        ),
        "allowed_scores": [0.0, 0.5, 1.0],
    },
    {
        "name": "การยกตัวอย่าง Algorithm (Algorithm Examples)",
        "score": 1.0,
        "description": (
            "ประเมินการยกตัวอย่าง Algorithm ตามที่โจทย์กำหนด (คะแนนเต็ม 1.00 คะแนน):\n"
            "• ได้ 1.00 คะแนน: มีการระบุชื่อ Algorithm ที่เกี่ยวข้องอย่างน้อย 1 ชื่อ (เช่น Merge Sort, Quick Sort, Bubble Sort, Insertion Sort, Heap Sort, Selection Sort หรือยกตัวอย่างฟังก์ชันโค้ดการทำงาน)\n"
            "• ได้ 0.00 คะแนน: ไม่มีการระบุชื่อ Algorithm หรือตัวอย่างใดๆ เลย"
        ),
        "allowed_scores": [0.0, 1.0],
    },
]

Q2_QUESTION_TEXT = "อธิบายว่าทำไม O(n log n) เหมาะกับข้อมูลใหญ่กว่า O(n^2) และยกตัวอย่าง Algorithm"

async def test_all():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["ชุดข้อสอบ_dataset"]
    rows = [r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value == 2]
    print(f"Testing Q2 with 2 natural sub-rubrics on {len(rows)} students...", flush=True)

    sem = asyncio.Semaphore(5)
    exact = 0
    within_05 = 0
    diffs = []

    async def grade_student(r):
        nonlocal exact, within_05
        sid = ws.cell(r, 1).value
        ans = str(ws.cell(r, 6).value or "")
        human = float(ws.cell(r, 7).value or 0.0)
        async with sem:
            res = await score_with_openai(
                question_text=Q2_QUESTION_TEXT,
                answer_text=ans,
                max_score=2.0,
                answer_key="",
                rubrics=Q2_TWO_RUBRICS,
                allowed_scores=[0.0, 0.5, 1.0, 1.5, 2.0],
                strict_rubric_enforcement=False,
            )
            ai = float(res["score"])
            diff = round(ai - human, 2)
            is_match = (diff == 0.0)
            if is_match:
                exact += 1
            if abs(diff) <= 0.5:
                within_05 += 1
            diffs.append((sid, human, ai, diff, ans[:50], res.get("teacher_feedback", "")[:80]))
            status = "MATCH" if is_match else f"DIFF={diff:+.1f}"
            print(f"[{sid}] Human: {human:.1f} | AI: {ai:.1f} | {status}", flush=True)

    await asyncio.gather(*[grade_student(r) for r in rows])

    mae = sum(abs(d[3]) for d in diffs) / len(diffs)
    print("\n" + "=" * 50, flush=True)
    print(f"RESULTS FOR QUESTION 2 (2 Sub-Rubrics):", flush=True)
    print(f"Exact Matches: {exact}/{len(rows)} ({exact/len(rows)*100:.2f}%)", flush=True)
    print(f"Within +/- 0.50 pt: {within_05}/{len(rows)} ({within_05/len(rows)*100:.2f}%)", flush=True)
    print(f"MAE: {mae:.4f}", flush=True)
    print("=" * 50, flush=True)

if __name__ == "__main__":
    asyncio.run(test_all())
