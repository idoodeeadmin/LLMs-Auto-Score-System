"""Blind pilot of the user-proposed four-part Q2 rubric (no workbook edits)."""

import asyncio
import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

# This module loads .env and supplies the word-counter shim needed by the
# bundled Python runtime; it does not run grading when imported.
from scripts.regrade_corrected_q2_q3_answers import OPENAI_MODEL, PROMPT_VERSION, score_with_openai

BOOK = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx"
OUT = ROOT / "docs_and_tests/q2_q3_review/q2_four_part_pilot.json"
RUBRIC = [
    {"name": "ระบุความเหมาะสมกับข้อมูลขนาดใหญ่", "score": 0.5, "allowed_scores": [0.0, 0.5],
     "description": "ให้ 0.5 เมื่อคำตอบสื่อว่า O(n log n) เหมาะกับข้อมูลขนาดใหญ่มากกว่า O(n^2); มิฉะนั้นให้ 0"},
    {"name": "เหตุผลหรือแนวโน้มเวลาที่ถูกทิศทาง", "score": 0.5, "allowed_scores": [0.0, 0.5],
     "description": "ให้ 0.5 เมื่ออธิบายเหตุผลว่าเวลาหรือจำนวนรอบของ O(n log n) เติบโตช้ากว่า/น้อยกว่า O(n^2) เมื่อ n เพิ่ม หรืออธิบายกลไกเปรียบเทียบได้ถูกทิศทาง; การบอกเพียงว่า 'เร็วกว่า' โดยไม่อธิบายเหตุผลยังไม่พอ; มิฉะนั้นให้ 0"},
    {"name": "ตัวอย่าง O(n^2)", "score": 0.5, "allowed_scores": [0.0, 0.5],
     "description": "ให้ 0.5 เมื่อยกตัวอย่าง O(n^2) ที่ถูกแนวคิด เช่น nested loop ที่วนตาม n สองชั้น, Bubble Sort, Selection Sort หรือ Insertion Sort; มิฉะนั้นให้ 0"},
    {"name": "ตัวอย่าง O(n log n)", "score": 0.5, "allowed_scores": [0.0, 0.5],
     "description": "ให้ 0.5 เมื่อยกตัวอย่าง O(n log n) ได้ถูกต้อง เช่น Merge Sort, Heap Sort หรือ Quick Sort กรณีเฉลี่ย หรือโค้ด/ขั้นตอนที่แสดงกลไก O(n log n) จริง; recursion หรือ factorial เพียงอย่างเดียวไม่ใช่ตัวอย่างที่ถูกต้อง; มิฉะนั้นให้ 0"},
]
KEY = "O(n log n) มีเวลาทำงานโตช้ากว่า O(n^2) เมื่อข้อมูลมีขนาดใหญ่ ตัวอย่าง O(n^2): nested loop ที่วนตาม n สองชั้น, Bubble Sort, Selection Sort. ตัวอย่าง O(n log n): Merge Sort, Heap Sort หรือ Quick Sort ในกรณีเฉลี่ย. การเรียกซ้ำหรือ factorial เพียงอย่างเดียวไม่ได้แปลว่าเป็น O(n log n)."


async def main():
    workbook = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    sheet = workbook["ชุดข้อสอบ_dataset"]
    items = []
    for row in range(40, 74):
        assert sheet.cell(row, 2).value == 2
        items.append({"sample_id": sheet.cell(row, 1).value, "row": row,
                      "question": sheet.cell(row, 4).value,
                      "answer": sheet.cell(row, 6).value,
                      "human_score": float(sheet.cell(row, 7).value),
                      "current_ai_score": float(sheet.cell(row, 8).value)})
    workbook.close()
    report = {"started_at": datetime.now(timezone.utc).isoformat(), "model": OPENAI_MODEL,
              "prompt_version": PROMPT_VERSION, "source_sha256": hashlib.sha256(BOOK.read_bytes()).hexdigest(),
              "rubric": RUBRIC, "results": []}
    semaphore = asyncio.Semaphore(4)

    async def grade(item):
        async with semaphore:
            result = await score_with_openai(question_text=item["question"], answer_text=item["answer"],
                                             max_score=2.0, answer_key=KEY, rubrics=RUBRIC,
                                             allowed_scores=[0.0, 0.5, 1.0, 1.5, 2.0],
                                             strict_rubric_enforcement=True)
            success = (not result.get("metrics", {}).get("manual_review_required", True)
                       and float(result.get("score", -1)) in [0.0, 0.5, 1.0, 1.5, 2.0]
                       and bool(result.get("teacher_feedback", "").strip()))
            return {**item, "success": success, "result": result}

    tasks = [asyncio.create_task(grade(item)) for item in items]
    for task in asyncio.as_completed(tasks):
        item = await task
        report["results"].append(item)
        OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(item["sample_id"], item["result"].get("score"), "success", item["success"], flush=True)
    print("report", OUT)


if __name__ == "__main__":
    asyncio.run(main())
