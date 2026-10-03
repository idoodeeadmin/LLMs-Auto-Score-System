"""Compare Q1 Excel answers with their original scored answer photos."""

from html import escape
from pathlib import Path

import openpyxl
from PIL import Image

ROOT = Path(__file__).resolve().parents[1]
NUMBERS = [
    2791, 2793, 2795, 2797, 2798, 2799, 2800, 2801, 2802, 2803, 2804,
    2808, 2811, 2812, 2813, 2814, 2815, 2816, 2818, 2819, 2821, 2822,
    2823, 2824, 2825, 2826, 2827, 2828, 2829, 2831, 2832, 2833, 2834,
    2835,
]
assert len(NUMBERS) == 34 and len(set(NUMBERS)) == 34
book = openpyxl.load_workbook(
    ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx", read_only=True, data_only=True
)
sheet = book["ชุดข้อสอบ_dataset"]
cards = []
preview_dir = ROOT / "docs_and_tests/assets/q1_scored_oriented"
preview_dir.mkdir(parents=True, exist_ok=True)
for index, image_number in enumerate(NUMBERS, 1):
    sample_id = f"DS-{index:03d}"
    row = index + 5
    assert sheet.cell(row, 1).value == sample_id
    answer = str(sheet.cell(row, 6).value or "")
    image_name = f"{sample_id}_IMG_{image_number}.jpg"
    image_path = ROOT / "public/photo_q1_original" / image_name
    heic_path = ROOT / "ชุดข้อสอบเก่า/Textชุดที่1" / f"IMG_{image_number}.HEIC"
    assert image_path.is_file() and heic_path.is_file(), image_name
    image_url = f"../public/photo_q1_original/{image_name}"
    preview_path = preview_dir / image_name
    with Image.open(image_path) as source:
        oriented = source.rotate(90, expand=True) if source.height > source.width else source.copy()
        oriented.save(preview_path, quality=90)
    preview_url = f"assets/q1_scored_oriented/{image_name}"
    answer_html = escape(answer).replace("\n", "<br>")
    search_text = escape(f"{sample_id} IMG_{image_number} {answer}", quote=True)
    cards.append(f"""
    <article class="card" data-search="{search_text}">
      <div class="card-head"><strong>{sample_id}</strong><span>IMG_{image_number}.HEIC</span>
        <a href="{image_url}" target="_blank" rel="noopener">เปิดภาพขนาดเต็ม ↗</a></div>
      <div class="comparison">
        <section><h2>ภาพต้นฉบับที่ยังมีคะแนนและรอยตรวจ</h2>
          <a class="photo" href="{image_url}" target="_blank" rel="noopener">
            <img src="{preview_url}" alt="ภาพต้นฉบับ {sample_id} หมุนให้อ่านตรง" loading="lazy"></a></section>
        <section><h2>ข้อความคำตอบใน Excel</h2><div class="answer">{answer_html}</div></section>
      </div>
    </article>""")

html = """<!doctype html><html lang="th"><head><meta charset="utf-8">
<meta name="viewport" content="width=device-width,initial-scale=1">
<title>เทียบคำตอบข้อ 1 กับภาพต้นฉบับ</title><style>
*{box-sizing:border-box}body{margin:0;background:#f6f8f7;color:#17352f;font:15px/1.65 Tahoma,Arial,sans-serif}
main{max-width:1280px;margin:auto;padding:28px 24px 80px}h1{font-size:28px;margin:0 0 5px}p{margin:0;color:#526c64}
.toolbar{display:flex;gap:14px;align-items:center;margin:18px 0;flex-wrap:wrap}
input{width:min(440px,100%);border:1px solid #bdcbc6;border-radius:8px;padding:11px 14px;font:inherit}
#count{color:#526c64}.note{padding:13px 16px;border-left:3px solid #277662;background:#edf5f1;margin:18px 0 26px}
.card{background:#fff;border:1px solid #dbe5e1;border-radius:12px;margin:18px 0;overflow:hidden}
.card-head{display:flex;gap:16px;align-items:center;padding:11px 18px;background:#f0f5f2;border-bottom:1px solid #dbe5e1}
.card-head a{margin-left:auto;color:#0b6756}.card-head span{color:#526c64}
.comparison{display:grid;grid-template-columns:1fr 1fr}.comparison section{min-width:0;padding:18px}
.comparison section+section{border-left:1px solid #dbe5e1}h2{font-size:14px;margin:0 0 12px}
.photo{height:500px;background:#e9edeb;border-radius:8px;display:flex;align-items:center;justify-content:center;overflow:hidden}
.photo img{max-width:100%;max-height:100%;object-fit:contain}
.answer{min-height:500px;border:1px solid #dbe5e1;border-radius:8px;padding:18px;overflow-wrap:anywhere}
[hidden]{display:none!important}@media(max-width:750px){main{padding:20px 14px}.comparison{grid-template-columns:1fr}
.comparison section+section{border-left:0;border-top:1px solid #dbe5e1}.photo,.answer{height:auto;min-height:220px}.photo img{max-height:420px}}
</style></head><body><main><h1>เทียบคำตอบข้อ 1 กับภาพต้นฉบับ</h1>
<p>คำตอบ 34 รายการจาก Excel ปัจจุบัน เทียบกับภาพกระดาษคำตอบในชุดข้อสอบเก่า</p>
<div class="toolbar"><input id="search" type="search" placeholder="ค้นหารหัส DS, ชื่อภาพ หรือข้อความคำตอบ"
aria-label="ค้นหารายการ"><span id="count">แสดง 34 จาก 34 รายการ</span></div>
<div class="note">ภาพด้านซ้ายเป็นภาพต้นฉบับก่อนลบคะแนนและรอยตรวจ ข้อความด้านขวาถอดเฉพาะคำ ตัวเลข และสัญลักษณ์จากคำตอบในภาพ โดยไม่เติมคำบรรยายรูป</div>
""" + "".join(cards) + """</main><script>
const search=document.getElementById('search'),cards=[...document.querySelectorAll('.card')],count=document.getElementById('count');
search.addEventListener('input',()=>{const q=search.value.trim().toLocaleLowerCase();let shown=0;
for(const card of cards){const visible=card.dataset.search.toLocaleLowerCase().includes(q);card.hidden=!visible;if(visible)shown++}
count.textContent=`แสดง ${shown} จาก ${cards.length} รายการ`});
</script></body></html>"""
output = ROOT / "docs_and_tests/text_image_comparison.html"
output.write_text(html, encoding="utf-8")
print(f"Wrote {output}: {len(cards)} scored source images")
