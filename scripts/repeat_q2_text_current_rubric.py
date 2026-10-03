"""Repeat Q2 text grading once and compare with saved workbook scores."""

import asyncio
import hashlib
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")
from scripts.regrade_corrected_q2_q3_answers import (  # noqa: E402
    OPENAI_MODEL,
    PROMPT_VERSION,
    score_with_openai,
)

BOOK = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx"
OUT = ROOT / "docs_and_tests/q2_q3_review" / os.getenv(
    "Q2_REPEAT_REPORT", "q2_text_repeat_current_rubric.json"
)
RUBRIC_VARIANT = os.getenv("Q2_RUBRIC_VARIANT", "current")
assert RUBRIC_VARIANT in ("current", "partial_credit", "user_four_level")
CANDIDATE_RUBRIC = (
    "เกณฑ์ให้คะแนนข้อ 2 (คะแนนเต็ม 2; เลือกได้เฉพาะ 0, 1, 1.5, 2):\n"
    "2.00: อธิบายถูกว่าเมื่อข้อมูลใหญ่ขึ้น O(n log n) มีจำนวนขั้นตอนเพิ่มช้ากว่า O(n^2) "
    "จึงเหมาะกว่า และมีตัวอย่างอัลกอริทึมหรือกลไกการทำงานที่ถูกต้องอย่างน้อยหนึ่งอย่าง\n"
    "1.50: เปรียบเทียบได้ถูกทิศทางว่า O(n log n) เหมาะกับข้อมูลใหญ่กว่า O(n^2) "
    "และมีเหตุผลบางส่วนหรือตัวอย่างอัลกอริทึม/โค้ดที่เกี่ยวข้องกับความซับซ้อนอย่างน้อยหนึ่งฝั่ง "
    "แม้รายละเอียดอื่นผิดหรือไม่ครบ แต่ต้องยังมีสาระถูกต้องพอให้เห็นความเข้าใจ เช่น อธิบายว่าลูปซ้อนทำงานหลายรอบ "
    "พร้อมพยายามเปรียบเทียบกับวิธีที่ทำงานน้อยกว่า แม้จัดประเภทโค้ดบางส่วนผิด\n"
    "1.00: บอกได้เพียงว่า O(n log n) เร็วกว่า/เหมาะกว่าโดยไม่มีเหตุผลหรือตัวอย่างที่ถูกต้องเพิ่ม "
    "หรือมีตัวอย่างที่ถูกต้องเพียงอย่างเดียวโดยไม่อธิบายความสัมพันธ์ของความซับซ้อน\n"
    "0.00: ไม่ตอบ ไม่เกี่ยวข้อง หรือไม่มีสาระที่ถูกต้องเกี่ยวกับการเปรียบเทียบและไม่มีตัวอย่างที่เกี่ยวข้องอย่างถูกต้อง\n"
    "การมีโค้ด recursive อย่างเดียวไม่พิสูจน์ว่าเป็น O(n log n); factorial แบบ fac(n-1) เป็น O(n). "
    "หากคำตอบมีทั้งส่วนถูกและผิด ให้พิจารณาส่วนที่ถูกต้องก่อน ไม่ตัดเป็นศูนย์เพราะข้อผิดพลาดเพียงจุดเดียว. "
    "ตัวอย่างที่ถูกต้องได้แก่ Merge Sort/Heap Sort สำหรับ O(n log n) และ Bubble Sort หรือ "
    "ลูปซ้อนที่แต่ละลูปวนตาม n สำหรับ O(n^2)."
)
USER_FOUR_LEVEL_RUBRIC = (
    "2.00 คะแนน\n"
    "อธิบายได้ถูกต้องว่าเมื่อขนาดข้อมูลเพิ่มขึ้น O(n log n) มีอัตราการเติบโตของจำนวนขั้นตอนต่ำกว่า O(n²) "
    "จึงเหมาะกับข้อมูลขนาดใหญ่มากกว่า และยกชื่อ Algorithm ที่เกี่ยวข้องได้ถูกต้อง เช่น Merge Sort หรือ Heap Sort\n"
    "อนุโลมการใช้ภาษาพูดหรือคำอธิบายที่ไม่เป็นทางการ หากสาระสำคัญถูกต้อง\n\n"
    "1.50 คะแนน\n"
    "แสดงความเข้าใจหลักเกี่ยวกับความแตกต่างระหว่าง O(n log n) และ O(n²) ได้ถูกต้อง "
    "และมีตัวอย่าง Algorithm หรือโค้ดประกอบ แต่คำอธิบายหรือตัวอย่างมีข้อผิดพลาด ไม่สมบูรณ์ "
    "หรือคลาดเคลื่อนบางส่วน โดยยังมีสาระที่ถูกต้องเพียงพอแสดงถึงความเข้าใจของผู้ตอบ\n\n"
    "1.00 คะแนน\n"
    "มีสาระสำคัญถูกต้องเพียงด้านใดด้านหนึ่ง เช่น\n"
    "- อธิบายได้ว่า O(n log n) เหมาะกับข้อมูลขนาดใหญ่มากกว่า O(n²) แต่ไม่มีตัวอย่าง Algorithm ที่ถูกต้อง\n"
    "- ระบุชื่อ Algorithm ที่มีความซับซ้อน O(n log n) ได้ถูกต้อง แต่ไม่มีคำอธิบายเหตุผลเปรียบเทียบที่มีสาระ\n"
    "- กล่าวเพียงว่าหนึ่งแบบ 'เร็วกว่า' หรือ 'มีประสิทธิภาพกว่า' โดยไม่ได้แสดงความเข้าใจเพิ่มเติมอย่างเพียงพอ\n\n"
    "0.00 คะแนน\n"
    "ไม่ตอบ ตอบไม่เกี่ยวข้อง หรือคำตอบไม่มีสาระที่ถูกต้องเกี่ยวกับการเปรียบเทียบ O(n log n) กับ O(n²) "
    "และไม่มีตัวอย่างที่เกี่ยวข้องอย่างถูกต้อง"
)
USER_FOUR_LEVEL_KEY = (
    "คำตอบควรแสดงความเข้าใจว่า เมื่อจำนวนข้อมูล n เพิ่มขึ้น O(n log n) จะมีจำนวนขั้นตอนการทำงานเพิ่มขึ้นช้ากว่า "
    "O(n²) อย่างมาก จึงเหมาะกับข้อมูลขนาดใหญ่มากกว่า\n\n"
    "ตัวอย่าง Algorithm ที่ยอมรับได้สำหรับ O(n log n) เช่น Merge Sort และ Heap Sort\n\n"
    "ตัวอย่างลักษณะการทำงานแบบ O(n²) เช่น การวนลูปซ้อนกันที่แต่ละลูปทำงานตามจำนวนข้อมูล n "
    "หรือ Algorithm เช่น Bubble Sort ในกรณีทั่วไป"
)
ANSWER_KEY = (
    "O(n log n) เติบโตช้ากว่า O(n^2) เมื่อข้อมูลมีขนาดใหญ่ เช่น Merge Sort หรือ "
    "Quick Sort เทียบกับ Bubble Sort หรือ Selection Sort"
)
ALLOWED = [0.0, 1.0, 1.5, 2.0]


async def main():
    wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    sheet = wb["ชุดข้อสอบ_dataset"]
    rubric_sheet = wb["Exam_Rubrics"]
    rubric_description = {
        "current": rubric_sheet.cell(7, 6).value,
        "partial_credit": CANDIDATE_RUBRIC,
        "user_four_level": USER_FOUR_LEVEL_RUBRIC,
    }[RUBRIC_VARIANT]
    answer_key = USER_FOUR_LEVEL_KEY if RUBRIC_VARIANT == "user_four_level" else ANSWER_KEY
    rubric = [{
        "name": rubric_sheet.cell(7, 4).value,
        "score": 2.0,
        "description": rubric_description,
        "allowed_scores": ALLOWED,
    }]
    if RUBRIC_VARIANT == "current":
        assert "คำอธิบายเพิ่มเติมในการแยกคะแนนบางส่วน" not in rubric[0]["description"]
    items = []
    for row_no in range(40, 74):
        row = [sheet.cell(row_no, col).value for col in range(1, 9)]
        assert row[0] == f"DS-{row_no - 5:03d}" and row[1] == 2
        items.append({
            "sample_id": row[0], "question": row[3], "answer": row[5],
            "human_score": float(row[6]), "saved_ai_score": float(row[7]),
        })
    wb.close()
    report = {
        "started_at": datetime.now(timezone.utc).isoformat(),
        "model": OPENAI_MODEL, "prompt_version": PROMPT_VERSION,
        "grading_function": "server.services.openai_grading.score_with_openai",
        "rubric_variant": RUBRIC_VARIANT,
        "source_sha256": hashlib.sha256(BOOK.read_bytes()).hexdigest(),
        "answer_key": answer_key, "rubric": rubric, "allowed_scores": ALLOWED,
        "input_modality": "text_only", "results": [],
    }
    semaphore = asyncio.Semaphore(3)

    async def grade(item):
        async with semaphore:
            try:
                result = await score_with_openai(
                    question_text=str(item["question"]), answer_text=str(item["answer"]),
                    max_score=2.0, answer_key=answer_key, rubrics=rubric,
                    allowed_scores=ALLOWED,
                )
                success = (
                    not result.get("metrics", {}).get("manual_review_required", False)
                    and float(result.get("score", -1)) in ALLOWED
                    and bool(result.get("teacher_feedback", "").strip())
                )
                error = None
            except Exception as exc:
                result, success, error = None, False, str(exc)
            return {**item, "success": success, "error": error, "result": result}

    tasks = [asyncio.create_task(grade(item)) for item in items]
    for task in asyncio.as_completed(tasks):
        item = await task
        report["results"].append(item)
        OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        score = item["result"]["score"] if item["result"] else None
        print(item["sample_id"], "saved", item["saved_ai_score"],
              "repeat", score, "success", item["success"], flush=True)

    report["results"].sort(key=lambda item: item["sample_id"])
    valid = [item for item in report["results"] if item["success"]]
    report["summary"] = {
        "completed": len(valid), "total": len(items),
        "same_as_saved": sum(x["result"]["score"] == x["saved_ai_score"] for x in valid),
        "changed_from_saved": sum(x["result"]["score"] != x["saved_ai_score"] for x in valid),
        "repeat_teacher_exact": sum(x["result"]["score"] == x["human_score"] for x in valid),
        "saved_teacher_exact": sum(x["saved_ai_score"] == x["human_score"] for x in valid),
    }
    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report["summary"]), flush=True)


if __name__ == "__main__":
    asyncio.run(main())
