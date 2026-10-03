import os
import sys
from pathlib import Path
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[1]

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3"):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), color)
        tblBorders.append(border)
    border = OxmlElement('w:insideV')
    border.set(qn('w:val'), 'none')
    tblBorders.append(border)
    tblPr.append(tblBorders)

def add_styled_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = "TH Sarabun New"
    run.bold = True
    if level == 0:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run.font.size = Pt(20)
    elif level == 1:
        run.font.size = Pt(18)
    elif level == 2:
        run.font.size = Pt(16)
    elif level == 3:
        run.font.size = Pt(15)
    return p

def add_body_paragraph(doc, text, bold_prefix="", indent=True):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    if indent:
        p.paragraph_format.first_line_indent = Inches(0.5)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "TH Sarabun New"
        r_pre.font.size = Pt(16)
        r_pre.bold = True
    r = p.add_run(text)
    r.font.name = "TH Sarabun New"
    r.font.size = Pt(16)
    return p

def create_table_header(table, headers, col_widths=None):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        for r in p.runs:
            r.font.name = "TH Sarabun New"
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = RGBColor(255, 255, 255)
        if col_widths and i < len(col_widths):
            hdr_cells[i].width = col_widths[i]

def add_table_row(table, row_data, is_zebra=False, col_widths=None, align_center_cols=[]):
    row_cells = table.add_row().cells
    bg_color = "F8FAFC" if is_zebra else "FFFFFF"
    for i, text in enumerate(row_data):
        row_cells[i].text = str(text)
        set_cell_background(row_cells[i], bg_color)
        set_cell_margins(row_cells[i], top=90, bottom=90, left=120, right=120)
        p = row_cells[i].paragraphs[0]
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if i in align_center_cols:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for r in p.runs:
            r.font.name = "TH Sarabun New"
            r.font.size = Pt(14)
        if col_widths and i < len(col_widths):
            row_cells[i].width = col_widths[i]

print("Docx generator helper initialized.")
