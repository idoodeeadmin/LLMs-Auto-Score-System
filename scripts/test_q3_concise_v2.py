import asyncio
import json
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q3_QUESTION_TEXT = "การใช้ลิ้งค์ลิสต์เป็นสแตกและคิวจะแตกต่างจากการใช้อาร์เรย์อย่างไร และมีข้อดีข้อเสียอย่างไร (1 คะแนน)"

Q3_ANSWER_KEY = (
    "เฉลยและเกณฑ์มาตรฐาน:\n"
    "1. ความแตกต่าง (Difference):\n"
    "   - Array: มีขนาดคงที่ (Fixed-size / Static) ต้องจองพื้นที่ล่วงหน้า การทำ Stack/Queue ต้องคุม index (top, front, rear)\n"
    "   - Linked List: มีขนาดปรับเปลี่ยนได้ตามจริง (Dynamic / ยืดหยุ่น) ไม่ต้องกำหนดขนาดล่วงหน้า จัดการผ่าน Node และ Pointer (head, tail)\n"
    "2. ข้อดี-ข้อเสีย (Pros & Cons):\n"
    "   - Linked List: ข้อดีคือไม่จำกัดขนาด ไม่เกิด Stack Overflow ได้ง่าย เพิ่ม/ลดโหนดได้สะดวก; "
    "ข้อเสียคือเข้าถึงข้อมูลช้ากว่า (Sequential access) และเปลืองหน่วยความจำเก็บพอยน์เตอร์\n"
    "   - Array: ข้อดีคือเข้าถึงข้อมูลได้รวดเร็วทันทีด้วย Index (O(1)); "
    "ข้อเสียคือขนาดคงที่เสี่ยงเกิด Overflow หากจองน้อยไป หรือเปลืองพื้นที่หากจองมากเกินไป"
)

Q3_RUBRIC_NAME = "การเปรียบเทียบ Linked List vs Array สำหรับ Stack และ Queue พร้อมข้อดีข้อเสีย"

Q3_RUBRIC_DESC = (
    "เกณฑ์การให้คะแนนสำหรับการเปรียบเทียบ Linked List vs Array สำหรับ Stack และ Queue (คะแนนเต็ม 1.00):\n\n"
    "• 1.00 คะแนน (สมบูรณ์): ตอบครบ 2 ส่วนหลักอย่างชัดเจน ได้แก่ (1) อธิบายความแตกต่างของโครงสร้าง/ขนาด (Dynamic vs Static/Fixed หรือไม่ต้องประกาศขนาดล่วงหน้า) และ (2) ระบุทั้ง 'ข้อดี' และ 'ข้อเสีย' อย่างชัดเจนเป็นหัวข้อ/ประเด็น (อนุโลมสำนวนภาษาตามระดับนักศึกษา ไม่ต้องเคร่งครัดศัพท์ pointer/index หากเนื้อหาสื่อข้อดีข้อเสียได้สมเหตุสมผล)\n"
    "• 0.75 คะแนน (ดีมาก): ตอบความแตกต่างได้ถูกต้อง แต่ระบุข้อดีหรือข้อเสียเพียงด้านเดียว (เช่น มีข้อดีแต่ไม่มีข้อเสีย) หรืออธิบายข้อดีข้อเสียกว้างๆ ในเชิงการเขียนโค้ด (เช่น เขียนง่ายกว่า/ยากกว่า)\n"
    "• 0.50 คะแนน (มาตรฐาน): ตอบถูกเพียงบางส่วน ได้แก่:\n"
    "  - เปรียบเทียบความแตกต่างได้ แต่เขียนรวมเป็นความเรียงเดียว ไม่ได้แจกแจงแยกข้อดีข้อเสียให้ชัดเจน\n"
    "  - หรือระบุข้อดี หรือข้อเสียเพียงอย่างเดียวสั้นๆ\n"
    "  - หรือเน้นอธิบายพฤติกรรม LIFO/FIFO ของ Stack และ Queue แต่ระบุข้อดีข้อเสียเพียงคร่าวๆ หรือสับสน\n"
    "• 0.25 คะแนน (เล็กน้อย): คำตอบสั้นมาก คลุมเครือ ไม่ระบุโครงสร้าง แต่มีคีย์เวิร์ดที่ถูกต้องตามหลักการเพียงจุดเดียว (เช่น ยืดหยุ่น หรือ static)\n"
    "• 0.00 คะแนน (ไม่ตรงประเด็น/ผิด): ตอบไม่ตรงประเด็น (เช่น ช่วยประหยัดพลังงาน, รถไฟ), ตอบสั้นเพียงประโยคเดียวโดยไม่มีการระบุข้อดีข้อเสียเลย, มีข้อความที่ผิดหลักการอย่างมีนัยสำคัญ (เช่น อ้างว่าข้อเสียคือไม่รู้ขนาด หรืออาร์เรย์เสี่ยง Error ง่ายกว่า), หรือไม่ตอบ\n\n"
    "(ต้องส่งกลับ teacher_feedback สำหรับผู้สอน และ student_feedback เป็นคำแนะนำสำหรับนักเรียน)"
)

Q3_RUBRICS = [
    {
        "name": Q3_RUBRIC_NAME,
        "score": 1.0,
        "description": Q3_RUBRIC_DESC,
    }
]

def calculate_qwk(y_true, y_pred, step=0.25, max_score=1.0):
    if len(y_true) != len(y_pred) or len(y_true) == 0: return 0.0
    if y_true == y_pred: return 1.0
    k = int(round(max_score / step)) + 1
    w = [[((i - j) ** 2) / ((k - 1) ** 2) for j in range(k)] for i in range(k)]
    cat_true = [min(k-1, max(0, int(round(v / step)))) for v in y_true]
    cat_pred = [min(k-1, max(0, int(round(v / step)))) for v in y_pred]
    o = [[0] * k for _ in range(k)]
    for t, p in zip(cat_true, cat_pred): o[t][p] += 1
    r_t = [sum(o[i][j] for j in range(k)) for i in range(k)]
    r_p = [sum(o[i][j] for i in range(k)) for j in range(k)]
    n = len(y_true)
    e = [[(r_t[i] * r_p[j]) / n for j in range(k)] for i in range(k)]
    num = sum(w[i][j] * o[i][j] for i in range(k) for j in range(k))
    den = sum(w[i][j] * e[i][j] for i in range(k) for j in range(k))
    return 1.0 - (num / den) if den != 0 else 1.0

async def main():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws_data = wb["ชุดข้อสอบ_dataset"]
    
    all_q3 = []
    for r in range(74, 108):
        sid = ws_data.cell(r, 1).value
        q_no = ws_data.cell(r, 2).value
        ans = str(ws_data.cell(r, 6).value or "")
        h = float(ws_data.cell(r, 7).value or 0.0)
        old_ai = float(ws_data.cell(r, 8).value or 0.0)
        all_q3.append({
            "row": r,
            "sample_id": sid,
            "student_ans": ans,
            "human_score": h,
            "old_ai_score": old_ai
        })
        
    print(f"=== TESTING CLEAN CONCISE V2 RUBRIC ON ALL 34 STUDENTS ===")
    sem = asyncio.Semaphore(6)
    
    async def grade_one(item):
        async with sem:
            res = await score_with_openai(
                question_text=Q3_QUESTION_TEXT,
                answer_text=item["student_ans"],
                max_score=1.0,
                answer_key=Q3_ANSWER_KEY,
                rubrics=Q3_RUBRICS,
                image_bytes_list=None,
                strict_rubric_enforcement=True,
            )
            new_ai = float(res.get("score", 0.0))
            diff = round(new_ai - item["human_score"], 2)
            match = (new_ai == item["human_score"])
            
            return {
                "row": item["row"],
                "sample_id": item["sample_id"],
                "student_ans": item["student_ans"],
                "human_score": item["human_score"],
                "old_ai_score": item["old_ai_score"],
                "new_ai_score": new_ai,
                "confidence": res.get("confidence", "high"),
                "diff": diff,
                "is_exact": match,
                "teacher_feedback": res.get("teacher_feedback", ""),
                "student_feedback": res.get("student_feedback", "")
            }

    tasks = [grade_one(item) for item in all_q3]
    results = await asyncio.gather(*tasks)
    results = sorted(results, key=lambda x: x["row"])
    
    print("\n--- RESULTS PER STUDENT (CONCISE V2) ---")
    for r in results:
        status = "🟢 EXACT" if r["is_exact"] else f"🟠 DIFF ({r['diff']:+0.2f})"
        print(f"{r['sample_id']} | Human: {r['human_score']:.2f} | AI: {r['new_ai_score']:.2f} -> {status}")
        
    y_true = [x["human_score"] for x in results]
    y_pred = [x["new_ai_score"] for x in results]
    
    exact = sum(1 for x in results if x["is_exact"])
    w025 = sum(1 for x in results if abs(x["diff"]) <= 0.2501)
    w050 = sum(1 for x in results if abs(x["diff"]) <= 0.5001)
    mae = sum(abs(x["diff"]) for x in results) / len(results)
    qwk = calculate_qwk(y_true, y_pred)
    
    print("\n" + "="*70)
    print("CONCISE V2 RUBRIC PERFORMANCE (N=34)")
    print(f"Exact Match:  {exact} / 34 ({exact/34*100:.1f}%) | MAE: {mae:.4f} | QWK: {qwk:.4f}")
    print(f"Within ≤0.25: {w025} / 34 ({w025/34*100:.1f}%)")
    print(f"Within ≤0.50: {w050} / 34 ({w050/34*100:.1f}%)")
    print("="*70)
    
    out_file = ROOT / "artifacts" / "q3_all34_concise_v2_results.json"
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)

if __name__ == "__main__":
    asyncio.run(main())
