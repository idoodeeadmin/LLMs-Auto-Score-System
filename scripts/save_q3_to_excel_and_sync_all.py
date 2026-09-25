import json
import os
import sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

EXCEL_DATASET = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
EXCEL_AI_RUBRICS = ROOT / "ชุดข้อสอบใหม่" / "เกณฑ์ตรวจสำหรับAI.xlsx"
RESULTS_JSON = ROOT / "artifacts" / "q3_all34_calibrated_v5_results.json"

from scripts.test_q3_calibrated_v5 import Q3_RUBRIC_NAME, Q3_RUBRIC_DESC

def main():
    if not RESULTS_JSON.exists():
        print(f"Results file not found: {RESULTS_JSON}")
        return
        
    with open(RESULTS_JSON, "r", encoding="utf-8") as f:
        results = json.load(f)
        
    print(f"Loaded {len(results)} evaluated items from {RESULTS_JSON.name}")
    
    # 1. Update ชุดข้อสอบ_dataset.xlsx
    wb = openpyxl.load_workbook(EXCEL_DATASET)
    ws_data = wb["ชุดข้อสอบ_dataset"]
    
    for item in results:
        r = item["row"]
        assert ws_data.cell(r, 1).value == item["sample_id"]
        assert ws_data.cell(r, 2).value == 3
        
        ws_data.cell(r, 8, item["new_ai_score"])
        ws_data.cell(r, 9, item["confidence"])
        
        # Dual feedback
        tf = item.get("teacher_feedback", "").strip()
        sf = item.get("student_feedback", "").strip()
        if tf and sf:
            combined = f"[สำหรับผู้สอน]\n{tf}\n\n[สำหรับนักเรียน]\n{sf}"
        else:
            combined = tf or sf
        ws_data.cell(r, 10, combined)
        
    print(f"Updated rows 74-107 in sheet 'ชุดข้อสอบ_dataset'")
    
    # 2. Update Exam_Rubrics sheet
    if "Exam_Rubrics" in wb.sheetnames:
        ws_rubric = wb["Exam_Rubrics"]
        # Find rows for Q3
        for r in range(4, ws_rubric.max_row + 1):
            q_val = ws_rubric.cell(r, 1).value
            if q_val == 3:
                ws_rubric.cell(r, 4, Q3_RUBRIC_NAME)
                ws_rubric.cell(r, 5, 1.0)
                ws_rubric.cell(r, 6, Q3_RUBRIC_DESC)
        print("Updated Q3 in sheet 'Exam_Rubrics'")
        
    wb.save(EXCEL_DATASET)
    print(f"Successfully saved {EXCEL_DATASET}")
    
    # 3. Update เกณฑ์ตรวจสำหรับAI.xlsx
    if EXCEL_AI_RUBRICS.exists():
        wb_ai = openpyxl.load_workbook(EXCEL_AI_RUBRICS)
        if "ข้อมูลสำหรับ API" in wb_ai.sheetnames:
            ws_ai = wb_ai["ข้อมูลสำหรับ API"]
            updated = 0
            for r in range(2, ws_ai.max_row + 1):
                q_text = str(ws_ai.cell(r, 1).value or "")
                if "ลิ้งค์ลิสต์" in q_text or "สแตกและคิว" in q_text:
                    ws_ai.cell(r, 2, Q3_RUBRIC_DESC)
                    updated += 1
            wb_ai.save(EXCEL_AI_RUBRICS)
            print(f"Updated {updated} rows in '{EXCEL_AI_RUBRICS.name}' (Sheet: ข้อมูลสำหรับ API)")

if __name__ == "__main__":
    main()
