import asyncio
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env")
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q2_RUBRIC_NAME = "ความซับซ้อน O(n log n) vs O(n^2) และตัวอย่างอัลกอริทึม"
Q2_RUBRIC_DESC = (
    "โจทย์: อธิบายว่าทำไม O(n log n) เหมาะกับข้อมูลใหญ่กว่า O(n^2) และยกตัวอย่าง Algorithm (คะแนนเต็ม 2.00 คะแนน)\n"
    "เกณฑ์การให้คะแนนแบ่งเป็น 2 กรณีตามการมีอยู่ของชื่อ Algorithm อย่างเคร่งครัด:\n\n"
    "กรณีที่ 1: มีการระบุชื่อ Algorithm ที่เกี่ยวข้องอย่างน้อย 1 ชื่อ (เช่น Merge Sort, Quick Sort, Bubble Sort, Insertion Sort, Heap Sort, Selection Sort เป็นต้น):\n"
    "• ได้ 2.00 คะแนน (เต็ม): อธิบายเหตุผลที่สื่อถึงความเข้าใจว่า O(n log n) ดีกว่า/เร็วกว่า/ทำงานน้อยกว่า/ลดเวลา/แบ่งครึ่ง/ลูปน้อยกว่า O(n^2) เมื่อข้อมูลมีขนาดใหญ่ และมีชื่อ Algorithm ที่เกี่ยวข้อง (ให้อนุโลมภาษาพูดและไม่หักคะแนนหากนิสิตใส่รายละเอียดเกินเช่น memory/recursive/CPU)\n"
    "• ได้ 1.50 คะแนน: มีชื่อ Algorithm ที่เกี่ยวข้อง แต่อธิบายเหตุผลสั้นมาก หรือคำอธิบายกำกวม\n"
    "• ได้ 1.00 คะแนน: มีชื่อ Algorithm ถูกต้อง แต่ไม่ได้อธิบายเหตุผลเลย\n\n"
    "กรณีที่ 2: ไม่มีการระบุชื่อ Algorithm ใดๆ เลย (ไม่มีชื่อ Merge, Quick, Bubble, Insertion, Sort ฯลฯ):\n"
    "• ได้ 1.00 คะแนน: อธิบายเหตุผลได้ว่า O(n log n) ดีกว่า/เร็วกว่า/ตัดครึ่ง/ลูปน้อยกว่า O(n^2) แต่ 'ไม่ได้ยกตัวอย่างชื่อ Algorithm เลย' (คะแนนเต็มของกรณีนี้คือ 1.00 คะแนนเท่านั้น ห้ามให้ 1.50 หรือ 2.00 เด็ดขาด)\n"
    "• ได้ 0.50 คะแนน: ตอบสั้นๆ แค่ว่าเร็วกว่าโดยไม่มีเหตุผลเปรียบเทียบและไม่มีตัวอย่าง\n"
    "• ได้ 0.00 คะแนน: ตอบผิดหลักการทั้งหมด ไม่ตอบ หรือตอบเรื่องลูกเต๋าโยนเหรียญ หรือเขียนเฉพาะตัวโจทย์ซ้ำ"
)

Q2_RUBRICS = [
    {
        "name": Q2_RUBRIC_NAME,
        "score": 2.0,
        "description": Q2_RUBRIC_DESC,
        "allowed_scores": [0.0, 0.5, 1.0, 1.5, 2.0],
    }
]

Q2_QUESTION_TEXT = "อธิบายว่าทำไม O(n log n) เหมาะกับข้อมูลใหญ่กว่า O(n^2) และยกตัวอย่าง Algorithm"
Q2_ANSWER_KEY = (
    "คำตอบที่สมบูรณ์ควรมี 2 ส่วน:\n"
    "1. เหตุผล: O(n log n) มีอัตราการเติบโตของเวลา/จำนวนรอบการทำงานช้ากว่า O(n^2) มากเมื่อขนาดข้อมูล (n) มีขนาดใหญ่ ทำให้ประหยัดเวลาและทรัพยากรมากกว่า O(n^2) ที่มักมีการวนลูปซ้อน 2 ชั้น\n"
    "2. ตัวอย่าง Algorithm:\n"
    "   - อัลกอริทึมกลุ่ม O(n log n) เช่น Merge Sort, Quick Sort (กรณีเฉลี่ย), Heap Sort\n"
    "   - อัลกอริทึมกลุ่ม O(n^2) เช่น Bubble Sort, Insertion Sort, Selection Sort"
)

async def main():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["ชุดข้อสอบ_dataset"]
    rows = [r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value == 2]
    assert len(rows) == 34
    
    print(f"Total students to grade for Q2: {len(rows)}")
    
    sem = asyncio.Semaphore(4)
    exact = 0
    within_05 = 0
    total_diff = 0.0

    async def grade_one(r):
        nonlocal exact, within_05, total_diff
        sid = ws.cell(r, 1).value
        ans = str(ws.cell(r, 6).value or "")
        human = float(ws.cell(r, 7).value or 0.0)
        async with sem:
            res = await score_with_openai(
                question_text=Q2_QUESTION_TEXT,
                answer_text=ans,
                max_score=2.0,
                answer_key="",
                rubrics=Q2_RUBRICS,
                allowed_scores=[0.0, 0.5, 1.0, 1.5, 2.0],
                strict_rubric_enforcement=True,
            )
            ai = float(res["score"])
            diff = round(ai - human, 2)
            is_match = (diff == 0.0)
            if is_match:
                exact += 1
            if abs(diff) <= 0.5:
                within_05 += 1
            total_diff += abs(diff)
            print(f"[{sid}] Human: {human:.1f} | AI: {ai:.1f} | {'MATCH' if is_match else f'DIFF={diff:+.1f}'}")
            return r, ai, res.get("confidence", "medium"), res.get("feedback", "")

    tasks = [grade_one(r) for r in rows]
    results = await asyncio.gather(*tasks)

    # Save to Excel
    for r, ai, conf, fb in results:
        ws.cell(r, 8, ai)
        ws.cell(r, 9, conf)
        ws.cell(r, 10, fb)

    # Update Exam_Rubrics Row 7
    ws_rubrics = wb["Exam_Rubrics"]
    ws_rubrics.cell(7, 4, Q2_RUBRIC_NAME)
    ws_rubrics.cell(7, 5, 2.0)
    ws_rubrics.cell(7, 6, Q2_RUBRIC_DESC)

    wb.save(excel_path)
    print(f"\n==========================================")
    print(f"Q2 Evaluation Complete!")
    print(f"Exact Matches: {exact}/34 ({exact/34*100:.2f}%)")
    print(f"Within +/-0.50: {within_05}/34 ({within_05/34*100:.2f}%)")
    print(f"MAE: {total_diff/34:.4f}")
    print(f"==========================================")

if __name__ == "__main__":
    asyncio.run(main())
