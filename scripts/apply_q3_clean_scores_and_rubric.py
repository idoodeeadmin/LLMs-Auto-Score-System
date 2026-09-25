import json
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[1]
RESULTS_FILE = ROOT / 'artifacts' / 'q3_clean_test_results.json'
EXCEL_FILE = ROOT / 'ชุดข้อสอบใหม่' / 'ชุดข้อสอบ_dataset.xlsx'
RUBRIC_AI_FILE = ROOT / 'ชุดข้อสอบใหม่' / 'เกณฑ์ตรวจสำหรับAI.xlsx'

PART1_DESC = (
    "ประเมินการเปรียบเทียบโครงสร้างระหว่าง Linked List กับ Array ในการทำ Stack/Queue (ให้ 0.50, 0.25 หรือ 0.00 คะแนน):\n"
    "• 0.50 คะแนน: อธิบายความแตกต่างของโครงสร้างได้ชัดเจน เช่น Array มีขนาดคงที่ (Fixed size) หรือต้องระบุขนาดล่วงหน้า จองพื้นที่ต่อเนื่อง "
    "ส่วน Linked List มีขนาดปรับเปลี่ยนได้แบบพลวัต (Dynamic size) ยืดหยุ่น หรือใช้พอยน์เตอร์เชื่อมโหนด\n"
    "• 0.25 คะแนน: อธิบายถูกเพียงเบื้องต้น หรือตอบถูกเพียงฝั่งใดฝั่งหนึ่ง\n"
    "• 0.00 คะแนน: ไม่ได้เปรียบเทียบระหว่าง Linked List กับ Array (เช่น ตอบเฉพาะการทำงานของ Stack/Queue หรือตอบไม่เกี่ยวกับโครงสร้างข้อมูล) หรือตอบผิดหลักการ"
)

PART2_DESC = (
    "ประเมินการระบุข้อดีและข้อเสียของการเลือกใช้ (ให้ 0.50, 0.25 หรือ 0.00 คะแนน):\n"
    "• 0.50 คะแนน: ระบุทั้งข้อดีและข้อเสียได้สมเหตุสมผล เช่น Linked List ไม่จำกัดขนาด/ไม่เกิด overflow แต่เข้าถึงช้ากว่า/เปลืองเนื้อที่ pointer หรือ Array เข้าถึงเร็ว O(1) แต่ขยายขนาดยาก/เสี่ยง overflow\n"
    "• 0.25 คะแนน: ระบุเฉพาะข้อดีอย่างเดียว หรือเฉพาะข้อเสียอย่างเดียว หรือข้อดีข้อเสียยังไม่ชัดเจน\n"
    "• 0.00 คะแนน: ไม่ได้ระบุข้อดีข้อเสีย หรือระบุผิดหลักการ"
)

COMBINED_AI_CRITERIA = (
    "เกณฑ์การประเมิน 2 ส่วนรวมกัน (คะแนนเต็ม 1.00 - คะแนนรวมย่อย 0.00, 0.25, 0.50, 0.75 หรือ 1.00):\n\n"
    "ส่วนที่ 1: ความแตกต่างเชิงโครงสร้าง (คะแนนเต็ม 0.50)\n"
    "• 0.50 คะแนน: อธิบายความแตกต่างของโครงสร้างได้ชัดเจน เช่น Array มีขนาดคงที่ (Fixed size) หรือต้องระบุขนาดล่วงหน้า จองพื้นที่ต่อเนื่อง "
    "ส่วน Linked List มีขนาดปรับเปลี่ยนได้แบบพลวัต (Dynamic size) ยืดหยุ่น หรือใช้พอยน์เตอร์เชื่อมโหนด\n"
    "• 0.25 คะแนน: อธิบายถูกเพียงเบื้องต้น หรือตอบถูกเพียงฝั่งใดฝั่งหนึ่ง\n"
    "• 0.00 คะแนน: ไม่ได้เปรียบเทียบระหว่าง Linked List กับ Array (เช่น ตอบเฉพาะการทำงานของ Stack/Queue หรือตอบไม่เกี่ยวกับโครงสร้างข้อมูล) หรือตอบผิดหลักการ\n\n"
    "ส่วนที่ 2: ข้อดีและข้อเสีย (คะแนนเต็ม 0.50)\n"
    "• 0.50 คะแนน: ระบุทั้งข้อดีและข้อเสียได้สมเหตุสมผล เช่น Linked List ไม่จำกัดขนาด/ไม่เกิด overflow แต่เข้าถึงช้ากว่า/เปลืองเนื้อที่ pointer หรือ Array เข้าถึงเร็ว O(1) แต่ขยายขนาดยาก/เสี่ยง overflow\n"
    "• 0.25 คะแนน: ระบุเฉพาะข้อดีอย่างเดียว หรือเฉพาะข้อเสียอย่างเดียว หรือข้อดีข้อเสียยังไม่ชัดเจน\n"
    "• 0.00 คะแนน: ไม่ได้ระบุข้อดีข้อเสีย หรือระบุผิดหลักการ\n\n"
    "(ต้องระบุ teacher_feedback แจกแจงจุดให้/หักคะแนนสำหรับผู้สอน และ student_feedback เป็นคำแนะนำพัฒนาการเรียนรู้สำหรับนักเรียน)"
)

def apply_q3():
    print(f"Loading results from {RESULTS_FILE.name}...")
    with open(RESULTS_FILE, encoding='utf-8') as f:
        data = json.load(f)

    results_map = {r['sample_id']: r for r in data['results']}
    print(f"Loaded {len(results_map)} student results for Question 3.")

    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb['ชุดข้อสอบ_dataset']

    thin_border = Border(
        left=Side(style='thin', color='00BFBFBF'),
        right=Side(style='thin', color='00BFBFBF'),
        top=Side(style='thin', color='00BFBFBF'),
        bottom=Side(style='thin', color='00BFBFBF')
    )
    font_aptos = Font(name='Aptos', size=10)
    align_center = Alignment(horizontal='center', vertical='top')
    align_feedback = Alignment(horizontal='left', vertical='top', wrap_text=True)

    # 1. Update ชุดข้อสอบ_dataset rows 74 to 107
    updated_count = 0
    for row in range(74, 108):
        sid = str(ws.cell(row=row, column=1).value or '').strip()
        qno = ws.cell(row=row, column=2).value
        if qno != 3 or not sid or sid not in results_map:
            continue

        res = results_map[sid]
        score = float(res['ai'])
        conf = str(res.get('confidence', 'high'))
        fb_teacher = res.get('teacher_feedback', '').strip()
        fb_student = res.get('student_feedback', '').strip()
        full_fb = f"[สำหรับผู้สอน]\n{fb_teacher}\n\n[สำหรับนักเรียน]\n{fb_student}"

        c8 = ws.cell(row=row, column=8, value=score)
        c8.font = font_aptos
        c8.alignment = align_center
        c8.border = thin_border

        c9 = ws.cell(row=row, column=9, value=conf)
        c9.font = font_aptos
        c9.alignment = align_center
        c9.border = thin_border

        c10 = ws.cell(row=row, column=10, value=full_fb)
        c10.font = font_aptos
        c10.alignment = align_feedback
        c10.border = thin_border

        updated_count += 1

    print(f"Updated {updated_count} student scores and dual feedbacks in 'ชุดข้อสอบ_dataset' (rows 74-107)")

    # 2. Update Exam_Rubrics Sheet: Rows 8 and 9
    if 'Exam_Rubrics' in wb.sheetnames:
        ws_rubric = wb['Exam_Rubrics']
        # Row 8: Part 1
        ws_rubric.cell(row=8, column=1, value=3)
        ws_rubric.cell(row=8, column=2, value='Linked List vs Array (Stack & Queue)')
        ws_rubric.cell(row=8, column=3, value=1.0)
        ws_rubric.cell(row=8, column=4, value='ความแตกต่างเชิงโครงสร้าง (Linked List vs Array)')
        ws_rubric.cell(row=8, column=5, value=0.5)
        ws_rubric.cell(row=8, column=6, value=PART1_DESC)

        # Row 9: Part 2
        ws_rubric.cell(row=9, column=1, value=3)
        ws_rubric.cell(row=9, column=2, value='Linked List vs Array (Stack & Queue)')
        ws_rubric.cell(row=9, column=3, value=1.0)
        ws_rubric.cell(row=9, column=4, value='การระบุข้อดีและข้อเสีย (Pros & Cons)')
        ws_rubric.cell(row=9, column=5, value=0.5)
        ws_rubric.cell(row=9, column=6, value=PART2_DESC)
        print("Updated Question 3 rubrics in 'Exam_Rubrics' sheet (Rows 8 and 9).")

    wb.save(EXCEL_FILE)
    print(f"Saved {EXCEL_FILE.name} successfully.")

    # 3. Update เกณฑ์ตรวจสำหรับAI.xlsx
    if RUBRIC_AI_FILE.exists():
        wb_ai = openpyxl.load_workbook(RUBRIC_AI_FILE)
        ws_ai = wb_ai.worksheets[0]
        ai_updated = 0
        for r in range(2, ws_ai.max_row + 1):
            q_val = str(ws_ai.cell(row=r, column=1).value or '')
            if 'ลิ้งค์ลิสต์' in q_val or 'สแตกและคิว' in q_val:
                ws_ai.cell(row=r, column=2, value=COMBINED_AI_CRITERIA)
                ai_updated += 1
        wb_ai.save(RUBRIC_AI_FILE)
        print(f"Updated {ai_updated} rows in '{RUBRIC_AI_FILE.name}'.")

    print("\nAll Q3 clean rubric updates completed successfully!")

if __name__ == '__main__':
    apply_q3()
