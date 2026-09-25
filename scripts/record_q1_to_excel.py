import json
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[1]
JSON_FILE = ROOT / 'artifacts' / 'q1-confidence-eval-20260922-151705' / 'q1_results.json'
EXCEL_FILE = ROOT / 'ชุดข้อสอบใหม่' / 'ชุดข้อสอบ_dataset.xlsx'

def main():
    if not JSON_FILE.exists():
        raise FileNotFoundError(f"JSON results not found at {JSON_FILE}")
    if not EXCEL_FILE.exists():
        raise FileNotFoundError(f"Excel file not found at {EXCEL_FILE}")

    data = json.loads(JSON_FILE.read_text(encoding='utf-8'))
    results_map = {r['student_id']: r for r in data['results']}
    print(f"Loaded {len(results_map)} results from {JSON_FILE.name}")

    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb['ชุดข้อสอบ_dataset']

    thin_border = Border(
        left=Side(style='thin', color='00BFBFBF'),
        right=Side(style='thin', color='00BFBFBF'),
        top=Side(style='thin', color='00BFBFBF'),
        bottom=Side(style='thin', color='00BFBFBF')
    )
    font_aptos = Font(name='Aptos', size=10)
    align_center = Alignment(horizontal='center', vertical='top')
    align_feedback = Alignment(horizontal='left', vertical='top', wrap_text=True)

    updated_count = 0
    for row in range(6, ws.max_row + 1):
        sid = ws.cell(row=row, column=1).value
        qno = ws.cell(row=row, column=2).value
        if qno != 1 or not sid:
            continue
        
        sid_clean = str(sid).strip()
        if sid_clean in results_map:
            res = results_map[sid_clean]
            
            # Col 8: ai_score
            c8 = ws.cell(row=row, column=8, value=float(res['ai_score']))
            c8.font = font_aptos
            c8.alignment = align_center
            c8.border = thin_border
            
            # Col 9: ai_confidence
            c9 = ws.cell(row=row, column=9, value=str(res['confidence']))
            c9.font = font_aptos
            c9.alignment = align_center
            c9.border = thin_border
            
            # Col 10: ai_feedback
            c10 = ws.cell(row=row, column=10, value=str(res['feedback']))
            c10.font = font_aptos
            c10.alignment = align_feedback
            c10.border = thin_border
            
            updated_count += 1
            print(f"Row {row} [{sid_clean}]: AI Score={res['ai_score']}, Conf={res['confidence']}")

    wb.save(EXCEL_FILE)
    print(f"\nSuccessfully recorded {updated_count} rows to {EXCEL_FILE}")

if __name__ == '__main__':
    main()
