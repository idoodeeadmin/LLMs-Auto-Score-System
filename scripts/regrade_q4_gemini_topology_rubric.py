"""Regrade all 34 Q4 answers using Gemini 3.8 Flash with the Topology-Aware Rubric.

Evaluates student Binary Search Tree drawings with tolerance for node crowding (13/58/15)
and geometric distortion, matching human teacher grading standards.
"""

import asyncio
import hashlib
import html
import io
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

from dotenv import load_dotenv
import openpyxl
from PIL import Image

load_dotenv(ROOT / ".env")
from server.services.gemini_grading import score_with_gemini, GEMINI_MODEL

BOOK = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx"
KEY_IMAGE_PATH = ROOT / "public/answer-keys/q4_bst_ground_truth.png"
OUT_DIR = ROOT / "docs_and_tests/gemini38_regrade"
OUT_JSON = OUT_DIR / "q4_topology_rubric_results.json"
OUT_HTML = OUT_DIR / "q4_topology_rubric_comparison.html"

Q4_TOPOLOGY_RUBRIC = [
    {
        "name": "ความถูกต้องของโครงสร้างต้นไม้ค้นหาทวิภาค (Binary Search Tree 12 โหนด)",
        "score": 1.0,
        "description": (
            "ประเมินภาพวาด Binary Search Tree จากข้อมูล: 9, 16, 10, 76, 5, 13, 58, 92, 11, 15, 80, 99 (คะแนนเต็ม 1.00 คะแนน, เลือกเฉพาะ 1.00 หรือ 0.00):\n\n"
            "• โครงสร้างความสัมพันธ์ที่ถูกต้อง (Parent-Child Topology):\n"
            "  - Root คือ 9 (ซ้าย 5, ขวา 16)\n"
            "  - ใต้ 16: ซ้ายคือ 10, ขวาคือ 76\n"
            "  - ใต้ 10: ขวาคือ 13 (ซ้ายว่าง)\n"
            "  - ใต้ 13: ซ้ายคือ 11, ขวาคือ 15\n"
            "  - ใต้ 76: ซ้ายคือ 58 (เป็น Leaf Node), ขวาคือ 92\n"
            "  - ใต้ 92: ซ้ายคือ 80, ขวาคือ 99\n\n"
            "• กฎการอนุโลมสำคัญ (ดูลำดับการเชื่อมต่อ ไม่ยึดติดมุมองศาเรขาคณิต):\n"
            "  1. จุดเบียดตรงกลาง (13 vs 58): ในข้อสอบนี้ โหนด 13 (ลูกขวาของ 10) และโหนด 58 (ลูกซ้ายของ 76) จะเดินทางเข้ามาอยู่ใกล้กันตรงกลางกระดาษ "
            "     ทำให้นักเรียนแทบทุกคนวาดเบียดกัน เส้นเชื่อมของ 15 อาจวาดดิ่งลงมาหรือเบี่ยงหลบ 58 ให้พิจารณาที่ 'เจตนาของเส้นเชื่อม' ว่า 13 เชื่อมไป 11 และ 15 "
            "     และ 76 เชื่อมไป 58 หากเชื่อมถูกต้องให้ 1.00 คะแนนเต็ม ห้ามตัดสินว่าผิดเพียงเพราะเส้นอยู่ใกล้กันหรือมุมเอียงไม่เหมือนภาพเฉลยคอมพิวเตอร์\n"
            "  2. ทิศทางภาพถ่าย (Orientation): หากภาพถ่ายเอียง ตะแคง 90 องศา หรือกลับหัว ให้สังเกตทิศทางของตัวเลขและปรับมุมมองในใจให้โหนดราก 9 อยู่ด้านบนสุดเสมอ\n"
            "  3. รอยร่างและวงกลมเกิน: ให้อนุโลมความสวยงาม ลายมือ รอยดินสอร่าง หรือวงกลมเปล่าส่วนเกิน ขอเพียงโหนดทั้ง 12 ตัวและเส้นเชื่อมหลักถูกต้อง\n"
            "  4. ให้ 0.00 คะแนนเฉพาะกรณี: โหนด Root ผิด (ไม่ใช่ 9), วางโหนดผิดหลัก BST อย่างแท้จริง (เช่น 16 ไปเป็น Root, เอา 13 ไปต่อใต้ 58 โดยตรง), หรือไม่วาดคำตอบ"
        ),
        "allowed_scores": [0.0, 1.0],
    }
]

Q4_QUESTION_TEXT = "จากข้อมูลต่อไปนี้จงนำไปสร้างเป็น Binary search tree (1 คะแนน)\n9 16 10 76 5 13 58 92 11 15 80 99"
Q4_ANSWER_KEY_TEXT = (
    "เฉลยโครงสร้าง Binary Search Tree (BST) ที่ถูกต้องสมบูรณ์:\n"
    "- ลำดับการนำเข้าข้อมูล: 9, 16, 10, 76, 5, 13, 58, 92, 11, 15, 80, 99\n"
    "- Root = 9 (ซ้าย 5, ขวา 16)\n"
    "- 16: ซ้าย 10, ขวา 76\n"
    "- 10: ขวา 13\n"
    "- 13: ซ้าย 11, ขวา 15\n"
    "- 76: ซ้าย 58, ขวา 92\n"
    "- 92: ซ้าย 80, ขวา 99\n"
    "- โหนด 13 และ 58 เบียดกันตรงกลาง ให้ยึดเจตนาเส้นเชื่อมเป็นสำคัญ"
)

def esc(v):
    return html.escape(str(v or ""))

def load_q4_items():
    wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    sheet = wb["ชุดข้อสอบ_dataset"]
    items = []
    for row in range(108, 142):
        sid = sheet.cell(row, 1).value
        idx = row - 107
        img_path = ROOT / "ชุดข้อสอบใหม่/photo_clean_ชุดที่1" / f"LINE_ALBUM_Photo1_260917_{idx}.jpg"
        assert img_path.is_file(), f"Missing image {img_path}"
        items.append({
            "row": row,
            "sample_id": sid,
            "question_no": 4,
            "img_path": img_path,
            "human_score": float(sheet.cell(row, 7).value),
            "previous_ai_score": float(sheet.cell(row, 8).value),
        })
    wb.close()
    return items

async def main():
    print(f"=== Starting Q4 Regrade with Topology-Aware Rubric ===")
    print(f"Model: {GEMINI_MODEL}")
    items = load_q4_items()
    print(f"Total items: {len(items)}")

    with open(KEY_IMAGE_PATH, "rb") as f:
        key_raw = f.read()

    sem = asyncio.Semaphore(4)
    results = []

    async def grade_item(item):
        async with sem:
            sid = item["sample_id"]
            img_bytes = item["img_path"].read_bytes()

            result = None
            for attempt in range(3):
                try:
                    result = await score_with_gemini(
                        question_text=Q4_QUESTION_TEXT,
                        answer_text="",
                        max_score=1.0,
                        answer_key=Q4_ANSWER_KEY_TEXT,
                        rubrics=Q4_TOPOLOGY_RUBRIC,
                        allowed_scores=[0.0, 1.0],
                        strict_rubric_enforcement=False,
                        image_bytes_list=[img_bytes],
                        image_mime_list=["image/jpeg"],
                        answer_key_image_bytes_list=[key_raw],
                        answer_key_image_mime_list=["image/png"],
                    )
                    if isinstance(result, dict) and "score" in result:
                        break
                except Exception as e:
                    print(f"[{sid}] Attempt {attempt+1} error: {e}")
                    if attempt < 2:
                        await asyncio.sleep(2 * (attempt + 1))

            score = float(result.get("score", 0.0)) if result else 0.0
            match = score == item["human_score"]
            status = "MATCH ✅" if match else "DIFF ❌"
            print(f"[{sid}] Human: {item['human_score']:.1f} | Gem: {score:.1f} | Prev: {item['previous_ai_score']:.1f} -> {status}", flush=True)

            return {
                "sample_id": sid,
                "row": item["row"],
                "human_score": item["human_score"],
                "previous_ai_score": item["previous_ai_score"],
                "gemini_score": score,
                "confidence": result.get("confidence", "medium") if result else "low",
                "teacher_feedback": result.get("teacher_feedback", "") if result else "",
                "student_feedback": result.get("student_feedback", "") if result else "",
                "image_rel": str(item["img_path"].relative_to(ROOT)).replace("\\", "/"),
                "match": match,
            }

    tasks = [asyncio.create_task(grade_item(item)) for item in items]
    results = await asyncio.gather(*tasks)

    # Sort by row
    results.sort(key=lambda x: x["row"])

    exact = sum(1 for r in results if r["match"])
    mae = sum(abs(r["gemini_score"] - r["human_score"]) for r in results) / len(results)
    prev_exact = sum(1 for r in results if r["previous_ai_score"] == r["human_score"])
    prev_mae = sum(abs(r["previous_ai_score"] - r["human_score"]) for r in results) / len(results)

    print("\n" + "="*50)
    print(f"Q4 Regrade Summary (34 items):")
    print(f"  Gemini 3.8 Flash (Topology Rubric): Exact {exact}/34 ({exact/34*100:.1f}%), MAE: {mae:.4f}")
    print(f"  Previous Baseline (Old Rubric):    Exact {prev_exact}/34 ({prev_exact/34*100:.1f}%), MAE: {prev_mae:.4f}")
    print("="*50)

    # Save JSON
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    report = {
        "timestamp": datetime.now(timezone.utc).isoformat(),
        "model": GEMINI_MODEL,
        "total": len(results),
        "exact_matches": exact,
        "match_rate": round(exact / len(results) * 100, 2),
        "mae": round(mae, 4),
        "results": results,
    }
    OUT_JSON.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"Saved JSON: {OUT_JSON}")

    # Build HTML comparison
    cards = []
    for r in results:
        badge = '<span class="badge match">ตรงผู้สอน ✅</span>' if r["match"] else '<span class="badge diff">ต่างจากผู้สอน ❌</span>'
        cards.append(f'''
        <article class="card" data-match="{int(r['match'])}">
          <div class="card-header">
            <h3>{esc(r['sample_id'])} {badge}</h3>
            <div class="scores">
              <span>ผู้สอน: <b>{r['human_score']:.1f}</b></span>
              <span>Gemini: <b class="score-gem">{r['gemini_score']:.1f}</b></span>
              <span>AI เดิม: <b>{r['previous_ai_score']:.1f}</b></span>
              <span>ความมั่นใจ: <b>{esc(r['confidence'])}</b></span>
            </div>
          </div>
          <div class="card-body">
            <div class="img-box">
              <h4>ภาพคำตอบผู้เรียน</h4>
              <img loading="lazy" src="../../{esc(r['image_rel'])}" alt="{esc(r['sample_id'])}">
            </div>
            <div class="fb-box">
              <h4>เหตุผลสำหรับผู้สอน (Teacher Feedback)</h4>
              <pre>{esc(r['teacher_feedback'])}</pre>
              <h4>คำแนะนำสำหรับนักเรียน (Student Feedback)</h4>
              <pre>{esc(r['student_feedback'])}</pre>
            </div>
          </div>
        </article>
        ''')

    html_content = f'''<!DOCTYPE html>
<html lang="th">
<head>
  <meta charset="utf-8">
  <meta name="viewport" content="width=device-width, initial-scale=1">
  <title>ผลตรวจข้อ 4 (BST 34 ข้อ) ด้วย Gemini 3.8 Flash (Topology-Aware Rubric)</title>
  <style>
    body {{ font-family: system-ui, 'Segoe UI', sans-serif; background: #0b1120; color: #f1f5f9; margin: 0; padding: 24px; line-height: 1.5; }}
    header {{ max-width: 1400px; margin: 0 auto 24px; background: #1e293b; padding: 24px; border-radius: 12px; border: 1px solid #334155; }}
    h1 {{ margin: 0 0 12px; font-size: 24px; color: #38bdf8; }}
    .stats-grid {{ display: grid; grid-template-columns: repeat(auto-fit, minmax(220px, 1fr)); gap: 16px; margin: 20px 0; }}
    .stat-card {{ background: #0f172a; padding: 16px; border-radius: 8px; border: 1px solid #334155; text-align: center; }}
    .stat-num {{ font-size: 32px; font-weight: bold; color: #38bdf8; }}
    .stat-label {{ font-size: 13px; color: #94a3b8; margin-top: 4px; }}
    .badge {{ display: inline-block; padding: 4px 10px; border-radius: 999px; font-size: 13px; font-weight: 600; margin-left: 8px; }}
    .badge.match {{ background: #065f46; color: #6ee7b7; }}
    .badge.diff {{ background: #881337; color: #fda4af; }}
    .filter-bar {{ margin-top: 16px; display: flex; gap: 12px; align-items: center; }}
    button {{ background: #334155; color: #f8fafc; border: 1px solid #475569; padding: 8px 16px; border-radius: 6px; cursor: pointer; font-weight: 500; }}
    button.active {{ background: #0284c7; border-color: #38bdf8; }}
    main {{ max-width: 1400px; margin: 0 auto; display: flex; flex-direction: column; gap: 20px; }}
    .card {{ background: #1e293b; border: 1px solid #334155; border-radius: 12px; padding: 20px; }}
    .card-header {{ display: flex; justify-content: space-between; align-items: center; border-bottom: 1px solid #334155; padding-bottom: 12px; margin-bottom: 16px; flex-wrap: wrap; gap: 10px; }}
    .card-header h3 {{ margin: 0; font-size: 18px; }}
    .scores {{ display: flex; gap: 16px; font-size: 14px; }}
    .score-gem {{ color: #38bdf8; }}
    .card-body {{ display: grid; grid-template-columns: 360px 1fr; gap: 20px; }}
    @media (max-width: 900px) {{ .card-body {{ grid-template-columns: 1fr; }} }}
    .img-box img {{ width: 100%; max-height: 480px; object-fit: contain; background: #0f172a; border-radius: 8px; border: 1px solid #334155; }}
    .fb-box h4 {{ margin: 0 0 6px; font-size: 14px; color: #94a3b8; }}
    .fb-box pre {{ background: #0f172a; padding: 12px; border-radius: 8px; border: 1px solid #334155; white-space: pre-wrap; font-family: inherit; font-size: 13.5px; line-height: 1.6; margin-bottom: 12px; color: #e2e8f0; }}
  </style>
</head>
<body>
  <header>
    <h1>ข้อ 4 (Binary Search Tree 34 ข้อ) — ผลตรวจด้วย Gemini 3.8 Flash (Topology-Aware Rubric)</h1>
    <p>เกณฑ์ฉบับปรับปรุง: อนุโลมการเบียดกันของโหนด 13, 15 และ 58 ตรงกลางกระดาษ และดูความสัมพันธ์ Parent-Child Topology เป็นหลัก</p>
    <div class="stats-grid">
      <div class="stat-card">
        <div class="stat-num">{exact} / 34</div>
        <div class="stat-label">ตรงผู้สอน (Exact Match: {exact/34*100:.1f}%)</div>
      </div>
      <div class="stat-card">
        <div class="stat-num">{mae:.3f}</div>
        <div class="stat-label">ค่าความคลาดเคลื่อนเฉลี่ย (MAE)</div>
      </div>
      <div class="stat-card">
        <div class="stat-num">{prev_exact} / 34</div>
        <div class="stat-label">เกณฑ์เดิมก่อนแก้ (Exact: {prev_exact/34*100:.1f}%)</div>
      </div>
    </div>
    <div class="filter-bar">
      <span>ตัวกรอง:</span>
      <button class="active" onclick="filterCards('all', this)">ทั้งหมด (34)</button>
      <button onclick="filterCards('diff', this)">เฉพาะต่างจากผู้สอน ({34 - exact})</button>
      <button onclick="filterCards('match', this)">เฉพาะตรงผู้สอน ({exact})</button>
    </div>
  </header>
  <main id="cards-container">
    {''.join(cards)}
  </main>
  <script>
    function filterCards(type, btn) {{
      document.querySelectorAll('.filter-bar button').forEach(b => b.classList.remove('active'));
      btn.classList.add('active');
      document.querySelectorAll('.card').forEach(card => {{
        const isMatch = card.dataset.match === '1';
        if (type === 'all') card.style.display = 'block';
        else if (type === 'diff') card.style.display = !isMatch ? 'block' : 'none';
        else if (type === 'match') card.style.display = isMatch ? 'block' : 'none';
      }});
    }}
  </script>
</body>
</html>'''

    OUT_HTML.write_text(html_content, encoding="utf-8")
    print(f"Saved HTML Review: {OUT_HTML}")

if __name__ == "__main__":
    asyncio.run(main())
