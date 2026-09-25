import json
import sys
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[1]

DATASET_PATH = ROOT / 'ชุดข้อสอบใหม่' / 'ชุดข้อสอบ_dataset.xlsx'
AI_RUBRIC_PATH = ROOT / 'ชุดข้อสอบใหม่' / 'เกณฑ์ตรวจสำหรับAI.xlsx'
RESULTS_PATH = ROOT / 'artifacts' / 'q2_short_2part_results.json'

data = json.loads(RESULTS_PATH.read_text(encoding='utf-8'))
results_map = {r['sample_id']: r for r in data['results']}

print(f"Loaded {len(results_map)} results from {RESULTS_PATH.name}")

# Concise 2-Part Rubric Text
CONCISE_RUBRIC_TEXT = (
    "ประเมินแยก 2 ส่วนแล้วรวมคะแนน (2.0, 1.5, 1.0, 0.5 หรือ 0.0 คะแนน):\n\n"
    "ส่วนที่ 1: การอธิบายเหตุผล (คะแนนเต็ม 1.5)\n"
    "• 1.5 คะแนน: อธิบายแก่นเหตุผลได้ (เช่น เวลา/รอบโตช้ากว่า, n² เป็นลูปซ้อนทำให้ช้า, n log n แบ่งข้อมูล/recursive) ให้อนุโลมภาษาพูด\n"
    "• 1.0 คะแนน: ตอบถูกเบื้องต้น เช่น ระบุว่าเร็วกว่า/เสถียรกว่า หรือมีตัวอย่างตัวเลขคำนวณเปรียบเทียบ\n"
    "• 0.0 คะแนน: ตอบผิดหลักการ หรือไม่ตอบ\n\n"
    "ส่วนที่ 2: ตัวอย่าง Algorithm (คะแนนเต็ม 0.5)\n"
    "• 0.5 คะแนน: ยกตัวอย่างอัลกอริทึมที่เกี่ยวข้องได้ถูกต้อง (เช่น Merge Sort, Quick Sort, Bubble Sort ฯลฯ)\n"
    "• 0.0 คะแนน: ไม่ได้ยกตัวอย่าง หรือยกตัวอย่างผิด"
)

# 1. Update ชุดข้อสอบ_dataset.xlsx
wb = openpyxl.load_workbook(DATASET_PATH)
ws = wb['ชุดข้อสอบ_dataset']

updated_scores = 0
for row_idx in range(6, ws.max_row + 1):
    q_no = ws.cell(row=row_idx, column=2).value
    if q_no != 2:
        continue
    sid = str(ws.cell(row=row_idx, column=1).value or '').strip()
    if sid in results_map:
        res = results_map[sid]
        score = float(res['ai'])
        conf = str(res.get('confidence') or 'medium')
        teacher_fb = str(res.get('teacher_feedback') or '').strip()
        student_fb = str(res.get('student_feedback') or '').strip()
        
        full_fb = f"[สำหรับผู้สอน]\n{teacher_fb}\n\n[สำหรับนักเรียน]\n{student_fb}"
        
        ws.cell(row=row_idx, column=8, value=score)
        ws.cell(row=row_idx, column=9, value=conf)
        ws.cell(row=row_idx, column=10, value=full_fb)
        updated_scores += 1

print(f"1. Updated {updated_scores} student scores & dual feedback in {DATASET_PATH.name}")

# Update Exam_Rubrics
if 'Exam_Rubrics' in wb.sheetnames:
    ws_rubrics = wb['Exam_Rubrics']
    ws_rubrics.cell(row=7, column=4, value="การอธิบายเหตุผลและตัวอย่างอัลกอริทึม (แยก 2 ส่วน)")
    ws_rubrics.cell(row=7, column=6, value=CONCISE_RUBRIC_TEXT)
    print("   Updated Exam_Rubrics Row 7.")

wb.save(DATASET_PATH)

# 2. Update เกณฑ์ตรวจสำหรับAI.xlsx
wb_ai = openpyxl.load_workbook(AI_RUBRIC_PATH)
ws_ai = wb_ai['ข้อมูลสำหรับ API']
updated_ai = 0
for r in range(2, ws_ai.max_row + 1):
    q_val = str(ws_ai.cell(r, 1).value or '')
    if 'O(n log n)' in q_val:
        ws_ai.cell(r, 2, value=CONCISE_RUBRIC_TEXT)
        updated_ai += 1
wb_ai.save(AI_RUBRIC_PATH)
print(f"2. Updated {updated_ai} rows in {AI_RUBRIC_PATH.name}")

# 3. Update write_rubrics_to_excel.py
SCRIPT_RUBRIC = ROOT / 'scripts' / 'write_rubrics_to_excel.py'
if SCRIPT_RUBRIC.exists():
    text = SCRIPT_RUBRIC.read_text(encoding='utf-8')
    text = text.replace(
        '(ระดับคะแนนที่ให้ได้คือ 2.0, 1.5, 1.0 หรือ 0.0 คะแนนเท่านั้น)',
        CONCISE_RUBRIC_TEXT
    )
    SCRIPT_RUBRIC.write_text(text, encoding='utf-8')
    print("3. Synchronized write_rubrics_to_excel.py")

print("\nAll files successfully updated with the concise 2-part rubric & scores!")
