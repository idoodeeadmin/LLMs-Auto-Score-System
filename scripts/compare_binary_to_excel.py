import sys
import json
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

with open('artifacts/q1_binary_test_results.json', encoding='utf-8') as f:
    d_bin = json.load(f)

wb = openpyxl.load_workbook('ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx', data_only=True)
ws = wb['ชุดข้อสอบ_dataset']

baseline = {ws.cell(r, 1).value: (float(ws.cell(r, 7).value), float(ws.cell(r, 8).value)) for r in range(6, 40)}

print('========================================================================')
print('COMPARISON: OLD EXCEL BASELINE vs NEW BINARY (0, 1, 2) REGRADE')
print('========================================================================')

diffs = []
same = []

for r in d_bin['results']:
    sid = r['sample_id']
    h, old_ai = baseline[sid]
    new_ai = float(r['ai'])
    if old_ai != new_ai:
        diffs.append({
            'sid': sid,
            'human': h,
            'old': old_ai,
            'new': new_ai,
            'ans': r['answer']
        })
    else:
        same.append(sid)

print(f"Total students: 34")
print(f"Identical scores: {len(same)}/34 ({len(same)/34*100:.1f}%)")
print(f"Different scores: {len(diffs)}/34 ({len(diffs)/34*100:.1f}%)")
print('------------------------------------------------------------------------')
for d in diffs:
    snippet = d['ans'].replace('\n', ' ')[:40]
    print(f"  {d['sid']}: Human={d['human']} | Old AI={d['old']} -> New AI={d['new']} | Ans: {snippet}")
