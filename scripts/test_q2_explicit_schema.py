import asyncio
import os
import sys
import json
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env")
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
import httpx
from server.services.openai_grading import (
    _get_openai_api_key,
    OPENAI_MODEL,
    OPENAI_RESPONSES_URL,
    _build_ssl_context,
    _extract_output_text,
)

excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
wb = openpyxl.load_workbook(excel_path, data_only=True)
ws = wb["ชุดข้อสอบ_dataset"]
rows = [r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value == 2]

prompt_template = """คุณคืออาจารย์ผู้เชี่ยวชาญในการตรวจข้อสอบวิชา Data Structures
โจทย์ข้อที่ 2: 'อธิบายว่าทำไม O(n log n) เหมาะกับข้อมูลใหญ่กว่า O(n^2) และยกตัวอย่าง Algorithm' (คะแนนเต็ม 2.00 คะแนน)

กรุณาประเมินคำตอบของนักเรียนอย่างเป็นกลางและให้คะแนนแยก 2 เกณฑ์อย่างเป็นอิสระต่อกัน:

เกณฑ์ที่ 1: เหตุผลเชิงเปรียบเทียบความซับซ้อน (คะแนนเต็ม 1.00 คะแนน)
- ได้ 1.00: อธิบายเหตุผลในเชิงกลไกการทำงานได้ถูกต้อง (เช่น O(n log n) เติบโตช้ากว่ามากเมื่อข้อมูลใหญ่, มีการตัดแบ่งครึ่งข้อมูล/Divide & Conquer, หรือ O(n^2) วนลูปซ้อนกันทำให้รอบการทำงานทวีคูณอย่างรวดเร็ว) ให้อนุโลมภาษาพูดและคำอธิบายเสริมเรื่อง memory/recursive/CPU
- ได้ 0.50: ตอบสั้นๆ แบบกว้างๆ หรือเป็นเหตุผลเชิงวงกลม เช่น ตอบแค่ว่า 'มีประสิทธิภาพดีกว่า' หรือ 'เร็วกว่า' โดยไม่อธิบายกลไกการวนลูปหรือการเติบโต
- ได้ 0.00: ไม่ตอบ หรือตอบผิดหลักการทั้งหมดอย่างชัดเจน (เช่น ตอบเรื่องลูกเต๋าโยนเหรียญ หรือเขียนเฉพาะตัวโจทย์ซ้ำ)

เกณฑ์ที่ 2: การยกตัวอย่าง Algorithm (คะแนนเต็ม 1.00 คะแนน)
- ได้ 1.00: มีการระบุชื่อ Algorithm ที่เกี่ยวข้องอย่างน้อย 1 ชื่อ (เช่น Merge Sort, Quick Sort, Bubble Sort, Insertion Sort, Heap Sort, Selection Sort หรือยกตัวอย่างฟังก์ชันโค้ดการทำงาน)
- ได้ 0.00: ไม่มีการระบุชื่อ Algorithm หรือตัวอย่างใดๆ เลย

คะแนนรวม (score) ต้องเท่ากับ reasoning_score + algorithm_score อย่างเคร่งครัด

คำตอบของนักเรียน:
{answer}
"""

schema = {
    "type": "object",
    "properties": {
        "reasoning_score": {"type": "number", "enum": [0.0, 0.5, 1.0], "description": "คะแนนส่วนเหตุผล"},
        "algorithm_score": {"type": "number", "enum": [0.0, 1.0], "description": "คะแนนส่วนชื่ออัลกอริทึม"},
        "score": {"type": "number", "enum": [0.0, 0.5, 1.0, 1.5, 2.0], "description": "คะแนนรวม = reasoning_score + algorithm_score"},
        "teacher_feedback": {"type": "string"},
        "student_feedback": {"type": "string"}
    },
    "required": ["reasoning_score", "algorithm_score", "score", "teacher_feedback", "student_feedback"],
    "additionalProperties": False
}

async def grade_one(client, ans):
    payload = {
        "model": OPENAI_MODEL,
        "instructions": "Evaluate student answer strictly by assigning reasoning_score (0.0, 0.5, 1.0) and algorithm_score (0.0, 1.0) and setting score = reasoning_score + algorithm_score.",
        "input": [{"role": "user", "content": [{"type": "input_text", "text": prompt_template.format(answer=ans)}]}],
        "reasoning": {"effort": "low"},
        "text": {"format": {"type": "json_schema", "name": "q2_grading", "strict": True, "schema": schema}},
        "max_output_tokens": 1000,
        "store": False
    }
    resp = await client.post(OPENAI_RESPONSES_URL, headers={"Authorization": f"Bearer {_get_openai_api_key()}", "Content-Type": "application/json"}, json=payload)
    resp.raise_for_status()
    body = resp.json()
    return json.loads(_extract_output_text(body))

async def main():
    async with httpx.AsyncClient(timeout=60.0, verify=_build_ssl_context()) as client:
        sem = asyncio.Semaphore(5)
        exact = 0
        within_05 = 0
        diffs = []

        async def worker(r):
            nonlocal exact, within_05
            sid = ws.cell(r, 1).value
            ans = str(ws.cell(r, 6).value or "")
            human = float(ws.cell(r, 7).value or 0.0)
            async with sem:
                res = await grade_one(client, ans)
                ai = float(res["score"])
                diff = round(ai - human, 2)
                is_match = (diff == 0.0)
                if is_match: exact += 1
                if abs(diff) <= 0.5: within_05 += 1
                diffs.append((sid, human, ai, diff, res["reasoning_score"], res["algorithm_score"]))
                status = "MATCH" if is_match else f"DIFF={diff:+.1f}"
                print(f"[{sid}] H={human:.1f} | AI={ai:.1f} (R={res['reasoning_score']}, A={res['algorithm_score']}) | {status}", flush=True)

        await asyncio.gather(*[worker(r) for r in rows])

        print("\n" + "="*50, flush=True)
        print(f"EXACT MATCHES: {exact}/34 ({exact/34*100:.2f}%)", flush=True)
        print(f"WITHIN 0.5: {within_05}/34 ({within_05/34*100:.2f}%)", flush=True)
        print(f"MAE: {sum(abs(d[3]) for d in diffs)/34:.4f}", flush=True)

if __name__ == "__main__":
    asyncio.run(main())
