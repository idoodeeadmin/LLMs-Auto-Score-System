import json
import sys
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')
ROOT = Path(__file__).resolve().parents[1]

EXCEL_PATH = ROOT / 'ชุดข้อสอบใหม่' / 'ชุดข้อสอบ_dataset.xlsx'
RESULTS_PATH = ROOT / 'artifacts' / 'q2_calibrated_5level_results.json'

if not RESULTS_PATH.exists():
    raise FileNotFoundError(f"Results file not found: {RESULTS_PATH}")

data = json.loads(RESULTS_PATH.read_text(encoding='utf-8'))
results_map = {r['sample_id']: r for r in data['results']}

print(f"Loaded {len(results_map)} results from {RESULTS_PATH.name}")

wb = openpyxl.load_workbook(EXCEL_PATH)
ws = wb['ชุดข้อสอบ_dataset']

updated_count = 0
for row_idx in range(6, ws.max_row + 1):
    q_no = ws.cell(row=row_idx, column=2).value
    if q_no != 2:
        continue
    sid = str(ws.cell(row=row_idx, column=1).value or '').strip()
    if sid in results_map:
        res = results_map[sid]
        score = float(res['ai'])
        conf = str(res.get('confidence') or 'medium')
        teacher_fb = str(res.get('teacher_feedback') or '').strip()
        student_fb = str(res.get('student_feedback') or '').strip()
        
        full_fb = f"[สำหรับผู้สอน]\n{teacher_fb}\n\n[สำหรับนักเรียน]\n{student_fb}"
        
        ws.cell(row=row_idx, column=8, value=score)
        ws.cell(row=row_idx, column=9, value=conf)
        ws.cell(row=row_idx, column=10, value=full_fb)
        updated_count += 1
        print(f"Row {row_idx}: {sid} -> Score={score:.1f}, Conf={conf}")

# Also ensure Exam_Rubrics Row 7 is synchronized
if 'Exam_Rubrics' in wb.sheetnames:
    ws_rubrics = wb['Exam_Rubrics']
    q2_rubric_5level = (
        "ประเมินตามระดับคะแนน 5 ระดับ (2.0, 1.5, 1.0, 0.5, 0.0) อย่างเคร่งครัด:\n"
        "• 2.0 คะแนน: อธิบายได้ว่าทำไม O(n log n) ดีกว่าเมื่อข้อมูลใหญ่ เช่น เร็วกว่า, จำนวนรอบ/การทำงานโตช้ากว่า, n² โตเร็วมาก และ มีตัวอย่าง Algorithm หรืออธิบายกลไกชัดเจน\n"
        "• 1.5 คะแนน: เข้าใจแก่นว่า O(n log n) มีประสิทธิภาพกว่า แต่คำอธิบายยังไม่ครบ/ไม่ยกตัวอย่าง หรือมีตัวอย่างแต่เหตุผลยังไม่ลงลึก\n"
        "• 1.0 คะแนน: รู้เพียงว่า O(n log n) “เร็วกว่า/ดีกว่า/ซ้ำซ้อนน้อยกว่า” แต่ไม่ได้อธิบายกลไกอย่างชัดเจน หรือคำอธิบายคลุมเครือ และไม่ได้ยกตัวอย่าง (หรือมีเพียงตัวอย่างคำนวณ)\n"
        "• 0.5 คะแนน: ระบุได้เพียงตัวอย่าง Algorithm ที่เกี่ยวข้องอย่างถูกต้อง หรือกล่าวถึงความซับซ้อน/ประสิทธิภาพเพียงฝั่งเดียวอย่างสั้นๆ โดยยังไม่สื่อชัดว่า O(n log n) เหมาะกับข้อมูลใหญ่กว่า O(n²) อย่างไร\n"
        "• 0.0 คะแนน: ไม่สามารถอธิบายความสัมพันธ์ของ O(n log n) กับ O(n²) ได้อย่างมีสาระ ตอบผิดหลักการสำคัญ หรือไม่ตอบ\n"
        "(ระดับคะแนนที่ให้ได้คือ 2.0, 1.5, 1.0, 0.5 หรือ 0.0 คะแนนเท่านั้น ห้ามให้เศษทศนิยมอื่น)"
    )
    ws_rubrics.cell(row=7, column=6, value=q2_rubric_5level)
    print("Exam_Rubrics Row 7 confirmed updated.")

wb.save(EXCEL_PATH)
print(f"\nSuccessfully updated {updated_count} student rows in {EXCEL_PATH.name}!")
