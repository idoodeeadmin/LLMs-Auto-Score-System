"""Read the four API input fields without exposing human scores or sample IDs."""
from pathlib import Path
import openpyxl

RUBRIC_FILE = Path(__file__).resolve().parents[1] / 'ชุดข้อสอบใหม่' / 'เกณฑ์ตรวจสำหรับAI.xlsx'

def load_proposed_rubric(question_no, max_score, question_text):
    wb = openpyxl.load_workbook(RUBRIC_FILE, read_only=True, data_only=True)
    try:
        sheet = wb['ข้อมูลสำหรับ API']
        if tuple(sheet.cell(1, col).value for col in range(1, 5)) != ('โจทย์', 'เกณฑ์', 'คำตอบ', 'แนวคำตอบ'):
            raise ValueError('Unexpected API input columns')
        for row in sheet.iter_rows(min_row=2, max_col=4, values_only=True):
            stored_question, criteria, _, _ = row
            if stored_question == question_text or str(stored_question).startswith(question_text + '\nภาพโจทย์:'):
                step = float(max_score) / 4  # score precision remains an API setting, not a worksheet column
                return [{'name': f'เกณฑ์ข้อ {question_no}', 'score': float(max_score),
                         'description': str(criteria)}], step
    finally:
        wb.close()
    raise ValueError(f'No proposed rubric for question {question_no}: {question_text}')
