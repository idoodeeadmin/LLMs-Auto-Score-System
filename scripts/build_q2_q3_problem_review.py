"""Show source images and grading disagreements for corrected answers."""

import html
import json
from pathlib import Path

import openpyxl

root = Path(__file__).resolve().parents[1]
out = root / "docs_and_tests/q2_q3_review/problem_cases.html"
report = json.loads((out.parent / "regrade_corrected_results.json").read_text(encoding="utf-8"))
book = openpyxl.load_workbook(root / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx", read_only=True, data_only=True)
sheet = book["ชุดข้อสอบ_dataset"]
images = {51: "IMG_2859.jpg", 55: "IMG_2863.jpg", 89: "IMG_2914.jpg", 101: "IMG_2926.jpg"}
parts = ['''<!doctype html><html lang="th"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>กรณีคะแนนไม่ตรงกับผู้สอน</title>
<style>body{font:16px/1.6 Tahoma,sans-serif;background:#f4f7f6;color:#152924;max-width:1380px;margin:0 auto;padding:24px}h1{margin:0 0 8px}p{margin:0 0 20px}.case{background:white;border:1px solid #d9e4df;border-radius:14px;padding:24px;margin:24px 0}.top{display:flex;justify-content:space-between;gap:16px;align-items:center}.badges{display:flex;gap:8px;flex-wrap:wrap}.badge{padding:5px 10px;border-radius:20px;background:#e9f1ee;font-weight:bold}.grid{display:grid;grid-template-columns:minmax(300px,1fr) minmax(400px,1fr);gap:24px;margin-top:18px}.photo{width:100%;max-height:800px;object-fit:contain;background:#ebeeed;border-radius:8px}pre{font:15px/1.6 Tahoma,sans-serif;white-space:pre-wrap;overflow-wrap:anywhere;background:#f6f8f7;padding:16px;border-radius:8px}h3{margin:14px 0 6px}@media(max-width:850px){.grid{grid-template-columns:1fr}}</style>
<h1>กรณีที่คะแนน AI หลังตรวจใหม่ไม่ตรงกับผู้สอน</h1><p>แสดงภาพคำตอบที่ลบคะแนนแล้ว ข้อความใน Excel และเหตุผลที่ AI รายงาน คะแนนผู้สอนไม่ได้ส่งให้ AI</p>''']
for item in report["items"]:
    n = int(item["sample_id"].split("-")[1])
    if n not in images:
        continue
    q = 2 if n < 69 else 3
    row = n + 5
    human = sheet.cell(row, 7).value
    ai = sheet.cell(row, 8).value
    if human == ai:
        continue
    old = item["prior_ai_score"]
    image = f"../../ชุดข้อสอบใหม่/photo_clean_text{q}/{images[n]}"
    answer = html.escape(str(sheet.cell(row, 6).value))
    feedback = html.escape(item["result"]["teacher_feedback"])
    parts.append(f'''<article class="case"><div class="top"><h2>{item["sample_id"]} · ข้อ {q}</h2><div class="badges"><span class="badge">ผู้สอน {human}</span><span class="badge">AI เดิม {old}</span><span class="badge">AI ใหม่ {ai}</span></div></div><div class="grid"><div><h3>ภาพคำตอบ</h3><img class="photo" src="{image}"></div><div><h3>ข้อความที่ถอด</h3><pre>{answer}</pre><h3>เหตุผลจาก AI</h3><pre>{feedback}</pre></div></div></article>''')
parts.append("</html>")
out.write_text("".join(parts), encoding="utf-8")
book.close()
