import json
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

with open('artifacts/q1_regrade_results.json', encoding='utf-8') as f:
    new_res = json.load(f)

wb = openpyxl.load_workbook('ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx', data_only=True)
ws = wb['ชุดข้อสอบ_dataset']

baseline = {}
for r in range(6, 40):
    sid = ws.cell(r, 1).value
    human = float(ws.cell(r, 7).value)
    ai_old = float(ws.cell(r, 8).value)
    baseline[sid] = {'human': human, 'old_ai': ai_old}

changes = []
for item in new_res['results']:
    sid = item['sample_id']
    h = baseline[sid]['human']
    old_ai = baseline[sid]['old_ai']
    new_ai = float(item['ai'])
    old_diff = round(old_ai - h, 2)
    new_diff = round(new_ai - h, 2)
    
    if old_ai == new_ai:
        status = 'SAME'
    elif abs(new_diff) < abs(old_diff):
        status = 'IMPROVED (+)'
    elif abs(new_diff) > abs(old_diff):
        status = 'REGRESSED (-)'
    else:
        status = 'CHANGED (=)'

    changes.append({
        'sid': sid,
        'human': h,
        'old_ai': old_ai,
        'new_ai': new_ai,
        'old_diff': old_diff,
        'new_diff': new_diff,
        'status': status,
        'answer': item['answer'],
        'teacher_feedback': item['teacher_feedback'],
        'student_feedback': item['student_feedback']
    })

old_exact = sum(1 for c in changes if c['old_diff'] == 0.0)
old_diff05 = sum(1 for c in changes if abs(c['old_diff']) <= 0.5)
old_mae = sum(abs(c['old_diff']) for c in changes) / len(changes)

new_exact = sum(1 for c in changes if c['new_diff'] == 0.0)
new_diff05 = sum(1 for c in changes if abs(c['new_diff']) <= 0.5)
new_mae = sum(abs(c['new_diff']) for c in changes) / len(changes)

print('========================================================================')
print('QUESTION 1 RE-GRADE METRICS COMPARISON (34 Students)')
print('========================================================================')
print(f'Old AI Baseline: Exact Match = {old_exact}/34 ({old_exact/34*100:.1f}%), Diff <= 0.5 = {old_diff05}/34 ({old_diff05/34*100:.1f}%), MAE = {old_mae:.4f}')
print(f'New AI Regrade : Exact Match = {new_exact}/34 ({new_exact/34*100:.1f}%), Diff <= 0.5 = {new_diff05}/34 ({new_diff05/34*100:.1f}%), MAE = {new_mae:.4f}')
print('------------------------------------------------------------------------')
print(f'Improved count: {sum(1 for c in changes if "IMPROVED" in c["status"])}')
print(f'Regressed count: {sum(1 for c in changes if "REGRESSED" in c["status"])}')
print(f'Same count: {sum(1 for c in changes if c["status"] == "SAME")}')
print('========================================================================\n')

print('ALL CHANGED SCORES:')
print(f'{"ID":<7} | {"Human":<5} | {"Old AI":<6} -> {"New AI":<6} | {"Status":<14} | Answer Snippet')
print('-'*80)
for c in changes:
    if c['old_ai'] != c['new_ai']:
        snippet = c['answer'].replace('\n', ' ')[:45]
        print(f"{c['sid']:<7} | {c['human']:<5.1f} | {c['old_ai']:<6.1f} -> {c['new_ai']:<6.1f} | {c['status']:<14} | {snippet}")

print('\n========================================================================')
print('DETAILED ANALYSIS OF IMPROVEMENTS AND REGRESSIONS:')
print('========================================================================')
for c in changes:
    if c['old_ai'] != c['new_ai']:
        print(f"\n[{c['sid']}] Human={c['human']} | Old AI={c['old_ai']} -> New AI={c['new_ai']} ({c['status']})")
        print(f"Answer: {c['answer']}")
        print(f"Teacher FB: {c['teacher_feedback']}")
