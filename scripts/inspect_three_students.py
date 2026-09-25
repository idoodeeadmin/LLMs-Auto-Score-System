import asyncio
import openpyxl
import sys
from pathlib import Path
ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")
from dotenv import load_dotenv
load_dotenv(ROOT / ".env")
from scripts.test_approach_a_strict import Q3_QUESTION_TEXT, Q3_ANSWER_KEY, RUBRICS
from server.services.openai_grading import score_with_openai

async def inspect():
    wb = openpyxl.load_workbook("ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx", data_only=True)
    ws = wb["ชุดข้อสอบ_dataset"]
    for sid_target in ["DS-070", "DS-072", "DS-073"]:
        for r in range(74, 84):
            if ws.cell(r, 1).value == sid_target:
                ans = str(ws.cell(r, 6).value or "")
                h = float(ws.cell(r, 7).value or 0.0)
                res = await score_with_openai(
                    question_text=Q3_QUESTION_TEXT,
                    answer_text=ans,
                    max_score=1.0,
                    answer_key=Q3_ANSWER_KEY,
                    rubrics=RUBRICS,
                    strict_rubric_enforcement=True
                )
                print(f"=== {sid_target} (Human: {h:.2f}, AI: {res.get('score')}) ===")
                print("Ans:", ans)
                print("TF:", res.get("teacher_feedback"))
                print()

if __name__ == "__main__":
    asyncio.run(inspect())
