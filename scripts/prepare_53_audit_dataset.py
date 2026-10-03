import os
import sys
import json
import openpyxl
from PIL import Image
import pillow_heif

sys.stdout.reconfigure(encoding='utf-8')
pillow_heif.register_heif_opener()

wb = openpyxl.load_workbook('ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx', data_only=True)
ws = wb['ชุดข้อสอบ_dataset']

out_img_dir = 'public/screenshots/audit_gallery'
os.makedirs(out_img_dir, exist_ok=True)

# Question topic titles
q_topics = {
    1: 'Row-major vs Column-major Order in 2D Array',
    2: 'Time Complexity: O(n log n) vs O(n²)',
    3: 'Linked List vs Array for Stack & Queue',
    4: 'Binary Search Tree (BST) Construction (12 Nodes)',
    5: 'Infix to Prefix and Postfix Expression Conversion',
    6: 'General Tree to Binary Tree (LCRS)'
}

# Image source files per question
q1_files = sorted(os.listdir('ชุดข้อสอบใหม่/photo_clean_text1')) # 34 files IMG_2791.jpg ..
q2_files = sorted(os.listdir('ชุดข้อสอบใหม่/photo_clean_text2'))
q3_files = sorted(os.listdir('ชุดข้อสอบใหม่/photo_clean_text3'))

mismatch_records = []

# Scan rows 6 to 209
for r in range(6, ws.max_row + 1):
    sid = ws.cell(row=r, column=1).value
    qno = ws.cell(row=r, column=2).value
    if not sid or not str(sid).startswith('DS-') or not qno:
        continue
    
    qno = int(qno)
    h_score = ws.cell(row=r, column=7).value
    ai_score = ws.cell(row=r, column=8).value
    if h_score is None or ai_score is None:
        continue
    
    h_val = float(h_score)
    ai_val = float(ai_score)
    
    # Check if mismatch
    if h_val != ai_val:
        std_idx = ((r - 6) % 34) # 0 to 33
        
        ROT_90_CCW_FILES_Q2 = {"IMG_2792.HEIC", "IMG_2868.HEIC", "IMG_2869.HEIC", "IMG_2870.HEIC"}
        
        # Raw file path
        if qno == 1:
            base_name = q1_files[std_idx].replace('.jpg', '.HEIC')
            raw_path = os.path.join('ชุดข้อสอบใหม่/โจทย์textชุดที่1', base_name)
        elif qno == 2:
            base_name = q2_files[std_idx].replace('.jpg', '.HEIC')
            raw_path = os.path.join('ชุดข้อสอบใหม่/โจทย์textชุดที่2', base_name)
        elif qno == 3:
            base_name = q3_files[std_idx].replace('.jpg', '.HEIC')
            raw_path = os.path.join('ชุดข้อสอบใหม่/โจทย์textชุดที่3', base_name)
        elif qno == 4:
            file_num = std_idx + 1
            raw_path = os.path.join('ชุดข้อสอบใหม่/photoชุดที่1', f'LINE_ALBUM_Photo1_260917_{file_num}.jpg')
        elif qno == 5:
            file_num = std_idx + 1
            raw_path = os.path.join('ชุดข้อสอบใหม่/photoชุดที่2', f'LINE_ALBUM_Photo2.1_260918_{file_num}.jpg')
        else:
            raw_path = None
            
        ans_text = str(ws.cell(row=r, column=6).value or '')
        confidence = str(ws.cell(row=r, column=9).value or 'high')
        feedback = str(ws.cell(row=r, column=10).value or '')
        
        # Process and save optimized image
        full_out_name = f'full_{sid}.jpg'
        full_out_path = os.path.join(out_img_dir, full_out_name)
        
        crop_out_name = f'crop_{sid}.jpg'
        crop_out_path = os.path.join(out_img_dir, crop_out_name)
        
        if os.path.exists(raw_path):
            try:
                im = Image.open(raw_path)
                
                # Precise orientation correction
                if qno == 1:
                    # Landscape orientation: only rotate if image is currently portrait
                    if im.size[0] < im.size[1]:
                        im = im.transpose(Image.Transpose.ROTATE_90)
                elif qno == 2:
                    if base_name in ROT_90_CCW_FILES_Q2:
                        im = im.transpose(Image.Transpose.ROTATE_90)
                elif qno == 4:
                    import cv2
                    import numpy as np
                    with open(raw_path, 'rb') as f:
                        cv_img = cv2.imdecode(np.frombuffer(f.read(), np.uint8), cv2.IMREAD_COLOR)
                    gray = cv2.cvtColor(cv_img, cv2.COLOR_BGR2GRAY)
                    left_dark = np.mean(gray[:, :150] < 120)
                    right_dark = np.mean(gray[:, -150:] < 120)
                    if left_dark > right_dark:
                        im = im.transpose(Image.Transpose.ROTATE_90)
                    else:
                        im = im.transpose(Image.Transpose.ROTATE_270)
                
                w, h = im.size
                max_dim = 1200
                if max(w, h) > max_dim:
                    scale = max_dim / float(max(w, h))
                    im_resized = im.resize((int(w * scale), int(h * scale)), Image.Resampling.LANCZOS)
                else:
                    im_resized = im.copy()
                
                im_resized.save(full_out_path, quality=90, optimize=True)
                
                # Create focused crop around likely score area
                rw, rh = im_resized.size
                if qno == 1:
                    # Score is in bottom right
                    crop_im = im_resized.crop((int(rw * 0.55), int(rh * 0.55), rw, rh))
                elif qno == 2:
                    # Score is often top right or bottom right
                    crop_im = im_resized.crop((int(rw * 0.45), 0, rw, int(rh * 0.45)))
                elif qno == 3:
                    # Score is bottom right
                    crop_im = im_resized.crop((int(rw * 0.5), int(rh * 0.6), rw, rh))
                elif qno == 4:
                    # Score is bottom left or bottom right
                    crop_im = im_resized.crop((0, int(rh * 0.5), rw, rh))
                elif qno == 5:
                    # Score is top right
                    crop_im = im_resized.crop((int(rw * 0.5), 0, rw, int(rh * 0.45)))
                else:
                    crop_im = im_resized
                    
                crop_im.save(crop_out_path, quality=90, optimize=True)
            except Exception as e:
                print(f"Error processing image for {sid}: {e}")
        else:
            print(f"Warning: {raw_path} does not exist!")

        # Categorize known special cases
        anomaly_tag = "ทั่วไป"
        if sid == "DS-031":
            anomaly_tag = "ตรวจพบ Excel บันทึกผิด (บนกระดาษแก้เป็น 2)"
        elif sid == "DS-121":
            anomaly_tag = "อาจารย์หักกิ่ง BST ผิด (กระดาษให้ 0.25, AI ให้ 1.00)"
        elif sid == "DS-122":
            anomaly_tag = "โครงสร้าง BST ซับซ้อน (กระดาษให้ 1.00, AI ให้ 0.00)"
        elif sid == "DS-007" or sid == "DS-015":
            anomaly_tag = "อาจารย์ให้คะแนนความพยายาม (AI ตรวจ 0)"
        elif sid == "DS-054":
            anomaly_tag = "สำนวนกวน/ติดตลก (AI ให้เต็ม 2)"
        elif sid == "DS-009":
            anomaly_tag = "อาจารย์มีเกณฑ์ในใจ (AI ให้เต็ม 2)"
        elif h_val < ai_val:
            anomaly_tag = "AI ให้คะแนนสูงกว่า Excel"
        else:
            anomaly_tag = "AI ให้คะแนนต่ำกว่า Excel"

        mismatch_records.append({
            'sid': sid,
            'qno': qno,
            'topic': q_topics.get(qno, ''),
            'excel_human_score': h_val,
            'ai_score': ai_val,
            'score_diff': round(ai_val - h_val, 2),
            'confidence': confidence,
            'ans_text': ans_text,
            'feedback': feedback,
            'anomaly_tag': anomaly_tag,
            'full_img_url': f'screenshots/audit_gallery/{full_out_name}',
            'crop_img_url': f'screenshots/audit_gallery/{crop_out_name}',
            'raw_file_source': os.path.basename(raw_path)
        })

print(f"Processed {len(mismatch_records)} mismatch records!")

# Save to JSON for UI
json_out_path = 'public/audit_53_mismatches.json'
with open(json_out_path, 'w', encoding='utf-8') as f:
    json.dump(mismatch_records, f, ensure_ascii=False, indent=2)

print(f"Saved JSON database to {json_out_path}")
