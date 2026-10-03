# -*- coding: utf-8 -*-
"""
Update Question 1 rubric in:
1. ชุดข้อสอบ_dataset.xlsx -> Sheet 'Exam_Rubrics' (Rows 5 & 6)
2. เกณฑ์ตรวจสำหรับAI.xlsx -> Sheet 'ข้อมูลสำหรับ API' (Rows where Q1 is tested)
"""
import sys
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[1]
EXCEL_DATASET = ROOT / 'ชุดข้อสอบใหม่' / 'ชุดข้อสอบ_dataset.xlsx'
EXCEL_AI_RUBRICS = ROOT / 'ชุดข้อสอบใหม่' / 'เกณฑ์ตรวจสำหรับAI.xlsx'

Q1_CONCISE_TEXT_API = (
    "ข้อ 1 (2 คะแนน):\n"
    "ประเมินความเข้าใจความต่างของ Row-major และ Column-major (คะแนนเต็ม 2.00 คะแนน):\n"
    "• 2.00 คะแนน: อธิบายความต่างได้ถูกต้องครบทั้ง 2 ฝั่ง (Row-major อิงตามแถว/แนวนอน และ Column-major อิงตามคอลัมน์/แนวตั้ง)\n"
    "• 1.00 คะแนน: อธิบายถูกต้องเพียงฝั่งเดียว หรือตอบสั้นเฉพาะความเข้าใจเบื้องต้น\n"
    "• 0.00 คะแนน: ตอบผิดทั้งหมด หรือไม่ตอบ\n"
    "(ระดับคะแนนที่ให้ได้คือ 2.0, 1.0 หรือ 0.0 คะแนนเท่านั้น)"
)

Q1_ROW5_DESC = (
    "อธิบายการจัดเก็บข้อมูลหรือเข้าถึงข้อมูลตามแถว (Row / แนวนอน) ได้ถูกต้อง ได้ 1.00 คะแนน (ตอบผิดหรือไม่ตอบได้ 0.00 คะแนน)"
)

Q1_ROW6_DESC = (
    "อธิบายการจัดเก็บข้อมูลหรือเข้าถึงข้อมูลตามคอลัมน์ (Column / แนวตั้ง) ได้ถูกต้อง ได้ 1.00 คะแนน (ตอบผิดหรือไม่ตอบได้ 0.00 คะแนน)\n\n"
    "(เกณฑ์ภาพรวม: อธิบายถูกทั้ง 2 ฝั่งได้ 2.00 คะแนน, อธิบายถูกฝั่งเดียวหรือระบุความเข้าใจเบื้องต้นได้ 1.00 คะแนน, ตอบผิดทั้งหมดหรือไม่ตอบได้ 0.00 คะแนน)"
)

def update_dataset_exam_rubrics():
    if not EXCEL_DATASET.exists():
        print(f"Error: {EXCEL_DATASET} not found")
        return
    
    wb = openpyxl.load_workbook(EXCEL_DATASET)
    if "Exam_Rubrics" not in wb.sheetnames:
        print("Error: Exam_Rubrics sheet not found")
        return
        
    ws = wb["Exam_Rubrics"]
    
    # Row 5 (Row-major)
    ws.cell(row=5, column=1, value=1)
    ws.cell(row=5, column=2, value="Row-major vs Column-major")
    ws.cell(row=5, column=3, value=2.0)
    ws.cell(row=5, column=4, value="Row-major")
    ws.cell(row=5, column=5, value=1.0)
    ws.cell(row=5, column=6, value=Q1_ROW5_DESC)
    
    # Row 6 (Column-major)
    ws.cell(row=6, column=1, value=1)
    ws.cell(row=6, column=2, value="Row-major vs Column-major")
    ws.cell(row=6, column=3, value=2.0)
    ws.cell(row=6, column=4, value="Column-major")
    ws.cell(row=6, column=5, value=1.0)
    ws.cell(row=6, column=6, value=Q1_ROW6_DESC)
    
    wb.save(EXCEL_DATASET)
    print(f"1. Successfully updated Exam_Rubrics (Rows 5 & 6) in '{EXCEL_DATASET.name}'")

def update_ai_rubrics_file():
    if not EXCEL_AI_RUBRICS.exists():
        print(f"Warning: {EXCEL_AI_RUBRICS} not found")
        return
        
    wb = openpyxl.load_workbook(EXCEL_AI_RUBRICS)
    if "ข้อมูลสำหรับ API" not in wb.sheetnames:
        print("Warning: 'ข้อมูลสำหรับ API' sheet not found")
        return
        
    ws = wb["ข้อมูลสำหรับ API"]
    updated = 0
    for r in range(2, ws.max_row + 1):
        q_val = str(ws.cell(row=r, column=1).value or '')
        if "Row-major" in q_val:
            ws.cell(row=r, column=2, value=Q1_CONCISE_TEXT_API)
            updated += 1
            
    wb.save(EXCEL_AI_RUBRICS)
    print(f"2. Successfully updated {updated} rows in '{EXCEL_AI_RUBRICS.name}' (Sheet: ข้อมูลสำหรับ API)")

if __name__ == '__main__':
    update_dataset_exam_rubrics()
    update_ai_rubrics_file()
