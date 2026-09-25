import asyncio
import json
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q6_RUBRIC_NAME = "ความถูกต้องของการแปลง General Tree เป็น Binary Tree (Left-Child Right-Sibling)"
Q6_RUBRIC_DESC = (
    "ประเมินตามโครงสร้างและความถูกต้องของการแปลง General Tree เป็น Binary Tree ตามหลักการ Left-Child Right-Sibling (LCRS) จากโจทย์ 10 โหนด (1 ถึง 10) ดังนี้:\n\n"
    "- ได้ 1.00 คะแนน: แปลงโครงสร้าง Binary Tree ได้ถูกต้องสมบูรณ์ครบทั้ง 10 โหนด ตามหลัก LCRS:\n"
    "  * Root คือ 1 โดยมีกิ่งซ้าย (Left Child) เป็น 2 และกิ่งขวาว่าง (ไม่มี sibling)\n"
    "  * โหนด 2 มีกิ่งซ้ายเป็น 5 (ลูกคนแรก) และกิ่งขวาเป็น 3 (พี่น้องถัดไป)\n"
    "  * ใต้กิ่ง 5: มีกิ่งขวาเป็น 6 -> โหนด 6 มีกิ่งขวาเป็น 7 (ทั้ง 5, 6, 7 ไม่มีกิ่งซ้าย)\n"
    "  * โหนด 3 ไม่มีกิ่งซ้าย และมีกิ่งขวาเป็น 4 (พี่น้องถัดไป)\n"
    "  * โหนด 4 มีกิ่งซ้ายเป็น 8 (ลูกคนแรก) และไม่มีกิ่งขวา\n"
    "  * ใต้กิ่ง 8: มีกิ่งขวาเป็น 9 -> โหนด 9 มีกิ่งขวาเป็น 10 (ทั้ง 8, 9, 10 ไม่มีกิ่งซ้าย)\n"
    "  (ให้อนุโลมความสวยงาม ความยาว หรือมุมเอียงของกิ่งในกระดาษเขียน หากมีเส้นเชื่อมครบ 9 เส้นตามลำดับดังกล่าว "
    "แม้มุมของ 4 ไปยัง 8 จะเอียงชันตามเนื้อที่กระดาษ ให้ถือว่าสื่อสารกิ่งซ้ายของ 4 ได้ถูกต้อง และได้ 1.00 คะแนนเต็ม)\n\n"
    "- ได้ 0.25 คะแนน: โครงสร้างแกนหลักของ LCRS ถูกต้องเกิน 70% (เช่น แกน 1->2->3->4 ถูกต้อง และกิ่งย่อย 5-6-7 ถูกต้อง) "
    "แต่มีข้อผิดพลาดเฉพาะที่ปลายกิ่งย่อยเพียง 1 จุด (เช่น ลืมใส่โหนด 10 หรือสลับเฉพาะ 9 กับ 10)\n\n"
    "- ได้ 0.00 คะแนน: ผิดพลาดอย่างมีนัยสำคัญ หรือผิดหลักการ LCRS เช่น:\n"
    "  * วางโหนด 3 เป็นกิ่งขวาของโหนด 1 โดยตรง (โหนด 1 ไม่มีพี่น้อง ห้ามมีกิ่งขวา โหนด 3 ต้องเป็นกิ่งขวาของ 2)\n"
    "  * วางโหนด 4 หรือ 3 เป็นกิ่งซ้ายของ 2 หรือนำ 4 ไปต่อกับ 1 โดยตรง\n"
    "  * วาดเป็นเส้นตรงซิกแซกเรียงเดี่ยว หรือเขียนเฉพาะตัวเลขเรียงกันโดยไม่วาดต้นไม้\n"
    "  * ไม่วาดคำตอบ หรือเขียนว่าแปลงไม่ได้\n"
    "  * โครงสร้างแตกกิ่งมั่ว ไม่มีความสัมพันธ์แบบ Left-Child Right-Sibling\n\n"
    "(ต้องระบุ teacher_feedback แจกแจงเหตุผลและจุดถูก/ผิดตามเกณฑ์ LCRS สำหรับผู้สอนอย่างละเอียด และระบุ student_feedback เป็นคำแนะนำพัฒนาการเรียนรู้สำหรับนักเรียน)"
)

Q6_RUBRICS = [
    {
        "name": Q6_RUBRIC_NAME,
        "score": 1.0,
        "description": Q6_RUBRIC_DESC,
    }
]

Q6_QUESTION_TEXT = (
    "จงแปลง tree ต่อไปนี้ให้เป็น Binary Tree (1 คะแนน)\n"
    "ข้อมูล General Tree: โหนด 1 เป็น Root มีลูก 3 ตัวคือ 2, 3, 4 เรียงจากซ้ายไปขวา; "
    "โหนด 2 มีลูก 3 ตัวคือ 5, 6, 7; โหนด 3 ไม่มีลูก (Leaf); โหนด 4 มีลูก 3 ตัวคือ 8, 9, 10 (รวม 10 โหนด)"
)

Q6_ANSWER_KEY = (
    "เฉลยโครงสร้าง Binary Tree ตามหลักการ Left-Child Right-Sibling (LCRS):\n"
    "- 1 (Root): กิ่งซ้ายคือ 2, กิ่งขวาว่าง (null)\n"
    "- 2: กิ่งซ้ายคือ 5, กิ่งขวาคือ 3\n"
    "- 5: กิ่งซ้ายว่าง, กิ่งขวาคือ 6\n"
    "- 6: กิ่งซ้ายว่าง, กิ่งขวาคือ 7\n"
    "- 7: เป็น Leaf (กิ่งซ้ายว่าง, กิ่งขวาว่าง)\n"
    "- 3: กิ่งซ้ายว่าง, กิ่งขวาคือ 4\n"
    "- 4: กิ่งซ้ายคือ 8, กิ่งขวาว่าง\n"
    "- 8: กิ่งซ้ายว่าง, กิ่งขวาคือ 9\n"
    "- 9: กิ่งซ้ายว่าง, กิ่งขวาคือ 10\n"
    "- 10: เป็น Leaf (กิ่งซ้ายว่าง, กิ่งขวาว่าง)\n"
    "รวมทั้งหมด 10 โหนด โดยกิ่งซ้ายแทน First Child และกิ่งขวาแทน Next Sibling"
)

KEY_IMG_PATH = ROOT / "public" / "answer-keys" / "q6-lcrs-binary-tree-answer-key.png"
key_img_bytes = KEY_IMG_PATH.read_bytes() if KEY_IMG_PATH.exists() else None

def calculate_qwk(y_true, y_pred, step=0.25, max_score=1.0):
    if len(y_true) != len(y_pred) or len(y_true) == 0:
        return 0.0
    if y_true == y_pred:
        return 1.0
    k = int(round(max_score / step)) + 1
    w = [[((i - j) ** 2) / ((k - 1) ** 2) for j in range(k)] for i in range(k)]
    cat_true = [int(round(v / step)) for v in y_true]
    cat_pred = [int(round(v / step)) for v in y_pred]
    o = [[0] * k for _ in range(k)]
    for t, p in zip(cat_true, cat_pred):
        o[t][p] += 1
    r_t = [sum(o[i][j] for j in range(k)) for i in range(k)]
    r_p = [sum(o[i][j] for i in range(k)) for j in range(k)]
    n = len(y_true)
    e = [[(r_t[i] * r_p[j]) / n for j in range(k)] for i in range(k)]
    num = sum(w[i][j] * o[i][j] for i in range(k) for j in range(k))
    den = sum(w[i][j] * e[i][j] for i in range(k) for j in range(k))
    if den == 0:
        return 1.0
    return 1.0 - (num / den)

async def main():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws_dataset = wb["ชุดข้อสอบ_dataset"]
    
    # First 10 students for Question 6: rows 176 to 185 (DS-171 to DS-180)
    rows = list(range(176, 186))
    assert len(rows) == 10
    
    to_evaluate = []
    for r in rows:
        sid = ws_dataset.cell(r, 1).value
        human_score = float(ws_dataset.cell(r, 7).value or 0.0)
        std_idx = r - 175
        img_name = f"LINE_ALBUM_Photo2.2_260918_{std_idx}.jpg"
        img_path = ROOT / "ชุดข้อสอบใหม่" / "photo_clean_ชุดที่3" / img_name
        to_evaluate.append({
            "row": r,
            "sample_id": sid,
            "student_idx": std_idx,
            "human_score": human_score,
            "img_name": img_name,
            "img_path": img_path
        })
        
    print(f"============================================================")
    print(f"Testing Question 6 (First 10 Students: DS-171 to DS-180)")
    print(f"Target Ground Truth Scores: {[x['human_score'] for x in to_evaluate]}")
    print(f"============================================================")
    
    sem = asyncio.Semaphore(4)
    
    async def grade_one(item):
        async with sem:
            img_bytes = item["img_path"].read_bytes()
            print(f"Grading {item['sample_id']} (idx={item['student_idx']:2d}) ...", flush=True)
            res = await score_with_openai(
                question_text=Q6_QUESTION_TEXT,
                answer_text="",
                max_score=1.0,
                answer_key=Q6_ANSWER_KEY,
                rubrics=Q6_RUBRICS,
                image_bytes_list=[img_bytes],
                image_mime_list=["image/jpeg"],
                answer_key_image_bytes_list=[key_img_bytes] if key_img_bytes else None,
                answer_key_image_mime_list=["image/png"] if key_img_bytes else None,
                strict_rubric_enforcement=True,
            )
            ai_score = float(res.get("score", 0.0))
            conf = res.get("confidence", "unknown")
            teacher_fb = res.get("teacher_feedback", "")
            student_fb = res.get("student_feedback", "")
            combined_fb = res.get("feedback", "")
            diff = round(ai_score - item["human_score"], 2)
            match = (ai_score == item["human_score"])
            
            result_item = {
                "row": item["row"],
                "sample_id": item["sample_id"],
                "student_idx": item["student_idx"],
                "img_name": item["img_name"],
                "human_score": item["human_score"],
                "ai_score": ai_score,
                "diff": diff,
                "match": match,
                "confidence": conf,
                "teacher_feedback": teacher_fb,
                "student_feedback": student_fb,
                "feedback": combined_fb,
            }
            status_str = "MATCH 🟢" if match else f"DIFF 🟠 ({diff:+.2f})"
            print(f"  -> {item['sample_id']} (idx={item['student_idx']:2d}): Human={item['human_score']:.2f} | AI={ai_score:.2f} | {status_str}")
            return result_item

    tasks = [grade_one(item) for item in to_evaluate]
    results = await asyncio.gather(*tasks)
    results.sort(key=lambda x: x["row"])
    
    # Save results to artifacts
    out_json = ROOT / "artifacts" / "q6_first10_test_results.json"
    out_json.parent.mkdir(parents=True, exist_ok=True)
    with open(out_json, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"\nSaved test results to {out_json}")
    
    # Statistics
    y_true = [x["human_score"] for x in results]
    y_pred = [x["ai_score"] for x in results]
    exacts = sum(1 for x in results if x["match"])
    mae = sum(abs(t - p) for t, p in zip(y_true, y_pred)) / len(y_true)
    qwk = calculate_qwk(y_true, y_pred, step=0.25, max_score=1.0)
    
    print("\n" + "="*80)
    print("EVALUATION SUMMARY REPORT FOR QUESTION 6 (FIRST 10 SAMPLES)")
    print("="*80)
    print(f"Total Evaluated: {len(results)}")
    print(f"Exact Matches:   {exacts} / {len(results)} ({exacts / len(results) * 100:.1f}%)")
    print(f"MAE:             {mae:.4f}")
    print(f"QWK (Kappa):     {qwk:.4f}")
    print("="*80)
    
    print(f"\n{'Sample ID':<10} | {'Idx':<4} | {'Human':<6} | {'AI':<6} | {'Diff':<7} | {'Status':<10} | {'Confidence':<10}")
    print("-" * 65)
    for x in results:
        status_text = "EXACT" if x["match"] else f"{x['diff']:+.2f}"
        print(f"{x['sample_id']:<10} | {x['student_idx']:<4d} | {x['human_score']:<6.2f} | {x['ai_score']:<6.2f} | {x['diff']:<+7.2f} | {status_text:<10} | {x['confidence']:<10}")
        
    print("\nSample Feedbacks:")
    for x in results[:3]:
        print(f"\n[{x['sample_id']} - Score: Human={x['human_score']}, AI={x['ai_score']}]")
        print(f"Teacher FB: {x['teacher_feedback'][:160]}...")
        print(f"Student FB: {x['student_feedback'][:160]}...")

if __name__ == "__main__":
    asyncio.run(main())
