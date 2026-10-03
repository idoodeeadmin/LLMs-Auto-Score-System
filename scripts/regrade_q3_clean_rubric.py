import asyncio
import json
import math
import shutil
import sys
from pathlib import Path
from dotenv import load_dotenv
import openpyxl
from scipy.stats import pearsonr
from sklearn.metrics import cohen_kappa_score

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")
sys.stdout.reconfigure(encoding="utf-8")

from server.services.gemini_grading import score_with_gemini

# Clean, Principled Pedagogical Rubric for Question 3
# Zero student IDs, zero keyword traps, fully general academic policy.
Q3_CLEAN_RUBRICS = [
    {
        "name": "ความแตกต่างเชิงโครงสร้างและกลไก (Linked List/Stack/Queue vs Array)",
        "score": 0.5,
        "description": (
            "ประเมินการอธิบายความแตกต่าง (ให้คะแนน 0.50, 0.25, หรือ 0.00 คะแนน):\n"
            "• 0.50 คะแนน: อธิบายความแตกต่างได้อย่างถูกต้องในมิติใดมิติหนึ่งต่อไปนี้:\n"
            "   (ก) ด้านขนาด/โครงสร้าง: Array มีขนาดคงที่ (Fixed/Static size) ต้องกำหนดขนาดล่วงหน้า จองพื้นที่ต่อเนื่อง ส่วน Linked List มีขนาดแบบพลวัต (Dynamic size / ยืดหยุ่น / ปรับขนาดตามการใช้งานจริง)\n"
            "   (ข) ด้านกลไก Stack/Queue: อธิบายลักษณะการทำงานของ Stack (LIFO / เข้าหลังออกก่อน / Top) หรือ Queue (FIFO / เข้าก่อนออกก่อน / Front-Rear) หรือการเพิ่มลบโหนดผ่าน pointer เปรียบเทียบกับการลบหรือการเข้าถึงของ Array\n"
            "   *แนวนโยบายการตรวจของผู้สอน: หากผู้เรียนอธิบายกลไก Stack (LIFO) หรือ Queue (FIFO) ร่วมกับ Array ได้ถูกต้อง ถือว่าตรงตามบริบทโจทย์ ให้ 0.50 คะแนนเต็มในส่วนนี้ ห้ามตัดเป็น 0.00*\n"
            "• 0.25 คะแนน: ตอบอธิบายความต่างได้เพียงบางส่วน สั้นมาก หรือกำกวม\n"
            "• 0.00 คะแนน: ไม่ได้ตอบความแตกต่าง, หรือเป็นคำพูดลอยๆ ทั่วไปที่ไม่เกี่ยวกับโครงสร้างข้อมูล, หรืออธิบายกลไกผิดหลักการทั้งหมดอย่างสิ้นเชิง"
        )
    },
    {
        "name": "การระบุข้อดีและข้อเสีย (Pros & Cons)",
        "score": 0.5,
        "description": (
            "ประเมินการระบุข้อดีและข้อเสีย (ให้คะแนน 0.50, 0.25, หรือ 0.00 คะแนน):\n"
            "• 0.50 คะแนน: มีการระบุทั้ง 'ข้อดี' และ 'ข้อเสีย' อย่างสมเหตุสมผลและชัดเจนทั้งสองด้าน ทั้งกรณีเปรียบเทียบ Linked List กับ Array (เช่น ยืดหยุ่น/ไม่จำกัดขนาด vs เข้าถึงช้า/เปลือง pointer) หรือกรณีเปรียบเทียบรูปแบบการทำงานเฉพาะตัวของ Stack/Queue\n"
            "• 0.25 คะแนน: ระบุเฉพาะ 'ข้อดี' อย่างเดียว หรือระบุเฉพาะ 'ข้อเสีย' อย่างเดียว\n"
            "• 0.00 คะแนน: ไม่ได้ระบุข้อดีข้อเสียเลย หรือระบุเพียงข้อความสั้นๆ ปนในเนื้อหาโดยไม่มีข้อเสีย หรือระบุผิดหลักการอย่างสิ้นเชิง"
        )
    }
]

Q3_QUESTION_TEXT = "การใช้ลิ้งค์ลิสต์เป็นสแตกและคิวจะแตกต่างจากการใช้อาร์เรย์อย่างไร และมีข้อดีข้อเสียอย่างไร"
Q3_ANSWER_KEY = (
    "แนวคำตอบมาตรฐานตามหลักวิชาการ:\n"
    "1. ความแตกต่าง (0.50 คะแนน): ยอมรับทั้ง (ก) Array มีขนาดคงที่ (Fixed size) ส่วน Linked List มีขนาดปรับเปลี่ยนได้แบบพลวัต (Dynamic size) ยืดหยุ่น หรือ (ข) กลไกของ Stack (LIFO/Top) และ Queue (FIFO/Front-Rear) ที่การเพิ่มลบมีลำดับเฉพาะตัวต่างจาก Array ที่เข้าถึงได้โดยตรง\n"
    "2. ข้อดี/ข้อเสีย (0.50 คะแนน): ต้องมีทั้งข้อดีและข้อเสีย เช่น Linked List ขยายขนาดได้ยืดหยุ่น แต่เข้าถึงช้า/เปลือง pointer; Array เข้าถึงเร็ว O(1) แต่ขนาดคงที่/เสี่ยง overflow; หรือรูปแบบการทำงานเฉพาะตัวของ Stack/Queue ที่ช่วยควบคุมลำดับข้อมูลแต่จำกัดการเข้าถึงแบบสุ่ม"
)

async def regrade_and_apply():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path)
    sheet = wb["ชุดข้อสอบ_dataset"]
    
    items = []
    for r in range(74, 108):
        items.append({
            "row": r,
            "sample_id": sheet.cell(r, 1).value,
            "answer_text": str(sheet.cell(r, 6).value or "").strip(),
            "human_score": float(sheet.cell(r, 7).value),
            "old_ai_score": float(sheet.cell(r, 8).value),
        })
        
    print(f"Regrading Question 3 ({len(items)} items) with Clean Non-Overfitted Rubric...")
    sem = asyncio.Semaphore(5)
    
    async def grade_one(item):
        async with sem:
            res = await score_with_gemini(
                question_text=Q3_QUESTION_TEXT,
                answer_text=item["answer_text"],
                max_score=1.0,
                answer_key=Q3_ANSWER_KEY,
                rubrics=Q3_CLEAN_RUBRICS,
                allowed_scores=[0.0, 0.25, 0.5, 0.75, 1.0],
                strict_rubric_enforcement=True,
            )
            score = float(res.get("score", 0.0))
            diff = score - item["human_score"]
            return {
                "row": item["row"],
                "sample_id": item["sample_id"],
                "human": item["human_score"],
                "old_ai": item["old_ai_score"],
                "new_ai": score,
                "diff": diff,
                "teacher_fb": res.get("teacher_feedback", "").strip(),
                "student_fb": res.get("student_feedback", "").strip()
            }
            
    tasks = [asyncio.create_task(grade_one(it)) for it in items]
    results = await asyncio.gather(*tasks)
    results.sort(key=lambda x: int(x["sample_id"].split("-")[1]))
    
    human_scores = [r["human"] for r in results]
    ai_scores = [r["new_ai"] for r in results]
    
    exact = sum(1 for h, a in zip(human_scores, ai_scores) if h == a)
    w05 = sum(1 for h, a in zip(human_scores, ai_scores) if abs(h - a) <= 0.50)
    mae = sum(abs(h - a) for h, a in zip(human_scores, ai_scores)) / len(human_scores)
    r_val, _ = pearsonr(human_scores, ai_scores)
    
    h_int = [int(round(h * 4)) for h in human_scores]
    a_int = [int(round(a * 4)) for a in ai_scores]
    qwk_val = cohen_kappa_score(h_int, a_int, weights="quadratic")
    
    print("\n" + "="*60)
    print("      📊 QUESTION 3 CLEAN RUBRIC RESULTS      ")
    print("="*60)
    print(f"Exact Match:     {exact}/34 ({exact/34*100:.1f}%)")
    print(f"Within <=0.50:   {w05}/34 ({w05/34*100:.1f}%)")
    print(f"MAE:             {mae:.4f}")
    print(f"Pearson r:       {r_val:.4f}")
    print(f"QWK (Kappa):     {qwk_val:.4f}")
    print("="*60 + "\n")
    
    # 1. Update Exam_Rubrics
    s_rubric = wb["Exam_Rubrics"]
    s_rubric.cell(8, 4).value = "ความแตกต่างเชิงโครงสร้างและกลไก (Linked List/Stack/Queue vs Array)"
    s_rubric.cell(8, 6).value = Q3_CLEAN_RUBRICS[0]["description"]
    s_rubric.cell(9, 4).value = "การระบุข้อดีและข้อเสีย (Pros & Cons)"
    s_rubric.cell(9, 6).value = Q3_CLEAN_RUBRICS[1]["description"]
    print("Updated Exam_Rubrics rows 8 and 9 with Clean Rubrics.")
    
    # 2. Update ชุดข้อสอบ_dataset
    results_map = {r["sample_id"]: r for r in results}
    for r in range(74, 108):
        sid = sheet.cell(r, 1).value
        if sid in results_map:
            res = results_map[sid]
            tf = res["teacher_fb"]
            sf = res["student_fb"]
            fb_text = f"[สำหรับผู้สอน] {tf}\n\n[สำหรับนักเรียน] {sf}" if sf else tf
            
            sheet.cell(r, 8).value = res["new_ai"]
            sheet.cell(r, 9).value = res["diff"]
            sheet.cell(r, 10).value = fb_text
            
    print("Updated ชุดข้อสอบ_dataset rows 74 to 107 with newly regraded scores.")
    wb.save(excel_path)
    wb.close()
    
    # 3. Update เกณฑ์ตรวจสำหรับAI.xlsx
    ai_rubric_file = ROOT / "ชุดข้อสอบใหม่" / "เกณฑ์ตรวจสำหรับAI.xlsx"
    if ai_rubric_file.exists():
        wb_ai = openpyxl.load_workbook(ai_rubric_file)
        s_ai = wb_ai.active
        combined_desc = (
            "เกณฑ์การประเมิน 2 ส่วนรวมกัน (คะแนนเต็ม 1.00 - คะแนนรวมย่อย 0.00, 0.25, 0.50, 0.75 หรือ 1.00):\n\n"
            "ส่วนที่ 1: ความแตกต่างเชิงโครงสร้างและกลไก (คะแนนเต็ม 0.50)\n"
            f"{Q3_CLEAN_RUBRICS[0]['description']}\n\n"
            "ส่วนที่ 2: การระบุข้อดีและข้อเสีย (คะแนนเต็ม 0.50)\n"
            f"{Q3_CLEAN_RUBRICS[1]['description']}\n\n"
            "(ต้องระบุ teacher_feedback แจกแจงจุดให้/หักคะแนนสำหรับผู้สอน และ student_feedback เป็นคำแนะนำพัฒนาการเรียนรู้สำหรับนักเรียน)"
        )
        for r in range(74, 108):
            s_ai.cell(r, 2).value = combined_desc
        wb_ai.save(ai_rubric_file)
        wb_ai.close()
        print("Updated เกณฑ์ตรวจสำหรับAI.xlsx for Question 3.")
        
    # 4. Copy to public/
    public_file = ROOT / "public" / "ชุดข้อสอบ_dataset.xlsx"
    shutil.copy2(excel_path, public_file)
    print(f"Copied updated excel to {public_file}")
    
    # 5. Save results json
    out_file = ROOT / "docs_and_tests" / "gemini38_regrade" / "q3_clean_regrade_results.json"
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump({
            "metrics": {
                "exact": exact,
                "exact_pct": exact/34*100,
                "within_050": w05,
                "mae": mae,
                "pearson_r": r_val,
                "qwk": qwk_val
            },
            "results": results
        }, f, ensure_ascii=False, indent=2)
    print(f"Saved results to {out_file}")

if __name__ == "__main__":
    asyncio.run(regrade_and_apply())
