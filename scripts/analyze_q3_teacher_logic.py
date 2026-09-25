import openpyxl
import sys
from collections import Counter
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

# 1. Check Rubrics in Exam_Rubrics for Question 3
ws_rubrics = wb["Exam_Rubrics"]
print("=== CURRENT EXAM_RUBRICS FOR QUESTION 3 ===")
for r in range(5, ws_rubrics.max_row + 1):
    q_no = ws_rubrics.cell(r, 1).value
    if q_no == 3:
        vals = [ws_rubrics.cell(r, c).value for c in range(1, 7)]
        print(f"Row {r}: Topic={vals[1]}, Max={vals[2]}, Name={vals[3]}, Score={vals[4]}")
        print(f"Description:\n{vals[5]}\n{'-'*60}")

# 2. Check Question 3 Submissions
ws_data = wb["ชุดข้อสอบ_dataset"]
q3_samples = []

for r in range(6, ws_data.max_row + 1):
    sid = ws_data.cell(r, 1).value
    q_no = ws_data.cell(r, 2).value
    if q_no == 3:
        q_content = ws_data.cell(r, 4).value
        std_ans = ws_data.cell(r, 6).value
        h = ws_data.cell(r, 7).value
        ai = ws_data.cell(r, 8).value
        fb = ws_data.cell(r, 10).value
        q3_samples.append({
            "row": r,
            "sample_id": sid,
            "q_content": q_content,
            "std_ans": str(std_ans or ""),
            "human_score": float(h) if h is not None else 0.0,
            "ai_score": float(ai) if ai is not None else 0.0,
            "diff": round(float(ai) - float(h), 2) if (ai is not None and h is not None) else 0.0,
            "feedback": str(fb or "")
        })

print(f"\nTotal Q3 Submissions: {len(q3_samples)}")
print(f"Question Prompt:\n{q3_samples[0]['q_content']}\n")

h_dist = Counter(s["human_score"] for s in q3_samples)
ai_dist = Counter(s["ai_score"] for s in q3_samples)
print("Human Score Distribution:", sorted(h_dist.items()))
print("AI Score Distribution:   ", sorted(ai_dist.items()))

exact_count = sum(1 for s in q3_samples if s["diff"] == 0.0)
print(f"Exact Matches: {exact_count} / {len(q3_samples)} ({exact_count/len(q3_samples)*100:.1f}%)")

print("\n" + "="*80)
print("ALL 34 SAMPLES (SORTED BY HUMAN SCORE THEN SAMPLE_ID):")
print("="*80)

for s in sorted(q3_samples, key=lambda x: (x["human_score"], x["sample_id"])):
    mark = "🟢 EXACT" if s["diff"] == 0.0 else f"🟠 DIFF {s['diff']:+0.2f}"
    print(f"[{s['sample_id']}] Human: {s['human_score']:.2f} | AI: {s['ai_score']:.2f} | {mark}")
    print(f"Student Answer:\n{s['std_ans']}")
    if s["diff"] != 0.0:
        print(f"AI Feedback:\n{s['feedback'][:200]}...")
    print("-" * 80)
