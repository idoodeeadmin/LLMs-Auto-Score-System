import asyncio
import json
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q3_QUESTION_TEXT = "การใช้ลิ้งค์ลิสต์เป็นสแตกและคิวจะแตกต่างจากการใช้อาร์เรย์อย่างไร และมีข้อดีข้อเสียอย่างไร (1 คะแนน)"

Q3_ANSWER_KEY = (
    "เฉลยมาตรฐาน:\n"
    "1. ความแตกต่างเชิงโครงสร้าง (Difference):\n"
    "   - Array: มีขนาดคงที่ (Static / Fixed-size) ต้องจองพื้นที่ล่วงหน้าต่อเนื่องในหน่วยความจำ\n"
    "   - Linked List: มีขนาดปรับเปลี่ยนได้ตามจริง (Dynamic / ยืดหยุ่น / ไม่จำกัดขนาดล่วงหน้า) เชื่อมต่อโหนดด้วยพอยน์เตอร์\n"
    "2. ข้อดีและข้อเสีย (Pros & Cons):\n"
    "   - Linked List: ข้อดีคือไม่จำกัดขนาด ไม่เกิด Stack Overflow ได้ง่าย เพิ่ม/ลดโหนดได้สะดวก; "
    "ข้อเสียคือเข้าถึงข้อมูลช้ากว่า (Sequential access ต้องท่องจาก Head) และเปลืองหน่วยความจำเก็บพอยน์เตอร์\n"
    "   - Array: ข้อดีคือเข้าถึงข้อมูลได้รวดเร็วทันทีด้วย Index (O(1)); "
    "ข้อเสียคือขนาดคงที่เสี่ยงเกิด Overflow หากจองน้อยไป หรือเปลืองพื้นที่หากจองมากเกินไป"
)

Q3_RUBRIC_NAME = "การเปรียบเทียบการใช้ Linked List vs Array สำหรับ Stack และ Queue พร้อมข้อดีข้อเสีย"

Q3_RUBRIC_DESC = (
    "ประเมินตามความถูกต้องและความสมบูรณ์ของการตอบคำถาม 2 ส่วนหลัก (ความแตกต่าง + ข้อดีข้อเสีย) รวม 1.00 คะแนน โดยแบ่งเกณฑ์การตัดสินออกเป็น 5 ระดับอย่างเคร่งครัดดังนี้:\n\n"
    "- ได้ 1.00 คะแนน (สมบูรณ์แบบ): \n"
    "  * ตอบความแตกต่างเชิงโครงสร้างได้ถูกต้องครบถ้วน (ระบุชัดเจนว่า Array มีขนาดคงที่/Static ส่วน Linked List มีขนาดปรับเปลี่ยนได้/Dynamic หรือขยายได้ตามต้องการ)\n"
    "  * และระบุข้อดีและข้อเสียที่ถูกต้องตามหลักการอย่างน้อย 2 ประเด็น (เช่น ข้อดี Linked List ไม่จำกัดขนาด/ไม่ล้น และข้อเสียคือเข้าถึงช้ากว่า หรือเปลืองเมมโมรี่เก็บ Pointer หรือเทียบกับ Array ที่เข้าถึงเร็วแต่ขนาดจำกัด)\n"
    "  (ให้อนุโลมสำนวนภาษาและคำศัพท์ หากผู้เรียนอธิบายสาระสำคัญเรื่องขนาดคงที่ vs ขยายได้ และมีข้อดีข้อเสียชัดเจน ให้ 1.00 คะแนนเต็ม ไม่จำเป็นต้องมีคำว่า pointer ก็ได้หากอธิบายกลไกได้ถูกต้อง)\n\n"
    "- ได้ 0.75 คะแนน (ดีมาก / ขาดประเด็นย่อย): \n"
    "  * อธิบายความแตกต่างเรื่องขนาด (Fixed vs Dynamic) ได้ถูกต้องชัดเจน\n"
    "  * แต่ระบุข้อดีหรือข้อเสียเพียงด้านเดียว (เช่น มีแต่ข้อดี ไม่มีข้อเสียเลย หรือมีแต่ข้อเสีย ไม่มีข้อดี หรือระบุข้อดีข้อเสียเฉพาะฝั่ง Linked List โดยไม่เปรียบเทียบกับ Array)\n"
    "  * หรือระบุข้อดีข้อเสียครบแต่มีจุดคลาดเคลื่อนเล็กน้อย (เช่น ระบุว่าอาร์เรย์ค้นหาช้า)\n\n"
    "- ได้ 0.50 คะแนน (ระดับมาตรฐาน / คนได้เยอะสุด): \n"
    "  * ตอบเฉพาะข้อดี หรือข้อเสียเพียงอย่างเดียวสั้นๆ หรือ\n"
    "  * ตอบเปรียบเทียบเฉพาะเรื่องขนาด (Fixed vs Dynamic) สั้นๆ โดยไม่ได้แยกแยะข้อดีข้อเสียอย่างเป็นรูปธรรม หรือ\n"
    "  * ตอบเน้นเฉพาะพฤติกรรมการทำงานของ Stack (LIFO) / Queue (FIFO) แต่ยังขาดการเปรียบเทียบเชิงลึกด้านหน่วยความจำ/การเข้าถึงข้อมูล หรือ\n"
    "  * เขียนเป็นความเรียงรวมๆ ไม่ได้แจกแจงข้อดีข้อเสีย แต่สื่อถึงความยืดหยุ่นของ Linked List หรือข้อจำกัดขนาดของ Array\n\n"
    "- ได้ 0.25 คะแนน (มีจุดถูกเล็กน้อย): \n"
    "  * เขียนสั้นมาก ไม่ระบุชื่อโครงสร้างให้ชัดเจน แต่มีคีย์เวิร์ดที่ถูกต้องตามหลักการเพียงจุดเดียว เช่น มีคำว่า 'ยืดหยุ่น' หรือ 'static' (เช่น ตอบเพียงบรรทัดเดียวว่า ยืดหยุ่น จัดการเป็นลำดับ มีพื้นที่จำกัด เป็น static)\n\n"
    "- ได้ 0.00 คะแนน (ผิดหรือไม่ตรงประเด็น): \n"
    "  * ตอบไม่ตรงประเด็นหรือตอบมั่ว (เช่น ตอบเรื่องประหยัดพลังงาน, เปรียบเทียบกับรถไฟ)\n"
    "  * เขียนเพียงประโยคเดียวสั้นๆ โดยไม่มีการตอบข้อดีหรือข้อเสียใดๆ เลย (เช่น เขียนแค่ว่า อาร์เรย์ต้องกำหนด Size ของ Array Linked List เลื่อน Size ได้ แค่นี้โดยไม่มีข้อดีข้อเสีย ให้ 0.00)\n"
    "  * เขียนข้อดีข้อเสียที่ผิดหลักการทางคอมพิวเตอร์อย่างมีนัยสำคัญ (เช่น บอกว่า Array เสี่ยง Error หรือไม่รู้ขนาดข้อมูล)\n"
    "  * ไม่ตอบ หรือตอบว่าทำไม่ได้\n\n"
    "(สำคัญ: ต้องส่งกลับ teacher_feedback แจกแจงจุดที่ได้/ตัดคะแนนตามระดับ 1.00, 0.75, 0.50, 0.25, 0.00 อย่างละเอียดสำหรับผู้สอน และ student_feedback เป็นคำแนะนำพัฒนาการเรียนรู้สำหรับนักเรียน)"
)

Q3_RUBRICS = [
    {
        "name": Q3_RUBRIC_NAME,
        "score": 1.0,
        "description": Q3_RUBRIC_DESC,
    }
]

async def main():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws_data = wb["ชุดข้อสอบ_dataset"]
    
    # Grab first 10 students of Q3: rows 74 to 83 (DS-069 to DS-078)
    first10 = []
    for r in range(74, 84):
        sid = ws_data.cell(r, 1).value
        q_no = ws_data.cell(r, 2).value
        assert q_no == 3
        ans = str(ws_data.cell(r, 6).value or "")
        h = float(ws_data.cell(r, 7).value or 0.0)
        old_ai = float(ws_data.cell(r, 8).value or 0.0)
        first10.append({
            "row": r,
            "sample_id": sid,
            "student_ans": ans,
            "human_score": h,
            "old_ai_score": old_ai
        })
        
    print(f"=== TESTING Q3 CALIBRATED 5-LEVEL RUBRIC ON FIRST 10 STUDENTS ===")
    print(f"Number of test samples: {len(first10)}")
    
    results = []
    exact_old = sum(1 for x in first10 if x["old_ai_score"] == x["human_score"])
    print(f"Previous Rubric Exact Matches: {exact_old} / 10 ({exact_old*10:.1f}%)\n")
    
    for item in first10:
        print(f"Grading {item['sample_id']} ...", end=" ", flush=True)
        res = await score_with_openai(
            question_text=Q3_QUESTION_TEXT,
            answer_text=item["student_ans"],
            max_score=1.0,
            answer_key=Q3_ANSWER_KEY,
            rubrics=Q3_RUBRICS,
            image_bytes_list=None,
            strict_rubric_enforcement=True,
        )
        new_ai = float(res.get("score", 0.0))
        diff = round(new_ai - item["human_score"], 2)
        match = (new_ai == item["human_score"])
        
        status = "🟢 EXACT" if match else f"🟠 DIFF ({diff:+0.2f})"
        print(f"Human: {item['human_score']:.2f} | New AI: {new_ai:.2f} (Old: {item['old_ai_score']:.2f}) -> {status}")
        
        results.append({
            "sample_id": item["sample_id"],
            "student_ans": item["student_ans"],
            "human_score": item["human_score"],
            "old_ai_score": item["old_ai_score"],
            "new_ai_score": new_ai,
            "diff": diff,
            "is_exact": match,
            "teacher_feedback": res.get("teacher_feedback", ""),
            "student_feedback": res.get("student_feedback", "")
        })
        
    exact_new = sum(1 for x in results if x["is_exact"])
    mae_new = sum(abs(x["diff"]) for x in results) / len(results)
    mae_old = sum(abs(x["old_ai_score"] - x["human_score"]) for x in results) / len(results)
    
    print("\n" + "="*70)
    print(f"PILOT RESULTS SUMMARY (10 STUDENTS):")
    print(f"Old Rubric Exact Matches: {exact_old} / 10 ({exact_old*10:.1f}%) | MAE: {mae_old:.4f}")
    print(f"New Rubric Exact Matches: {exact_new} / 10 ({exact_new*10:.1f}%) | MAE: {mae_new:.4f}")
    print("="*70)
    
    # Save pilot results
    out_file = ROOT / "artifacts" / "q3_first10_pilot_results.json"
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"Saved pilot results to {out_file}")

if __name__ == "__main__":
    asyncio.run(main())
