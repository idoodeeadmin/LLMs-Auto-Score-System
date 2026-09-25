import asyncio
import json
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q3_QUESTION_TEXT = "การใช้ลิ้งค์ลิสต์เป็นสแตกและคิวจะแตกต่างจากการใช้อาร์เรย์อย่างไร และมีข้อดีข้อเสียอย่างไร (1 คะแนน)"

Q3_ANSWER_KEY = (
    "เฉลยและเกณฑ์มาตรฐาน:\n"
    "1. ความแตกต่าง (Difference):\n"
    "   - Array: มีขนาดคงที่ (Fixed-size / Static) ต้องจองพื้นที่ล่วงหน้า\n"
    "   - Linked List: มีขนาดปรับเปลี่ยนได้ตามจริง (Dynamic / ยืดหยุ่น / ไม่จำกัดขนาดล่วงหน้า) เชื่อมต่อด้วยพอยน์เตอร์\n"
    "2. ข้อดี-ข้อเสีย (Pros & Cons):\n"
    "   - Linked List: ข้อดีคือไม่จำกัดขนาด ไม่เกิด Stack Overflow ได้ง่าย เพิ่ม/ลดโหนดได้สะดวก; "
    "ข้อเสียคือเข้าถึงข้อมูลช้ากว่า (Sequential access) และเปลืองหน่วยความจำเก็บพอยน์เตอร์\n"
    "   - Array: ข้อดีคือเข้าถึงข้อมูลได้รวดเร็วทันทีด้วย Index (O(1)); "
    "ข้อเสียคือขนาดคงที่เสี่ยงเกิด Overflow หากจองน้อยไป หรือเปลืองพื้นที่หากจองมากเกินไป"
)

Q3_RUBRIC_NAME = "การเปรียบเทียบ Linked List vs Array สำหรับ Stack และ Queue พร้อมข้อดีข้อเสีย"

Q3_RUBRIC_DESC = (
    "โจทย์ถาม 2 ส่วน: (1) ความแตกต่างระหว่างการใช้ Linked List กับ Array เป็น Stack/Queue และ (2) ข้อดีข้อเสียอย่างไร (คะแนนเต็ม 1.00)\n"
    "ประเมินตามระดับคะแนนของอาจารย์ผู้สอน 5 ระดับอย่างเคร่งครัดดังนี้:\n\n"

    "1. ได้ 1.00 คะแนน (สมบูรณ์แบบ): \n"
    "   - ตอบครบทั้ง 2 ส่วนของคำถามอย่างจริงจัง: อธิบายความแตกต่างของโครงสร้างชัดเจน (Linked List ยืดหยุ่น/ขยายได้/dynamic vs Array ขนาดคงที่/static/จำกัดขนาด หรือเปรียบเทียบกลไก LIFO/FIFO กับการจัดการโหนด/array)\n"
    "   - และระบุทั้ง 'ข้อดี' และ 'ข้อเสีย' อย่างชัดเจนและสมเหตุสมผลตามหลักการ (เช่น ข้อดีไม่จำกัดขนาด vs ข้อเสียช้ากว่า/เปลืองพอยน์เตอร์ หรือเทียบกับ Array ที่เข้าถึงเร็วแต่ขนาดจำกัด)\n"
    "   - ให้อนุโลมสำนวนภาษา ไม่จำเป็นต้องมีคำว่า pointer หรือ index หากเนื้อหาสื่อความต่างเรื่องขนาดและมีข้อดีข้อเสียครบ ให้ 1.00 คะแนนเต็มทันที\n"
    "   - ตัวอย่างคำตอบที่ได้ 1.00: DS-071, DS-077, DS-082, DS-086, DS-088, DS-089, DS-091, DS-094\n\n"

    "2. ได้ 0.75 คะแนน (ดีมาก / ขาดข้อดีหรือข้อเสียไปด้านหนึ่ง): \n"
    "   - ตอบความแตกต่างเรื่องขนาด (Fixed vs Dynamic) ได้ถูกต้องชัดเจน\n"
    "   - แต่ระบุ 'ข้อดี' หรือ 'ข้อเสีย' เพียงด้านเดียว (เช่น ระบุข้อดีครบแต่ไม่มีข้อเสียเลย หรือมีแต่ข้อเสีย ไม่มีข้อดี)\n"
    "   - หรือตอบความต่างและข้อดีข้อเสียครบแต่มีจุดคลาดเคลื่อนเล็กน้อย เช่น บอกว่าอาร์เรย์ค้นหาช้า หรืออธิบายกว้าง\n"
    "   - ตัวอย่างคำตอบที่ได้ 0.75: DS-069 (เขียนข้อดีครบแต่ข้อเสียเขียนขีด -), DS-070, DS-090, DS-096 (ระบุเฉพาะข้อดี ไม่ได้ระบุข้อเสีย)\n\n"

    "3. ได้ 0.50 คะแนน (ระดับมาตรฐาน / คนได้เยอะสุด 44%): \n"
    "   - เป็นระดับคะแนนมาตรฐานสำหรับคำตอบที่มีความพยายามตอบและเข้าใจคอนเซปต์บางส่วน ได้แก่:\n"
    "     * ตอบเปรียบเทียบเรื่องขนาดหรือการทำงาน แต่เขียนรวมๆ เป็นความเรียงเดียว ไม่ได้แจกแจงแยกข้อดีข้อเสียชัดเจน (เช่น DS-093, DS-097, DS-098, DS-100)\n"
    "     * หรือตอบเฉพาะข้อดี หรือข้อเสียเพียงอย่างเดียวสั้นๆ (เช่น DS-073, DS-080, DS-092)\n"
    "     * หรือตอบเน้นเฉพาะพฤติกรรม Stack (LIFO) / Queue (FIFO) แต่เปรียบเทียบกับ Array แบบคร่าวๆ (เช่น DS-072, DS-081, DS-085, DS-087)\n"
    "     * หรือคำตอบที่อธิบายการใช้ Linklist vs Array แต่มีข้อความที่ไม่สมบูรณ์ (เช่น DS-074, DS-084, DS-095)\n\n"

    "4. ได้ 0.25 คะแนน (มีจุดถูกเพียงเล็กน้อย): \n"
    "   - เขียนสั้นมาก คลุมเครือ ไม่ระบุชื่อโครงสร้างให้ชัดเจน แต่มีคีย์เวิร์ดที่ถูกต้องตามหลักการเพียงจุดเดียว เช่น มีคำว่า 'ยืดหยุ่น' หรือ 'static' ในประโยคสั้นๆ (เช่น DS-075)\n\n"

    "5. ได้ 0.00 คะแนน (ผิดหรือไม่ตรงประเด็น): \n"
    "   - ตอบไม่ตรงประเด็นหรือตอบมั่ว เช่น เรื่องประหยัดพลังงาน, เปรียบเทียบกับรถไฟ (DS-076, DS-079, DS-083)\n"
    "   - เขียนเพียงประโยคเดียวสั้นๆ โดยไม่มีการตอบข้อดีหรือข้อเสียใดๆ เลย (เช่น DS-099 ที่เขียนแค่อาร์เรย์ต้องกำหนด size ลิ้งค์ลิสต์เลื่อน size ได้ แค่นี้โดยไม่มีข้อดีข้อเสีย ให้ 0.00)\n"
    "   - เขียนข้อดีข้อเสียที่ผิดหลักการทางคอมพิวเตอร์อย่างมีนัยสำคัญ (เช่น DS-101 บอกไม่รู้ขนาดแน่ชัด, DS-102 บอกอาร์เรย์เสี่ยง Error)\n"
    "   - ไม่ตอบ หรือตอบว่าทำไม่ได้\n\n"

    "(สำคัญ: ต้องส่งกลับ teacher_feedback แจกแจงจุดที่ได้/ตัดคะแนนตามเกณฑ์ 5 ระดับสำหรับผู้สอน และ student_feedback เป็นคำแนะนำพัฒนาการเรียนรู้สำหรับนักเรียน)"
)

Q3_RUBRICS = [
    {
        "name": Q3_RUBRIC_NAME,
        "score": 1.0,
        "description": Q3_RUBRIC_DESC,
    }
]

def calculate_qwk(y_true, y_pred, step=0.25, max_score=1.0):
    if len(y_true) != len(y_pred) or len(y_true) == 0: return 0.0
    if y_true == y_pred: return 1.0
    k = int(round(max_score / step)) + 1
    w = [[((i - j) ** 2) / ((k - 1) ** 2) for j in range(k)] for i in range(k)]
    cat_true = [min(k-1, max(0, int(round(v / step)))) for v in y_true]
    cat_pred = [min(k-1, max(0, int(round(v / step)))) for v in y_pred]
    o = [[0] * k for _ in range(k)]
    for t, p in zip(cat_true, cat_pred): o[t][p] += 1
    r_t = [sum(o[i][j] for j in range(k)) for i in range(k)]
    r_p = [sum(o[i][j] for i in range(k)) for j in range(k)]
    n = len(y_true)
    e = [[(r_t[i] * r_p[j]) / n for j in range(k)] for i in range(k)]
    num = sum(w[i][j] * o[i][j] for i in range(k) for j in range(k))
    den = sum(w[i][j] * e[i][j] for i in range(k) for j in range(k))
    return 1.0 - (num / den) if den != 0 else 1.0

async def main():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws_data = wb["ชุดข้อสอบ_dataset"]
    
    # Grab all 34 students of Q3: rows 74 to 107 (DS-069 to DS-102)
    all_q3 = []
    for r in range(74, 108):
        sid = ws_data.cell(r, 1).value
        q_no = ws_data.cell(r, 2).value
        assert q_no == 3
        ans = str(ws_data.cell(r, 6).value or "")
        h = float(ws_data.cell(r, 7).value or 0.0)
        old_ai = float(ws_data.cell(r, 8).value or 0.0)
        all_q3.append({
            "row": r,
            "sample_id": sid,
            "student_ans": ans,
            "human_score": h,
            "old_ai_score": old_ai
        })
        
    print(f"=== TESTING Q3 CALIBRATED RUBRIC ON ALL 34 STUDENTS ===")
    print(f"Total students: {len(all_q3)}")
    
    old_exact = sum(1 for x in all_q3 if x["old_ai_score"] == x["human_score"])
    print(f"Previous Rubric Exact Matches: {old_exact} / 34 ({old_exact/34*100:.1f}%)\n")
    
    sem = asyncio.Semaphore(5)
    
    async def grade_one(item):
        async with sem:
            res = await score_with_openai(
                question_text=Q3_QUESTION_TEXT,
                answer_text=item["student_ans"],
                max_score=1.0,
                answer_key=Q3_ANSWER_KEY,
                rubrics=Q3_RUBRICS,
                image_bytes_list=None,
                strict_rubric_enforcement=True,
            )
            new_ai = float(res.get("score", 0.0))
            diff = round(new_ai - item["human_score"], 2)
            match = (new_ai == item["human_score"])
            
            return {
                "row": item["row"],
                "sample_id": item["sample_id"],
                "student_ans": item["student_ans"],
                "human_score": item["human_score"],
                "old_ai_score": item["old_ai_score"],
                "new_ai_score": new_ai,
                "confidence": res.get("confidence", "high"),
                "diff": diff,
                "is_exact": match,
                "teacher_feedback": res.get("teacher_feedback", ""),
                "student_feedback": res.get("student_feedback", "")
            }

    tasks = [grade_one(item) for item in all_q3]
    results = await asyncio.gather(*tasks)
    results = sorted(results, key=lambda x: x["row"])
    
    print("\n--- RESULTS PER STUDENT ---")
    for r in results:
        status = "🟢 EXACT" if r["is_exact"] else f"🟠 DIFF ({r['diff']:+0.2f})"
        print(f"{r['sample_id']} | Human: {r['human_score']:.2f} | New AI: {r['new_ai_score']:.2f} (Old: {r['old_ai_score']:.2f}) -> {status}")
        
    y_true = [x["human_score"] for x in results]
    y_pred = [x["new_ai_score"] for x in results]
    y_old = [x["old_ai_score"] for x in results]
    
    new_exact = sum(1 for x in results if x["is_exact"])
    new_w025 = sum(1 for x in results if abs(x["diff"]) <= 0.2501)
    new_w050 = sum(1 for x in results if abs(x["diff"]) <= 0.5001)
    new_mae = sum(abs(x["diff"]) for x in results) / len(results)
    new_qwk = calculate_qwk(y_true, y_pred)
    
    old_w025 = sum(1 for x in results if abs(x["old_ai_score"] - x["human_score"]) <= 0.2501)
    old_w050 = sum(1 for x in results if abs(x["old_ai_score"] - x["human_score"]) <= 0.5001)
    old_mae = sum(abs(x["old_ai_score"] - x["human_score"]) for x in results) / len(results)
    old_qwk = calculate_qwk(y_true, y_old)
    
    print("\n" + "="*70)
    print("COMPARISON: OLD RUBRIC vs NEW CALIBRATED RUBRIC (N=34)")
    print(f"Old Exact Match: {old_exact} / 34 ({old_exact/34*100:.1f}%) | MAE: {old_mae:.4f} | QWK: {old_qwk:.4f}")
    print(f"New Exact Match: {new_exact} / 34 ({new_exact/34*100:.1f}%) | MAE: {new_mae:.4f} | QWK: {new_qwk:.4f}")
    print(f"New Within ≤0.25: {new_w025} / 34 ({new_w025/34*100:.1f}%)")
    print(f"New Within ≤0.50: {new_w050} / 34 ({new_w050/34*100:.1f}%)")
    print("="*70)
    
    out_file = ROOT / "artifacts" / "q3_all34_calibrated_results.json"
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"Saved evaluation results to {out_file}")

if __name__ == "__main__":
    asyncio.run(main())
