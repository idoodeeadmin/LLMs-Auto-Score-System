"""Grade the 34 cleaned Q2 or Q3 answer images without changing the workbook."""

import asyncio
import hashlib
import json
import os
import sys
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")
from scripts.regrade_corrected_q2_q3_answers import (  # noqa: E402
    OPENAI_MODEL,
    PROMPT_VERSION,
    score_with_openai,
)

BOOK = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx"
QUESTION = int(os.getenv("IMAGE_QUESTION", "2"))
assert QUESTION in (2, 3)
TEXT_RECHECK = os.getenv("TEXT_RECHECK") == "1"
IMAGE_DIR = ROOT / f"ชุดข้อสอบใหม่/photo_clean_text{QUESTION}"
OUT = ROOT / (
    f"docs_and_tests/q2_q3_review/q{QUESTION}_text_repeat_current_rubric.json"
    if TEXT_RECHECK else
    f"docs_and_tests/q2_q3_review/q{QUESTION}_clean_image_current_rubric_results.json"
)
ANSWER_KEY = {
    2: "O(n log n) เติบโตช้ากว่า O(n^2) เมื่อข้อมูลมีขนาดใหญ่ เช่น Merge Sort หรือ Quick Sort เทียบกับ Bubble Sort หรือ Selection Sort",
    3: "Array ต้องกำหนดขนาดล่วงหน้าและใช้ตำแหน่ง index ส่วน Linked List ปรับขนาดได้และใช้โหนดเชื่อมกัน ข้อดีของ Linked List คือเพิ่มหรือลบข้อมูลได้โดยไม่ต้องจองขนาดคงที่ ข้อเสียคือใช้หน่วยความจำสำหรับตัวชี้และเข้าถึงตำแหน่งโดยตรงไม่ได้",
}[QUESTION]
ALLOWED = [0.0, 1.0, 1.5, 2.0] if QUESTION == 2 else [0.0, 0.25, 0.5, 0.75, 1.0]
MAX_SCORE = 2.0 if QUESTION == 2 else 1.0


async def main():
    wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    sheet = wb["ชุดข้อสอบ_dataset"]
    rubric_sheet = wb["Exam_Rubrics"]
    rubric = [{
        "name": rubric_sheet.cell(row, 4).value,
        "score": float(rubric_sheet.cell(row, 5).value),
        "description": rubric_sheet.cell(row, 6).value,
        "allowed_scores": ALLOWED if QUESTION == 2 else [0.0, 0.25, 0.5],
    } for row in ((7,) if QUESTION == 2 else (8, 9))]
    images = sorted(IMAGE_DIR.glob("IMG_*.jpg"))
    assert len(images) == 34, f"Expected 34 cleaned images, got {len(images)}"
    if QUESTION == 2:
        assert "คำอธิบายเพิ่มเติมในการแยกคะแนนบางส่วน" not in rubric[0]["description"]
    items = []
    first_sample = 35 if QUESTION == 2 else 69
    first_row = first_sample + 5
    for index, row_no in enumerate(range(first_row, first_row + 34)):
        row = [sheet.cell(row_no, col).value for col in range(1, 9)]
        assert row[0] == f"DS-{first_sample + index:03d}" and row[1] == QUESTION
        assert row[4] == "text" and isinstance(row[6], (int, float))
        items.append({
            "sample_id": row[0], "question": row[3], "answer": row[5],
            "human_score": float(row[6]),
            "text_ai_score": float(row[7]), "image": images[index].name,
            "image_sha256": hashlib.sha256(images[index].read_bytes()).hexdigest(),
        })
    wb.close()
    report = {
        "started_at": datetime.now(timezone.utc).isoformat(),
        "model": OPENAI_MODEL,
        "prompt_version": PROMPT_VERSION,
        "question_no": QUESTION,
        "grading_function": "server.services.openai_grading.score_with_openai",
        "source_sha256": hashlib.sha256(BOOK.read_bytes()).hexdigest(),
        "image_directory": str(IMAGE_DIR),
        "answer_key": ANSWER_KEY,
        "rubric": rubric,
        "allowed_scores": ALLOWED,
        "input_modality": "text_only" if TEXT_RECHECK else "image_only",
        "results": [],
    }
    semaphore = asyncio.Semaphore(3)

    async def grade(item):
        async with semaphore:
            path = IMAGE_DIR / item["image"]
            try:
                result = await score_with_openai(
                    question_text=str(item["question"]),
                    answer_text=str(item["answer"]) if TEXT_RECHECK else "",
                    max_score=MAX_SCORE,
                    answer_key=ANSWER_KEY,
                    rubrics=rubric,
                    allowed_scores=ALLOWED,
                    image_bytes_list=None if TEXT_RECHECK else [path.read_bytes()],
                    image_mime_list=None if TEXT_RECHECK else ["image/jpeg"],
                )
                success = (
                    not result.get("metrics", {}).get("manual_review_required", False)
                    and float(result.get("score", -1)) in ALLOWED
                    and bool(result.get("teacher_feedback", "").strip())
                )
                error = None
            except Exception as exc:
                result, success, error = None, False, str(exc)
            return {**item, "success": success, "error": error, "result": result}

    tasks = [asyncio.create_task(grade(item)) for item in items]
    for task in asyncio.as_completed(tasks):
        item = await task
        report["results"].append(item)
        OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        score = item["result"]["score"] if item["result"] else None
        print(item["sample_id"], "score", score, "success", item["success"], flush=True)

    report["results"].sort(key=lambda item: item["sample_id"])
    valid = [item for item in report["results"] if item["success"]]
    modality = "text" if TEXT_RECHECK else "image"
    report["summary"] = {
        "completed": len(valid), "total": len(items),
        f"{modality}_teacher_exact": sum(x["result"]["score"] == x["human_score"] for x in valid),
        f"{modality}_saved_text_exact": sum(x["result"]["score"] == x["text_ai_score"] for x in valid),
        f"{modality}_teacher_mae": (
            sum(abs(x["result"]["score"] - x["human_score"]) for x in valid) / len(valid)
            if valid else None
        ),
    }
    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report["summary"]), flush=True)


if __name__ == "__main__":
    asyncio.run(main())
