# -*- coding: utf-8 -*-
import sys
import os
import json
import asyncio
from pathlib import Path
from dotenv import load_dotenv

load_dotenv()
sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.append(str(ROOT / '.venv/Lib/site-packages'))

from server.services.openai_grading import score_with_openai
from scripts.run_dataset_benchmark import calculate_mae, calculate_qwk

Q1_RUBRIC = [
    {
        "name": "Row-major vs Column-major",
        "score": 2.0,
        "description": (
            "2.00 คะแนน (เต็ม):\n"
            "อธิบายความแตกต่างครบทั้ง 2 ฝั่ง โดยต้องระบุ \"วิธีเรียกใช้งาน / ลำดับการเข้าถึง Array\" หรือ \"ยกตัวอย่างที่แสดงให้เห็นความต่างในการเรียก\" "
            "(เช่น ระบุว่า Row-major เรียกแถวก่อนหลัก vs Column-major เรียกหลักก่อนแถว, หรือแสดงตัวอย่างลำดับดัชนี/สูตร/การเข้าถึงข้อมูลที่สะท้อนความต่างชัดเจน)\n\n"
            "1.00 คะแนน:\n"
            "ตอบได้เพียงมโนทัศน์กว้างๆ ว่าเป็น \"แนวนอน vs. แนวตั้ง\" (หรือบอกแค่นิยาม Row/Column หรือเก็บแบบแถว vs คอลัมน์) โดยไม่ได้อธิบายวิธีเรียกหรือไม่มีตัวอย่างการเข้าถึงข้อมูล "
            "หรืออธิบายวิธีเรียกถูกต้องเพียงฝั่งเดียว\n\n"
            "0.00 คะแนน:\n"
            "ไม่ตอบ หรือตอบไม่ถูกต้องตามหลักการทั้งหมดอย่างสิ้นเชิง"
        ),
        "allowed_scores": [0.0, 1.0, 2.0]
    }
]

ANSWER_KEY = (
    "Row-major คือการจัดเก็บหรือเข้าถึงข้อมูลตามแถว (แนวนอน) โดยเรียกแถวก่อนหลัก "
    "ส่วน Column-major คือการจัดเก็บหรือเข้าถึงข้อมูลตามคอลัมน์ (แนวตั้ง) โดยเรียกหลักก่อนแถว"
)

# Load the verified 34 clean items
import openpyxl
wb = openpyxl.load_workbook('ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx', data_only=True)
ws = wb['ชุดข้อสอบ_dataset']
q1_clean_files = sorted(os.listdir('ชุดข้อสอบใหม่/photo_clean_text1'))

file_to_hscore = {}
file_to_ans = {}
for r in range(6, 40):
    idx = r - 6
    if idx < len(q1_clean_files):
        fn = q1_clean_files[idx].replace('.jpg', '.HEIC')
        h = ws.cell(row=r, column=7).value
        ans = ws.cell(row=r, column=6).value
        file_to_hscore[fn] = float(h) if h is not None else 0.0
        file_to_ans[fn] = ans

data = json.load(open('artifacts/q1_all_42_scan.json', encoding='utf-8'))
bad_files = {
    'IMG_2791.HEIC', 'IMG_2797.HEIC', 'IMG_2802.HEIC', 'IMG_2810.HEIC',
    'IMG_2815.HEIC', 'IMG_2822.HEIC', 'IMG_2825.HEIC', 'IMG_2839.HEIC',
}
candidates = [d for d in data if d['filename'] not in bad_files]

items = []
for i, c in enumerate(candidates):
    fn = c['filename']
    if fn in file_to_hscore:
        score = file_to_hscore[fn]
        ans = file_to_ans[fn]
    else:
        raw_s = c.get('teacher_score')
        score = 2.0 if (raw_s is None or raw_s > 2.0) else float(raw_s)
        ans = c.get('answer_text')

    items.append({
        'sample_id': f"DS-{i+1:03d}",
        'filename': fn,
        'student_answer': str(ans or '').strip(),
        'human_score': score
    })

async def main():
    print(f"Testing All 34 Clean Q1 Samples with AI Grading...\n")
    sem = asyncio.Semaphore(5)

    async def grade_item(it):
        async with sem:
            res = await score_with_openai(
                question_text="อธิบายความต่างของ Row-major vs Column-major",
                answer_text=it['student_answer'],
                max_score=2.0,
                answer_key=ANSWER_KEY,
                rubrics=Q1_RUBRIC,
                allowed_scores=[0.0, 1.0, 2.0]
            )
            s = float(res['score'])
            h = float(it['human_score'])
            diff = round(s - h, 2)
            match_str = "MATCH" if diff == 0.0 else f"DIFF ({diff:+.1f})"
            print(f"{it['sample_id']} ({it['filename']}) | Human: {h:.1f} | AI: {s:.1f} -> {match_str}", flush=True)
            return {
                'sample_id': it['sample_id'],
                'filename': it['filename'],
                'student_answer': it['student_answer'],
                'human': h,
                'ai': s,
                'diff': diff,
                'confidence': res.get('confidence', 'medium'),
                'teacher_feedback': res.get('teacher_feedback', ''),
                'student_feedback': res.get('student_feedback', '')
            }

    results = await asyncio.gather(*(grade_item(it) for it in items))
    results.sort(key=lambda x: x['sample_id'])

    exact = sum(1 for r in results if r['diff'] == 0.0)
    diff05 = sum(1 for r in results if abs(r['diff']) <= 0.5)
    truth = [r['human'] for r in results]
    predicted = [r['ai'] for r in results]
    mae = calculate_mae(truth, predicted)
    qwk = calculate_qwk(truth, predicted, step=1.0, max_score=2.0)

    out_file = ROOT / 'artifacts' / 'q1_34_clean_results.json'
    out_file.write_text(json.dumps({
        'total': len(results),
        'exact': exact,
        'diff05': diff05,
        'mae': mae,
        'qwk': qwk,
        'results': results
    }, ensure_ascii=False, indent=2), encoding='utf-8')

    print("\n" + "="*70)
    print("34 CLEAN Q1 BENCHMARK SUMMARY:")
    print(f"Exact Match: {exact}/{len(results)} ({exact/len(results)*100:.1f}%)")
    print(f"Diff <= 0.5: {diff05}/{len(results)} ({diff05/len(results)*100:.1f}%)")
    print(f"MAE:         {mae:.4f}")
    print(f"QWK:         {qwk:.4f}")
    print("="*70)

if __name__ == '__main__':
    asyncio.run(main())
