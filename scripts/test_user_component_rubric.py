import asyncio
import json
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q3_QUESTION_TEXT = "การใช้ลิ้งค์ลิสต์เป็นสแตกและคิวจะแตกต่างจากการใช้อาร์เรย์อย่างไร และมีข้อดีข้อเสียอย่างไร (1 คะแนน)"

Q3_ANSWER_KEY = (
    "เฉลยและแนวคำตอบ:\n"
    "1. ความแตกต่าง: Linked List ขนาดปรับเปลี่ยนได้ (Dynamic) ใช้โหนด/พอยน์เตอร์ ส่วน Array ขนาดคงที่ (Static/Fixed) จองพื้นที่ล่วงหน้า\n"
    "2. ข้อดี: Linked List ไม่จำกัดขนาด / Array เข้าถึงข้อมูลได้เร็ว\n"
    "3. ข้อเสีย: Linked List เปลืองหน่วยความจำเก็บพอยน์เตอร์ เข้าถึงข้อมูลช้ากว่า / Array ขนาดคงที่ เสี่ยง Overflow หรือเปลืองพื้นที่"
)

# เกณฑ์ตามที่ผู้ใช้กำหนด (แยก 3 องค์ประกอบ รวม 1.00 คะแนน พร้อมหลักการอนุโลมปล่อยผ่านหากไม่อธิบายมั่ว)
USER_COMPONENT_RUBRIC_NAME = "การประเมินแยก 3 องค์ประกอบ (ความแตกต่าง 0.50 + ข้อดี 0.25 + ข้อเสีย 0.25)"

USER_COMPONENT_RUBRIC_DESC = (
    "เกณฑ์การให้คะแนนรวม 1.00 คะแนน พิจารณาแยก 3 ส่วนอย่างเป็นอิสระต่อกัน แล้วนำคะแนนมารวมกัน:\n"
    "1. ส่วนความแตกต่าง (เต็ม 0.50 คะแนน):\n"
    "   - ให้ 0.50 คะแนน: มีการอธิบายความแตกต่างระหว่าง Linked List กับ Array (เช่น เรื่องขนาดคงที่ vs ปรับได้, การจองพื้นที่, หรือกลไกการทำงาน Stack/Queue เทียบกับ Array) ได้พอเข้าใจ โดยให้อนุโลมปล่อยผ่านหากไม่ได้เข้าเกณฑ์อธิบายมั่วหรือตอบนอกเรื่อง\n"
    "   - ให้ 0.00 คะแนน: ตอบมั่ว ตอบไม่ตรงประเด็น (เช่น ช่วยประหยัดพลังงาน, รถไฟ) หรือไม่ได้อธิบายความแตกต่าง\n\n"
    "2. ส่วนข้อดี (เต็ม 0.25 คะแนน):\n"
    "   - ให้ 0.25 คะแนน: มีการระบุข้อดี ได้พอเข้าใจ โดยให้อนุโลมปล่อยผ่านหากไม่ได้เข้าเกณฑ์อธิบายมั่ว\n"
    "   - ให้ 0.00 คะแนน: ไม่ได้ระบุข้อดี หรือระบุมั่ว/ผิดเพี้ยนสิ้นเชิง\n\n"
    "3. ส่วนข้อเสีย (เต็ม 0.25 คะแนน):\n"
    "   - ให้ 0.25 คะแนน: มีการระบุข้อเสีย ได้พอเข้าใจ โดยให้อนุโลมปล่อยผ่านหากไม่ได้เข้าเกณฑ์อธิบายมั่ว\n"
    "   - ให้ 0.00 คะแนน: ไม่ได้ระบุข้อเสีย (เช่น ใส่เครื่องหมายขีด - หรือละไว้) หรือระบุมั่ว/ผิดเพี้ยนสิ้นเชิง\n\n"
    "ผลรวมคะแนนจะเป็น: 0.00, 0.25, 0.50, 0.75 หรือ 1.00 คะแนน ตามผลรวมของทั้ง 3 ส่วน\n"
    "(ใน teacher_feedback ให้แจกแจงผลการตรวจทั้ง 3 ส่วนว่าส่วนใดได้/ไม่ได้คะแนนเพราะเหตุใด และ student_feedback ให้คำแนะนำการปรับปรุง)"
)

RUBRICS = [
    {
        "name": USER_COMPONENT_RUBRIC_NAME,
        "score": 1.0,
        "description": USER_COMPONENT_RUBRIC_DESC,
    }
]

async def test_pilot_10():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["ชุดข้อสอบ_dataset"]
    
    first10 = []
    for r in range(74, 84):
        sid = ws.cell(r, 1).value
        ans = str(ws.cell(r, 6).value or "")
        h = float(ws.cell(r, 7).value or 0.0)
        first10.append({"sid": sid, "ans": ans, "human": h})
        
    print(f"=== ทดสอบเกณฑ์ใหม่ของคุณ (0.50 ความแตกต่าง + 0.25 ข้อดี + 0.25 ข้อเสีย) บน 10 ข้อแรก ===")
    
    sem = asyncio.Semaphore(5)
    async def grade_one(item):
        async with sem:
            res = await score_with_openai(
                question_text=Q3_QUESTION_TEXT,
                answer_text=item["ans"],
                max_score=1.0,
                answer_key=Q3_ANSWER_KEY,
                rubrics=RUBRICS,
                strict_rubric_enforcement=True,
            )
            score = float(res.get("score", 0.0))
            diff = round(score - item["human"], 2)
            match = (score == item["human"])
            status = "🟢 ตรงเป๊ะ" if match else f"🟠 ต่าง ({diff:+0.2f})"
            print(f"{item['sid']} | อาจารย์: {item['human']:.2f} | AI: {score:.2f} | {status}")
            return {
                "sid": item["sid"],
                "human": item["human"],
                "ai": score,
                "diff": diff,
                "match": match,
                "tf": res.get("teacher_feedback", ""),
                "sf": res.get("student_feedback", "")
            }

    tasks = [grade_one(x) for x in first10]
    results = await asyncio.gather(*tasks)
    
    exact = sum(1 for x in results if x["match"])
    print(f"\nผลลัพธ์ 10 ข้อแรก: ตรงเป๊ะ {exact}/10 ({exact*10}%)")

if __name__ == "__main__":
    asyncio.run(test_pilot_10())
