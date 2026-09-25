import asyncio
import json
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q3_QUESTION_TEXT = "การใช้ลิ้งค์ลิสต์เป็นสแตกและคิวจะแตกต่างจากการใช้อาร์เรย์อย่างไร และมีข้อดีข้อเสียอย่างไร (1 คะแนน)"

Q3_ANSWER_KEY = (
    "เฉลยและแนวคำตอบมาตรฐาน:\n"
    "1. ความแตกต่าง (Difference):\n"
    "   - Array: มีขนาดคงที่ (Fixed-size / Static) ต้องจองพื้นที่ล่วงหน้า การทำ Stack/Queue ต้องคุม index (top, front, rear)\n"
    "   - Linked List: มีขนาดปรับเปลี่ยนได้ตามจริง (Dynamic / ยืดหยุ่น) ไม่ต้องกำหนดขนาดล่วงหน้า จัดการผ่าน Node และ Pointer (head, tail)\n"
    "2. ข้อดี-ข้อเสีย (Pros & Cons):\n"
    "   - Linked List: ข้อดีคือไม่จำกัดขนาด ไม่เกิด Stack Overflow ได้ง่าย เพิ่ม/ลดโหนดได้สะดวก; "
    "ข้อเสียคือเข้าถึงข้อมูลช้ากว่า (Sequential access) และเปลืองหน่วยความจำเก็บพอยน์เตอร์\n"
    "   - Array: ข้อดีคือเข้าถึงข้อมูลได้รวดเร็วทันทีด้วย Index (O(1)); "
    "ข้อเสียคือขนาดคงที่เสี่ยงเกิด Overflow หากจองน้อยไป หรือเปลืองพื้นที่หากจองมากเกินไป"
)

# เกณฑ์กระชับ เป็นกลาง ฉบับมาตรฐาน (ไม่มีรหัสนิสิตและไม่มีข้อความเฉพาะบุคคล)
Q3_RUBRIC_NAME = "การเปรียบเทียบ Linked List vs Array สำหรับ Stack และ Queue พร้อมข้อดีข้อเสีย"

Q3_RUBRIC_DESC = (
    "เกณฑ์การประเมิน 5 ระดับ (คะแนนเต็ม 1.00):\n"
    "• 1.00 คะแนน (สมบูรณ์): ตอบครบ 2 ส่วนหลักอย่างชัดเจน ได้แก่ (1) อธิบายความแตกต่างของโครงสร้าง/ขนาด (Dynamic vs Static/Fixed หรือไม่ต้องประกาศขนาดล่วงหน้า) หรือเชื่อมโยงการทำงานของ Stack/Queue และ (2) ระบุทั้ง 'ข้อดี' และ 'ข้อเสีย' อย่างครบถ้วน (อนุโลมสำนวนภาษาตามระดับนักศึกษา ไม่ต้องเคร่งครัดศัพท์ pointer/index หากเนื้อหาสื่อข้อดีข้อเสียได้สมเหตุสมผล)\n"
    "• 0.75 คะแนน (ดีมาก): ตอบความแตกต่างได้ถูกต้อง แต่ระบุข้อดีหรือข้อเสียเพียงด้านเดียว (เช่น มีข้อดีแต่ไม่มีข้อเสีย) หรืออธิบายข้อดีข้อเสียกว้างๆ ในเชิงการเขียนโปรแกรม\n"
    "• 0.50 คะแนน (มาตรฐาน): ตอบถูกเพียงบางส่วน ได้แก่ (1) เปรียบเทียบโครงสร้างทั่วไปแต่ไม่ได้แจกแจงแยกข้อดีข้อเสีย, (2) เขียนรวมเป็นความเรียงเดียว, (3) ตอบเฉพาะข้อดีหรือข้อเสียเพียงด้านเดียวสั้นๆ, หรือ (4) อธิบายเฉพาะพฤติกรรม LIFO/FIFO ของ Stack/Queue โดยข้อดีข้อเสียไม่ชัดเจน\n"
    "• 0.25 คะแนน (เล็กน้อย): คำตอบสั้นมาก คลุมเครือ ไม่ระบุชื่อโครงสร้าง แต่มีคีย์เวิร์ดที่ถูกต้องตามหลักการเพียงจุดเดียว (เช่น ยืดหยุ่น หรือ static)\n"
    "• 0.00 คะแนน (ไม่ได้คะแนน): ตอบไม่ตรงประเด็น (เช่น ช่วยประหยัดพลังงาน, รถไฟ), ตอบสั้นเพียงประโยคเดียวโดยไม่มีข้อดีข้อเสีย, ระบุข้อดีข้อเสียที่ผิดหลักการอย่างมีนัยสำคัญ, หรือไม่ตอบ\n\n"
    "(ต้องระบุ teacher_feedback แจกแจงจุดให้/หักคะแนนสำหรับผู้สอน และ student_feedback เป็นคำแนะนำพัฒนาการเรียนรู้สำหรับนักเรียน)"
)

Q3_RUBRICS = [
    {
        "name": Q3_RUBRIC_NAME,
        "score": 1.0,
        "description": Q3_RUBRIC_DESC,
    }
]

def calculate_qwk(y_true, y_pred, step=0.25, max_score=1.0):
    if len(y_true) != len(y_pred) or len(y_true) == 0: return 0.0
    if y_true == y_pred: return 1.0
    k = int(round(max_score / step)) + 1
    w = [[((i - j) ** 2) / ((k - 1) ** 2) for j in range(k)] for i in range(k)]
    cat_true = [min(k-1, max(0, int(round(v / step)))) for v in y_true]
    cat_pred = [min(k-1, max(0, int(round(v / step)))) for v in y_pred]
    o = [[0] * k for _ in range(k)]
    for t, p in zip(cat_true, cat_pred): o[t][p] += 1
    r_t = [sum(o[i][j] for j in range(k)) for i in range(k)]
    r_p = [sum(o[i][j] for i in range(k)) for j in range(k)]
    n = len(y_true)
    e = [[(r_t[i] * r_p[j]) / n for j in range(k)] for i in range(k)]
    num = sum(w[i][j] * o[i][j] for i in range(k) for j in range(k))
    den = sum(w[i][j] * e[i][j] for i in range(k) for j in range(k))
    return 1.0 - (num / den) if den != 0 else 1.0

def calc_pearson(x, y):
    n = len(x)
    if n == 0: return 0.0
    mx = sum(x) / n
    my = sum(y) / n
    cov = sum((a - mx) * (b - my) for a, b in zip(x, y))
    vx = sum((a - mx) ** 2 for a in x)
    vy = sum((b - my) ** 2 for b in y)
    if vx == 0 or vy == 0: return 0.0
    return cov / ((vx * vy) ** 0.5)

async def main():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    excel_ai_rubrics = ROOT / "ชุดข้อสอบใหม่" / "เกณฑ์ตรวจสำหรับAI.xlsx"
    
    wb = openpyxl.load_workbook(excel_path)
    ws_data = wb["ชุดข้อสอบ_dataset"]
    
    # 1. โหลดผลตรวจ 10 คนแรกที่เพิ่งตรวจไป
    first10_file = ROOT / "artifacts" / "q3_first10_concise_test_results.json"
    all_results = []
    if first10_file.exists():
        with open(first10_file, "r", encoding="utf-8") as f:
            all_results = json.load(f)
        print(f"โหลดผลตรวจ 10 คนแรกเดิมเรียบร้อย ({len(all_results)} ข้อ)")
    
    # 2. ดึง 24 คนที่เหลือ: แถว 84 ถึง 107 (DS-079 ถึง DS-102)
    remaining_24 = []
    for r in range(84, 108):
        sid = ws_data.cell(r, 1).value
        q_no = ws_data.cell(r, 2).value
        assert q_no == 3
        ans = str(ws_data.cell(r, 6).value or "")
        h = float(ws_data.cell(r, 7).value or 0.0)
        old_ai = float(ws_data.cell(r, 8).value or 0.0)
        remaining_24.append({
            "row": r,
            "sample_id": sid,
            "student_ans": ans,
            "human_score": h,
            "old_ai_score": old_ai
        })
        
    print(f"\n=== กำลังส่งตรวจ 24 ข้อที่เหลือ (DS-079 ถึง DS-102) ===")
    print(f"จำนวนที่ส่งตรวจ: {len(remaining_24)} ข้อ\n")
    
    sem = asyncio.Semaphore(6)
    
    async def grade_one(item):
        async with sem:
            res = await score_with_openai(
                question_text=Q3_QUESTION_TEXT,
                answer_text=item["student_ans"],
                max_score=1.0,
                answer_key=Q3_ANSWER_KEY,
                rubrics=Q3_RUBRICS,
                image_bytes_list=None,
                strict_rubric_enforcement=True,
            )
            new_ai = float(res.get("score", 0.0))
            diff = round(new_ai - item["human_score"], 2)
            match = (new_ai == item["human_score"])
            
            tf = res.get("teacher_feedback", "").strip()
            sf = res.get("student_feedback", "").strip()
            
            res_item = {
                "row": item["row"],
                "sample_id": item["sample_id"],
                "student_ans": item["student_ans"],
                "human_score": item["human_score"],
                "old_ai_score": item["old_ai_score"],
                "new_ai_score": new_ai,
                "confidence": res.get("confidence", "high"),
                "diff": diff,
                "is_exact": match,
                "teacher_feedback": tf,
                "student_feedback": sf
            }
            status = "🟢 ตรงเป๊ะ" if match else f"🟠 ต่าง ({diff:+0.2f})"
            print(f"{item['sample_id']} | อาจารย์: {item['human_score']:.2f} | AI: {new_ai:.2f} | {status}")
            return res_item

    tasks = [grade_one(item) for item in remaining_24]
    new_results = await asyncio.gather(*tasks)
    all_results.extend(new_results)
    all_results = sorted(all_results, key=lambda x: x["row"])
    
    print("\n" + "=" * 85)
    print(f"{'Sample ID':<10} | {'คะแนนอาจารย์':<12} | {'คะแนน AI':<10} | {'ผลต่าง':<8} | {'สถานะ':<10}")
    print("-" * 85)
    for r in all_results:
        status = "🟢 ตรงเป๊ะ" if r["is_exact"] else f"🟠 ต่าง ({r['diff']:+0.2f})"
        print(f"{r['sample_id']:<10} | {r['human_score']:<12.2f} | {r['new_ai_score']:<10.2f} | {r['diff']:<+8.2f} | {status}")
    print("=" * 85)
    
    # สถิติภาพรวม 34 คน
    y_true = [x["human_score"] for x in all_results]
    y_pred = [x["new_ai_score"] for x in all_results]
    exact = sum(1 for x in all_results if x["is_exact"])
    w025 = sum(1 for x in all_results if abs(x["diff"]) <= 0.2501)
    w050 = sum(1 for x in all_results if abs(x["diff"]) <= 0.5001)
    mae = sum(abs(x["diff"]) for x in all_results) / len(all_results)
    qwk = calculate_qwk(y_true, y_pred)
    pearson_r = calc_pearson(y_true, y_pred)
    
    print("\n📊 สรุปผลการประเมิน ข้อ 3 ทั้งหมด 34 คน ด้วยเกณฑ์กระชับฉบับมาตรฐาน:")
    print(f"• ตรงกันเป๊ะ (Exact Match): {exact} / 34 ({exact / 34 * 100:.1f}%)")
    print(f"• ต่างไม่เกิน ±0.25 pt:       {w025} / 34 ({w025 / 34 * 100:.1f}%)")
    print(f"• ต่างไม่เกิน ±0.50 pt:       {w050} / 34 ({w050 / 34 * 100:.1f}%)")
    print(f"• MAE เฉลี่ย:                {mae:.4f} คะแนน")
    print(f"• Pearson Correlation (r):   {pearson_r:.4f}")
    print(f"• QWK (Quadratic Weighted):  {qwk:.4f}")
    print("=" * 85)
    
    # 3. บันทึกผลลง Excel: ชุดข้อสอบ_dataset.xlsx
    for res in all_results:
        r = res["row"]
        ws_data.cell(r, 8, res["new_ai_score"])
        ws_data.cell(r, 9, res["confidence"])
        
        tf = res.get("teacher_feedback", "").strip()
        sf = res.get("student_feedback", "").strip()
        combined_fb = f"[สำหรับผู้สอน]\n{tf}\n\n[สำหรับนักเรียน]\n{sf}" if (tf and sf) else (tf or sf)
        ws_data.cell(r, 10, combined_fb)
        
    print(f"อัปเดตแถว 74-107 ในชีต 'ชุดข้อสอบ_dataset' เรียบร้อย")
    
    # อัปเดตชีต Exam_Rubrics
    if "Exam_Rubrics" in wb.sheetnames:
        ws_rubric = wb["Exam_Rubrics"]
        for r in range(4, ws_rubric.max_row + 1):
            q_val = ws_rubric.cell(r, 1).value
            if q_val == 3:
                ws_rubric.cell(r, 4, Q3_RUBRIC_NAME)
                ws_rubric.cell(r, 5, 1.0)
                ws_rubric.cell(r, 6, Q3_RUBRIC_DESC)
        print("อัปเดตเกณฑ์ข้อ 3 ในชีต 'Exam_Rubrics' เรียบร้อย")
        
    wb.save(excel_path)
    print(f"บันทึกไฟล์ Excel สำเร็จ: {excel_path.name}")
    
    # อัปเดตเกณฑ์ตรวจสำหรับAI.xlsx
    if excel_ai_rubrics.exists():
        wb_ai = openpyxl.load_workbook(excel_ai_rubrics)
        if "ข้อมูลสำหรับ API" in wb_ai.sheetnames:
            ws_ai = wb_ai["ข้อมูลสำหรับ API"]
            updated = 0
            for r in range(2, ws_ai.max_row + 1):
                q_text = str(ws_ai.cell(r, 1).value or "")
                if "ลิ้งค์ลิสต์" in q_text or "สแตกและคิว" in q_text:
                    ws_ai.cell(r, 2, Q3_RUBRIC_DESC)
                    updated += 1
            wb_ai.save(excel_ai_rubrics)
            print(f"อัปเดต {updated} แถวใน 'เกณฑ์ตรวจสำหรับAI.xlsx'")
            
    # บันทึก JSON artifact
    out_file = ROOT / "artifacts" / "q3_all34_concise_final_results.json"
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(all_results, f, ensure_ascii=False, indent=2)
    print(f"บันทึกไฟล์ JSON สรุป 34 คนลงใน {out_file.name}")

if __name__ == "__main__":
    asyncio.run(main())
