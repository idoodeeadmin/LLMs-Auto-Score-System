"""Apply a complete, successful Q1 grading report without losing workbook images."""

import hashlib
import json
import os
import sys
import zipfile
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx"
REPORT = Path(sys.argv[1]).resolve()
data = json.loads(REPORT.read_text(encoding="utf-8"))
assert data["source_sha256"] == hashlib.sha256(BOOK.read_bytes()).hexdigest(), "Workbook changed since grading"
results = data["results"]
assert len(results) == 34 and len({item["sample_id"] for item in results}) == 34
assert all(item["success"] and not item["result"]["metrics"]["manual_review_required"] for item in results)

workbook = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
sheet = workbook["ชุดข้อสอบ_dataset"]
updates = {}
for item in results:
    number = int(item["sample_id"].split("-")[1])
    row = number + 5
    assert 1 <= number <= 34 and sheet.cell(row, 1).value == item["sample_id"]
    assert sheet.cell(row, 6).value == item["answer"]
    assert sheet.cell(row, 7).value == item["human_score"]
    result = item["result"]
    score = float(result["score"])
    assert score in (0.0, 1.0, 2.0)
    assert result["confidence"] in ("high", "medium", "low")
    assert result["teacher_feedback"].strip() and result["student_feedback"].strip()
    updates[row] = (score, result["confidence"], result["feedback"])
workbook.close()

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ET.register_namespace("", NS)
with zipfile.ZipFile(BOOK) as source:
    root = ET.fromstring(source.read("xl/worksheets/sheet1.xml"))
    cells = {cell.attrib["r"]: cell for cell in root.iter(f"{{{NS}}}c")}
    for row, (score, confidence, feedback) in updates.items():
        for column, value in (("H", score), ("I", confidence), ("J", feedback)):
            cell = cells[f"{column}{row}"]
            for child in list(cell):
                if child.tag in {f"{{{NS}}}v", f"{{{NS}}}is", f"{{{NS}}}f"}:
                    cell.remove(child)
            if column == "H":
                cell.attrib.pop("t", None)
                ET.SubElement(cell, f"{{{NS}}}v").text = str(value)
            else:
                cell.attrib["t"] = "inlineStr"
                inline = ET.SubElement(cell, f"{{{NS}}}is")
                text = ET.SubElement(inline, f"{{{NS}}}t")
                text.attrib["{http://www.w3.org/XML/1998/namespace}space"] = "preserve"
                text.text = value
    xml = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with NamedTemporaryFile(dir=BOOK.parent, suffix=".xlsx", delete=False) as temporary:
        temporary_name = temporary.name
    with zipfile.ZipFile(temporary_name, "w") as target:
        for entry in source.infolist():
            target.writestr(entry, xml if entry.filename == "xl/worksheets/sheet1.xml" else source.read(entry.filename))

try:
    with zipfile.ZipFile(temporary_name) as archive:
        assert len([name for name in archive.namelist() if name.startswith("xl/media/")]) == 34
    check = openpyxl.load_workbook(temporary_name, read_only=True, data_only=True)
    try:
        sheet = check["ชุดข้อสอบ_dataset"]
        assert all(tuple(sheet.cell(row, column).value for column in (8, 9, 10)) == values
                   for row, values in updates.items())
    finally:
        check.close()
    os.replace(temporary_name, BOOK)
finally:
    if os.path.exists(temporary_name):
        os.unlink(temporary_name)

print(f"Updated {len(updates)} Q1 AI results in {BOOK}")
