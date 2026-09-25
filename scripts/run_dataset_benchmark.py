import os
import sys
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
sys.path.insert(0, os.path.abspath("."))
import json
import hashlib
from pathlib import Path
import asyncio
from typing import List, Dict, Any, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from dotenv import load_dotenv

sys.stdout.reconfigure(encoding='utf-8')
load_dotenv()

from server.services.openai_grading import score_with_openai

DATASET_EXCEL = os.path.join("ชุดข้อสอบใหม่", "ชุดข้อสอบ_dataset.xlsx")
_rubric_file = Path(__file__).resolve().parents[1] / 'ชุดข้อสอบใหม่' / 'เกณฑ์ตรวจสำหรับAI.xlsx'
_rubric_hash = hashlib.sha256(_rubric_file.read_bytes() + os.getenv('OPENAI_MODEL', 'gpt-5.6-luna').encode()).hexdigest()[:12]
CACHE_FILE = os.path.join("artifacts", f"benchmark_cache_rubric_{_rubric_hash}.json")
REPORT_EXCEL = os.path.join("artifacts", "LLM_AutoScore_Benchmark_204_Final_Report.xlsx")

EXAM_QUESTIONS = {
    1: {
        "question_no": 1,
        "type": "text",
        "topic": "Row-major vs Column-major",
        "max_score": 2.0,
        "question_text": "อธิบายความต่างของ Row-major vs Column-major",
        "answer_key": "Row-major คือการจัดเก็บข้อมูล ลำดับการเรียง หรือคำนวณตำแหน่ง address โดยอิงตามแถว (Row หรือแนวนอน/แกน X) ส่วน Column-major คือการจัดเก็บข้อมูล ลำดับการเรียง หรือคำนวณตำแหน่ง address โดยอิงตามคอลัมน์ (Column หรือแนวตั้ง/แกน Y) ต่างกันที่ลำดับและมิติการเรียงข้อมูลในหน่วยความจำ",
        "rubrics": [
            {"name": "Row-major", "score": 1.0, "description": "อธิบาย Row-major ว่าเป็นการจัดเก็บข้อมูล ลำดับการเรียง หรือการหาตำแหน่ง address ตามแนวแถว (Row/แนวนอน/แกน X หรือรูปแบบดัชนี [i][j])"},
            {"name": "Column-major", "score": 1.0, "description": "อธิบาย Column-major ว่าเป็นการจัดเก็บข้อมูล ลำดับการเรียง หรือการหาตำแหน่ง address ตามแนวคอลัมน์ (Column/แนวตั้ง/แกน Y หรือรูปแบบดัชนี [j][i])"}
        ]
    },
    2: {
        "question_no": 2,
        "type": "text",
        "topic": "O(n log n) vs O(n^2) Complexity",
        "max_score": 2.0,
        "question_text": "อธิบายว่าทำไม O(n log n) เหมาะกับข้อมูลใหญ่กว่า O(n^2) และยกตัวอย่าง Algorithm",
        "answer_key": "O(n log n) มีอัตราการเติบโตของเวลาในการทำงาน (Growth rate) ช้ากว่า O(n^2) มากเมื่อข้อมูลมีขนาดใหญ่ขึ้น ทำให้ใช้เวลาประมวลผลน้อยกว่า ตัวอย่าง O(n log n) เช่น Merge Sort, Quick Sort, Heap Sort และ O(n^2) เช่น Bubble Sort, Selection Sort, Insertion Sort",
        "rubrics": [
            {
                "name": "การเปรียบเทียบ O(n log n) vs O(n^2) และตัวอย่าง Algorithm",
                "score": 2.0,
                "description": (
                    "ประเมินตามระดับคะแนนดังต่อไปนี้อย่างเคร่งครัด:\n"
                    "- 2.0 คะแนน: อธิบายได้ว่าทำไม O(n log n) ดีกว่าเมื่อข้อมูลใหญ่ เช่น เร็วกว่า, จำนวนรอบ/การทำงานโตช้ากว่า, n² โตเร็วมาก และ มีตัวอย่าง Algorithm หรืออธิบายละเอียดพอ\n"
                    "- 1.5 คะแนน: เข้าใจแก่นว่า O(n log n) มีประสิทธิภาพกว่า แต่คำอธิบายยังไม่ครบ/มีส่วนคลาดเคลื่อน/ไม่ยกตัวอย่าง หรือมีตัวอย่างแต่เหตุผลไม่แข็งมาก\n"
                    "- 1.0 คะแนน: รู้เพียงว่า O(n log n) “เร็วกว่า/ดีกว่า/ซ้ำซ้อนน้อยกว่า” แต่ไม่ได้อธิบายว่าทำไมอย่างชัดเจน หรือคำอธิบายคลุมเครือ\n"
                    "- 0.5 คะแนน: ระบุได้เพียงตัวอย่าง Algorithm ที่เกี่ยวข้องอย่างถูกต้อง หรือกล่าวถึงความซับซ้อนเพียงฝั่งเดียว โดยยังไม่สื่อชัดว่า O(n log n) เหมาะกับข้อมูลใหญ่กว่า O(n²) อย่างไร และไม่มีข้อความที่ขัดกับหลักการสำคัญ\n"
                    "- 0.0 คะแนน: ไม่สามารถอธิบายความสัมพันธ์ของ O(n log n) กับ O(n²) ได้อย่างมีสาระ หรือตอบผิดหลักการ\n"
                    "**ระดับคะแนนที่ให้ได้คือ 2.0, 1.5, 1.0, 0.5 หรือ 0.0 คะแนนเท่านั้น ห้ามให้คะแนนเป็นเศษทศนิยมอื่น**"
                )
            }
        ]
    },
    3: {
        "question_no": 3,
        "type": "text",
        "topic": "Linked List vs Array (Stack & Queue)",
        "max_score": 1.0,
        "question_text": "การใช้ลิ้งค์ลิสต์เป็นสแตกและคิวจะแตกต่างจากการใช้อาร์เรย์อย่างไร และมีข้อดีข้อเสียอย่างไร",
        "answer_key": "ความแตกต่าง: Array มีขนาดคงที่ (Fixed size) จองพื้นที่ต่อเนื่อง ส่วน Linked List มีขนาดปรับเปลี่ยนได้แบบพลวัต (Dynamic size) ใช้พอยน์เตอร์ชี้โหนดถัดไป\nข้อดีข้อเสีย: Array เข้าถึงข้อมูลเร็วแบบ O(1) แต่เสี่ยงต่อ Overflow/ขยายขนาดยาก; Linked List ยืดหยุ่นไม่จำกัดขนาด ไม่เกิด overflow แต่ใช้เนื้อที่เพิ่มสำหรับ Pointer และการเข้าถึงข้อมูลช้ากว่า",
        "rubrics": [
            {
                "name": "ความแตกต่างเชิงโครงสร้าง (Comparison)",
                "score": 0.5,
                "description": (
                    "อธิบายความต่างระหว่าง Array และ Linked List ในการนำมาทำเป็น Stack/Queue:\n"
                    "- 0.50 คะแนน: อธิบายชัดเจนว่า Array มีขนาดคงที่ (Fixed size) หรือจองพื้นที่ต่อเนื่อง ส่วน Linked List มีขนาดปรับเปลี่ยนได้ (Dynamic size) หรือใช้พอยน์เตอร์เชื่อมโหนด\n"
                    "- 0.25 คะแนน: อธิบายถูกเพียงบางส่วน เช่น ตอบแค่ว่า Array ต้องระบุขนาด หรือบอกเรื่อง index แต่ไม่อธิบาย Linked List\n"
                    "- 0.00 คะแนน: ไม่ได้เปรียบเทียบ หรือตอบผิดหลักการ\n"
                    "**เกณฑ์นี้ให้คะแนนเป็น 0.50, 0.25 หรือ 0.00 คะแนนเท่านั้น**"
                )
            },
            {
                "name": "ข้อดีและข้อเสีย (Pros & Cons)",
                "score": 0.5,
                "description": (
                    "ระบุข้อดีและข้อเสียของการใช้ Linked List เทียบกับ Array:\n"
                    "- 0.50 คะแนน: มีทั้งข้อดีและข้อเสีย เช่น Linked List ไม่จำกัดขนาด/ไม่เกิด overflow แต่เข้าถึงข้อมูลช้ากว่า/เปลืองเนื้อที่ pointer หรือ Array เข้าถึงเร็ว O(1) แต่เสี่ยง overflow/จำกัดขนาด\n"
                    "- 0.25 คะแนน: ระบุเฉพาะข้อดีอย่างเดียว หรือเฉพาะข้อเสียอย่างเดียว หรือข้อดีข้อเสียยังไม่ชัดเจน\n"
                    "- 0.00 คะแนน: ไม่ได้ระบุข้อดีข้อเสีย หรือระบุผิด\n"
                    "**เกณฑ์นี้ให้คะแนนเป็น 0.50, 0.25 หรือ 0.00 คะแนนเท่านั้น**"
                )
            }
        ]
    },
    4: {
        "question_no": 4,
        "type": "img",
        "topic": "Binary Search Tree Construction",
        "max_score": 1.0,
        "question_text": "จากข้อมูลต่อไปนี้จงนำไปสร้างเป็น Binary search tree: 9 16 10 76 5 13 58 92 11 15 80 99",
        "answer_key": "Binary Search Tree: รากคือ 9\n- กิ่งซ้ายของ 9: 5\n- กิ่งขวาของ 9: 16\n  - กิ่งซ้ายของ 16: 10 (กิ่งขวาของ 10: 13 -> กิ่งซ้ายของ 13: 11, กิ่งขวาของ 13: 15)\n  - กิ่งขวาของ 16: 76 (กิ่งซ้ายของ 76: 58, กิ่งขวาของ 76: 92 -> กิ่งซ้ายของ 92: 80, กิ่งขวาของ 92: 99)",
        "rubrics": [
            {"name": "โครงสร้าง Binary Search Tree", "score": 1.0, "description": "สร้าง Binary Search Tree จากข้อมูล 9, 16, 10, 76, 5, 13, 58, 92, 11, 15, 80, 99 ได้ถูกต้องครบทั้ง 12 โหนดตามหลัก BST โดยให้อนุโลมลายมือ รอยร่าง หรือความยาวกิ่ง (เกณฑ์ให้คะแนนเฉพาะ 1.0 หรือ 0.0 เท่านั้น)"}
        ]
    },
    5: {
        "question_no": 5,
        "type": "img",
        "topic": "1D Array Representation of BST",
        "max_score": 1.0,
        "question_text": "จากทรีในข้อ 4 นิสิตสามารถนำทรีนี้มาจากโดยใช้ Array 1 มิติ ได้ จงแสดงค่าตัวเลขใน Array (Array ไม่พอ เพิ่มเติมได้)",
        "answer_key": "การแทนที่ Tree ใน Array 1 มิติ (เมื่อ index เริ่มต้นที่ 1): root อยู่ที่ 1, left child อยู่ที่ 2i, right child อยู่ที่ 2i+1\nแสดงตำแหน่งข้อมูลตัวเลขลงในช่อง Array ถูกต้องตามตำแหน่งโหนด",
        "rubrics": [
            {"name": "การแทนค่าใน Array 1 มิติ", "score": 1.0, "description": "แสดงค่าตัวเลขลงในช่อง Array 1 มิติ ได้ถูกต้องตามตำแหน่งของโครงสร้าง Tree ในข้อ 4"}
        ]
    },
    6: {
        "question_no": 6,
        "type": "img",
        "topic": "General Tree to Binary Tree Conversion",
        "max_score": 1.0,
        "question_text": "จงแปลง tree ต่อไปนี้ให้เป็น Binary Tree (ดูรูปภาพโจทย์ประกอบ)",
        "answer_key": "การแปลง General Tree เป็น Binary Tree ด้วยหลักการ Left-Child Right-Sibling (ลูกคนแรกเป็นกิ่งซ้าย, พี่น้องลำดับถัดไปเป็นกิ่งขวา)",
        "rubrics": [
            {"name": "การแปลง Tree เป็น Binary Tree", "score": 1.0, "description": "แปลงโครงสร้างตามหลักการ Left-Child Right-Sibling ได้ถูกต้อง"}
        ]
    }
}

def load_dataset() -> List[Dict[str, Any]]:
    wb = openpyxl.load_workbook(DATASET_EXCEL, data_only=True)
    ws = wb['ชุดข้อสอบ_dataset']
    items = []
    
    # Q6 question image
    q6_question_img = os.path.join("ชุดข้อสอบใหม่", "โจทphoto3", "LINE_ALBUM_โจทphoto6_260918_1.jpg")

    for r in range(6, ws.max_row + 1):
        sid = ws.cell(row=r, column=1).value
        if not sid:
            continue
        q_no = int(ws.cell(row=r, column=2).value)
        q_type = str(ws.cell(row=r, column=3).value or "text")
        q_content = ws.cell(row=r, column=4).value or EXAM_QUESTIONS[q_no]["question_text"]
        ans_type = str(ws.cell(row=r, column=5).value or "text")
        std_ans = ws.cell(row=r, column=6).value or ""
        h_score = float(ws.cell(row=r, column=7).value or 0.0)
        
        # Determine student index (1-34)
        std_idx = ((r - 6) % 34) + 1
        
        # Only cleaned answer sheets are valid blind-test inputs; masks are not answers.
        img_path = None
        if ans_type == "img":
            if q_no == 4:
                img_path = os.path.join("ชุดข้อสอบใหม่", "photo_clean_ชุดที่1", f"LINE_ALBUM_Photo1_260917_{std_idx}.jpg")
            elif q_no == 5:
                img_path = os.path.join("ชุดข้อสอบใหม่", "photo_clean_ชุดที่2", f"LINE_ALBUM_Photo2.1_260918_{std_idx}.jpg")
            elif q_no == 6:
                img_path = os.path.join("ชุดข้อสอบใหม่", "photo_clean_ชุดที่3", f"LINE_ALBUM_Photo2.2_260918_{std_idx}.jpg")
            if not img_path or not os.path.isfile(img_path):
                raise FileNotFoundError(f"Cleaned answer image missing for {sid}: {img_path}")

        q_img_path = q6_question_img if q_no == 6 else None

        items.append({
            "row": r,
            "sample_id": sid,
            "student_idx": std_idx,
            "question_no": q_no,
            "question_type": q_type,
            "question_content": q_content,
            "answer_type": ans_type,
            "student_answer": std_ans,
            "human_score": h_score,
            "student_img_path": img_path,
            "question_img_path": q_img_path
        })
    return items

def calculate_qwk(y_true: List[float], y_pred: List[float], step: float = 0.25, max_score: float = 2.0) -> float:
    if len(y_true) != len(y_pred) or len(y_true) == 0:
        return 0.0
    if y_true == y_pred:
        return 1.0

    k = int(round(max_score / step)) + 1
    w = [[((i - j) ** 2) / ((k - 1) ** 2) for j in range(k)] for i in range(k)]

    O = [[0] * k for _ in range(k)]
    hist_true = [0] * k
    hist_pred = [0] * k

    for t, p in zip(y_true, y_pred):
        ti = max(0, min(k - 1, int(round(t / step))))
        pi = max(0, min(k - 1, int(round(p / step))))
        O[ti][pi] += 1
        hist_true[ti] += 1
        hist_pred[pi] += 1

    n = len(y_true)
    E = [[(hist_true[i] * hist_pred[j]) / n for j in range(k)] for i in range(k)]

    num = sum(w[i][j] * O[i][j] for i in range(k) for j in range(k))
    den = sum(w[i][j] * E[i][j] for i in range(k) for j in range(k))

    if den == 0:
        return 1.0 if num == 0 else 0.0
    return round(1.0 - (num / den), 4)

def calculate_mae(y_true: List[float], y_pred: List[float]) -> float:
    if not y_true or len(y_true) != len(y_pred):
        return 0.0
    return round(sum(abs(t - p) for t, p in zip(y_true, y_pred)) / len(y_true), 4)

def calculate_confusion_matrix(y_true: List[float], y_pred: List[float], labels: List[float]) -> Dict[str, Any]:
    grid = [[0] * len(labels) for _ in range(len(labels))]
    label_to_idx = {round(l, 2): idx for idx, l in enumerate(labels)}

    for t, p in zip(y_true, y_pred):
        # find closest label
        t_closest = min(labels, key=lambda x: abs(x - t))
        p_closest = min(labels, key=lambda x: abs(x - p))
        ti = label_to_idx[round(t_closest, 2)]
        pi = label_to_idx[round(p_closest, 2)]
        grid[ti][pi] += 1

    return {
        "labels": labels,
        "matrix": grid,
        "totals_human": [sum(grid[r]) for r in range(len(labels))],
        "totals_ai": [sum(grid[r][c] for r in range(len(labels))) for c in range(len(labels))]
    }

def get_agreement_interpretation(qwk: float) -> Dict[str, str]:
    if qwk < 0.0:
        return {"level_en": "Poor", "level_th": "ความสอดคล้องต่ำมาก (Poor Agreement)", "badge_color": "red"}
    elif qwk <= 0.20:
        return {"level_en": "Slight", "level_th": "สอดคล้องกันเล็กน้อย (Slight Agreement)", "badge_color": "amber"}
    elif qwk <= 0.40:
        return {"level_en": "Fair", "level_th": "สอดคล้องกันพอใช้ (Fair Agreement)", "badge_color": "yellow"}
    elif qwk <= 0.60:
        return {"level_en": "Moderate", "level_th": "สอดคล้องกันปานกลาง (Moderate Agreement)", "badge_color": "blue"}
    elif qwk <= 0.80:
        return {"level_en": "Substantial", "level_th": "สอดคล้องกันมาก (Substantial Agreement)", "badge_color": "indigo"}
    else:
        return {"level_en": "Almost Perfect", "level_th": "สอดคล้องกันเกือบสมบูรณ์แบบ (Almost Perfect Agreement)", "badge_color": "emerald"}

async def grade_single_item(item: Dict[str, Any], sem: asyncio.Semaphore) -> Dict[str, Any]:
    async with sem:
        q_meta = EXAM_QUESTIONS[item["question_no"]]
        q_text = item["question_content"]
        ans_text = item["student_answer"]
        max_score = q_meta["max_score"]
        ans_key = q_meta["answer_key"]
        from scripts.benchmark_rubrics import load_proposed_rubric
        rubrics, score_step = load_proposed_rubric(item["question_no"], max_score, q_text)

        # Read images if applicable
        img_bytes_list = None
        img_mime_list = None
        if item["student_img_path"] and os.path.exists(item["student_img_path"]):
            with open(item["student_img_path"], "rb") as f:
                img_bytes_list = [f.read()]
            img_mime_list = ["image/jpeg"]

        q_img_bytes_list = None
        q_img_mime_list = None
        if item["question_img_path"] and os.path.exists(item["question_img_path"]):
            with open(item["question_img_path"], "rb") as f:
                q_img_bytes_list = [f.read()]
            q_img_mime_list = ["image/jpeg"]

        result = await score_with_openai(
            question_text=q_text,
            answer_text=ans_text,
            max_score=max_score,
            answer_key=ans_key,
            rubrics=rubrics,
            image_bytes_list=img_bytes_list,
            image_mime_list=img_mime_list,
            q_image_bytes_list=q_img_bytes_list,
            q_image_mime_list=q_img_mime_list,
            score_step=score_step,
            allowed_scores=[0, 1, 2] if item["question_no"] == 1 else ([0, 0.5, 1, 1.5, 2] if item["question_no"] == 2 else ([0.0, 0.25, 0.5, 0.75, 1.0] if item["question_no"] == 3 else ([0, 1] if item["question_no"] == 4 else None))),
        )

        return {
            "sample_id": item["sample_id"],
            "row": item["row"],
            "question_no": item["question_no"],
            "answer_type": item["answer_type"],
            "student_idx": item["student_idx"],
            "human_score": item["human_score"],
            "ai_score": result.get("score", 0.0),
            "ai_confidence": result.get("confidence", "medium"),
            "ai_feedback": result.get("feedback", ""),
            "transcription": result.get("transcription", ""),
            "student_img_path": item["student_img_path"]
        }

async def run_benchmark(max_items: Optional[int] = None, concurrency: int = 3):
    os.makedirs("artifacts", exist_ok=True)
    items = load_dataset()
    if max_items:
        items = items[:max_items]

    print(f"============================================================")
    print(f" Starting Dataset Benchmark Evaluation (Total: {len(items)} samples)")
    print(f" Concurrency: {concurrency} workers")
    print(f"============================================================")

    # Load cache if exists
    cache = {}
    if os.path.exists(CACHE_FILE):
        try:
            with open(CACHE_FILE, "r", encoding="utf-8") as f:
                cache = json.load(f)
            print(f"Loaded {len(cache)} existing cached results from {CACHE_FILE}")
        except Exception:
            cache = {}

    sem = asyncio.Semaphore(concurrency)
    results = []

    async def process_item(item, idx, total):
        sid = item["sample_id"]
        if sid in cache:
            cached_res = cache[sid]
            print(f"[{idx}/{total}] Cached: {sid} (Q{item['question_no']} std#{item['student_idx']}) Human={item['human_score']} AI={cached_res['ai_score']}")
            return cached_res

        print(f"[{idx}/{total}] Scoring: {sid} (Q{item['question_no']} std#{item['student_idx']} {item['answer_type']}) ...")
        res = await grade_single_item(item, sem)
        cache[sid] = res
        # Save cache periodically
        with open(CACHE_FILE, "w", encoding="utf-8") as f:
            json.dump(cache, f, ensure_ascii=False, indent=2)
        print(f"[{idx}/{total}] Done: {sid} -> AI Score: {res['ai_score']} (Human: {item['human_score']}, Conf: {res['ai_confidence']})")
        return res

    tasks = [process_item(item, i + 1, len(items)) for i, item in enumerate(items)]
    results = await asyncio.gather(*tasks)

    print("\nAll items graded! Calculating benchmark metrics...")
    
    # Update Excel dataset with AI results
    wb = openpyxl.load_workbook(DATASET_EXCEL)
    ws = wb['ชุดข้อสอบ_dataset']
    for res in results:
        r = res["row"]
        ws.cell(row=r, column=8, value=res["ai_score"])
        ws.cell(row=r, column=9, value=res["ai_confidence"])
        ws.cell(row=r, column=10, value=res["ai_feedback"])
        # If student answer was empty for image, write image filename
        if res.get("student_img_path") and not ws.cell(row=r, column=6).value:
            ws.cell(row=r, column=6, value=os.path.basename(res["student_img_path"]))
    wb.save(DATASET_EXCEL)
    print(f"Updated {DATASET_EXCEL} with AI scores, confidences, and feedback.")
    try:
        from scripts.format_dataset_excel import format_dataset_excel
        format_dataset_excel(DATASET_EXCEL)
    except Exception as fe:
        print(f"Format notice: {fe}")

    # Calculate overall metrics
    y_true = [r["human_score"] for r in results]
    y_pred = [r["ai_score"] for r in results]

    overall_qwk = calculate_qwk(y_true, y_pred, step=0.25, max_score=2.0)
    overall_mae = calculate_mae(y_true, y_pred)
    exact_count = sum(1 for t, p in zip(y_true, y_pred) if abs(t - p) < 1e-4)
    adj_count = sum(1 for t, p in zip(y_true, y_pred) if abs(t - p) <= 0.25 + 1e-4)
    exact_rate = round(exact_count / len(results) * 100, 2)
    adj_rate = round(adj_count / len(results) * 100, 2)
    agreement = get_agreement_interpretation(overall_qwk)

    # Modality comparison
    text_results = [r for r in results if r["answer_type"] == "text"]
    img_results = [r for r in results if r["answer_type"] == "img"]

    def calc_subset(sub):
        if not sub:
            return {"count": 0, "qwk": 0.0, "mae": 0.0, "exact_pct": 0.0, "adj_pct": 0.0, "agreement": get_agreement_interpretation(0.0)}
        st = [r["human_score"] for r in sub]
        sp = [r["ai_score"] for r in sub]
        q = calculate_qwk(st, sp, step=0.25, max_score=2.0)
        m = calculate_mae(st, sp)
        ex = sum(1 for t, p in zip(st, sp) if abs(t - p) < 1e-4)
        ad = sum(1 for t, p in zip(st, sp) if abs(t - p) <= 0.25 + 1e-4)
        return {
            "count": len(sub),
            "qwk": q,
            "mae": m,
            "exact_pct": round(ex / len(sub) * 100, 2),
            "adj_pct": round(ad / len(sub) * 100, 2),
            "agreement": get_agreement_interpretation(q)
        }

    text_metrics = calc_subset(text_results)
    img_metrics = calc_subset(img_results)

    # Per question breakdown
    per_q = []
    for q_no in sorted(EXAM_QUESTIONS.keys()):
        q_sub = [r for r in results if r["question_no"] == q_no]
        m = calc_subset(q_sub)
        m["question_no"] = q_no
        m["topic"] = EXAM_QUESTIONS[q_no]["topic"]
        m["max_score"] = EXAM_QUESTIONS[q_no]["max_score"]
        m["type"] = EXAM_QUESTIONS[q_no]["type"]
        per_q.append(m)

    # Confusion matrix on discrete labels: 0, 0.25, 0.5, 0.75, 1.0, 1.25, 1.5, 1.75, 2.0
    labels = [0.0, 0.25, 0.5, 0.75, 1.0, 1.25, 1.5, 1.75, 2.0]
    cm = calculate_confusion_matrix(y_true, y_pred, labels)

    summary_report = {
        "total_samples": len(results),
        "overall_qwk": overall_qwk,
        "overall_mae": overall_mae,
        "exact_agreement_rate": exact_rate,
        "adjacent_agreement_rate": adj_rate,
        "agreement": agreement,
        "modality_comparison": {
            "text": text_metrics,
            "image": img_metrics
        },
        "per_question": per_q,
        "confusion_matrix": cm
    }

    # Save summary JSON
    summary_path = os.path.join("artifacts", "benchmark_summary_metrics.json")
    with open(summary_path, "w", encoding="utf-8") as f:
        json.dump(summary_report, f, ensure_ascii=False, indent=2)

    # Generate styled Excel Report
    generate_excel_report(summary_report, results)

    print("\n============================================================")
    print("                 BENCHMARK EVALUATION SUMMARY                ")
    print("============================================================")
    print(f"Total Samples Evaluated: {len(results)}")
    print(f"Quadratic Weighted Kappa (QWK): {overall_qwk:.4f} ({agreement['level_th']})")
    print(f"Mean Absolute Error (MAE): {overall_mae:.4f} points")
    print(f"Exact Agreement Rate: {exact_rate}%")
    print(f"Adjacent Agreement Rate (±0.25): {adj_rate}%")
    print("------------------------------------------------------------")
    print(f"Text Modality (n={text_metrics['count']}): QWK={text_metrics['qwk']:.4f}, MAE={text_metrics['mae']:.4f}")
    print(f"Image Modality (n={img_metrics['count']}): QWK={img_metrics['qwk']:.4f}, MAE={img_metrics['mae']:.4f}")
    print("------------------------------------------------------------")
    for qm in per_q:
        print(f"Q{qm['question_no']} ({qm['type']}, max={qm['max_score']}): QWK={qm['qwk']:.4f}, MAE={qm['mae']:.4f}, Exact={qm['exact_pct']}% | {qm['topic']}")
    print("============================================================")
    print(f"Excel Report Saved: {REPORT_EXCEL}")
    print(f"Summary JSON Saved: {summary_path}")

    return summary_report

def generate_excel_report(summary: Dict[str, Any], results: List[Dict[str, Any]]):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    title_font = Font(name="Segoe UI", size=14, bold=True, color="1F4E79")
    subtitle_font = Font(name="Segoe UI", size=10, italic=True, color="595959")
    header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
    bold_font = Font(name="Segoe UI", size=10, bold=True)
    regular_font = Font(name="Segoe UI", size=10)
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    subhead_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    zebra_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    # 1. Summary Sheet
    ws_sum = wb.create_sheet("Benchmark_Summary")
    ws_sum["A1"] = "รายงานผลการประเมินความแม่นยำระบบตรวจข้อสอบอัตโนมัติ (Model Benchmark Evaluation)"
    ws_sum["A1"].font = title_font
    ws_sum["A2"] = "วิชา Data Structures | การวัดผลด้วย Quadratic Weighted Kappa (QWK), MAE และ Confusion Matrix"
    ws_sum["A2"].font = subtitle_font

    ws_sum["A4"] = "ตัวชี้วัดหลัก (Key Evaluation Indicators)"
    ws_sum["A4"].font = Font(name="Segoe UI", size=11, bold=True, color="1F4E79")

    metrics_rows = [
        ("จำนวนตัวอย่างทั้งหมด (Total Samples)", f"{summary['total_samples']} ตัวอย่าง (34 นิสิต × 6 ข้อ)"),
        ("Quadratic Weighted Kappa (QWK)", f"{summary['overall_qwk']:.4f}"),
        ("ระดับความสอดคล้อง (Interpretation)", f"{summary['agreement']['level_th']}"),
        ("Mean Absolute Error (MAE)", f"{summary['overall_mae']:.4f} คะแนน"),
        ("ความสอดคล้องสมบูรณ์ (Exact Agreement Rate)", f"{summary['exact_agreement_rate']:.2f}%"),
        ("ความสอดคล้องใกล้เคียง (Adjacent Agreement Rate ±0.25)", f"{summary['adjacent_agreement_rate']:.2f}%")
    ]
    for r_idx, (label, val) in enumerate(metrics_rows, 5):
        ws_sum.cell(row=r_idx, column=1, value=label).font = bold_font
        ws_sum.cell(row=r_idx, column=2, value=val).font = regular_font

    # Modality comparison table
    ws_sum["A12"] = "การเปรียบเทียบตามรูปแบบคำตอบ (Modality Comparison)"
    ws_sum["A12"].font = Font(name="Segoe UI", size=11, bold=True, color="1F4E79")
    mod_headers = ["Modality", "จำนวนคำตอบ", "QWK", "ระดับความสอดคล้อง", "MAE (คะแนน)", "Exact Match (%)", "Adjacent Match (%)"]
    for c_idx, h in enumerate(mod_headers, 1):
        c = ws_sum.cell(row=13, column=c_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    txt_m = summary["modality_comparison"]["text"]
    img_m = summary["modality_comparison"]["image"]
    mod_data = [
        ["Text-based (ข้อความพิมพ์)", txt_m["count"], f"{txt_m['qwk']:.4f}", txt_m["agreement"]["level_th"], f"{txt_m['mae']:.4f}", f"{txt_m['exact_pct']}%", f"{txt_m['adj_pct']}%"],
        ["Image-based (ลายมือ/แผนภาพ)", img_m["count"], f"{img_m['qwk']:.4f}", img_m["agreement"]["level_th"], f"{img_m['mae']:.4f}", f"{img_m['exact_pct']}%", f"{img_m['adj_pct']}%"]
    ]
    for r_idx, row_vals in enumerate(mod_data, 14):
        for c_idx, val in enumerate(row_vals, 1):
            c = ws_sum.cell(row=r_idx, column=c_idx, value=val)
            c.font = regular_font
            c.border = thin_border
            if c_idx in [2, 3, 5, 6, 7]:
                c.alignment = Alignment(horizontal="center", vertical="center")

    # Per Question table
    ws_sum["A18"] = "ผลการวิเคราะห์แยกรายข้อ (Per-Question Breakdown Q1–Q6)"
    ws_sum["A18"].font = Font(name="Segoe UI", size=11, bold=True, color="1F4E79")
    q_headers = ["ข้อที่", "หัวข้อโจทย์", "ประเภท", "คะแนนเต็ม", "จำนวน", "QWK", "ระดับความสอดคล้อง", "MAE", "Exact Match (%)"]
    for c_idx, h in enumerate(q_headers, 1):
        c = ws_sum.cell(row=19, column=c_idx, value=h)
        c.font = header_font
        c.fill = subhead_fill
        c.alignment = Alignment(horizontal="center", vertical="center")

    for r_idx, qm in enumerate(summary["per_question"], 20):
        q_row = [
            f"ข้อ {qm['question_no']}",
            qm["topic"],
            "ข้อความ" if qm["type"] == "text" else "รูปภาพ",
            qm["max_score"],
            qm["count"],
            f"{qm['qwk']:.4f}",
            qm["agreement"]["level_th"],
            f"{qm['mae']:.4f}",
            f"{qm['exact_pct']}%"
        ]
        for c_idx, val in enumerate(q_row, 1):
            c = ws_sum.cell(row=r_idx, column=c_idx, value=val)
            c.font = regular_font
            c.border = thin_border
            if c_idx in [1, 3, 4, 5, 6, 8, 9]:
                c.alignment = Alignment(horizontal="center", vertical="center")

    # Column widths
    ws_sum.column_dimensions["A"].width = 30
    ws_sum.column_dimensions["B"].width = 35
    ws_sum.column_dimensions["C"].width = 16
    ws_sum.column_dimensions["D"].width = 12
    ws_sum.column_dimensions["E"].width = 12
    ws_sum.column_dimensions["F"].width = 14
    ws_sum.column_dimensions["G"].width = 40
    ws_sum.column_dimensions["H"].width = 14
    ws_sum.column_dimensions["I"].width = 18

    # 2. Raw Dataset Sheet
    ws_data = wb.create_sheet("Dataset_QWK_204")
    ws_data.row_dimensions[1].height = 28.0
    data_headers = [
        ("sample_id", 14),
        ("question_no", 12),
        ("answer_type", 16),
        ("student_idx", 12),
        ("human_score", 14),
        ("ai_score", 14),
        ("difference", 14),
        ("ai_confidence", 16),
        ("student_img_path", 35),
        ("ai_feedback", 85),
        ("transcription", 45)
    ]
    for c_idx, (h, w) in enumerate(data_headers, 1):
        c = ws_data.cell(row=1, column=c_idx, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center")
        col_letter = get_column_letter(c_idx)
        ws_data.column_dimensions[col_letter].width = w

    for r_idx, r in enumerate(results, 2):
        is_zebra = (r_idx % 2 == 1)
        diff = round(abs(r["human_score"] - r["ai_score"]), 4)
        fb_text = str(r.get("ai_feedback") or "")
        tr_text = str(r.get("transcription") or "")
        
        # Calculate needed lines for row height
        max_lines = 1
        for text_val, col_w in [(fb_text, 85), (tr_text, 45)]:
            if text_val:
                lines = sum(max(1, math.ceil(len(p) / max(20, int(col_w * 0.85)))) for p in text_val.split("\n"))
                max_lines = max(max_lines, lines)
        ws_data.row_dimensions[r_idx].height = max(28.0, max_lines * 22.0 + 8.0)

        row_vals = [
            r["sample_id"], r["question_no"], r["answer_type"], r["student_idx"],
            r["human_score"], r["ai_score"], diff, r["ai_confidence"],
            r.get("student_img_path") or "", fb_text, tr_text
        ]
        for c_idx, val in enumerate(row_vals, 1):
            c = ws_data.cell(row=r_idx, column=c_idx, value=val)
            c.font = regular_font
            c.border = thin_border
            if is_zebra:
                c.fill = zebra_fill
            if c_idx in [1, 2, 3, 4, 5, 6, 7, 8]:
                c.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            else:
                c.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    wb.save(REPORT_EXCEL)
    print(f"Generated Excel report: {REPORT_EXCEL}")

if __name__ == "__main__":
    import argparse
    parser = argparse.ArgumentParser()
    parser.add_argument("--limit", type=int, default=None, help="Limit number of items to score")
    parser.add_argument("--workers", type=int, default=3, help="Number of concurrent scoring workers")
    args = parser.parse_args()

    asyncio.run(run_benchmark(max_items=args.limit, concurrency=args.workers))
