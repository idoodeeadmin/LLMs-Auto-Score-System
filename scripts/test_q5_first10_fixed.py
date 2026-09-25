import asyncio
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q5_RUBRIC = [
    {
        "name": "ความถูกต้องของการแทนค่า Binary Search Tree ใน Array 1 มิติ",
        "score": 1.0,
        "description": (
            "ประเมินเป็น 2 ส่วนแล้วนำคะแนนมารวมกันอย่างเคร่งครัด ห้ามให้คะแนนปลอบใจ:\n"
            "ส่วนที่ 1: ตรวจช่วง index 0–19 (คะแนนเต็ม 0.75 คะแนน)\n"
            "- ได้ 0.75 คะแนน: หากระบุตำแหน่ง 0=9, 1=5, 2=16, 5=10, 6=76, 12=13, 13=58, 14=92 และเว้นตำแหน่งอื่นว่าง ได้ถูกต้องครบถ้วนทั้งหมด "
            "(หรือหากนิสิตแปลงตำแหน่งสอดคล้องถูกต้องตามโครงสร้าง Tree ข้อ 4 ของตนเอง)\n"
            "- ได้ 0.25 คะแนน: หากระบุโครงสร้างหลักถูก แต่มีจุดผิดพลาดเพียง 1 ตำแหน่งพอดี\n"
            "- ได้ 0.00 คะแนน: หากผิดพลาดตั้งแต่ 2 ตำแหน่งขึ้นไป หรือเขียนตัวเลขเรียงติดกันในช่อง 0–11 โดยไม่เว้นตำแหน่งตามสูตร Tree ให้ 0 คะแนนในส่วนนี้ทันที\n"
            "ส่วนที่ 2: ตรวจช่วง index 20–30 (คะแนนเต็ม 0.25 คะแนน)\n"
            "- ได้เพิ่ม 0.25 คะแนน: หากระบุค่า 11, 15, 80, 99 (ที่ index 25, 26, 29, 30) และช่องว่างอื่นถูกต้องครบถ้วนทั้งหมด\n"
            "- ได้ 0.00 คะแนน: หากผิดพลาดแม้แต่ตำแหน่งเดียว หรือไม่มีการเขียนแสดงช่วง index 20–30"
        ),
    }
]

Q5_QUESTION_TEXT = (
    "จากทรีในข้อ 4 นิสิตสามารถนำทรีนี้มาจากโดยใช้ Array 1 มิติ ได้ จงแสดงค่าตัวเลขใน Array (Array ไม่พอ เพิ่มเติมได้)\n"
    "บริบทของ Tree ที่โจทย์ข้อ 4 อ้างถึง: จงสร้าง Binary Search Tree จากข้อมูล: 9, 16, 10, 76, 5, 13, 58, 92, 11, 15, 80, 99"
)

Q5_ANSWER_KEY = (
    "เฉลยการแทน Binary Search Tree ใน Array (เริ่ม index ที่ 0): "
    "0=9, 1=5, 2=16, 5=10, 6=76, 12=13, 13=58, 14=92, "
    "25=11, 26=15, 29=80, 30=99; ตำแหน่งอื่นว่าง. "
    "หากเริ่ม index ที่ 1 ให้เลื่อนตำแหน่งทั้งหมดเพิ่ม 1 "
    "(1=9, 2=5, 3=16, 6=10, 7=76, 13=13, 14=58, 15=92, 26=11, 27=15, 30=80, 31=99) แต่ค่ากับโครงสร้างต้องตรงกันทุกจุด. "
    "หมายเหตุ: นิสิตสามารถแปลงตามทรีที่ตนเองวาดไว้ในข้อ 4 ได้อย่างสมเหตุสมผล หากตำแหน่งลูกสอดคล้องกับสูตร 2i+1, 2i+2"
)

KEY_IMG_PATH = ROOT / "public" / "answer-keys" / "q5-array-bst-answer-key.png"
key_img_bytes = KEY_IMG_PATH.read_bytes() if KEY_IMG_PATH.exists() else None

async def main():
    wb = openpyxl.load_workbook(ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx")
    sheet = wb["ชุดข้อสอบ_dataset"]
    
    rows_to_test = list(range(142, 152))
    print(f"=== Grading Q5 First 10 Samples (Rows 142-151) via System API (Refined Rubric) ===")
    
    results = []
    
    for r in rows_to_test:
        sample_id = sheet.cell(r, 1).value
        human_score = float(sheet.cell(r, 7).value or 0.0)
        std_idx = r - 141
        img_name = f"LINE_ALBUM_Photo2.1_260918_{std_idx}.jpg"
        img_path = ROOT / "ชุดข้อสอบใหม่" / "photo_clean_ชุดที่2" / img_name
        
        img_bytes = img_path.read_bytes()
        
        print(f"Grading {sample_id} ({img_name}) ...", flush=True)
        res = await score_with_openai(
            question_text=Q5_QUESTION_TEXT,
            answer_text="",
            max_score=1.0,
            answer_key=Q5_ANSWER_KEY,
            rubrics=Q5_RUBRIC,
            image_bytes_list=[img_bytes],
            image_mime_list=["image/jpeg"],
            answer_key_image_bytes_list=[key_img_bytes] if key_img_bytes else None,
            answer_key_image_mime_list=["image/png"] if key_img_bytes else None,
            strict_rubric_enforcement=True,
        )
        
        ai_score = float(res.get("score", 0.0))
        confidence = res.get("confidence", "unknown")
        feedback = res.get("feedback", "")
        transcription = res.get("transcription", "")
        diff = round(ai_score - human_score, 2)
        match = (ai_score == human_score)
        
        item_res = {
            "row": r,
            "sample_id": sample_id,
            "img_name": img_name,
            "human_score": human_score,
            "ai_score": ai_score,
            "diff": diff,
            "match": match,
            "confidence": confidence,
            "feedback": feedback,
            "transcription": transcription,
        }
        results.append(item_res)
        print(f"  -> {sample_id}: Human={human_score:.2f} | AI={ai_score:.2f} | Diff={diff:+.2f} | Match={'MATCH' if match else 'DIFF'} | Conf={confidence}")
        
    print("\n" + "="*80)
    print("SUMMARY RESULTS (First 10 of Question 5 - Refined Rubric):")
    print("="*80)
    exact_matches = sum(1 for x in results if x["match"])
    print(f"Exact Match: {exact_matches}/{len(results)} ({exact_matches/len(results)*100:.1f}%)")
    
    print(f"\n{'Sample ID':<10} | {'Human':<6} | {'AI':<6} | {'Diff':<6} | {'Status':<8} | {'Feedback'}")
    print("-" * 110)
    for x in results:
        status = "EXACT" if x["match"] else f"{x['diff']:+.2f}"
        fb_short = (x['feedback'][:65] + "...") if len(x['feedback']) > 65 else x['feedback']
        print(f"{x['sample_id']:<10} | {x['human_score']:<6.2f} | {x['ai_score']:<6.2f} | {x['diff']:<+6.2f} | {status:<8} | {fb_short}")
        
    import json
    out_file = ROOT / "artifacts" / "q5_first10_refined_rubric_results.json"
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)
    print(f"\nSaved to {out_file}")

if __name__ == "__main__":
    asyncio.run(main())
