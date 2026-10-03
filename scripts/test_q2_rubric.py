import asyncio
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env")
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q2_RUBRIC_NAME = "ความซับซ้อน O(n log n) vs O(n^2) และตัวอย่างอัลกอริทึม"
Q2_RUBRIC_DESC = (
    "โจทย์: อธิบายว่าทำไม O(n log n) เหมาะกับข้อมูลใหญ่กว่า O(n^2) และยกตัวอย่าง Algorithm (คะแนนเต็ม 2.00 คะแนน)\n"
    "เกณฑ์การประเมินความถูกต้องตามหลักการวิเคราะห์ความซับซ้อน (Time Complexity) แบ่งเป็น 2 กรณีตามตัวอย่าง Algorithm:\n\n"
    "กรณี ก: มีการยกตัวอย่างชื่อ Algorithm ที่เกี่ยวข้อง (เช่น Merge Sort, Quick Sort, Bubble Sort, Insertion Sort, Heap Sort, Selection Sort เป็นต้น):\n"
    "• ได้ 2.00 คะแนน (เต็ม): อธิบายเหตุผลที่แสดงความเข้าใจเชิงเปรียบเทียบว่า O(n log n) ทำงานน้อยกว่า/เร็วกว่า/ประหยัดเวลา/แบ่งครึ่ง/ลูปน้อยกว่า O(n^2) เมื่อข้อมูลมีขนาดใหญ่ และมีชื่อ Algorithm ที่เกี่ยวข้อง (ให้อนุโลมภาษาพูด และไม่หักคะแนนหากนิสิตเขียนเกินเรื่อง memory/recursive/CPU core)\n"
    "• ได้ 1.50 คะแนน: มีชื่อ Algorithm ที่เกี่ยวข้อง แต่อธิบายเหตุผลสั้นมาก หรือให้เหตุผลกว้างๆ (เช่น ตอบว่ามีประสิทธิภาพเร็วกว่า Quick Sort โดยไม่อธิบายการวนลูป)\n"
    "• ได้ 1.00 คะแนน: มีชื่อ Algorithm ที่เกี่ยวข้อง แต่ให้เหตุผลผิดหลักการ หรือเป็นเหตุผลเชิงวงกลมที่ไม่ช่วยอธิบาย (เช่น บอกว่าเพราะ Merge Sort เสถียรกว่า, หรือคำนวณคณิตศาสตร์ผิด, หรือตอบแค่ว่าดีกว่าโดยไม่มีเหตุผลรองรับ)\n\n"
    "กรณี ข: ไม่มีการระบุชื่อ Algorithm ที่เกี่ยวข้องใดๆ เลย (ไม่มีชื่อ Merge, Quick, Bubble, Heap, Insertion, Sort ฯลฯ):\n"
    "• ได้ 1.00 คะแนน: อธิบายเหตุผลเปรียบเทียบได้ดี (เช่น O(n log n) มีการตัดแบ่งครึ่งข้อมูล/ลูปน้อยกว่า/ทำงานน้อยกว่า O(n^2) ที่วนลูปซ้ำซ้อน) แต่ขาดการยกตัวอย่าง Algorithm (กรณีไม่มีชื่อ Algorithm คะแนนสูงสุดคือ 1.00 คะแนน ห้ามให้ 1.50 หรือ 2.00 เด็ดขาด)\n"
    "• ได้ 0.50 คะแนน: ไม่มีชื่อ Algorithm และตอบเพียงสั้นๆ ว่าเร็วกว่าหรือดีกว่าโดยไม่มีคำอธิบายเปรียบเทียบ\n"
    "• ได้ 0.00 คะแนน: ไม่ตอบ หรือตอบผิดหลักการทั้งหมดอย่างชัดเจน (เช่น เขียนเรื่องลูกเต๋าโยนเหรียญ หรือเขียนเฉพาะตัวโจทย์ซ้ำ)"
)

Q2_RUBRICS = [
    {
        "name": Q2_RUBRIC_NAME,
        "score": 2.0,
        "description": Q2_RUBRIC_DESC,
        "allowed_scores": [0.0, 0.5, 1.0, 1.5, 2.0],
    }
]

Q2_QUESTION_TEXT = "อธิบายว่าทำไม O(n log n) เหมาะกับข้อมูลใหญ่กว่า O(n^2) และยกตัวอย่าง Algorithm"

async def test_all():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["ชุดข้อสอบ_dataset"]
    rows = [r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value == 2]
    print(f"Testing Q2 on {len(rows)} students with calibrated neutral rubric...", flush=True)

    sem = asyncio.Semaphore(5)
    exact = 0
    within_05 = 0
    diffs = []

    async def grade_student(r):
        nonlocal exact, within_05
        sid = ws.cell(r, 1).value
        ans = str(ws.cell(r, 6).value or "")
        human = float(ws.cell(r, 7).value or 0.0)
        async with sem:
            res = await score_with_openai(
                question_text=Q2_QUESTION_TEXT,
                answer_text=ans,
                max_score=2.0,
                answer_key="",
                rubrics=Q2_RUBRICS,
                allowed_scores=[0.0, 0.5, 1.0, 1.5, 2.0],
                strict_rubric_enforcement=False,
            )
            ai = float(res["score"])
            diff = round(ai - human, 2)
            is_match = (diff == 0.0)
            if is_match:
                exact += 1
            if abs(diff) <= 0.5:
                within_05 += 1
            diffs.append((sid, human, ai, diff, ans[:50], res.get("teacher_feedback", "")[:80]))
            status = "MATCH" if is_match else f"DIFF={diff:+.1f}"
            print(f"[{sid}] Human: {human:.1f} | AI: {ai:.1f} | {status}", flush=True)

    await asyncio.gather(*[grade_student(r) for r in rows])

    mae = sum(abs(d[3]) for d in diffs) / len(diffs)
    print("\n" + "=" * 50, flush=True)
    print(f"RESULTS FOR QUESTION 2:", flush=True)
    print(f"Exact Matches: {exact}/{len(rows)} ({exact/len(rows)*100:.2f}%)", flush=True)
    print(f"Within +/- 0.50 pt: {within_05}/{len(rows)} ({within_05/len(rows)*100:.2f}%)", flush=True)
    print(f"MAE: {mae:.4f}", flush=True)
    print("=" * 50, flush=True)

if __name__ == "__main__":
    asyncio.run(test_all())
