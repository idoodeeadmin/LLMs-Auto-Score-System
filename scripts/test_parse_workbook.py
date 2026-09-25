import openpyxl
import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

# 1. Parse Sheet 'ชุดข้อสอบ_dataset'
ws_data = wb["ชุดข้อสอบ_dataset"]
dataset_items = []
for r in range(6, ws_data.max_row + 1):
    sid = ws_data.cell(r, 1).value
    if not sid:
        continue
    q_no = ws_data.cell(r, 2).value
    q_type = ws_data.cell(r, 3).value
    q_content = ws_data.cell(r, 4).value
    ans_type = ws_data.cell(r, 5).value
    ans_val = ws_data.cell(r, 6).value
    h_score = ws_data.cell(r, 7).value
    ai_score = ws_data.cell(r, 8).value
    ai_conf = ws_data.cell(r, 9).value
    ai_fb = ws_data.cell(r, 10).value
    
    # Extract image path if hyperlink formula
    img_path = None
    if ans_type == "img" or "Photo" in str(ans_val):
        match = re.search(r'"([^"]+)"', str(ans_val))
        if match:
            img_path = match.group(1)
        elif str(ans_val).startswith("photo_clean"):
            img_path = str(ans_val)
        elif str(ans_val).startswith("LINE_ALBUM"):
            # determine folder
            if q_no == 4: img_path = f"photo_clean_ชุดที่1/{ans_val}"
            elif q_no == 5: img_path = f"photo_clean_ชุดที่2/{ans_val}"
            elif q_no == 6: img_path = f"photo_clean_ชุดที่3/{ans_val}"
            
    h_val = float(h_score) if (h_score is not None and str(h_score).strip() != "") else None
    ai_val = float(ai_score) if (ai_score is not None and str(ai_score).strip() != "") else None
    
    diff = round(ai_val - h_val, 2) if (ai_val is not None and h_val is not None) else None
    is_exact = (diff == 0.0) if diff is not None else None
    
    dataset_items.append({
        "row": r,
        "sample_id": str(sid),
        "question_no": int(q_no) if q_no else None,
        "question_type": str(q_type or ""),
        "question_content": str(q_content or ""),
        "answer_type": str(ans_type or ""),
        "student_answer": str(ans_val or ""),
        "image_path": img_path,
        "human_score": h_val,
        "ai_score": ai_val,
        "diff": diff,
        "is_exact": is_exact,
        "ai_confidence": str(ai_conf or ""),
        "ai_feedback": str(ai_fb or "")
    })

# 2. Parse Sheet 'Exam_Rubrics'
ws_rubrics = wb["Exam_Rubrics"]
rubric_items = []
for r in range(5, ws_rubrics.max_row + 1):
    q_no = ws_rubrics.cell(r, 1).value
    if q_no is None:
        continue
    topic = ws_rubrics.cell(r, 2).value
    max_s = ws_rubrics.cell(r, 3).value
    name = ws_rubrics.cell(r, 4).value
    score = ws_rubrics.cell(r, 5).value
    desc = ws_rubrics.cell(r, 6).value
    
    rubric_items.append({
        "question_no": int(q_no),
        "topic": str(topic or ""),
        "max_score": float(max_s) if max_s is not None else None,
        "name": str(name or ""),
        "score": float(score) if score is not None else None,
        "description": str(desc or "")
    })

# 3. Parse Sheet 'Template_Example'
ws_tmpl = wb["Template_Example"]
template_items = []
for r in range(6, ws_tmpl.max_row + 1):
    sid = ws_tmpl.cell(r, 1).value
    if not sid:
        continue
    template_items.append({
        "sample_id": str(sid),
        "question_no": ws_tmpl.cell(r, 2).value,
        "question_type": ws_tmpl.cell(r, 3).value,
        "question_content": ws_tmpl.cell(r, 4).value,
        "answer_type": ws_tmpl.cell(r, 5).value,
        "student_answer": ws_tmpl.cell(r, 6).value,
        "human_score": ws_tmpl.cell(r, 7).value,
        "ai_score": ws_tmpl.cell(r, 8).value,
        "ai_confidence": ws_tmpl.cell(r, 9).value,
        "ai_feedback": ws_tmpl.cell(r, 10).value
    })

print(f"Extracted {len(dataset_items)} dataset items, {len(rubric_items)} rubrics, {len(template_items)} template rows.")
