# -*- coding: utf-8 -*-
"""
Regrade all 34 answers of Question 1 using the approved concise rubric,
and save the new AI scores, confidence, and dual-perspective feedbacks
directly into sheet 'ชุดข้อสอบ_dataset' of 'ชุดข้อสอบ_dataset.xlsx'.
"""
import asyncio
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
load_dotenv(ROOT / ".env")
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q1_RUBRIC_NAME = "Row-major vs Column-major"
Q1_RUBRIC_DESC = (
    "ประเมินความเข้าใจความต่างของ Row-major และ Column-major (คะแนนเต็ม 2.00 คะแนน):\n"
    "• 2.00 คะแนน: อธิบายความต่างได้ถูกต้องครบทั้ง 2 ฝั่ง (Row-major อิงตามแถว/แนวนอน/แกน X และ Column-major อิงตามคอลัมน์/แนวตั้ง/แกน Y หรืออิงการจัดเก็บ/การหา address ของแต่ละมิติ)\n"
    "• 1.00 คะแนน: อธิบายถูกต้องเพียงฝั่งเดียว หรือตอบสั้นเฉพาะความเข้าใจเบื้องต้น\n"
    "• 0.00 คะแนน: ตอบผิดทั้งหมด หรือไม่ตอบ\n"
    "(ระดับคะแนนที่ให้ได้คือ 2.0, 1.0 หรือ 0.0 คะแนนเท่านั้น)"
)

Q1_RUBRICS = [
    {
        "name": Q1_RUBRIC_NAME,
        "score": 2.0,
        "description": Q1_RUBRIC_DESC,
        "allowed_scores": [0.0, 1.0, 2.0],
    }
]

Q1_QUESTION_TEXT = "อธิบายความต่างของ Row-major vs Column-major"
Q1_ANSWER_KEY = (
    "Row-major คือการจัดเก็บหรือเข้าถึงข้อมูลตามแถว (แนวนอน) "
    "ส่วน Column-major คือการจัดเก็บหรือเข้าถึงข้อมูลตามคอลัมน์ (แนวตั้ง)"
)

def calc_qwk(y_t, y_p, max_s=2.0, step=1.0):
    k = int(round(max_s / step)) + 1
    w = [[((i - j) ** 2) / ((k - 1) ** 2) for j in range(k)] for i in range(k)]
    cat_t = [min(k - 1, max(0, int(round(v / step)))) for v in y_t]
    cat_p = [min(k - 1, max(0, int(round(v / step)))) for v in y_p]
    n = len(y_t)
    if n == 0:
        return 0.0
    o = [[0] * k for _ in range(k)]
    for t, p in zip(cat_t, cat_p):
        o[t][p] += 1
    hist_t = [0] * k
    hist_p = [0] * k
    for t, p in zip(cat_t, cat_p):
        hist_t[t] += 1
        hist_p[p] += 1
    e = [[(hist_t[i] * hist_p[j]) / n for j in range(k)] for i in range(k)]
    num = sum(w[i][j] * o[i][j] for i in range(k) for j in range(k))
    den = sum(w[i][j] * e[i][j] for i in range(k) for j in range(k))
    return 1.0 - (num / den) if den != 0 else 1.0

async def main():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path)
    ws = wb["ชุดข้อสอบ_dataset"]
    
    rows = [r for r in range(2, ws.max_row + 1) if ws.cell(r, 2).value == 1]
    assert len(rows) == 34, f"Expected 34 rows for Q1, found {len(rows)}"
    
    print(f"Total students to grade for Q1: {len(rows)}")
    
    sem = asyncio.Semaphore(5)
    exact = 0
    within_05 = 0
    total_diff = 0.0
    y_true = []
    y_pred = []
    discrepancies = []

    async def grade_one(r):
        nonlocal exact, within_05, total_diff
        sid = str(ws.cell(r, 1).value)
        ans = str(ws.cell(r, 6).value or "")
        human = float(ws.cell(r, 7).value or 0.0)
        
        async with sem:
            res = await score_with_openai(
                question_text=Q1_QUESTION_TEXT,
                answer_text=ans,
                max_score=2.0,
                answer_key=Q1_ANSWER_KEY,
                rubrics=Q1_RUBRICS,
                allowed_scores=[0.0, 1.0, 2.0],
            )
            ai = float(res["score"])
            diff = round(ai - human, 2)
            is_match = (diff == 0.0)
            if is_match:
                exact += 1
            else:
                discrepancies.append((sid, human, ai, diff, ans))
                
            if abs(diff) <= 0.5:
                within_05 += 1
            total_diff += abs(diff)
            y_true.append(human)
            y_pred.append(ai)
            
            fb = res.get("feedback", "")
            conf = res.get("confidence", "high")
            
            print(f"[{sid}] Human: {human:.1f} | AI: {ai:.1f} | {'MATCH' if is_match else f'DIFF={diff:+.1f}'}")
            return r, ai, conf, fb

    tasks = [grade_one(r) for r in rows]
    results = await asyncio.gather(*tasks)

    # Save results to Excel
    for r, ai, conf, fb in results:
        ws.cell(r, 8, ai)
        ws.cell(r, 9, conf)
        ws.cell(r, 10, fb)

    wb.save(excel_path)
    print(f"\nSuccessfully saved new Q1 AI scores to {excel_path.name} (Sheet: ชุดข้อสอบ_dataset)!")
    
    qwk = calc_qwk(y_true, y_pred, max_s=2.0, step=1.0)
    mae = total_diff / 34.0
    
    print("\n==========================================")
    print("Q1 REGRADING COMPLETED & SAVED TO EXCEL:")
    print(f"Exact Matches:      {exact}/34 ({exact/34*100:.2f}%)")
    print(f"Within +/-0.50:     {within_05}/34 ({within_05/34*100:.2f}%)")
    print(f"MAE:                {mae:.4f}")
    print(f"QWK:                {qwk:.4f}")
    print("==========================================")
    print(f"Mismatches ({len(discrepancies)} items):")
    for sid, h, a, d, ans in discrepancies:
        print(f"  {sid} | Human: {h:.1f} | AI: {a:.1f} | diff={d:+.1f} | ans: {ans[:40]}")

if __name__ == "__main__":
    asyncio.run(main())
