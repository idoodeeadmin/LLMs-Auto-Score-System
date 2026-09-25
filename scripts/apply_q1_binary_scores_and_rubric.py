import json
import sys
from pathlib import Path
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

sys.stdout.reconfigure(encoding='utf-8')

ROOT = Path(__file__).resolve().parents[1]
RESULTS_FILE = ROOT / 'artifacts' / 'q1_binary_test_results.json'
EXCEL_FILE = ROOT / 'ชุดข้อสอบใหม่' / 'ชุดข้อสอบ_dataset.xlsx'
RUBRIC_AI_FILE = ROOT / 'ชุดข้อสอบใหม่' / 'เกณฑ์ตรวจสำหรับAI.xlsx'

def apply_q1():
    print(f"Loading results from {RESULTS_FILE.name}...")
    with open(RESULTS_FILE, encoding='utf-8') as f:
        data = json.load(f)

    results_map = {r['sample_id']: r for r in data['results']}
    print(f"Loaded {len(results_map)} student results.")

    wb = openpyxl.load_workbook(EXCEL_FILE)
    ws = wb['ชุดข้อสอบ_dataset']

    thin_border = Border(
        left=Side(style='thin', color='00BFBFBF'),
        right=Side(style='thin', color='00BFBFBF'),
        top=Side(style='thin', color='00BFBFBF'),
        bottom=Side(style='thin', color='00BFBFBF')
    )
    font_aptos = Font(name='Aptos', size=10)
    align_center = Alignment(horizontal='center', vertical='top')
    align_feedback = Alignment(horizontal='left', vertical='top', wrap_text=True)

    # 1. Update ชุดข้อสอบ_dataset rows 6 to 39
    updated_count = 0
    for row in range(6, 40):
        sid = str(ws.cell(row=row, column=1).value or '').strip()
        qno = ws.cell(row=row, column=2).value
        if qno != 1 or not sid or sid not in results_map:
            continue

        res = results_map[sid]
        score = float(res['ai'])
        conf = str(res.get('confidence', 'high'))
        fb_teacher = res.get('teacher_feedback', '').strip()
        fb_student = res.get('student_feedback', '').strip()
        full_fb = f"[สำหรับผู้สอน]\n{fb_teacher}\n\n[สำหรับนักเรียน]\n{fb_student}"

        c8 = ws.cell(row=row, column=8, value=score)
        c8.font = font_aptos
        c8.alignment = align_center
        c8.border = thin_border

        c9 = ws.cell(row=row, column=9, value=conf)
        c9.font = font_aptos
        c9.alignment = align_center
        c9.border = thin_border

        c10 = ws.cell(row=row, column=10, value=full_fb)
        c10.font = font_aptos
        c10.alignment = align_feedback
        c10.border = thin_border

        updated_count += 1

    print(f"Updated {updated_count} student scores and dual feedbacks in 'ชุดข้อสอบ_dataset'")

    # 2. Update Exam_Rubrics
    if 'Exam_Rubrics' in wb.sheetnames:
        ws_rubric = wb['Exam_Rubrics']
        r5_desc = (
            "อธิบาย Row-major ว่าเป็นการจัดเก็บข้อมูล ลำดับการเรียง หรือการหาตำแหน่ง address ตามแนวแถว "
            "(Row/แนวนอน/แกน X หรือรูปแบบดัชนี [i][j]) ได้ 1.0 คะแนน (ตอบผิดหรือไม่ตอบได้ 0.0 คะแนน)"
        )
        r6_desc = (
            "อธิบาย Column-major ว่าเป็นการจัดเก็บข้อมูล ลำดับการเรียง หรือการหาตำแหน่ง address ตามแนวคอลัมน์ "
            "(Column/แนวตั้ง/แกน Y หรือรูปแบบดัชนี [j][i]) ได้ 1.0 คะแนน (ตอบผิดหรือไม่ตอบได้ 0.0 คะแนน)\n\n"
            "(คะแนนรวมคือผลรวมของ 2 ส่วน = 0.0, 1.0 หรือ 2.0 คะแนนเท่านั้น)"
        )
        ws_rubric.cell(row=5, column=6, value=r5_desc)
        ws_rubric.cell(row=6, column=6, value=r6_desc)
        print("Updated Question 1 rubrics in 'Exam_Rubrics' sheet.")

    wb.save(EXCEL_FILE)
    print(f"Saved {EXCEL_FILE.name} successfully.")

    # 3. Update เกณฑ์ตรวจสำหรับAI.xlsx
    if RUBRIC_AI_FILE.exists():
        wb_ai = openpyxl.load_workbook(RUBRIC_AI_FILE)
        ws_ai = wb_ai.worksheets[0]
        q1_ai_criteria = (
            "ข้อ 1 (2 คะแนน): ประเมิน 2 ด้านย่อย ด้านละ 1.0 คะแนน (คะแนนรวมเป็น 2.0, 1.0 หรือ 0.0 เท่านั้น):\n"
            "- Row-major (1.0 คะแนน): อธิบายว่าจัดเก็บข้อมูล ลำดับการเรียง หรือหา address ตามแนวแถว (Row/แนวนอน/แกน X หรือดัชนี [i][j])\n"
            "- Column-major (1.0 คะแนน): อธิบายว่าจัดเก็บข้อมูล ลำดับการเรียง หรือหา address ตามแนวคอลัมน์ (Column/แนวตั้ง/แกน Y หรือดัชนี [j][i])\n"
            "(ยอมรับคำอธิบายภาษาพูด เช่น คิดแถวก่อนหลัก/หลักก่อนแถว, ซ้ายไปขวา/บนลงล่าง, อิงแกน X vs แกน Y)"
        )
        ai_updated = 0
        for r in range(2, ws_ai.max_row + 1):
            q_val = ws_ai.cell(row=r, column=1).value
            if q_val and 'Row-major' in str(q_val):
                ws_ai.cell(row=r, column=2, value=q1_ai_criteria)
                ai_updated += 1
        wb_ai.save(RUBRIC_AI_FILE)
        print(f"Updated {ai_updated} rows in '{RUBRIC_AI_FILE.name}'.")

    print("\nAll Q1 updates completed successfully!")

if __name__ == '__main__':
    apply_q1()
