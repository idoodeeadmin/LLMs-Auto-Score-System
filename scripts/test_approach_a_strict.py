import asyncio
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q3_QUESTION_TEXT = "การใช้ลิ้งค์ลิสต์เป็นสแตกและคิวจะแตกต่างจากการใช้อาร์เรย์อย่างไร และมีข้อดีข้อเสียอย่างไร (1 คะแนน)"

Q3_ANSWER_KEY = (
    "เฉลยและแนวคำตอบ:\n"
    "1. ความแตกต่าง (0.50 คะแนน): Linked List ขนาดปรับเปลี่ยนได้ตามจริง (Dynamic) ไม่ต้องกำหนดขนาดล่วงหน้า ส่วน Array ขนาดคงที่ (Static/Fixed) จองพื้นที่ล่วงหน้า\n"
    "2. ข้อดี (0.25 คะแนน): Linked List ขนาดไม่จำกัด ไม่เกิด overflow / Array เข้าถึงข้อมูลได้เร็ว\n"
    "3. ข้อเสีย (0.25 คะแนน): Linked List เปลืองเมมโมรี่พอยน์เตอร์ เข้าถึงช้ากว่า / Array ขนาดคงที่เสี่ยง overflow หรือเปลืองเนื้อที่"
)

# เกณฑ์แยก 3 ส่วน แบบกำหนดจุดตรวจเด็ดขาดและห้ามนับซ้ำ
APPROACH_A_RUBRIC_NAME = "การประเมินแยก 3 องค์ประกอบแบบรัดกุม (ความแตกต่าง 0.50 + ข้อดี 0.25 + ข้อเสีย 0.25)"

APPROACH_A_RUBRIC_DESC = (
    "เกณฑ์การประเมินคะแนนเต็ม 1.00 คะแนน โดยให้ตรวจและรวมคะแนนจาก 3 องค์ประกอบแยกจากกันอย่างเด็ดขาด "
    "(คะแนนรวมต้องเป็น 0.00, 0.25, 0.50, 0.75 หรือ 1.00 เท่านั้น):\n\n"
    "1. ส่วนความแตกต่าง (เต็ม 0.50 คะแนน):\n"
    "   - ได้ 0.50 คะแนน: อธิบายความแตกต่างของโครงสร้างข้อมูล (เช่น Linked List ปรับเปลี่ยนขนาดได้ ยืดหยุ่น ไม่จำกัดขนาด ส่วน Array มีขนาดคงที่ แน่นอน ต้องจองพื้นที่ล่วงหน้า) หรืออธิบายกลไก Stack/Queue (LIFO/FIFO เทียบกับ Array) ได้พอเข้าใจ\n"
    "   - ได้ 0.00 คะแนน: ไม่ได้เขียนเปรียบเทียบความแตกต่าง หรือตอบนอกเรื่องวิชา Data Structures สิ้นเชิง (เช่น ช่วยประหยัดพลังงาน, รถไฟ)\n\n"
    "2. ส่วนข้อดี (เต็ม 0.25 คะแนน):\n"
    "   - ได้ 0.25 คะแนน: มีการระบุ 'ข้อดี' เพิ่มเติม (เช่น เขียนหัวข้อ 'ข้อดี' หรือระบุประโยชน์จุดเด่นชัดเจน เช่น ไม่จำกัดขนาด, ยืดหยุ่น, ค้นหาเร็ว, จัดการข้อมูลตามลำดับได้ดี) พอเข้าใจ\n"
    "   - ได้ 0.00 คะแนน: ไม่ได้ระบุข้อดี หรือระบุมั่ว/ตอบนอกเรื่อง\n\n"
    "3. ส่วนข้อเสีย (เต็ม 0.25 คะแนน):\n"
    "   - ได้ 0.25 คะแนน: มีการระบุ 'ข้อเสีย' เพิ่มเติม (เช่น เขียนหัวข้อ 'ข้อเสีย' หรือระบุจุดด้อยชัดเจน เช่น index out of bounds, เปลืองพื้นที่, ซับซ้อนสับสน, เข้าถึงช้า) พอเข้าใจ\n"
    "   - ได้ 0.00 คะแนน: ไม่ได้ระบุข้อเสีย (เช่น ใส่เครื่องหมายขีด - หรือละไว้ไม่เขียนถึงข้อเสีย) หรือตอบมั่ว/ตอบนอกเรื่อง\n\n"
    "🚨 **กฎเหล็กการตรวจ (Strict Constraints)**:\n"
    "1. **ห้ามนับความแตกต่างซ้ำเป็นข้อดีข้อเสียเด็ดขาด**: หากนักเรียนเขียนอธิบายความแตกต่างเป็นประโยคเดียว (เช่น 'ลิ้งค์ลิสต์เพิ่มลดขนาดได้ แต่อาร์เรย์ต้องกำหนดขนาดล่วงหน้า') โดยไม่ได้เขียนแจกแจงแยกหัวข้อข้อดีและข้อเสีย ให้นักเรียนได้เฉพาะคะแนนส่วนที่ 1 (0.50) เท่านั้น ห้ามเหมาให้ส่วนข้อดีและข้อเสียเด็ดขาด\n"
    "2. **นักเรียนที่เขียนเฉพาะข้อดีข้อเสียสั้นๆ แต่ไม่ได้อธิบายความแตกต่าง**: จะไม่ได้คะแนนส่วนที่ 1 (0.50) แต่จะได้เฉพาะคะแนนส่วนข้อดี (0.25) หรือข้อเสีย (0.25) เท่าที่มีจริง\n"
    "3. **คำตอบที่ตอบนอกเรื่อง (เช่น ช่วยประหยัดพลังงาน)**: ต้องได้ 0.00 คะแนนในทุกส่วนทันที ห้ามให้คะแนนอนุโลมใดๆ\n\n"
    "(ใน teacher_feedback ให้แจกแจงผลคะแนนทั้ง 3 ส่วน: ความแตกต่าง (0.50/0.00), ข้อดี (0.25/0.00), ข้อเสีย (0.25/0.00) และผลรวมคะแนนอย่างชัดเจน)"
)

RUBRICS = [
    {
        "name": APPROACH_A_RUBRIC_NAME,
        "score": 1.0,
        "description": APPROACH_A_RUBRIC_DESC,
    }
]

async def test():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["ชุดข้อสอบ_dataset"]
    
    first10 = []
    for r in range(74, 84):
        sid = ws.cell(r, 1).value
        ans = str(ws.cell(r, 6).value or "")
        h = float(ws.cell(r, 7).value or 0.0)
        first10.append({"sid": sid, "ans": ans, "human": h})
        
    print(f"=== กำลังทดสอบแนวทาง A (เกณฑ์ 3 ส่วนแบบรัดกุม ห้ามนับซ้ำ) บน 10 ข้อแรก ===")
    
    sem = asyncio.Semaphore(5)
    async def grade_one(item):
        async with sem:
            res = await score_with_openai(
                question_text=Q3_QUESTION_TEXT,
                answer_text=item["ans"],
                max_score=1.0,
                answer_key=Q3_ANSWER_KEY,
                rubrics=RUBRICS,
                strict_rubric_enforcement=True,
            )
            s = float(res.get("score", 0.0))
            diff = round(s - item["human"], 2)
            match = (s == item["human"])
            icon = "🟢 ตรงเป๊ะ" if match else f"🟠 ต่าง ({diff:+0.2f})"
            tf = res.get("teacher_feedback", "").replace("\n", " ")
            print(f"{item['sid']} | อาจารย์: {item['human']:.2f} | AI: {s:.2f} | {icon}")
            return {"sid": item["sid"], "human": item["human"], "ai": s, "diff": diff, "match": match, "tf": tf}

    tasks = [grade_one(x) for x in first10]
    results = await asyncio.gather(*tasks)
    exact = sum(1 for r in results if r["match"])
    w025 = sum(1 for r in results if abs(r["diff"]) <= 0.2501)
    print("\n" + "="*60)
    print(f"สรุป 10 ข้อแรก แนวทาง A:")
    print(f"• ตรงกันเป๊ะ (Exact Match): {exact} / 10 ({exact*10}%)")
    print(f"• ต่างไม่เกิน ±0.25 pt:       {w025} / 10 ({w025*10}%)")
    print("="*60)

if __name__ == "__main__":
    asyncio.run(test())
