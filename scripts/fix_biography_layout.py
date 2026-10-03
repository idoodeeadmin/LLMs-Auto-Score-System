from pathlib import Path

from docx import Document
from docx.enum.text import WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


SOURCE = Path(r"C:\Users\idood\Downloads\LLMs-Auto-Score-System-main\docs_and_tests\LLMs_Auto_Score_Systems_Pro1_1-5.docx")
OUTPUT = Path(r"C:\Users\idood\Downloads\LLMs-Auto-Score-System-main\docs_and_tests\LLMs_Auto_Score_Systems_Pro1_1-5_edited.docx")


def set_cell_margins(cell, top=0, bottom=0, start=70, end=70):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("bottom", bottom), ("start", start), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


document = Document(SOURCE)

# The final two tables contain the two author biographies. Compact only these
# tables so all biography details fit on one page without altering the rest of
# the thesis layout.
for table in document.tables[-2:]:
    for row in table.rows:
        for cell in row.cells:
            set_cell_margins(cell)
            for paragraph in cell.paragraphs:
                fmt = paragraph.paragraph_format
                fmt.space_before = Pt(0)
                fmt.space_after = Pt(0)
                fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
                for run in paragraph.runs:
                    run.font.size = Pt(14)

# Tighten only the two biography subheadings.
for paragraph in document.paragraphs:
    if paragraph.text.strip() in {
        "ประวัติย่อผู้จัดทำโครงงานคนที่ 1",
        "ประวัติย่อผู้จัดทำโครงงานคนที่ 2",
    }:
        paragraph.paragraph_format.space_before = Pt(2)
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.keep_with_next = True
        for run in paragraph.runs:
            run.font.size = Pt(15)

# Ask Word to refresh TOC and page-reference fields when the edited file opens.
settings = document.settings._element
update = settings.find(qn("w:updateFields"))
if update is None:
    update = OxmlElement("w:updateFields")
    settings.append(update)
update.set(qn("w:val"), "true")

document.save(OUTPUT)
print(OUTPUT)
