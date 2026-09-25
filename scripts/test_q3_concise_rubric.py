import asyncio
import json
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q3_QUESTION_TEXT = "การใช้ลิ้งค์ลิสต์เป็นสแตกและคิวจะแตกต่างจากการใช้อาร์เรย์อย่างไร และมีข้อดีข้อเสียอย่างไร (1 คะแนน)"

Q3_ANSWER_KEY = (
    "1. ความแตกต่าง (Difference):\n"
    "   - Array: ขนาดคงที่ (Fixed-size / Static) ต้องกำหนดขนาดล่วงหน้า การทำ Stack/Queue ต้องคุม index\n"
    "   - Linked List: ขนาดยืดหยุ่นปรับเปลี่ยนได้ตามจริง (Dynamic) ไม่ต้องกำหนดขนาดล่วงหน้า เชื่อมต่อด้วย Node และ Pointer\n"
    "2. ข้อดี-ข้อเสีย (Pros & Cons):\n"
    "   - Linked List: ข้อดีคือไม่จำกัดขนาด ไม่เกิด Stack Overflow ง่าย เพิ่มลดโหนดสะดวก; ข้อเสียคือเข้าถึงข้อมูลช้ากว่า (Sequential access) และเปลืองหน่วยความจำเก็บ Pointer\n"
    "   - Array: ข้อดีคือเข้าถึงข้อมูลได้รวดเร็วทันทีด้วย Index (O(1)); ข้อเสียคือขนาดคงที่เสี่ยงเกิด Overflow หรือเปลืองพื้นที่หากจองเกิน"
)

# กระชับ ชัดเจน เป็นธรรมชาติ ไม่ระบุชื่อคน/ID และไม่มีกฎยิบย่อย
Q3_CONCISE_RUBRIC_NAME = "การเปรียบเทียบ Linked List vs Array สำหรับ Stack และ Queue พร้อมข้อดีข้อเสีย"

Q3_CONCISE_RUBRIC_DESC = (
    "โจทย์ถาม 2 ส่วนหลัก: (1) ความแตกต่างระหว่าง Linked List กับ Array ใน Stack/Queue และ (2) ข้อดีข้อเสีย (คะแนนเต็ม 1.00)\n\n"
    "เกณฑ์การให้คะแนนตามระดับความสมบูรณ์ของคำตอบ:\n"
    "- 1.00 คะแนน: ตอบครบทั้ง 2 ส่วนหลักอย่างชัดเจน ได้แก่ (1) อธิบายความแตกต่างของโครงสร้าง/ขนาด (เช่น Dynamic ขยายได้ vs Static ขนาดคงที่ หรือเชื่อมโยงการทำงาน Stack/Queue) และ (2) ระบุทั้ง 'ข้อดี' และ 'ข้อเสีย' อย่างสมเหตุสมผลตามระดับนักศึกษา (อนุโลมสำนวนภาษา ไม่จำเป็นต้องระบุคำว่า pointer หรือ index หากเนื้อหาสื่อข้อดีข้อเสียได้ชัดเจน)\n"
    "- 0.75 คะแนน: ตอบความแตกต่างได้ถูกต้อง และระบุข้อดี แต่ขาดข้อเสีย (หรือระบุข้อเสียแต่ขาดข้อดี) หรืออธิบายข้อดีข้อเสียกว้างๆ ในเชิงการเขียนโปรแกรม\n"
    "- 0.50 คะแนน: ตอบถูกเพียงส่วนใดส่วนหนึ่ง เช่น เปรียบเทียบความแตกต่างเรื่องขนาดได้แต่ไม่มีการแจกแจงข้อดีข้อเสียที่ชัดเจน (เขียนรวมเป็นความเรียงเดียว), หรือตอบเฉพาะพฤติกรรม LIFO/FIFO โดยอธิบายข้อดีข้อเสียเพียงคร่าวๆ, หรือระบุข้อดี/ข้อเสียเพียงด้านเดียวสั้นๆ\n"
    "- 0.25 คะแนน: ตอบสั้นมาก คลุมเครือ หรือไม่ระบุชื่อโครงสร้าง แต่มีคีย์เวิร์ดที่ถูกต้องตามหลักการเพียงจุดเดียว (เช่น คำว่า ยืดหยุ่น หรือ static)\n"
    "- 0.00 คะแนน: ตอบไม่ตรงประเด็น (เช่น เรื่องประหยัดพลังงาน), หรือเขียนเพียงประโยคเดียวโดยไม่มีข้อดีข้อเสีย, หรือตอบผิดหลักการอย่างมีนัยสำคัญ, หรือไม่ตอบ\n\n"
    "(ต้องระบุ teacher_feedback สำหรับผู้สอน และ student_feedback เป็นคำแนะนำสำหรับนักเรียน)"
)

Q3_RUBRICS = [
    {
        "name": Q3_CONCISE_RUBRIC_NAME,
        "score": 1.0,
        "description": Q3_CONCISE_RUBRIC_DESC,
    }
]

def calculate_qwk(y_true, y_pred, step=0.25, max_score=1.0):
    if len(y_true) != len(y_pred) or len(y_true) == 0: return 0.0
    if y_true == y_pred: return 1.0
    k = int(round(max_score / step)) + 1
    w = [[((i - j) ** 2) / ((k - 1) ** 2) for j in range(k)] for i in range(k)]
    cat_true = [min(k-1, max(0, int(round(v / step)))) for v in y_true]
    cat_pred = [min(k-1, max(0, int(round(v / step)))) for v in y_pred]
    o = [[0] * k for _ in range(k)]
    for t, p in zip(cat_true, cat_pred): o[t][p] += 1
    r_t = [sum(o[i][j] for j in range(k)) for i in range(k)]
    r_p = [sum(o[i][j] for i in range(k)) for j in range(k)]
    n = len(y_true)
    e = [[(r_t[i] * r_p[j]) / n for j in range(k)] for i in range(k)]
    num = sum(w[i][j] * o[i][j] for i in range(k) for j in range(k))
    den = sum(w[i][j] * e[i][j] for i in range(k) for j in range(k))
    return 1.0 - (num / den) if den != 0 else 1.0

async def main():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws_data = wb["ชุดข้อสอบ_dataset"]
    
    all_q3 = []
    for r in range(74, 108):
        sid = ws_data.cell(r, 1).value
        q_no = ws_data.cell(r, 2).value
        ans = str(ws_data.cell(r, 6).value or "")
        h = float(ws_data.cell(r, 7).value or 0.0)
        old_ai = float(ws_data.cell(r, 8).value or 0.0)
        all_q3.append({
            "row": r,
            "sample_id": sid,
            "student_ans": ans,
            "human_score": h,
            "old_ai_score": old_ai
        })
        
    print(f"=== TESTING CONCISE & NATURAL Q3 RUBRIC ON ALL 34 STUDENTS ===")
    sem = asyncio.Semaphore(6)
    
    async def grade_one(item):
        async with sem:
            res = await score_with_openai(
                question_text=Q3_QUESTION_TEXT,
                answer_text=item["student_ans"],
                max_score=1.0,
                answer_key=Q3_ANSWER_KEY,
                rubrics=Q3_RUBRICS,
                image_bytes_list=None,
                strict_rubric_enforcement=True,
            )
            new_ai = float(res.get("score", 0.0))
            diff = round(new_ai - item["human_score"], 2)
            match = (new_ai == item["human_score"])
            
            return {
                "row": item["row"],
                "sample_id": item["sample_id"],
                "student_ans": item["student_ans"],
                "human_score": item["human_score"],
                "old_ai_score": item["old_ai_score"],
                "new_ai_score": new_ai,
                "confidence": res.get("confidence", "high"),
                "diff": diff,
                "is_exact": match,
                "teacher_feedback": res.get("teacher_feedback", ""),
                "student_feedback": res.get("student_feedback", "")
            }

    tasks = [grade_one(item) for item in all_q3]
    results = await asyncio.gather(*tasks)
    results = sorted(results, key=lambda x: x["row"])
    
    print("\n--- RESULTS PER STUDENT (CONCISE RUBRIC) ---")
    for r in results:
        status = "🟢 EXACT" if r["is_exact"] else f"🟠 DIFF ({r['diff']:+0.2f})"
        print(f"{r['sample_id']} | Human: {r['human_score']:.2f} | AI: {r['new_ai_score']:.2f} -> {status}")
        
    y_true = [x["human_score"] for x in results]
    y_pred = [x["new_ai_score"] for x in results]
    
    exact = sum(1 for x in results if x["is_exact"])
    w025 = sum(1 for x in results if abs(x["diff"]) <= 0.2501)
    w050 = sum(1 for x in results if abs(x["diff"]) <= 0.5001)
    mae = sum(abs(x["diff"]) for x in results) / len(results)
    qwk = calculate_qwk(y_true, y_pred)
    
    print("\n" + "="*70)
    print("CONCISE RUBRIC PERFORMANCE (N=34)")
    print(f"Exact Match:  {exact} / 34 ({exact/34*100:.1f}%) | MAE: {mae:.4f} | QWK: {qwk:.4f}")
    print(f"Within ≤0.25: {w025} / 34 ({w025/34*100:.1f}%)")
    print(f"Within ≤0.50: {w050} / 34 ({w050/34*100:.1f}%)")
    print("="*70)
    
    out_file = ROOT / "artifacts" / "q3_all34_concise_results.json"
    with open(out_file, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)

if __name__ == "__main__":
    asyncio.run(main())
