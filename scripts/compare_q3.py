import sys
import json
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

with open('artifacts/q3_clean_test_results.json', encoding='utf-8') as f:
    d = json.load(f)

wb = openpyxl.load_workbook('ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx', data_only=True)
ws = wb['ชุดข้อสอบ_dataset']

excel_scores = {ws.cell(r, 1).value: (float(ws.cell(r, 7).value), float(ws.cell(r, 8).value)) for r in range(74, 108)}

old_exact = sum(1 for sid, (h, ai) in excel_scores.items() if h == ai)
old_mae = sum(abs(h - ai) for sid, (h, ai) in excel_scores.items()) / 34

print('========================================================================')
print('QUESTION 3: OVERFITTED BASELINE vs CLEAN NATURAL RUBRIC')
print('========================================================================')
print(f"Old Excel Baseline (Overfitted): Exact = {old_exact}/34 (47.1%), MAE = {old_mae:.4f}")
print(f"New Clean Rubric (AI คิดเอง)    : Exact = {d['exact']}/34 ({d['exact']/34*100:.1f}%), Tol<=0.25 = {d['diff025']}/34 ({d['diff025']/34*100:.1f}%), MAE = {d['mae']:.4f}, QWK = {d['qwk']:.4f}")
print('------------------------------------------------------------------------')

changes = []
for r in d['results']:
    sid = r['sample_id']
    h, old_ai = excel_scores[sid]
    new_ai = r['ai']
    if old_ai != new_ai:
        changes.append((sid, h, old_ai, new_ai, r['answer']))

print(f"Changed scores: {len(changes)}/34\n")
for sid, h, old, new, ans in changes:
    snippet = ans.replace('\n', ' ')[:45]
    if new == h:
        status = 'MATCH (+)'
    elif abs(new - h) < abs(old - h):
        status = 'CLOSER (+)'
    elif abs(new - h) > abs(old - h):
        status = 'REGRESSED (-)'
    else:
        status = 'CHANGED (=)'
    print(f"  {sid} | Human={h:<4} | Old={old:<4} -> New={new:<4} | {status:<13} | {snippet}")
