"""Grade all 204 current workbook answers with Gemini and the production template.

Reads the authoritative workbook and writes a resumable JSON report. It never
modifies the source workbook or sends the teacher's score to the model.
"""

import asyncio
import hashlib
import json
import os
import re
import sys
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from dotenv import load_dotenv

load_dotenv(ROOT / ".env")
# Reuse the local word-count shim; the bundled Python and repository Pydantic
# wheel target different interpreter versions.
import scripts.regrade_corrected_q2_q3_answers  # noqa: E402,F401
from server.services.gemini_grading import GEMINI_MODEL, score_with_gemini  # noqa: E402
from server.services.openai_grading import PROMPT_VERSION  # noqa: E402

BOOK = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx"
OUT_DIR = ROOT / "docs_and_tests/gemini38_regrade"
ONLY_IDS = {part.strip() for part in os.getenv("GEMINI_GRADE_ONLY_IDS", "").split(",") if part.strip()}
ONLY_QUESTION = int(os.getenv("GEMINI_GRADE_ONLY_QUESTION", "0"))
assert ONLY_QUESTION in range(7) and not (ONLY_IDS and ONLY_QUESTION)
REPORT_TAG = os.getenv("GEMINI_GRADE_REPORT_TAG", "new_prompt")
assert re.fullmatch(r"[a-zA-Z0-9_-]+", REPORT_TAG)
OUT = OUT_DIR / ("pilot.json" if ONLY_IDS else f"q{ONLY_QUESTION}_{REPORT_TAG}_results.json" if ONLY_QUESTION else "all_204_results.json")

ALLOWED = {
    1: [0.0, 1.0, 2.0],
    2: [0.0, 1.0, 1.5, 2.0],
    3: [0.0, 0.25, 0.5, 0.75, 1.0],
    4: [0.0, 0.5, 1.0],
    5: [0.0, 0.25, 0.5, 1.0],
    6: [0.0, 1.0],
}

ANSWER_KEYS = {
    1: "Row-major จัดเก็บหรือเข้าถึงข้อมูลตามแถว (แนวนอน) ส่วน Column-major ตามคอลัมน์ (แนวตั้ง)",
    2: "เมื่อข้อมูลใหญ่ O(n log n) มีจำนวนขั้นตอนเติบโตช้ากว่า O(n²) เช่น Merge Sort หรือ Heap Sort; O(n²) เช่น Bubble Sort หรือลูปซ้อนกัน",
    3: "Array ใช้ตำแหน่ง index และมักกำหนดขนาดล่วงหน้า ส่วน Linked List เชื่อมโหนดและปรับขนาดได้; ข้อดีคือเพิ่มลบและขยายขนาดได้ยืดหยุ่น ข้อเสียคือใช้ pointer เพิ่มและเข้าถึงตำแหน่งโดยตรงไม่ได้",
    4: "BST จากลำดับที่ให้: 9 เป็นราก มี 5 ทางซ้ายและ 16 ทางขวา; 16 มี 10 ซ้าย 76 ขวา; 10 มี 13 ขวา; 13 มี 11 ซ้าย 15 ขวา; 76 มี 58 ซ้าย 92 ขวา; 92 มี 80 ซ้าย 99 ขวา",
    5: "Prefix: + A * B - C / D * F 2; Postfix: A B C D F 2 * / - * +. พิจารณาวิธีทำและคำตอบสุดท้ายตามเกณฑ์",
    6: "แปลง General Tree ตาม Left-Child Right-Sibling: 1 ซ้าย 2; 2 ซ้าย 5 ขวา 3; 5 ขวา 6; 6 ขวา 7; 3 ขวา 4; 4 ซ้าย 8; 8 ขวา 9; 9 ขวา 10",
}

KEY_IMAGES = {
    4: ROOT / "public/answer-keys/q4_bst_ground_truth.png",
    5: ROOT / "public/answer-keys/q5-infix-prefix-postfix-answer-key.png",
    6: ROOT / "public/answer-keys/q6-lcrs-binary-tree-answer-key.png",
}
QUESTION_IMAGE = ROOT / "ชุดข้อสอบใหม่/โจทphoto3/LINE_ALBUM_โจทphoto6_260918_1.jpg"
IMAGE_DIRS = {
    4: ("photo_clean_ชุดที่1", "LINE_ALBUM_Photo1_260917_"),
    5: ("photo_clean_ชุดที่2", "LINE_ALBUM_Photo2.1_260918_"),
    6: ("photo_clean_ชุดที่3", "LINE_ALBUM_Photo2.2_260918_"),
}


def image_info(path):
    raw = path.read_bytes()
    return {"path": str(path.relative_to(ROOT)), "sha256": hashlib.sha256(raw).hexdigest()}, raw


def load_inputs():
    workbook = openpyxl.load_workbook(BOOK, read_only=False, data_only=False)
    sheet = workbook["ชุดข้อสอบ_dataset"]
    rubric_sheet = workbook["Exam_Rubrics"]
    rubrics = {}
    max_scores = {}
    for row in range(5, 13):
        question = int(rubric_sheet.cell(row, 1).value)
        max_scores[question] = float(rubric_sheet.cell(row, 3).value)
        rubrics.setdefault(question, []).append({
            "name": str(rubric_sheet.cell(row, 4).value),
            "score": float(rubric_sheet.cell(row, 5).value),
            "description": str(rubric_sheet.cell(row, 6).value),
            "allowed_scores": ([0.0, 1.0] if question == 1 else [0.0, 0.25, 0.5] if question == 3 else ALLOWED[question]),
        })
    assert set(rubrics) == set(range(1, 7))
    for question, criteria in rubrics.items():
        assert abs(sum(x["score"] for x in criteria) - max_scores[question]) < 1e-8

    items = []
    for row in range(6, 210):
        sample_id = sheet.cell(row, 1).value
        question = int(sheet.cell(row, 2).value)
        index = (row - 6) % 34 + 1
        assert sample_id == f"DS-{row-5:03d}" and question == (row - 6) // 34 + 1
        answer_type = sheet.cell(row, 5).value
        assert answer_type == ("text" if question <= 3 else "img")
        answer = sheet.cell(row, 6).value
        image_path = None
        if question <= 3:
            assert isinstance(answer, str) and answer.strip()
        else:
            directory, prefix = IMAGE_DIRS[question]
            image_path = ROOT / "ชุดข้อสอบใหม่" / directory / f"{prefix}{index}.jpg"
            assert image_path.is_file(), image_path
            formula = str(answer)
            assert re.search(r'HYPERLINK\("([^\"]+)', formula)
            assert image_path.name in formula, (sample_id, formula)
        question_text = str(sheet.cell(row, 4).value or "")
        if question == 6:
            assert "HYPERLINK" in question_text and QUESTION_IMAGE.is_file()
            question_text = "จงแปลง General Tree ในภาพโจทย์เป็น Binary Tree ตามหลัก Left-Child Right-Sibling"
        items.append({
            "row": row, "sample_id": sample_id, "question_no": question,
            "question_text": question_text, "answer_text": answer if question <= 3 else "",
            "answer_image": image_path, "human_score": float(sheet.cell(row, 7).value),
            "previous_ai_score": float(sheet.cell(row, 8).value),
        })
    workbook.close()
    assert len(items) == 204 and Counter(x["question_no"] for x in items) == {q: 34 for q in range(1, 7)}
    return items, rubrics, max_scores


def valid_result(result, question):
    if not isinstance(result, dict):
        return False
    metrics = result.get("metrics") or {}
    return (result.get("score") in ALLOWED[question]
            and result.get("confidence") in {"high", "medium", "low"}
            and bool(str(result.get("teacher_feedback") or "").strip())
            and bool(str(result.get("student_feedback") or "").strip())
            and metrics.get("provider") == "gemini"
            and metrics.get("model") == GEMINI_MODEL
            and "answer_word_count" in metrics)


async def main():
    assert GEMINI_MODEL == "gemini-3.8-flash", f"Expected gemini-3.8-flash, got {GEMINI_MODEL}"
    all_items, rubrics, max_scores = load_inputs()
    if ONLY_QUESTION:
        items = [x for x in all_items if x["question_no"] == ONLY_QUESTION]
        assert len(items) == 34
    elif ONLY_IDS:
        items = [x for x in all_items if x["sample_id"] in ONLY_IDS]
        assert len(items) == len(ONLY_IDS)
    else:
        items = all_items
    for path in KEY_IMAGES.values():
        assert path.is_file(), path
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    source_hash = hashlib.sha256(BOOK.read_bytes()).hexdigest()
    prompt_hash = hashlib.sha256((ROOT / "server/services/openai_grading.py").read_bytes()
                                 + (ROOT / "server/services/gemini_grading.py").read_bytes()).hexdigest()
    report = {
        "started_at": datetime.now(timezone.utc).isoformat(), "provider": "gemini", "model": GEMINI_MODEL,
        "thinking_level": "low", "prompt_version": PROMPT_VERSION,
        "grading_function": "server.services.gemini_grading.score_with_gemini",
        "source_workbook": str(BOOK.relative_to(ROOT)), "source_sha256": source_hash,
        "grading_code_sha256": prompt_hash,
        "blind_to_human_score": True, "rubrics": rubrics, "answer_keys": ANSWER_KEYS,
        "allowed_scores": ALLOWED, "results": [],
    }
    if OUT.exists() and not ONLY_IDS:
        existing = json.loads(OUT.read_text(encoding="utf-8"))
        assert (existing["source_sha256"] == source_hash and existing["model"] == GEMINI_MODEL
                and existing.get("grading_code_sha256") == prompt_hash)
        report = existing
    complete = {x["sample_id"] for x in report["results"] if x.get("success")}
    semaphore = asyncio.Semaphore(3)

    async def grade(item):
        async with semaphore:
            question = item["question_no"]
            kwargs = dict(question_text=item["question_text"], answer_text=item["answer_text"],
                          max_score=max_scores[question], answer_key=ANSWER_KEYS[question],
                          rubrics=rubrics[question], allowed_scores=ALLOWED[question],
                          strict_rubric_enforcement=question >= 4)
            image_metadata = {}
            if item["answer_image"]:
                image_metadata["answer_image"], raw = image_info(item["answer_image"])
                kwargs.update(image_bytes_list=[raw], image_mime_list=["image/jpeg"])
                image_metadata["answer_key_image"], key_raw = image_info(KEY_IMAGES[question])
                kwargs.update(answer_key_image_bytes_list=[key_raw], answer_key_image_mime_list=["image/png"])
            if question == 6:
                image_metadata["question_image"], q_raw = image_info(QUESTION_IMAGE)
                kwargs.update(q_image_bytes_list=[q_raw], q_image_mime_list=["image/jpeg"])
            result = None
            for attempt in range(3):
                result = await score_with_gemini(**kwargs)
                if valid_result(result, question):
                    break
                if attempt < 2:
                    await asyncio.sleep(3 * (attempt + 1))
            success = valid_result(result, question)
            return {
                "sample_id": item["sample_id"], "row": item["row"], "question_no": question,
                "answer_type": "text" if question <= 3 else "img",
                "answer_sha256": hashlib.sha256(item["answer_text"].encode("utf-8")).hexdigest() if question <= 3 else None,
                "image_inputs": image_metadata, "human_score": item["human_score"],
                "previous_ai_score": item["previous_ai_score"], "success": success, "result": result,
            }

    for question in range(1, 7):
        todo = [x for x in items if x["question_no"] == question and x["sample_id"] not in complete]
        print(f"Q{question}: {len(todo)} remaining", flush=True)
        tasks = [asyncio.create_task(grade(item)) for item in todo]
        for task in asyncio.as_completed(tasks):
            entry = await task
            report["results"] = [x for x in report["results"] if x["sample_id"] != entry["sample_id"]]
            report["results"].append(entry)
            report["results"].sort(key=lambda x: x["row"])
            OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
            print(entry["sample_id"], "score", entry["result"]["score"],
                  "confidence", entry["result"]["confidence"], "valid", entry["success"], flush=True)
        if any(not x["success"] for x in report["results"] if x["question_no"] == question):
            print(f"Q{question} contains failures; retry them before using this report", flush=True)
    report["completed_at"] = datetime.now(timezone.utc).isoformat()
    report["summary"] = {}
    for question in range(1, 7):
        rows = [x for x in report["results"] if x["question_no"] == question and x["success"]]
        report["summary"][str(question)] = {
            "completed": len(rows), "total": 34,
            "exact": sum(x["human_score"] == x["result"]["score"] for x in rows),
            "changed_from_previous_ai": sum(x["previous_ai_score"] != x["result"]["score"] for x in rows),
            "mae": (sum(abs(x["human_score"] - x["result"]["score"]) for x in rows) / len(rows)) if rows else None,
        }
    OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report["summary"], ensure_ascii=False), flush=True)
    print("REPORT", OUT, flush=True)


if __name__ == "__main__":
    asyncio.run(main())
