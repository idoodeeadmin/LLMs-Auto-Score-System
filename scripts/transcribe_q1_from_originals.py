"""Apply photo-checked Q1 transcriptions to only F6:F39 of the dataset workbook.

The workbook contains embedded photos; patch its worksheet XML rather than
round-tripping the package through an editor that could discard drawings.
"""

from pathlib import Path
from tempfile import NamedTemporaryFile
import json
import os
import re
import zipfile
from xml.etree import ElementTree as ET

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx"
PHOTOS = ROOT / "public/photo_q1_original"
OCR = {
    item["filename"]: item["verbatim_student_answer"]
    for item in json.loads(
        (ROOT / "artifacts/q1_42_precision_extracted.json").read_text(encoding="utf-8")
    )
}

# Each override below was checked against the scored, pre-cleaning photo.
# Transcribe visible words, numbers and symbols; do not narrate drawings.
OVERRIDES = {
    1: "Row จะนับเป็นแถวจากซ้ายไปขวา\n0 1 2 3\n4 5 6 7\n8 9 10 11\nส่วน Column จากบนลงล่าง\n0 4 8\n1 5 9\n2 6 10\n3 7 11",
    2: "ต่างกันที่การเรียงลำดับข้อมูล เช่น\n1 2 3 4\n5 6 7 8\nvs\n1 3 5 7\n2 4 6 8",
    4: "Row-major คือ Array ที่มีลักษณะเป็นแนวนอน\n0 1 2 3 4\nColumn-major คือ Array ที่เก็บข้อมูลในลักษณะแนวตั้ง\n0\n1\n2\n3\n4\nทั้งสองต่างกันที่เก็บข้อมูลลักษณะแถวและคอลัมน์",
    5: "คำนวณ ตำแหน่ง แนวตั้ง แนวนอน",
    8: "Row จะไล่จาก ซ้ายไปขวา\nColumn จะไล่จากบนลงล่าง\nสูตรที่ใช้ต่างกันเล็กน้อย\nหากใช้ Array 2D จะต้องคำนวณ Index 2 ตัว\nเพื่อให้ Row, Column ทำงานควบคู่กัน\nRow →\nCol ↓",
    9: "Row-major เริ่มเรียงจากแถวแกน x ก่อน เช่น\n1, 2, 3, 4, 5  R1\n6, 7, 8, 9, 10  R2\nColumn-major เริ่มเรียงจากแกน y ก่อน เช่น\nC1  C2  C3  C4  C5\n1   2   3   4   5\n6   7   8   9   10",
    10: "Row-major จะเรียกแถวก่อนหลัก\nColumn-major จะเรียกหลักก่อนแถว",
    11: "ถ้าเป็น Row-major จะนำข้อมูลจากแถว\nถ้าเป็น Column-major จะนำข้อมูลจากด้านบนทั้งหมด ซึ่งไม่เหมือนกัน",
    14: "Row-major และ Column-major ต่างกันตรงการคำนวณหา Address เนื่องจาก Row major จัดเก็บเป็น Row และ Column-major จัดเก็บ Col ซึ่งคือ Array 2D",
    15: "Row-major → 1D นับเป็นแถว\n2D นับ แถว,หลัก\nนับเป็นแนวนอน ซ้ายไปขวา\nColumn-major → 1D นับเป็นหลัก\n2D นับเป็น หลัก,แถว\nนับแนวตั้ง บนลงล่าง\nA[2:3]",
    17: "Arr[3][4]\nRow จะเรียงแถวก่อน\nRow →\n1 2 3 4\n5 6 7 8\n9 10 11 12\nColumn จะเรียงหลักก่อน\nCol ↓\n1 4 7 10\n2 5 8 11\n3 6 9 12",
    19: "Row major คือ การคำนวณจาก แถว\nColumn-major คือการคำนวณจากหลักตัวตั้ง\nความแตกต่างคือ การจัดเนื้อที่ของค่าจากแนวนอน และแนวตั้ง",
    20: "Row-major เป็นแนวนอน\ncolumn-major เป็นแนวตั้ง ทั้งคู่ก็ไม่เหมือนกัน",
    21: "Row-Major vs Colum-major\nRow-major คือ การนำ Array แบบการกำหนดตาม\nต่างจาก การนำ Addres แบบ Column-major ที่มีการนำ Addes แบบ\nการดูและกำหนดแบบ Colum",
    22: "Row-major\n0 1 2 3\n4 5 6 7\n8 9 10 11\na(2,1) = 9\nColumn-major\n0 3 6 9\n1 4 7 10\n2 5 8 11\na(2,1) = 7\nRow จะนับแถวก่อนหลัก a(แถว,หลัก)\nColumn จะนับหลักก่อนแถว a(หลัก,แถว)",
    23: "Row-major\nเป็นการจัดแถวในการเก็บ Array เป็นหลัก เช่น\nA[2,1] = 8\nColumn-major\nเป็นการจัดหลัก (Column) ของอาเรย์เป็นหลัก เช่น A[2,1] = 5\n4 2 6 9\n0 1 5 3\n2 8 11 12",
    24: "Row major เป็นการจัดเก็บ Array แบบ แถว แต่ใช้กับ Array 1D\nColumn major เป็นการจัดเก็บ Array ร่วมกับ Row คือ Array 2D",
    25: "Row-major จะนับจากบนลงล่าง เช่น [0][0], [1][0]\n[0][1], [1][1]\n[0][2], [1][2]\nColumn-major จะนับจากซ้ายไปขวา เช่น [0][0], [0][1], [0][2]\n[1][0], [1][1], [1][2]",
    29: "Row-major\n1 2 3 4\n5 6 7 8\n9 10 11 12\nนับจากซ้ายไปขวา\nColumn-major\n1 4 7 10\n2 5 8 11\n3 6 9 12\nนับจากบนลงล่าง",
    30: "Row-major จะเก็บข้อมูลในแนวนอนก่อนจากซ้ายไปขวาแล้วแถวถัดมา\nColumn-major จะเก็บข้อมูลในแนวตั้งก่อนจากบนลงล่างแล้วซ้ายไปขวา",
    31: "Row-major หาแนวนอน\nColumn-major หาแนวตั้ง\nแตกต่างกันที่สูตร\nRow-major = B + W[(U2-L2+1)(i-L1) + (j-L2)]\nColumn-major = B + W[(U1-L1+1)(j-L2) + (i-L1)]",
    32: "แต่ละภาษาจะใช้ Array ต่างกันออกไป เช่น\njava จะใช้ Row-major\nRow-major คือ A[แถว][หลัก] = A[0][1] = -7\nColumn-major คือ A[หลัก][แถว] = A[0][1] = 4\n6 -7 15\n4 1 -9\n8 9 27",
    33: "Row-major คือการดูแถวก่อนค่อยดูคอลัมน์ หรือแถวคือแกนหลัก คอลัมน์คือแกนรอง (แถว,คอลัมน์)\nColumn-major คือการดูคอลัมน์ก่อนค่อยดูแถว หรือคอลัมน์คือแกนหลัก แถวคือแกนรอง (คอลัมน์,แถว)",
    34: "row-major => บังคับอาเรย์เป็นแถวในแนวนอน\nวิธีการคำนวณคือ B + W[(U2-L2+1)(I-L1) + (J-L2)]\ncolumn-major => บังคับอาเรย์เป็นแนวตั้ง\nวิธีการคำนวณคือ B + W[(U1-L1+1)(J-L2) + (I-L1)]\nซึ่งต่างกันที่การเก็บข้อมูลหรือเรียงของสูตรในการคำนวณค่าดัชนี\n(U1-L1+1)(I-L1)   (U2-L2+1)(J-L2)\n↓ Col                 ↓ ROW",
}

book = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
sheet = book["ชุดข้อสอบ_dataset"]
photos = sorted(PHOTOS.glob("DS-*_IMG_*.jpg"))
assert len(photos) == 34
new_text = {}
for number, photo in enumerate(photos, 1):
    sample_id = f"DS-{number:03d}"
    assert photo.name.startswith(sample_id + "_")
    assert sheet.cell(number + 5, 1).value == sample_id
    original_name = re.search(r"(IMG_\d+\.jpg)$", photo.name).group(1)
    text = OVERRIDES.get(number, OCR[original_name])
    assert text.strip() and "คะแนน" not in text
    new_text[number + 5] = text
book.close()

NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ET.register_namespace("", NS)
with zipfile.ZipFile(BOOK) as src:
    xml = ET.fromstring(src.read("xl/worksheets/sheet1.xml"))
    cells = {cell.attrib.get("r"): cell for cell in xml.iter(f"{{{NS}}}c")}
    for row, answer in new_text.items():
        cell = cells[f"F{row}"]
        cell.attrib["t"] = "inlineStr"
        for child in list(cell):
            if child.tag in {f"{{{NS}}}v", f"{{{NS}}}is", f"{{{NS}}}f"}:
                cell.remove(child)
        inline = ET.SubElement(cell, f"{{{NS}}}is")
        text_node = ET.SubElement(inline, f"{{{NS}}}t")
        text_node.attrib["{http://www.w3.org/XML/1998/namespace}space"] = "preserve"
        text_node.text = answer
    updated_xml = ET.tostring(xml, encoding="utf-8", xml_declaration=True)
    with NamedTemporaryFile(dir=BOOK.parent, suffix=".xlsx", delete=False) as temp:
        temp_name = temp.name
    with zipfile.ZipFile(temp_name, "w") as dst:
        for entry in src.infolist():
            data = updated_xml if entry.filename == "xl/worksheets/sheet1.xml" else src.read(entry.filename)
            dst.writestr(entry, data)

try:
    check = openpyxl.load_workbook(temp_name, read_only=True, data_only=True)
    target = check["ชุดข้อสอบ_dataset"]
    assert all(target.cell(row, 6).value == answer for row, answer in new_text.items())
    check.close()
    with zipfile.ZipFile(temp_name) as verify:
        assert len([name for name in verify.namelist() if name.startswith("xl/media/")]) == 34
    os.replace(temp_name, BOOK)
finally:
    if os.path.exists(temp_name):
        os.unlink(temp_name)

print(f"Updated {len(new_text)} Q1 answers in {BOOK}")
