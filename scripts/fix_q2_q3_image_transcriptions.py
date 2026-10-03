"""Correct verified Q2/Q3 transcriptions while preserving workbook media."""

import json
import os
import shutil
import zipfile
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET

import openpyxl


ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx"
BACKUP = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset_before_q2_q3_text_fixes.xlsx"
MANIFEST = ROOT / "docs_and_tests/q2_q3_review/corrections.json"
updates = {int(k): v for k, v in json.loads(MANIFEST.read_text(encoding="utf-8")).items()}

workbook = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
sheet = workbook["ชุดข้อสอบ_dataset"]
for number, answer in updates.items():
    assert 35 <= number <= 102
    assert sheet.cell(number + 5, 1).value == f"DS-{number:03d}"
    assert answer.strip() and answer != sheet.cell(number + 5, 6).value
workbook.close()

if not BACKUP.exists():
    shutil.copy2(BOOK, BACKUP)

ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ET.register_namespace("", ns)
with zipfile.ZipFile(BOOK) as source:
    root = ET.fromstring(source.read("xl/worksheets/sheet1.xml"))
    cells = {cell.attrib["r"]: cell for cell in root.iter(f"{{{ns}}}c")}
    for number, answer in updates.items():
        cell = cells[f"F{number + 5}"]
        for child in list(cell):
            if child.tag in {f"{{{ns}}}v", f"{{{ns}}}is", f"{{{ns}}}f"}:
                cell.remove(child)
        cell.attrib["t"] = "inlineStr"
        inline = ET.SubElement(cell, f"{{{ns}}}is")
        text = ET.SubElement(inline, f"{{{ns}}}t")
        text.attrib["{http://www.w3.org/XML/1998/namespace}space"] = "preserve"
        text.text = answer
    xml = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with NamedTemporaryFile(dir=BOOK.parent, suffix=".xlsx", delete=False) as temp:
        temp_name = temp.name
    with zipfile.ZipFile(temp_name, "w") as target:
        for entry in source.infolist():
            target.writestr(entry, xml if entry.filename == "xl/worksheets/sheet1.xml" else source.read(entry.filename))

try:
    with zipfile.ZipFile(temp_name) as archive:
        with zipfile.ZipFile(BACKUP) as original:
            for name in original.namelist():
                if name != "xl/worksheets/sheet1.xml":
                    assert archive.read(name) == original.read(name), name
    check = openpyxl.load_workbook(temp_name, read_only=True, data_only=True)
    try:
        sheet = check["ชุดข้อสอบ_dataset"]
        assert all(sheet.cell(number + 5, 6).value == value for number, value in updates.items())
    finally:
        check.close()
    os.replace(temp_name, BOOK)
finally:
    if os.path.exists(temp_name):
        os.unlink(temp_name)

print(f"Corrected {len(updates)} answers; backup created")
