"""Clarify partial credit for Q2 without changing workbook media or scores."""

import os
import shutil
import zipfile
from pathlib import Path
from tempfile import NamedTemporaryFile
from xml.etree import ElementTree as ET

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx"
BACKUP = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset_before_q2_rubric_clarification.xlsx"
NOTE = (
    "\n\nคำอธิบายเพิ่มเติมในการแยกคะแนนบางส่วน: การเรียกซ้ำ (recursive) เพียงอย่างเดียว"
    "ไม่ได้หมายความว่าอัลกอริทึมนั้นเป็น O(n log n) และลูปซ้อนต้องพิจารณาจำนวนรอบตามขนาดข้อมูลด้วย "
    "อย่างไรก็ตาม หากคำตอบยกตัวอย่างโค้ดเรียกซ้ำควบคู่กับลูปซ้อนเพื่อเปรียบเทียบวิธีทำงาน "
    "แม้จับคู่ตัวอย่างกับระดับ Big O ไม่แม่น แต่ยังสื่อความพยายามเปรียบเทียบจำนวนรอบ/ขั้นตอน "
    "ให้พิจารณา 1.50 คะแนนตามระดับอนุโลม ไม่ตัดเป็น 0.00 เพียงเพราะจัดประเภทตัวอย่างผิด "
    "ส่วน 0.00 ใช้เมื่อไม่พบทั้งเหตุผลเปรียบเทียบและตัวอย่างที่เกี่ยวข้องกับโจทย์เลย"
)

wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
old = wb["Exam_Rubrics"].cell(7, 6).value
assert "ระดับคะแนน 0.0, 1.0, 1.5, 2.0" in old
assert "คำอธิบายเพิ่มเติมในการแยกคะแนนบางส่วน" not in old
wb.close()
if not BACKUP.exists():
    shutil.copy2(BOOK, BACKUP)

ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
ET.register_namespace("", ns)
with zipfile.ZipFile(BOOK) as source:
    root = ET.fromstring(source.read("xl/worksheets/sheet3.xml"))
    cell = next(c for c in root.iter(f"{{{ns}}}c") if c.attrib.get("r") == "F7")
    for child in list(cell):
        if child.tag in {f"{{{ns}}}v", f"{{{ns}}}is", f"{{{ns}}}f"}:
            cell.remove(child)
    cell.attrib["t"] = "inlineStr"
    inline = ET.SubElement(cell, f"{{{ns}}}is")
    text = ET.SubElement(inline, f"{{{ns}}}t")
    text.attrib["{http://www.w3.org/XML/1998/namespace}space"] = "preserve"
    text.text = old + NOTE
    xml = ET.tostring(root, encoding="utf-8", xml_declaration=True)
    with NamedTemporaryFile(dir=BOOK.parent, suffix=".xlsx", delete=False) as temp:
        temp_name = temp.name
    with zipfile.ZipFile(temp_name, "w") as target:
        for entry in source.infolist():
            target.writestr(entry, xml if entry.filename == "xl/worksheets/sheet3.xml" else source.read(entry.filename))

try:
    with zipfile.ZipFile(temp_name) as current, zipfile.ZipFile(BACKUP) as original:
        for name in original.namelist():
            if name != "xl/worksheets/sheet3.xml":
                assert current.read(name) == original.read(name), name
    check = openpyxl.load_workbook(temp_name, read_only=True, data_only=True)
    try:
        assert check["Exam_Rubrics"].cell(7, 6).value == old + NOTE
    finally:
        check.close()
    os.replace(temp_name, BOOK)
finally:
    if os.path.exists(temp_name):
        os.unlink(temp_name)
print("Clarified Q2 rubric; scores unchanged")
