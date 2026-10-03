"""Regrade only corrected Q2/Q3 text answers with the production template."""

import asyncio
import hashlib
import json
import os
import sys
import re
import subprocess
import types
from datetime import datetime, timezone
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

from dotenv import load_dotenv

load_dotenv(ROOT / ".env")

# The bundled Python can read the workbook, but the repository venv's compiled
# pydantic_core targets another Python version. Supply the grading function's
# word-count dependency without importing FastAPI/Pydantic.
policy = types.ModuleType("server.exam_policy")
policy.MAX_ANSWER_WORDS = 300
def count_answer_words(value):
    try:
        node = Path.home() / ".cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node.exe"
        result = subprocess.run([str(node), str(ROOT / "server/services/word_counter.mjs")],
                                input=value, capture_output=True, text=True, encoding="utf-8", timeout=3, check=True)
        return int(result.stdout.strip())
    except Exception:
        return len(re.findall(r"[\u0E00-\u0E7F]+|[a-zA-Z0-9_]+", value))
policy.count_answer_words = count_answer_words
sys.modules["server.exam_policy"] = policy

from server.services.openai_grading import OPENAI_MODEL, PROMPT_VERSION, score_with_openai

BOOK = ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx"
strict_one = os.getenv("STRICT_DS057") == "1"
pilot_51 = os.getenv("PILOT_DS051") == "1"
OUT = ROOT / ("docs_and_tests/q2_q3_review/regrade_ds051_clarified_rubric.json" if pilot_51
              else "docs_and_tests/q2_q3_review/regrade_ds057_strict.json" if strict_one
              else "docs_and_tests/q2_q3_review/regrade_corrected_results.json")
IDS = (51,) if pilot_51 else (57,) if strict_one else (51, 55, 57, 89, 101)


async def main():
    wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
    sheet = wb["ชุดข้อสอบ_dataset"]
    rubric_sheet = wb["Exam_Rubrics"]
    rubric = {
        2: [{"name": rubric_sheet.cell(7, 4).value, "score": 2.0, "description": rubric_sheet.cell(7, 6).value,
             "allowed_scores": [0.0, 1.0, 1.5, 2.0]}],
        3: [{"name": rubric_sheet.cell(row, 4).value, "score": 0.5, "description": rubric_sheet.cell(row, 6).value,
             "allowed_scores": [0.0, 0.25, 0.5]} for row in (8, 9)],
    }
    answer_key = {
        2: "O(n log n) เติบโตช้ากว่า O(n^2) เมื่อข้อมูลมีขนาดใหญ่ เช่น Merge Sort หรือ Quick Sort เทียบกับ Bubble Sort หรือ Selection Sort",
        3: "Array ต้องกำหนดขนาดล่วงหน้าและใช้ตำแหน่ง index ส่วน Linked List ปรับขนาดได้และใช้โหนดเชื่อมกัน ข้อดีของ Linked List คือเพิ่มหรือลบข้อมูลได้โดยไม่ต้องจองขนาดคงที่ ข้อเสียคือใช้หน่วยความจำสำหรับตัวชี้และเข้าถึงตำแหน่งโดยตรงไม่ได้",
    }
    items = []
    for number in IDS:
        row = number + 5
        q = int(sheet.cell(row, 2).value)
        assert sheet.cell(row, 1).value == f"DS-{number:03d}" and q in (2, 3)
        items.append({"sample_id": f"DS-{number:03d}", "row": row, "question_no": q,
                      "question": sheet.cell(row, 4).value, "answer": sheet.cell(row, 6).value,
                      "prior_ai_score": sheet.cell(row, 8).value})
    report = {"timestamp": datetime.now(timezone.utc).isoformat(), "model": OPENAI_MODEL,
              "prompt_version": PROMPT_VERSION, "source_sha256": hashlib.sha256(BOOK.read_bytes()).hexdigest(),
              "rubrics": rubric, "items": []}
    wb.close()
    for item in items:
        q = item["question_no"]
        allowed = [0.0, 1.0, 1.5, 2.0] if q == 2 else [0.0, 0.25, 0.5, 0.75, 1.0]
        result = await score_with_openai(question_text=item["question"], answer_text=item["answer"],
                                         max_score=2.0 if q == 2 else 1.0, answer_key=answer_key[q],
                                         rubrics=rubric[q], allowed_scores=allowed,
                                         strict_rubric_enforcement=strict_one or pilot_51)
        success = (not result.get("metrics", {}).get("manual_review_required", True)
                   and float(result.get("score", -1)) in allowed
                   and bool(result.get("teacher_feedback", "").strip())
                   and bool(result.get("student_feedback", "").strip()))
        report["items"].append({**item, "success": success, "result": result})
        OUT.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
        print(item["sample_id"], "new", result.get("score"), "old", item["prior_ai_score"],
              "success", success, flush=True)
    print("report", OUT)


if __name__ == "__main__":
    asyncio.run(main())
