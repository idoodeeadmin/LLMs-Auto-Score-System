import sys
import os
import json
from pathlib import Path
import openpyxl

sys.stdout.reconfigure(encoding="utf-8")

ROOT = Path(__file__).resolve().parents[1]
RESULTS_JSON = ROOT / "public/q4_regrade_visual_key_results.json"

with open(RESULTS_JSON, "r", encoding="utf-8") as f:
    regrade_data = json.load(f)

results_map = {r["sample_id"]: r for r in regrade_data["results"]}
print(f"Loaded {len(results_map)} regrade results from {RESULTS_JSON.name}")

target_files = [
    ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx",
    ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset_formatted.xlsx",
    ROOT / "ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset_จัดรูปแบบใหม่.xlsx",
]

report_file = ROOT / "artifacts/LLM_AutoScore_Benchmark_204_Final_Report.xlsx"
if report_file.exists():
    target_files.append(report_file)

for excel_path in target_files:
    if not excel_path.exists():
        print(f"Skipping missing file: {excel_path}")
        continue

    print(f"\nProcessing: {excel_path.relative_to(ROOT)}")
    wb = openpyxl.load_workbook(excel_path)
    
    # Identify sheet name
    sheet_name = None
    if "ชุดข้อสอบ_dataset" in wb.sheetnames:
        sheet_name = "ชุดข้อสอบ_dataset"
    elif "Dataset_QWK_204" in wb.sheetnames:
        sheet_name = "Dataset_QWK_204"
    else:
        sheet_name = wb.sheetnames[0]
        
    ws = wb[sheet_name]
    print(f"  Working on sheet: {sheet_name}")

    # Determine column indexes from header row (row 4)
    sid_col = 1
    qno_col = 2
    ai_score_col = 8
    ai_conf_col = 9
    ai_fb_col = 10

    for c in range(1, 15):
        val = str(ws.cell(row=4, column=c).value or "").strip().lower()
        if val == "sample_id":
            sid_col = c
        elif val == "question_no":
            qno_col = c
        elif val == "ai_score":
            ai_score_col = c
        elif val in ("ai_confidence", "confidence"):
            ai_conf_col = c
        elif val in ("ai_feedback", "feedback"):
            ai_fb_col = c

    updated_count = 0
    for r in range(5, ws.max_row + 1):
        q_no = ws.cell(row=r, column=qno_col).value
        sid = str(ws.cell(row=r, column=sid_col).value or "").strip()
        
        if q_no == 4 and sid in results_map:
            res = results_map[sid]
            new_score = float(res["ai"])
            new_conf = str(res.get("confidence", "high"))
            tfb = res.get("teacher_feedback", "").strip()
            sfb = res.get("student_feedback", "").strip()
            
            combined_fb = ""
            if tfb and sfb:
                combined_fb = f"[สำหรับผู้สอน]\n{tfb}\n\n[สำหรับนักเรียน]\n{sfb}"
            elif tfb:
                combined_fb = f"[สำหรับผู้สอน]\n{tfb}"
            elif sfb:
                combined_fb = f"[สำหรับนักเรียน]\n{sfb}"

            ws.cell(row=r, column=ai_score_col).value = new_score
            ws.cell(row=r, column=ai_conf_col).value = new_conf
            ws.cell(row=r, column=ai_fb_col).value = combined_fb
            updated_count += 1

    wb.save(excel_path)
    print(f"  Successfully updated {updated_count} student answers in {excel_path.name}")

print("\nAll target Excel files updated successfully!")
