import os
import sys
import math
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

sys.stdout.reconfigure(encoding='utf-8')

SOURCE_PATH = os.path.join("ชุดข้อสอบใหม่", "ชุดข้อสอบ_dataset.xlsx")
NEW_PATH_TH = os.path.join("ชุดข้อสอบใหม่", "ชุดข้อสอบ_dataset_จัดรูปแบบใหม่.xlsx")
NEW_PATH_EN = os.path.join("ชุดข้อสอบใหม่", "ชุดข้อสอบ_dataset_formatted.xlsx")

def create_formatted_dataset():
    print(f"Reading from {SOURCE_PATH}...")
    src_wb = openpyxl.load_workbook(SOURCE_PATH, data_only=True)
    
    new_wb = openpyxl.Workbook()
    new_wb.remove(new_wb.active) # remove default sheet

    # Styling definitions
    font_name = "Segoe UI"
    header_font = Font(name=font_name, size=10, bold=True, color="FFFFFF")
    subhead_font = Font(name=font_name, size=9, italic=True, color="E2E8F0")
    data_font = Font(name=font_name, size=10, color="1E293B")
    data_font_bold = Font(name=font_name, size=10, bold=True, color="0F172A")
    
    header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Royal Dark Blue
    subhead_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid") # Vibrant Navy Blue
    zebra_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")   # Subtle soft grey-blue
    
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1")
    )

    col_widths = {
        1: 14,   # A: sample_id
        2: 12,   # B: question_no
        3: 16,   # C: question_type
        4: 55,   # D: question_content
        5: 16,   # E: answer_type
        6: 72,   # F: student_answer
        7: 16,   # G: human_score
        8: 14,   # H: ai_score
        9: 16,   # I: ai_confidence
        10: 88   # J: ai_feedback
    }

    # 1. Main Sheet: ชุดข้อสอบ_dataset
    ws_src = src_wb['ชุดข้อสอบ_dataset']
    ws_new = new_wb.create_sheet(title="ชุดข้อสอบ_dataset")
    ws_new.views.sheetView[0].showGridLines = True

    # Set column widths
    for col_idx, width in col_widths.items():
        col_letter = get_column_letter(col_idx)
        ws_new.column_dimensions[col_letter].width = width

    # Row 1-3 empty or title
    for r in range(1, 4):
        ws_new.row_dimensions[r].height = 15.0

    # Row 4: English Column Keys
    ws_new.row_dimensions[4].height = 32.0
    for c in range(1, 11):
        cell = ws_new.cell(row=4, column=c, value=ws_src.cell(row=4, column=c).value)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Row 5: Thai Descriptions
    ws_new.row_dimensions[5].height = 28.0
    for c in range(1, 11):
        cell = ws_new.cell(row=5, column=c, value=ws_src.cell(row=5, column=c).value)
        cell.font = subhead_font
        cell.fill = subhead_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = thin_border

    # Data Rows (6 to max_row)
    for r in range(6, ws_src.max_row + 1):
        sid = ws_src.cell(row=r, column=1).value
        if not sid:
            continue
        
        q_no = ws_src.cell(row=r, column=2).value
        q_type = ws_src.cell(row=r, column=3).value or "text"
        q_content = ws_src.cell(row=r, column=4).value
        ans_type = ws_src.cell(row=r, column=5).value or "text"
        std_ans = ws_src.cell(row=r, column=6).value
        h_score = ws_src.cell(row=r, column=7).value
        ai_score = ws_src.cell(row=r, column=8).value
        ai_conf = ws_src.cell(row=r, column=9).value
        ai_feedback = ws_src.cell(row=r, column=10).value

        # Calculate student index (1-34)
        std_idx = ((r - 6) % 34) + 1

        # Populate missing values for image questions
        if not q_content and q_no == 6:
            q_content = "จงแปลง tree ต่อไปนี้ให้เป็น Binary Tree (1 คะแนน)\n[รูปภาพโจทย์: โจทphoto3/LINE_ALBUM_โจทphoto6_260918_1.jpg]"

        if ans_type == "img":
            if q_no == 4:
                std_ans = f'=HYPERLINK("photo_clean_ชุดที่1/LINE_ALBUM_Photo1_260917_{std_idx}.jpg", "เปิดภาพคำตอบที่ลบคะแนนแล้ว")'
            elif q_no == 5:
                std_ans = f'=HYPERLINK("photo_clean_ชุดที่2/LINE_ALBUM_Photo2.1_260918_{std_idx}.jpg", "เปิดภาพคำตอบที่ลบคะแนนแล้ว")'
            elif q_no == 6:
                std_ans = f'=HYPERLINK("photo_clean_ชุดที่3/LINE_ALBUM_Photo2.2_260918_{std_idx}.jpg", "เปิดภาพคำตอบที่ลบคะแนนแล้ว")'

        row_vals = [sid, q_no, q_type, q_content, ans_type, std_ans, h_score, ai_score, ai_conf, ai_feedback]

        # Calculate required lines for dynamic row height
        max_lines = 1
        for c_idx, w in [(4, col_widths[4]), (6, col_widths[6]), (10, col_widths[10])]:
            val_str = str(row_vals[c_idx - 1] or "")
            if not val_str:
                continue
            char_limit = max(20, int(w * 0.82))
            lines = 0
            for para in val_str.split("\n"):
                p_len = len(para)
                if p_len == 0:
                    lines += 1
                else:
                    lines += max(1, math.ceil(p_len / char_limit))
            max_lines = max(max_lines, lines)

        calc_height = max(28.0, max_lines * 22.0 + 10.0)
        ws_new.row_dimensions[r].height = round(calc_height, 1)

        is_zebra = (r % 2 == 1)
        for c_idx, val in enumerate(row_vals, 1):
            cell = ws_new.cell(row=r, column=c_idx, value=val)
            cell.font = data_font_bold if c_idx in [1, 7, 8] else data_font
            cell.border = thin_border
            if is_zebra:
                cell.fill = zebra_fill

            if c_idx in [1, 2, 3, 5, 7, 8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
            else:
                cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

    # 2. Copy Template_Example sheet if exists
    if "Template_Example" in src_wb.sheetnames:
        ws_tpl_src = src_wb["Template_Example"]
        ws_tpl_new = new_wb.create_sheet(title="Template_Example")
        ws_tpl_new.views.sheetView[0].showGridLines = True

        for col_idx, width in col_widths.items():
            col_letter = get_column_letter(col_idx)
            ws_tpl_new.column_dimensions[col_letter].width = width

        for r in range(1, ws_tpl_src.max_row + 1):
            for c in range(1, ws_tpl_src.max_column + 1):
                val = ws_tpl_src.cell(row=r, column=c).value
                cell = ws_tpl_new.cell(row=r, column=c, value=val)
                if r == 4:
                    cell.font = header_font
                    cell.fill = header_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    ws_tpl_new.row_dimensions[r].height = 30.0
                elif r == 5:
                    cell.font = subhead_font
                    cell.fill = subhead_fill
                    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                    ws_tpl_new.row_dimensions[r].height = 26.0
                elif r >= 6 and val is not None:
                    cell.font = data_font
                    cell.border = thin_border
                    if c in [1, 2, 3, 5, 7, 8, 9]:
                        cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
                    else:
                        cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
                    ws_tpl_new.row_dimensions[r].height = 45.0

    # Save both Thai and standard name for convenience
    new_wb.save(NEW_PATH_TH)
    new_wb.save(NEW_PATH_EN)
    print(f"Created new file 1: {NEW_PATH_TH}")
    print(f"Created new file 2: {NEW_PATH_EN}")

if __name__ == "__main__":
    create_formatted_dataset()
