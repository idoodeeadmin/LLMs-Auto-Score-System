import sys
import io
import re
import os
import openpyxl

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

wb = openpyxl.load_workbook('ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx')
sheet = wb['ชุดข้อสอบ_dataset']

for r in range(142, 152):
    sample_id = sheet.cell(r, 1).value
    val = sheet.cell(r, 6).value
    h_score = sheet.cell(r, 7).value
    match = re.search(r'"([^"]+)"', str(val))
    path = match.group(1) if match else str(val)
    
    # Try different potential locations
    p1 = os.path.join('ชุดข้อสอบใหม่', path)
    p2 = path
    p3 = os.path.join('ชุดข้อสอบใหม่/ชุดที่2', path)
    
    actual_path = None
    for p in [p1, p2, p3]:
        if os.path.exists(p):
            actual_path = p
            break
            
    print(f"Row {r} | ID: {sample_id} | Path in sheet: {path} | Found at: {actual_path} | Human: {h_score}")
