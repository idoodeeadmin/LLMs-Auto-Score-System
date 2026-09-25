import asyncio
import os
import sys
from pathlib import Path
from dotenv import load_dotenv

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")
sys.stdout.reconfigure(encoding="utf-8")

import openpyxl
from server.services.openai_grading import score_with_openai

Q3_QUESTION_TEXT = "การใช้ลิ้งค์ลิสต์เป็นสแตกและคิวจะแตกต่างจากการใช้อาร์เรย์อย่างไร และมีข้อดีข้อเสียอย่างไร (1 คะแนน)"

Q3_ANSWER_KEY = (
    "เฉลยและแนวคำตอบ:\n"
    "1. ความแตกต่าง (0.50 คะแนน): Linked List ขนาดปรับเปลี่ยนได้ตามจริง (Dynamic) ไม่ต้องกำหนดขนาดล่วงหน้า ส่วน Array ขนาดคงที่ (Static/Fixed) จองพื้นที่ล่วงหน้า\n"
    "2. ข้อดี (0.25 คะแนน): Linked List ขนาดไม่จำกัด ไม่เกิด overflow / Array เข้าถึงข้อมูลได้เร็ว\n"
    "3. ข้อเสีย (0.25 คะแนน): Linked List เปลืองเมมโมรี่พอยน์เตอร์ เข้าถึงช้ากว่า / Array ขนาดคงที่เสี่ยง overflow หรือเปลืองเนื้อที่"
)

USER_COMPONENT_RUBRIC_NAME = "เกณฑ์แยก 3 ส่วน: ความแตกต่าง (0.50) + ข้อดี (0.25) + ข้อเสีย (0.25)"

USER_COMPONENT_RUBRIC_DESC = (
    "เกณฑ์การให้คะแนนรวม 1.00 คะแนน ตัดสินโดยการรวมคะแนนจาก 3 ส่วนที่นักเรียนเขียนตอบ (คะแนนสุดท้ายต้องเป็น 0.00, 0.25, 0.50, 0.75 หรือ 1.00):\n\n"
    "ส่วนที่ 1: การอธิบายความแตกต่าง (เต็ม 0.50 คะแนน)\n"
    "- ได้ 0.50 คะแนน: มีการเขียนอธิบายความแตกต่างของโครงสร้างข้อมูลได้พอเข้าใจ (เช่น Linked List ปรับเปลี่ยนขนาดได้ ยืดหยุ่น ไม่จำกัดขนาด ส่วน Array มีขนาดคงที่ แน่นอน ต้องจองพื้นที่ล่วงหน้า หรืออธิบายกลไก Stack/Queue) โดยให้อนุโลมปล่อยผ่านหากไม่ได้ตอบมั่ว\n"
    "- ได้ 0.00 คะแนน: ไม่ได้เขียนอธิบายความแตกต่าง หรือตอบมั่วออกนอกเรื่อง (เช่น ประหยัดพลังงาน, รถไฟ)\n\n"
    "ส่วนที่ 2: การระบุข้อดี (เต็ม 0.25 คะแนน)\n"
    "- ได้ 0.25 คะแนน: มีการระบุข้อดี (เช่น เขียนหัวข้อ 'ข้อดี' หรือระบุจุดเด่น เช่น ค้นหาเร็ว ไม่จำกัดขนาด ยืดหยุ่น ใช้งานได้หลากหลาย) ได้พอเข้าใจ โดยให้อนุโลมปล่อยผ่านหากไม่ได้ตอบมั่ว\n"
    "- ได้ 0.00 คะแนน: ไม่ได้ระบุข้อดี หรือตอบมั่ว\n\n"
    "ส่วนที่ 3: การระบุข้อเสีย (เต็ม 0.25 คะแนน)\n"
    "- ได้ 0.25 คะแนน: มีการระบุข้อเสีย (เช่น เขียนหัวข้อ 'ข้อเสีย' หรือระบุจุดด้อย เช่น index out of bounds, ซับซ้อนสับสน, เข้าถึงช้า, เปลืองพื้นที่) ได้พอเข้าใจ โดยให้อนุโลมปล่อยผ่านหากไม่ได้ตอบมั่ว\n"
    "- ได้ 0.00 คะแนน: ไม่ได้ระบุข้อเสีย (เช่น เขียนเครื่องหมายขีด - หรือไม่ได้เขียนถึงข้อเสียเลย) หรือตอบมั่ว\n\n"
    "**ข้อกำหนดสำคัญ**:\n"
    "1. ห้ามนำข้อความในส่วนความแตกต่างมาเหมาให้คะแนนข้อดีและข้อเสียซ้ำซ้อนโดยอัตโนมัติ หากนักเรียนเขียนเฉพาะความแตกต่างแต่ไม่ได้แจกแจงข้อดีข้อเสีย ให้ได้เฉพาะคะแนนส่วนที่ 1 (0.50)\n"
    "2. หากนักเรียนเขียนเฉพาะข้อดีและข้อเสียสั้นๆ แต่ไม่ได้อธิบายความแตกต่างของโครงสร้างข้อมูล จะไม่ได้คะแนนส่วนที่ 1 (0.50) แต่จะได้คะแนนเฉพาะส่วนข้อดีหรือข้อเสียที่มี (0.25)"
)

RUBRICS = [
    {
        "name": USER_COMPONENT_RUBRIC_NAME,
        "score": 1.0,
        "description": USER_COMPONENT_RUBRIC_DESC,
    }
]

async def test():
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["ชุดข้อสอบ_dataset"]
    
    first10 = []
    for r in range(74, 84):
        sid = ws.cell(r, 1).value
        ans = str(ws.cell(r, 6).value or "")
        h = float(ws.cell(r, 7).value or 0.0)
        first10.append({"sid": sid, "ans": ans, "human": h})
        
    print(f"=== ทดสอบเกณฑ์แยก 3 ส่วน (ความแตกต่าง 0.50 + ข้อดี 0.25 + ข้อเสีย 0.25) บน 10 ข้อแรก ===")
    
    sem = asyncio.Semaphore(5)
    async def grade_one(item):
        async with sem:
            res = await score_with_openai(
                question_text=Q3_QUESTION_TEXT,
                answer_text=item["ans"],
                max_score=1.0,
                answer_key=Q3_ANSWER_KEY,
                rubrics=RUBRICS,
                strict_rubric_enforcement=True,
            )
            s = float(res.get("score", 0.0))
            diff = round(s - item["human"], 2)
            match = (s == item["human"])
            icon = "🟢 ตรงเป๊ะ" if match else f"🟠 ต่าง ({diff:+0.2f})"
            print(f"{item['sid']} | อาจารย์: {item['human']:.2f} | AI: {s:.2f} | {icon}")
            return match

    tasks = [grade_one(x) for x in first10]
    matches = await asyncio.gather(*tasks)
    exact = sum(1 for m in matches if m)
    print(f"\nผลรวม 10 ข้อแรก: ตรงเป๊ะ {exact}/10 ({exact*10}%)")

if __name__ == "__main__":
    asyncio.run(test())
