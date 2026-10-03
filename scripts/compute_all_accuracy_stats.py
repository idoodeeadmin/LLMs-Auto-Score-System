import openpyxl
import math
import json
import sys
from pathlib import Path

sys.stdout.reconfigure(encoding="utf-8")
ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
ws = wb["ชุดข้อสอบ_dataset"]

def calc_qwk(y_t, y_p, max_s=1.0, step=0.25):
    if len(y_t) != len(y_p) or len(y_t) == 0:
        return 0.0
    if y_t == y_p:
        return 1.0
    k = int(round(max_s / step)) + 1
    w = [[((i - j) ** 2) / ((k - 1) ** 2) for j in range(k)] for i in range(k)]
    cat_t = [min(k - 1, max(0, int(round(v / step)))) for v in y_t]
    cat_p = [min(k - 1, max(0, int(round(v / step)))) for v in y_p]
    o = [[0] * k for _ in range(k)]
    for t, p in zip(cat_t, cat_p):
        o[t][p] += 1
    r_t = [sum(o[i][j] for j in range(k)) for i in range(k)]
    r_p = [sum(o[i][j] for i in range(k)) for j in range(k)]
    n = len(y_t)
    e = [[(r_t[i] * r_p[j]) / n for j in range(k)] for i in range(k)]
    num = sum(w[i][j] * o[i][j] for i in range(k) for j in range(k))
    den = sum(w[i][j] * e[i][j] for i in range(k) for j in range(k))
    return 1.0 - (num / den) if den != 0 else 1.0

def calc_pearson(x, y):
    n = len(x)
    if n == 0:
        return 0.0
    mean_x = sum(x) / n
    mean_y = sum(y) / n
    cov = sum((a - mean_x) * (b - mean_y) for a, b in zip(x, y))
    var_x = sum((a - mean_x) ** 2 for a in x)
    var_y = sum((b - mean_y) ** 2 for b in y)
    denom = math.sqrt(var_x * var_y)
    return cov / denom if denom != 0 else 0.0

q_meta = {
    1: {"topic": "Row-major vs Column-major", "type": "Text (ข้อความ)", "max_s": 2.0},
    2: {"topic": "O(n log n) vs O(n^2) Complexity", "type": "Text (ข้อความ)", "max_s": 2.0},
    3: {"topic": "Linked List vs Array (Stack/Queue)", "type": "Text (ข้อความ)", "max_s": 1.0},
    4: {"topic": "Binary Search Tree Construction", "type": "Vision (ภาพลายมือ)", "max_s": 1.0},
    5: {"topic": "Infix to Prefix & Postfix Expression", "type": "Vision (ภาพลายมือ)", "max_s": 1.0},
    6: {"topic": "General Tree to Binary Tree (LCRS)", "type": "Vision (ภาพลายมือ)", "max_s": 1.0},
}

q_data = {q: {"h": [], "a": []} for q in range(1, 7)}
all_h, all_a = [], []

for r in range(6, ws.max_row + 1):
    sid = ws.cell(r, 1).value
    if not sid:
        continue
    q_no = int(ws.cell(r, 2).value)
    h = float(ws.cell(r, 7).value or 0.0)
    a = float(ws.cell(r, 8).value or 0.0)
    
    q_data[q_no]["h"].append(h)
    q_data[q_no]["a"].append(a)
    all_h.append(h)
    all_a.append(a)

per_question = {}
for q, d in sorted(q_data.items()):
    h, a = d["h"], d["a"]
    n = len(h)
    exact = sum(1 for x, y in zip(h, a) if round(x, 2) == round(y, 2))
    within_025 = sum(1 for x, y in zip(h, a) if abs(x - y) <= 0.2501)
    within_050 = sum(1 for x, y in zip(h, a) if abs(x - y) <= 0.5001)
    mae = sum(abs(x - y) for x, y in zip(h, a)) / n
    rmse = math.sqrt(sum((x - y) ** 2 for x, y in zip(h, a)) / n)
    r = calc_pearson(h, a)
    max_s = q_meta[q]["max_s"]
    qwk = calc_qwk(h, a, max_s=max_s, step=0.25)
    
    per_question[q] = {
        "question_no": q,
        "topic": q_meta[q]["topic"],
        "type": q_meta[q]["type"],
        "max_score": max_s,
        "n": n,
        "exact_count": exact,
        "exact_pct": round((exact / n) * 100, 2),
        "within_025_count": within_025,
        "within_025_pct": round((within_025 / n) * 100, 2),
        "within_050_count": within_050,
        "within_050_pct": round((within_050 / n) * 100, 2),
        "mae": round(mae, 4),
        "rmse": round(rmse, 4),
        "pearson_r": round(r, 4),
        "qwk": round(qwk, 4),
    }

# Subgroups
text_h = q_data[1]["h"] + q_data[2]["h"] + q_data[3]["h"]
text_a = q_data[1]["a"] + q_data[2]["a"] + q_data[3]["a"]
img_h = q_data[4]["h"] + q_data[5]["h"] + q_data[6]["h"]
img_a = q_data[4]["a"] + q_data[5]["a"] + q_data[6]["a"]

def group_stats(h, a, name):
    n = len(h)
    exact = sum(1 for x, y in zip(h, a) if round(x, 2) == round(y, 2))
    within_025 = sum(1 for x, y in zip(h, a) if abs(x - y) <= 0.2501)
    within_050 = sum(1 for x, y in zip(h, a) if abs(x - y) <= 0.5001)
    mae = sum(abs(x - y) for x, y in zip(h, a)) / n
    rmse = math.sqrt(sum((x - y) ** 2 for x, y in zip(h, a)) / n)
    r = calc_pearson(h, a)
    return {
        "name": name,
        "n": n,
        "exact_count": exact,
        "exact_pct": round((exact / n) * 100, 2),
        "within_025_count": within_025,
        "within_025_pct": round((within_025 / n) * 100, 2),
        "within_050_count": within_050,
        "within_050_pct": round((within_050 / n) * 100, 2),
        "mae": round(mae, 4),
        "rmse": round(rmse, 4),
        "pearson_r": round(r, 4),
    }

text_stats = group_stats(text_h, text_a, "กลุ่มข้อสอบข้อเขียน Text (ข้อ 1-3)")
img_stats = group_stats(img_h, img_a, "กลุ่มข้อสอบรูปภาพลายมือ Vision (ข้อ 4-6)")
total_stats = group_stats(all_h, all_a, "ภาพรวมทั้งชุดข้อสอบ (ข้อ 1-6 รวม 204 ข้อ)")

summary_report = {
    "per_question": per_question,
    "text_group": text_stats,
    "image_group": img_stats,
    "overall": total_stats,
}

out_json = ROOT / "artifacts" / "accuracy_summary_1_to_6.json"
out_json.parent.mkdir(parents=True, exist_ok=True)
with open(out_json, "w", encoding="utf-8") as f:
    json.dump(summary_report, f, ensure_ascii=False, indent=2)

print(json.dumps(summary_report, ensure_ascii=False, indent=2))
