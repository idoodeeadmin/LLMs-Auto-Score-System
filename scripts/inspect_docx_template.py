# -*- coding: utf-8 -*-
import sys
import docx
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

DOCX_PATH = "docs_and_tests/LLMs_Auto_Score_Systems_Pro1_1-3.docx"

doc = docx.Document(DOCX_PATH)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print(f"Total sections: {len(doc.sections)}")

print("\n=== SECTIONS & MARGINS ===")
for i, s in enumerate(doc.sections):
    print(f"Section {i+1}:")
    print(f"  Page: {s.page_width.cm:.2f} x {s.page_height.cm:.2f} cm")
    print(f"  Margins: Top={s.top_margin.cm:.2f} cm, Bottom={s.bottom_margin.cm:.2f} cm, Left={s.left_margin.cm:.2f} cm, Right={s.right_margin.cm:.2f} cm")
    print(f"  Header distance: {s.header_distance.cm:.2f} cm, Footer distance: {s.footer_distance.cm:.2f} cm")
    print(f"  Different first page header: {s.different_first_page_header_footer}")

print("\n=== STYLES DEFINED ===")
for s in doc.styles:
    if s.type == docx.enum.style.WD_STYLE_TYPE.PARAGRAPH and s.builtin is False:
        print(f"Custom Paragraph Style: {s.name}")

print("\n=== SAMPLE HEADINGS AND BODY PARAGRAPHS ===")
captions = []
headings = []
for idx, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    if any(txt.startswith(k) for k in ['บทที่', '1.', '2.', '3.', 'ตารางที่', 'รูปที่', 'ภาพที่', 'ภาพประกอบที่']):
        font_name = None
        font_size = None
        bold = None
        if p.runs:
            font_name = p.runs[0].font.name
            font_size = p.runs[0].font.size.pt if p.runs[0].font.size else None
            bold = p.runs[0].bold
        p_info = {
            'idx': idx,
            'text': txt[:70],
            'style': p.style.name if p.style else None,
            'font': font_name,
            'size': font_size,
            'bold': bold,
            'align': str(p.alignment),
            'first_line_indent': f"{p.paragraph_format.first_line_indent.cm:.2f} cm" if p.paragraph_format.first_line_indent else None,
            'space_before': f"{p.paragraph_format.space_before.pt:.1f} pt" if p.paragraph_format.space_before else None,
            'space_after': f"{p.paragraph_format.space_after.pt:.1f} pt" if p.paragraph_format.space_after else None,
            'line_spacing': p.paragraph_format.line_spacing
        }
        if any(txt.startswith(k) for k in ['ตารางที่', 'รูปที่', 'ภาพที่', 'ภาพประกอบที่']):
            captions.append(p_info)
        elif any(txt.startswith(k) for k in ['บทที่', '1.', '2.', '3.']):
            headings.append(p_info)

print(f"\nFound {len(headings)} heading paragraphs. Showing first 15 and last 10:")
for h in headings[:15]:
    print(" ", h)
print("  ...")
for h in headings[-10:]:
    print(" ", h)

print(f"\nFound {len(captions)} caption paragraphs. Showing samples:")
for c in captions[:10]:
    print(" ", c)
if len(captions) > 10:
    print("  ...")
    for c in captions[-5:]:
        print(" ", c)

print("\n=== LAST 20 PARAGRAPHS IN DOCUMENT ===")
for idx in range(max(0, len(doc.paragraphs)-20), len(doc.paragraphs)):
    p = doc.paragraphs[idx]
    if p.text.strip():
        print(f"[{idx}] style={p.style.name if p.style else None}: {p.text.strip()[:80]}")
