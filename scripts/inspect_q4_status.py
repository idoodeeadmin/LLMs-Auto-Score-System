import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('ชุดข้อสอบใหม่/ชุดข้อสอบ_dataset.xlsx', data_only=True)
ws = wb['ชุดข้อสอบ_dataset']

q4_data = []
for r in range(108, 142):
    sid = ws.cell(r, 1).value
    ans = ws.cell(r, 6).value
    h = float(ws.cell(r, 7).value or 0)
    ai = float(ws.cell(r, 8).value or 0)
    conf = ws.cell(r, 9).value
    fb = ws.cell(r, 10).value
    q4_data.append({'row': r, 'sid': sid, 'h': h, 'ai': ai, 'conf': conf, 'fb': fb, 'ans': ans})

exact = sum(1 for x in q4_data if x['h'] == x['ai'])
diff025 = sum(1 for x in q4_data if abs(x['h'] - x['ai']) <= 0.25)
diff05 = sum(1 for x in q4_data if abs(x['h'] - x['ai']) <= 0.5)
mae = sum(abs(x['h'] - x['ai']) for x in q4_data) / len(q4_data)

print('========================================================================')
print('QUESTION 4 STATUS IN EXCEL (34 Students)')
print('========================================================================')
print(f'Total students: {len(q4_data)}')
print(f'Exact Match:    {exact}/{len(q4_data)} ({exact/len(q4_data)*100:.1f}%)')
print(f'Diff <= 0.25:   {diff025}/{len(q4_data)} ({diff025/len(q4_data)*100:.1f}%)')
print(f'Diff <= 0.50:   {diff05}/{len(q4_data)} ({diff05/len(q4_data)*100:.1f}%)')
print(f'MAE:            {mae:.4f}')

from collections import Counter
print(f'Human Distribution: {Counter(x["h"] for x in q4_data)}')
print(f'AI Distribution   : {Counter(x["ai"] for x in q4_data)}')
print('------------------------------------------------------------------------')

mismatches = [x for x in q4_data if x['h'] != x['ai']]
print(f'Mismatches ({len(mismatches)} students):')
for m in mismatches:
    delta = m['ai'] - m['h']
    print(f"  {m['sid']} | Human={m['h']:.2f} | AI={m['ai']:.2f} | Diff={delta:+.2f}")
    fb_teacher = str(m['fb']).split('[สำหรับนักเรียน]')[0].replace('[สำหรับผู้สอน]', '').strip().replace('\n', ' ')
    print(f"    Feedback: {fb_teacher[:140]}...")
