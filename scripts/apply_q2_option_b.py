import json
import shutil
import sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding="utf-8")

def update_q2_option_b():
    results_path = ROOT / "docs_and_tests" / "gemini38_regrade" / "q2_option_b_results.json"
    with open(results_path, encoding="utf-8") as f:
        q2_results = json.load(f)
    
    results_map = {r["sample_id"]: r for r in q2_results}
    
    excel_path = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
    wb = openpyxl.load_workbook(excel_path)
    
    # 1. Update Exam_Rubrics Row 7
    s_rubric = wb["Exam_Rubrics"]
    q2_rubric_desc = (
        "โจทย์: อธิบายว่าทำไม O(n log n) เหมาะกับข้อมูลใหญ่กว่า O(n^2) และยกตัวอย่าง Algorithm (คะแนนเต็ม 2.00 คะแนน)\n"
        "เกณฑ์การประเมินตามแนวนโยบายการตรวจของผู้สอน (Pedagogical Policy) สำหรับข้อสอบอัตนัยระดับปริญญาตรี (ระดับคะแนน: 2.0, 1.5, 1.0, หรือ 0.0):\n\n"
        "• 2.00 คะแนน (คะแนนเต็ม - สมบูรณ์เชิงมโนทัศน์):\n"
        "  1. มีการระบุชื่อ Algorithm ที่สอดคล้องอย่างถูกต้อง (เช่น Merge Sort, Quick Sort, หรือเทียบกับ Bubble/Insertion/Selection Sort ฝั่ง O(n^2))\n"
        "  2. มีคำอธิบายเปรียบเทียบว่า O(n log n) เร็วกว่า / มีขั้นตอนการทำงานน้อยกว่า / อัตราการเติบโตช้ากว่า O(n^2) เมื่อข้อมูลมีขนาดใหญ่\n"
        "  *แนวนโยบายการตรวจของผู้สอน: อาจารย์เน้นประเมินความเข้าใจเชิงมโนทัศน์ (Conceptual Understanding) เป็นหลัก ยอมรับการอธิบายด้วยภาษาธรรมชาติของผู้เรียน ให้อนุโลมคะแนนเต็มได้แม้ผู้เรียนจะมีการอธิบายบริบทการทำงานเสริมเรื่องหน่วยความจำ ระบบประมวลผล หรือความสัมพันธ์เชิงโครงสร้าง ขอเพียงมีชื่ออัลกอริทึมและใจความหลักเรื่องประสิทธิภาพความเร็วครบถ้วน*\n"
        "  *หรือ อธิบายกลไก Divide & Conquer (การแบ่งย่อยปัญหาทำงานทีละส่วนทำให้ประมวลผลเร็วกว่าลูปซ้ำซ้อน) ได้อย่างชัดเจนและสมบูรณ์*\n\n"
        "• 1.50 คะแนน (ระดับดี - มีสาระสำคัญและตัวอย่างชัดเจน):\n"
        "  - มีการเขียนโค้ดตัวอย่างประกอบ (เช่น โค้ด recursive function หรือการเปรียบเทียบลูป) เพื่อแสดงความพยายามอธิบายกลไกการทำงานเชิงปฏิบัติ\n"
        "  - มีคำอธิบายเชิงกลไกเรื่องการทำงานไม่ซ้ำซ้อน vs ลูปซ้อนลูป (Nested loops) อย่างชัดเจน แม้ไม่ได้ระบุชื่อ Algorithm เฉพาะเจาะจง\n"
        "  - มีชื่อ Algorithm แต่คำอธิบายสั้นกระชับมาก หรือมีจุดคลาดเคลื่อนทางเทคนิคเล็กน้อย\n\n"
        "• 1.00 คะแนน (ระดับพื้นฐาน - ตอบถูกเพียงด้านใดด้านหนึ่ง):\n"
        "  - ตอบเฉพาะเหตุผลว่าเร็วกว่า/ดีกว่า หรือแสดงการคำนวณตัวเลข/ตารางเปรียบเทียบ แต่ไม่มีชื่อ Algorithm ประกอบ\n"
        "  - มีชื่อ Algorithm แต่เขียนเพียงทวนคำถามซ้ำว่ามีประสิทธิภาพดีกว่าโดยไม่อธิบายเหตุผลรองรับ\n"
        "  - ตอบสั้นมาก แต่พอมีสาระสำคัญที่ถูกทางบางส่วน (เช่น ตอบเรื่องการลดลูปซ้ำซ้อน)\n\n"
        "• 0.00 คะแนน (ไม่ผ่านเกณฑ์):\n"
        "  - ไม่ตอบ หรือตอบผิดหลักการทั้งหมดอย่างสิ้นเชิง ไม่มีความเกี่ยวข้องกับเนื้อหาข้อสอบ"
    )
    s_rubric.cell(7, 4).value = "ความซับซ้อน O(n log n) vs O(n^2) และตัวอย่างอัลกอริทึม"
    s_rubric.cell(7, 6).value = q2_rubric_desc
    print("Updated Exam_Rubrics row 7 (Question 2).")
    
    # 2. Update ชุดข้อสอบ_dataset rows 40 to 73
    s_data = wb["ชุดข้อสอบ_dataset"]
    updated_count = 0
    for r in range(40, 74):
        sid = s_data.cell(r, 1).value
        if sid in results_map:
            res = results_map[sid]
            new_score = res["new_ai"]
            diff = res["diff"]
            tf = res.get("teacher_fb", "").strip()
            sf = res.get("student_fb", "").strip()
            
            fb_text = f"[สำหรับผู้สอน] {tf}\n\n[สำหรับนักเรียน] {sf}" if sf else tf
            
            s_data.cell(r, 8).value = new_score
            s_data.cell(r, 9).value = diff
            s_data.cell(r, 10).value = fb_text
            updated_count += 1
            
    print(f"Updated {updated_count} rows in ชุดข้อสอบ_dataset (Rows 40-73).")
    
    wb.save(excel_path)
    wb.close()
    print(f"Saved changes to {excel_path}")
    
    # 3. Update เกณฑ์ตรวจสำหรับAI.xlsx
    ai_rubric_file = ROOT / "ชุดข้อสอบใหม่" / "เกณฑ์ตรวจสำหรับAI.xlsx"
    if ai_rubric_file.exists():
        wb_ai = openpyxl.load_workbook(ai_rubric_file)
        s_ai = wb_ai.active
        for r in range(40, 74):
            s_ai.cell(r, 2).value = q2_rubric_desc
        wb_ai.save(ai_rubric_file)
        wb_ai.close()
        print("Updated เกณฑ์ตรวจสำหรับAI.xlsx for Question 2.")
        
    # 4. Copy to public/
    public_file = ROOT / "public" / "ชุดข้อสอบ_dataset.xlsx"
    shutil.copy2(excel_path, public_file)
    print(f"Copied updated excel to {public_file}")

if __name__ == "__main__":
    update_q2_option_b()
