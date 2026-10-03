import asyncio
import json
import os
import sys
from pathlib import Path
from dotenv import load_dotenv
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
load_dotenv(ROOT / ".env")
sys.stdout.reconfigure(encoding="utf-8")

from server.services.openai_grading import score_with_openai
from scripts.run_dataset_benchmark import calculate_qwk, calculate_mae

EXCEL_PATH = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"

# 1. Q1 Config
Q1_RUBRIC = [
    {
        "name": "Row-major",
        "score": 1.0,
        "description": "อธิบาย Row-major ว่าเป็นการจัดเก็บข้อมูล ลำดับการเรียง หรือการหาตำแหน่ง address ตามแนวแถว (Row/แนวนอน/แกน X/หรือ [i][j])",
        "allowed_scores": [0.0, 1.0]
    },
    {
        "name": "Column-major",
        "score": 1.0,
        "description": "อธิบาย Column-major ว่าเป็นการจัดเก็บข้อมูล ลำดับการเรียง หรือการหาตำแหน่ง address ตามแนวคอลัมน์ (Column/แนวตั้ง/แกน Y/หรือ [j][i])",
        "allowed_scores": [0.0, 1.0]
    }
]
Q1_ANSWER_KEY = (
    "Row-major คือการจัดเก็บข้อมูล ลำดับการเรียง หรือคำนวณตำแหน่ง address โดยอิงตามแถว (Row หรือแนวนอน/แกน X) "
    "ส่วน Column-major คือการจัดเก็บข้อมูล ลำดับการเรียง หรือคำนวณตำแหน่ง address โดยอิงตามคอลัมน์ (Column หรือแนวตั้ง/แกน Y) "
    "ต่างกันที่ลำดับและมิติการเรียงข้อมูลในหน่วยความจำ"
)

# 2. Q2 Config
Q2_RUBRIC = [
    {
        "name": "การเปรียบเทียบ O(n log n) vs O(n^2) และตัวอย่าง Algorithm",
        "score": 2.0,
        "description": (
            "ประเมินตามระดับคะแนน 5 ระดับ (2.0, 1.5, 1.0, 0.5, 0.0) อย่างเคร่งครัด:\n"
            "• 2.0 คะแนน: อธิบายได้ว่าทำไม O(n log n) ดีกว่าเมื่อข้อมูลใหญ่ เช่น เร็วกว่า, จำนวนรอบ/การทำงานโตช้ากว่า, n² โตเร็วมาก และ มีตัวอย่าง Algorithm หรืออธิบายกลไกชัดเจน\n"
            "• 1.5 คะแนน: เข้าใจแก่นว่า O(n log n) มีประสิทธิภาพกว่า แต่คำอธิบายยังไม่ครบ/ไม่ยกตัวอย่าง หรือมีตัวอย่างแต่เหตุผลยังไม่ลงลึก\n"
            "• 1.0 คะแนน: รู้เพียงว่า O(n log n) “เร็วกว่า/ดีกว่า/ซ้ำซ้อนน้อยกว่า” แต่ไม่ได้อธิบายกลไกอย่างชัดเจน หรือคำอธิบายคลุมเครือ และไม่ได้ยกตัวอย่าง (หรือมีเพียงตัวอย่างคำนวณ)\n"
            "• 0.5 คะแนน: ระบุได้เพียงตัวอย่าง Algorithm ที่เกี่ยวข้องอย่างถูกต้อง หรือกล่าวถึงความซับซ้อน/ประสิทธิภาพเพียงฝั่งเดียวอย่างสั้นๆ โดยยังไม่สื่อชัดว่า O(n log n) เหมาะกับข้อมูลใหญ่กว่า O(n²) อย่างไร\n"
            "• 0.0 คะแนน: ไม่สามารถอธิบายความสัมพันธ์ของ O(n log n) กับ O(n²) ได้อย่างมีสาระ ตอบผิดหลักการสำคัญ หรือไม่ตอบ\n"
            "(ระดับคะแนนที่ให้ได้คือ 2.0, 1.5, 1.0, 0.5 หรือ 0.0 คะแนนเท่านั้น ห้ามให้เศษทศนิยมอื่น)"
        ),
        "allowed_scores": [0.0, 0.5, 1.0, 1.5, 2.0]
    }
]
Q2_ANSWER_KEY = (
    "O(n log n) มีอัตราการเติบโตของเวลาในการทำงาน (Growth rate) ช้ากว่า O(n^2) มากเมื่อข้อมูลมีขนาดใหญ่ขึ้น ทำให้ใช้เวลาประมวลผลน้อยกว่า "
    "ตัวอย่าง O(n log n) เช่น Merge Sort, Quick Sort, Heap Sort และ O(n^2) เช่น Bubble Sort, Selection Sort, Insertion Sort"
)

# 3. Q3 Config
from scripts.test_q3_calibrated_v5 import Q3_RUBRIC_NAME, Q3_RUBRIC_DESC, Q3_ANSWER_KEY
Q3_RUBRIC = [
    {
        "name": Q3_RUBRIC_NAME,
        "score": 1.0,
        "description": Q3_RUBRIC_DESC,
        "allowed_scores": [0.0, 0.25, 0.5, 0.75, 1.0]
    }
]

async def grade_items(items, q_no, rubric, answer_key, max_score, allowed_scores, sem):
    results = []
    async def grade_one(it):
        async with sem:
            res = await score_with_openai(
                question_text=it["question_content"],
                answer_text=it["student_answer"],
                max_score=max_score,
                answer_key=answer_key,
                rubrics=rubric,
                allowed_scores=allowed_scores
            )
            s = float(res.get("score", 0.0))
            h = float(it["human_score"])
            diff = round(s - h, 2)
            match = "MATCH" if diff == 0.0 else f"diff={diff:+.2f}"
            print(f"[{it['sample_id']}] Human: {h:.2f} | AI: {s:.2f} | {match}", flush=True)
            return {
                "row": it["row"],
                "sample_id": it["sample_id"],
                "human": h,
                "ai": s,
                "diff": diff,
                "confidence": res.get("confidence", "high"),
                "teacher_feedback": res.get("teacher_feedback", "").strip(),
                "student_feedback": res.get("student_feedback", "").strip(),
                "metrics": res.get("metrics", {})
            }

    graded = await asyncio.gather(*(grade_one(it) for it in items))
    return sorted(graded, key=lambda x: x["row"])

async def main():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb["ชุดข้อสอบ_dataset"]
    
    # Load all rows 6 to 107 (Questions 1, 2, 3)
    q_data = {1: [], 2: [], 3: []}
    for r in range(6, 108):
        q_no = ws.cell(r, 2).value
        if q_no in [1, 2, 3]:
            q_data[q_no].append({
                "row": r,
                "sample_id": ws.cell(r, 1).value,
                "question_content": ws.cell(r, 4).value,
                "student_answer": ws.cell(r, 6).value,
                "human_score": ws.cell(r, 7).value,
                "ai_score": ws.cell(r, 8).value
            })

    sem = asyncio.Semaphore(4)

    all_results = {}
    print("\n=================== Grading Question 1 ===================")
    all_results[1] = await grade_items(q_data[1], 1, Q1_RUBRIC, Q1_ANSWER_KEY, 2.0, [0.0, 1.0, 2.0], sem)

    print("\n=================== Grading Question 2 ===================")
    all_results[2] = await grade_items(q_data[2], 2, Q2_RUBRIC, Q2_ANSWER_KEY, 2.0, [0.0, 0.5, 1.0, 1.5, 2.0], sem)

    print("\n=================== Grading Question 3 ===================")
    all_results[3] = await grade_items(q_data[3], 3, Q3_RUBRIC, Q3_ANSWER_KEY, 1.0, [0.0, 0.25, 0.5, 0.75, 1.0], sem)

    # Save to Excel
    print("\nWriting AI scores & feedback to Excel...")
    for q_no, results in all_results.items():
        for res in results:
            r = res["row"]
            ws.cell(r, 8, res["ai"])
            ws.cell(r, 9, res["confidence"])
            
            tf = res["teacher_feedback"]
            sf = res["student_feedback"]
            combined = f"[สำหรับผู้สอน]\n{tf}\n\n[สำหรับนักเรียน]\n{sf}" if (tf and sf) else (tf or sf)
            ws.cell(r, 10, combined)

    wb.save(EXCEL_PATH)
    print(f"Saved {EXCEL_PATH}")

    # Summary Statistics
    print("\n=================== BENCHMARK SUMMARY ===================")
    for q_no in [1, 2, 3]:
        res_list = all_results[q_no]
        human_scores = [r["human"] for r in res_list]
        ai_scores = [r["ai"] for r in res_list]
        exact = sum(1 for r in res_list if r["diff"] == 0.0)
        diff_05 = sum(1 for r in res_list if abs(r["diff"]) <= 0.5)
        mae = calculate_mae(human_scores, ai_scores)
        qwk = calculate_qwk(human_scores, ai_scores)
        print(f"Question {q_no} (N={len(res_list)}):")
        print(f"  Exact Match: {exact}/{len(res_list)} ({exact/len(res_list)*100:.1f}%)")
        print(f"  Diff <= 0.5: {diff_05}/{len(res_list)} ({diff_05/len(res_list)*100:.1f}%)")
        print(f"  MAE: {mae:.4f}")
        print(f"  QWK: {qwk:.4f}")

if __name__ == "__main__":
    asyncio.run(main())
