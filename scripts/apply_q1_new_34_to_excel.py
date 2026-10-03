import shutil
import sys
from pathlib import Path
import openpyxl
import json

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.stdout.reconfigure(encoding='utf-8')

EXCEL_PATH = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset.xlsx"
BACKUP_PATH = ROOT / "ชุดข้อสอบใหม่" / "ชุดข้อสอบ_dataset_backup_before_q1_34.xlsx"
JSON_PATH = ROOT / "scratch" / "system_graded_34.json"

def main():
    print(f"1. Creating backup: {BACKUP_PATH.name}...")
    shutil.copyfile(EXCEL_PATH, BACKUP_PATH)

    print("2. Loading graded data...")
    with open(JSON_PATH, "r", encoding="utf-8") as f:
        data = json.load(f)

    if len(data) != 34:
        raise ValueError(f"Expected 34 items, got {len(data)}")

    print(f"3. Loading workbook: {EXCEL_PATH}...")
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb["ชุดข้อสอบ_dataset"]

    print("4. Updating rows 6 to 39 for Question 1...")
    for i, item in enumerate(data):
        row = i + 6
        expected_id = f"DS-{i+1:03d}"
        
        ws.cell(row, 1, expected_id)
        ws.cell(row, 2, 1)
        ws.cell(row, 3, "text")
        ws.cell(row, 4, "อธิบายความต่างของ Row-major vs Column-major")
        ws.cell(row, 5, "text")
        ws.cell(row, 6, item["student_answer"])
        ws.cell(row, 7, float(item["human_score"]))
        ws.cell(row, 8, float(item["ai_score"]))
        ws.cell(row, 9, item.get("confidence", "high"))
        
        fb = item.get("feedback", "")
        if not fb:
            tf = item.get("teacher_feedback", "")
            sf = item.get("student_feedback", "")
            fb = f"[สำหรับผู้สอน]\n{tf}\n\n[สำหรับนักเรียน]\n{sf}"
        ws.cell(row, 10, fb.strip())

    print(f"5. Saving workbook: {EXCEL_PATH}...")
    wb.save(EXCEL_PATH)
    print("Update complete!")

    # Verify check
    print("\n6. Verifying updated Excel...")
    wb_check = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    ws_check = wb_check["ชุดข้อสอบ_dataset"]
    
    # Check first 3 and last 3 of Q1
    for r in [6, 7, 8, 37, 38, 39]:
        sid = ws_check.cell(r, 1).value
        qno = ws_check.cell(r, 2).value
        ans = (ws_check.cell(r, 6).value or "").replace("\n", " ")[:40]
        h = ws_check.cell(r, 7).value
        ai = ws_check.cell(r, 8).value
        conf = ws_check.cell(r, 9).value
        print(f"Row {r:02d} | {sid} | Q{qno} | H:{h} | AI:{ai} | {conf} | Ans: {ans}")

    # Check row 40 (Q2 first item) is untouched
    print("\nCheck Row 40 (Question 2 untouched):")
    print(f"Row 40 | {ws_check.cell(40, 1).value} | Q{ws_check.cell(40, 2).value} | H:{ws_check.cell(40, 7).value} | AI:{ws_check.cell(40, 8).value}")

if __name__ == "__main__":
    main()
